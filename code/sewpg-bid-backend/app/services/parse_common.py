"""解析链路通用件：取消检查、进度心跳、协程桥接、解析产物目录与文本归一。

来源：parsing-01 拆分，自 app/services/parsing.py 逐字搬迁，实现与行为不变；
符号经 parsing.py 门面 re-export，外部仍按 `app.services.parsing.<符号>` 访问。
"""
from __future__ import annotations

import asyncio
import re
import threading
import time
from pathlib import Path
from typing import Any, Callable

from docx import Document

from app.core.config import settings
from app.services.bid_parse_cancel import ParseCancelledError


def _raise_if_parse_cancelled(cancel_check: Callable[[], bool] | None) -> None:
    if cancel_check is not None and cancel_check():
        raise ParseCancelledError("解析已取消。")


def parsed_project_dir(project_id: str) -> Path:
    path = settings.parsed_dir / project_id
    path.mkdir(parents=True, exist_ok=True)
    return path


def parsed_appendix_path(project_id: str) -> Path:
    return settings.parsed_dir / project_id / "s1_appendices"


def _normalize_text(raw_text: str) -> str:
    lines = [line.rstrip() for line in raw_text.replace("\r\n", "\n").split("\n")]
    compact: list[str] = []
    previous_blank = False
    for line in lines:
        blank = not line.strip()
        if blank and previous_blank:
            continue
        compact.append(line)
        previous_blank = blank
    return "\n".join(compact).strip()


def _run_with_progress_heartbeat(
    operation: Callable[[], Any],
    *,
    heartbeat: Callable[[dict[str, Any]], None],
    interval_seconds: float = 10.0,
    cancel_check: Callable[[], bool] | None = None,
) -> Any:
    result: Any = None
    error: BaseException | None = None
    done = threading.Event()

    def run() -> None:
        nonlocal result, error
        try:
            result = operation()
        except BaseException as exc:  # pragma: no cover - re-raised in caller thread
            error = exc
        finally:
            done.set()

    thread = threading.Thread(target=run, daemon=True, name="parse-progress-heartbeat")
    thread.start()
    interval = max(0.001, float(interval_seconds))
    started_at = time.monotonic()
    heartbeat_index = 0
    while not done.wait(interval):
        _raise_if_parse_cancelled(cancel_check)
        heartbeat_index += 1
        heartbeat(
            {
                "heartbeat": True,
                "heartbeatIndex": heartbeat_index,
                "elapsedSeconds": max(0, round(time.monotonic() - started_at)),
            }
        )
    thread.join()
    _raise_if_parse_cancelled(cancel_check)
    if error:
        raise error
    return result


def _run_coroutine_blocking(coro: Any) -> Any:
    """同步上下文执行协程的本模块既有桥接（OCR 识别与 opencode 引擎调用共用）。

    无线程外事件循环时直接 asyncio.run；已在事件循环里（如本地内联执行解析任务
    跑在请求线程上）则开新线程跑独立循环并 join——保持原同步阻塞语义。
    """
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        return asyncio.run(coro)

    result: Any = None
    error: BaseException | None = None

    def run() -> None:
        nonlocal result, error
        try:
            result = asyncio.run(coro)
        except BaseException as exc:  # pragma: no cover - defensive bridge
            error = exc

    thread = threading.Thread(target=run, daemon=True, name="parse-blocking-bridge")
    thread.start()
    thread.join()
    if error:
        raise error
    return result


# parsing-01 接缝收口：docx 区段/商务格式区识别原在 parse_appendix，parse_business_fields
# 曾靠运行时注入模块全局引用；现统一下沉到本模块，两侧正常 import。
# WORD_NAMESPACE 原在 parse_extract，一并下沉（parse_extract 自本模块 import 后 re-export）。


WORD_NAMESPACE = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _looks_like_toc_or_directory_line(text: str) -> bool:
    normalized = re.sub(r"\s+", " ", str(text or "").strip().lstrip("#").strip())
    if not normalized:
        return False
    if any(token in normalized for token in ("……", "....", "----")):
        return True
    if re.search(r"(?:\s|\.{2,}|…{2,})\d{1,4}$", normalized):
        return True
    return False


def _looks_like_business_template_body_sentence(text: str) -> bool:
    normalized = str(text or "").strip().lstrip("#").strip()
    if not normalized:
        return False
    if any(mark in normalized for mark in ("。", "；", ";")):
        return True
    return any(
        token in normalized
        for token in (
            "我方",
            "本公司",
            "投标人承诺",
            "按招标文件",
            "按照招标文件",
            "愿意",
            "承担",
            "提交",
            "提供",
        )
    )


def _is_business_template_section_heading(text: str) -> bool:
    normalized = str(text or "").strip().lstrip("#").strip()
    if not normalized:
        return False
    if _looks_like_toc_or_directory_line(normalized):
        return False
    if _looks_like_business_template_body_sentence(normalized):
        return False
    if any(token in normalized for token in BUSINESS_FORMAT_START_KEYWORDS):
        return True
    return bool(re.match(r"^第[六6]章", normalized)) and any(
        token in normalized for token in ("格式", "投标文件", "响应文件", "商务文件")
    )


def _is_business_format_end_heading(text: str) -> bool:
    normalized = str(text or "").strip().lstrip("#").strip()
    if not normalized or _looks_like_toc_or_directory_line(normalized):
        return False
    return any(keyword in normalized for keyword in BUSINESS_FORMAT_END_KEYWORDS)


def _block_heading_rank(block: dict[str, Any], text: str) -> int | None:
    raw_level = block.get("headingLevel")
    if isinstance(raw_level, int):
        return raw_level
    normalized = str(text or "").strip().lstrip("#").strip()
    if re.match(r"^第[一二三四五六七八九十0-9]+章", normalized):
        return 0
    if re.match(r"^(?:[一二三四五六七八九十0-9]+[、.．]|[（(][一二三四五六七八九十0-9]+[）)])", normalized):
        return 1
    if bool(block.get("isLikelyHeading")):
        return 2
    return None


def _is_business_major_section_heading(text: str) -> bool:
    normalized = str(text or "").strip().lstrip("#").strip()
    return bool(re.match(r"^第[一二三四五六七八九十0-9]+章", normalized))


def _detect_business_format_regions(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    regions: list[dict[str, Any]] = []
    active: dict[str, Any] | None = None
    for index, block in enumerate(blocks):
        if block.get("type") != "paragraph":
            continue
        text = str(block.get("text") or "").strip()
        if not text:
            continue
        rank = _block_heading_rank(block, text)
        is_start = _is_business_template_section_heading(text)
        if is_start:
            if active is not None:
                active["end"] = max(active["start"] + 1, index)
                regions.append(active)
            active = {
                "start": index,
                "end": len(blocks),
                "title": text,
                "rank": rank if rank is not None else 0,
            }
            continue
        if active is None:
            continue
        active_rank = int(active.get("rank") if isinstance(active.get("rank"), int) else 0)
        same_or_higher_heading = rank is not None and rank <= active_rank
        if same_or_higher_heading and (_is_business_format_end_heading(text) or _is_business_major_section_heading(text)):
            active["end"] = index
            regions.append(active)
            active = None
    if active is not None:
        regions.append(active)
    return regions


def _docx_paragraph_text(element: Any) -> str:
    return "".join(node.text or "" for node in element.iter(f"{WORD_NAMESPACE}t")).strip()


def _docx_xml_attr(element: Any, name: str) -> str:
    if element is None:
        return ""
    return str(element.get(f"{WORD_NAMESPACE}{name}") or element.get(name) or "").strip()


def _docx_paragraph_xml_outline_level(element: Any) -> int | None:
    p_pr = element.find(f"{WORD_NAMESPACE}pPr")
    if p_pr is None:
        return None
    outline = p_pr.find(f"{WORD_NAMESPACE}outlineLvl")
    if outline is not None:
        raw = _docx_xml_attr(outline, "val")
        if raw.isdigit():
            return int(raw)
    p_style = p_pr.find(f"{WORD_NAMESPACE}pStyle")
    raw_style = _docx_xml_attr(p_style, "val")
    match = re.search(r"(?:Heading|标题)\s*([1-9])", raw_style, flags=re.IGNORECASE)
    if match:
        return max(0, int(match.group(1)) - 1)
    return None


def _docx_paragraph_alignment_value(element: Any, paragraph: Any) -> str:
    alignment = getattr(paragraph, "alignment", None) if paragraph is not None else None
    if alignment is not None:
        value = getattr(alignment, "value", alignment)
        if str(value) in {"1", "CENTER", "WD_ALIGN_PARAGRAPH.CENTER"}:
            return "center"
    p_pr = element.find(f"{WORD_NAMESPACE}pPr")
    jc = p_pr.find(f"{WORD_NAMESPACE}jc") if p_pr is not None else None
    raw = _docx_xml_attr(jc, "val").lower()
    return raw


def _docx_paragraph_style_level(style_name: str) -> int | None:
    normalized = str(style_name or "").strip()
    match = re.search(r"(?:Heading|标题)\s*([1-9])", normalized, flags=re.IGNORECASE)
    if match:
        return max(0, int(match.group(1)) - 1)
    return None


def _docx_paragraph_metadata(element: Any, paragraph: Any, text: str) -> dict[str, Any]:
    style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
    outline_level = _docx_paragraph_xml_outline_level(element)
    style_level = _docx_paragraph_style_level(style_name)
    heading_level = outline_level if outline_level is not None else style_level
    alignment = _docx_paragraph_alignment_value(element, paragraph)
    runs = list(getattr(paragraph, "runs", []) or []) if paragraph is not None else []
    non_empty_runs = [run for run in runs if str(getattr(run, "text", "") or "").strip()]
    bold_runs = [run for run in non_empty_runs if bool(getattr(getattr(run, "font", None), "bold", False) or getattr(run, "bold", False))]
    font_sizes = []
    for run in non_empty_runs:
        size = getattr(getattr(run, "font", None), "size", None)
        if size is not None and getattr(size, "pt", None):
            font_sizes.append(float(size.pt))
    bold_ratio = (len(bold_runs) / len(non_empty_runs)) if non_empty_runs else 0.0
    normalized = re.sub(r"\s+", "", str(text or ""))
    is_short = 0 < len(normalized) <= 48
    is_likely_heading = bool(
        heading_level is not None
        or style_name.lower().startswith("heading")
        or "标题" in style_name
        or (alignment == "center" and is_short)
        or (bold_ratio >= 0.6 and is_short)
    )
    return {
        "styleName": style_name,
        "outlineLevel": outline_level,
        "styleLevel": style_level,
        "headingLevel": heading_level,
        "alignment": alignment,
        "isCentered": alignment == "center",
        "boldRatio": round(bold_ratio, 3),
        "fontSizeMax": max(font_sizes) if font_sizes else None,
        "isLikelyHeading": is_likely_heading,
    }


def _docx_table_rows(table: Any) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def _docx_table_has_cell_merges(table_element: Any) -> bool:
    return (
        table_element.find(f".//{WORD_NAMESPACE}gridSpan") is not None
        or table_element.find(f".//{WORD_NAMESPACE}vMerge") is not None
    )


def _iter_docx_blocks(path: Path) -> list[dict[str, Any]]:
    """Walk the docx body and yield {paragraph,table} blocks.

    Each block records ``body_index`` — its position among ``body.iterchildren()`` —
    so callers can map a block back to the original ``<w:p>`` / ``<w:tbl>`` element
    when slicing the source docx (see ``_slice_appendix_from_source``)."""

    # 注意：business_material_splitter._iter_docx_blocks 是另一份变体（按 Document 对象遍历、
    # 字段为 kind/style/element，服务素材切分）；本版按路径遍历、带 body_index/元数据，
    # 服务附表识别与源文件切片。两者 block schema 不同，不合并，改动时互相看一眼。

    doc = Document(str(path))
    tables = iter(doc.tables)
    paragraphs = iter(doc.paragraphs)
    blocks: list[dict[str, Any]] = []
    for body_index, child in enumerate(doc.element.body.iterchildren()):
        if child.tag == f"{WORD_NAMESPACE}p":
            paragraph = next(paragraphs, None)
            text = _docx_paragraph_text(child)
            blocks.append(
                {
                    "type": "paragraph",
                    "text": text,
                    "body_index": body_index,
                    **_docx_paragraph_metadata(child, paragraph, text),
                }
            )
        elif child.tag == f"{WORD_NAMESPACE}tbl":
            table = next(tables, None)
            if table is not None:
                blocks.append(
                    {
                        "type": "table",
                        "rows": _docx_table_rows(table),
                        "body_index": body_index,
                        "hasCellMerges": _docx_table_has_cell_merges(child),
                    }
                )
    return blocks


BUSINESS_FORMAT_START_KEYWORDS = (
    "投标文件格式",
    "响应文件格式",
    "商务文件格式",
    "商务标格式",
    "商务响应文件格式",
)


BUSINESS_FORMAT_END_KEYWORDS = (
    "评标办法",
    "评审办法",
    "合同条款",
    "技术要求",
    "技术规范",
    "供货要求",
    "用户需求",
)
