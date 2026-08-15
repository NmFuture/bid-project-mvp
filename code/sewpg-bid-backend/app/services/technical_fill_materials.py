"""技术标 AI 填写的素材备料：下载落地、Excel 原件回源、PDF OCR sidecar、待插入嵌入备料。

从 technical_gap_ai_fill 拆出（纯搬迁，不改实现与签名）；门面
app.services.technical_gap_ai_fill re-export 本模块全部符号，
app.services.technical_gap_ai_fill.<符号> 保持可解析、可 patch。
"""
from __future__ import annotations

import asyncio
import os
import re
import threading
from pathlib import Path
from typing import Any

from app.services.identity import build_project_material_scope
from app.services.minio_client import minio_client
from app.services.ocr_service import ocr_service
from app.services.file_utils import safe_filename
from app.services.peripheral import PeripheralError
from app.services.technical_material_store import technical_material_store
from app.services.turbine_models import project_turbine_model


# 共享的专用事件循环：一键填写并发化后，多个 worker 线程都会跑到异步素材下载。
# 原来每次调用 asyncio.run 各开新 loop——共享 AsyncEngine 的 asyncpg 连接
# 在「前一线程的 loop 上创建、后一线程的 loop 上等待」时会永久挂起（实测并发批
# 填时正文模板下载卡死 2 小时+，select 空转）。改成所有调用都提交到同一个常驻
# loop（run_coroutine_threadsafe），异步原语的 loop 亲和性天然一致。
_SHARED_ASYNC_LOOP: asyncio.AbstractEventLoop | None = None
_SHARED_ASYNC_LOOP_LOCK = threading.Lock()


def _shared_async_loop() -> asyncio.AbstractEventLoop:
    global _SHARED_ASYNC_LOOP
    with _SHARED_ASYNC_LOOP_LOCK:
        if _SHARED_ASYNC_LOOP is None or _SHARED_ASYNC_LOOP.is_closed():
            loop = asyncio.new_event_loop()
            threading.Thread(
                target=loop.run_forever,
                daemon=True,
                name="technical-ai-fill-async-loop",
            ).start()
            _SHARED_ASYNC_LOOP = loop
        return _SHARED_ASYNC_LOOP


def _run_async(awaitable: Any) -> Any:
    try:
        loop = asyncio.get_running_loop()
    except RuntimeError:
        return asyncio.run_coroutine_threadsafe(awaitable, _shared_async_loop()).result()

    if loop is _SHARED_ASYNC_LOOP:
        raise RuntimeError(
            "_run_async was called from the shared event loop's own thread. "
            "Wrap the calling sync code with asyncio.to_thread or run it in a worker thread."
        )
    return asyncio.run_coroutine_threadsafe(awaitable, _shared_async_loop()).result()


def _object_items(value: Any) -> list[dict[str, Any]]:
    return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []


def _string_items(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    result: list[str] = []
    seen: set[str] = set()
    for item in value:
        text = str(item or "").strip()
        if not text or text in seen:
            continue
        seen.add(text)
        result.append(text)
    return result


def _material_key(material: dict[str, Any]) -> str:
    return str(material.get("id") or material.get("materialId") or material.get("path") or material.get("docx") or "").strip()


def _material_summary(material: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": str(material.get("id") or material.get("materialId") or ""),
        "name": str(material.get("name") or material.get("title") or material.get("fileName") or ""),
        "path": str(material.get("path") or material.get("docx") or ""),
        "folderPath": str(material.get("folderPath") or ""),
        "materialTier": str(material.get("materialTier") or material.get("materialScope") or ""),
        "usage": str(material.get("usage") or ""),
        "source": str(material.get("source") or ""),
        "hasCleanedWord": bool(material.get("hasCleanedWord")),
        "cleanedFileName": str(material.get("cleanedFileName") or ""),
        "turbineModelLabel": str(material.get("turbineModelLabel") or ""),
        "matchReason": str(material.get("matchReason") or ""),
        "turbineFit": str(material.get("turbineFit") or ""),
    }


def _dedupe_material_summaries(materials: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for material in materials:
        key = _material_key(material)
        if not key or key in seen:
            continue
        seen.add(key)
        result.append(_material_summary(material))
    return result


def effective_source_routing(item: dict[str, Any], appendix_task: dict[str, Any]) -> dict[str, Any]:
    """附表来源规则：任务级优先，目录项级兜底（与矩阵写入位置一致）。"""
    for carrier in (appendix_task, item):
        routing = carrier.get("sourceRouting") if isinstance(carrier.get("sourceRouting"), dict) else {}
        if routing.get("source") == "appendix_source_matrix":
            return routing
    return {}


def _project_tender_documents_for_fill(
    project: dict[str, Any],
    appendix_task: dict[str, Any],
) -> list[dict[str, Any]]:
    routing = appendix_task.get("sourceRouting") if isinstance(appendix_task.get("sourceRouting"), dict) else {}
    if not routing.get("useTenderParseFields"):
        return []
    # 完整文档记录（含 sourcePath/textPath）由解析链路写入 project.parse_storage.documents；
    # project.parse_result 里只有摘要（sourceFiles），旧数据可能带 documents，均作兜底。
    parse_storage = project.get("parse_storage") if isinstance(project.get("parse_storage"), dict) else {}
    documents = parse_storage.get("documents")
    if not isinstance(documents, list):
        parse_result = project.get("parse_result") if isinstance(project.get("parse_result"), dict) else {}
        documents = parse_result.get("documents")
        if not isinstance(documents, list):
            documents = parse_result.get("sourceFiles")
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for document in _object_items(documents):
        if document.get("status") == "failed":
            continue
        source_path = Path(str(document.get("sourcePath") or ""))
        text_path = Path(str(document.get("textPath") or ""))
        supported_source = source_path.suffix.lower() in {".docx", ".pdf", ".xlsx", ".xlsm", ".txt", ".md"}
        readable_source = (
            source_path
            if supported_source and source_path.is_file()
            else text_path
            if text_path.is_file()
            else None
        )
        if readable_source is None:
            continue
        document_id = str(document.get("id") or "").strip()
        key = document_id or str(readable_source.resolve())
        if key in seen:
            continue
        seen.add(key)
        nav_path = Path(str(document.get("documentNavPath") or ""))
        result.append(
            {
                "id": document_id or f"TENDER-{len(result) + 1}",
                "name": str(document.get("name") or document.get("fileName") or source_path.name or readable_source.name),
                "sourcePath": str(readable_source),
                "originalSourcePath": str(source_path) if source_path.is_file() else "",
                "textPath": str(text_path) if text_path.is_file() else "",
                "ocrTextPath": str(text_path) if text_path.is_file() else "",
                "documentNavPath": str(nav_path) if nav_path.is_file() else "",
                "pageCount": document.get("pageCount") or 0,
                "sourceType": "project_tender_document",
                "parsedTextFallback": readable_source == text_path and readable_source != source_path,
            }
        )
    return result


def _selected_reference_material_ids(
    item: dict[str, Any],
    appendix_task: dict[str, Any],
    data: dict[str, Any],
) -> list[str]:
    if "referenceMaterialIds" in data:
        return _string_items(data.get("referenceMaterialIds"))

    if effective_source_routing(item, appendix_task):
        # 有来源规则时严格执行：只用规则命中的素材，命中为空就是空，
        # 禁止回退目录项通用 matchedMaterials（否则会绕过规则填错来源）。
        return _string_items([
            _material_key(material)
            for material in (
                _object_items(appendix_task.get("recommendedMaterials"))
                + _object_items(item.get("sourceRoutedMaterials"))
            )
        ])

    matched = [_material_key(material) for material in _object_items(item.get("matchedMaterials"))]
    matched = [item for item in matched if item]
    if matched:
        return matched

    recommended = _string_items([
        _material_key(material)
        for material in _object_items(appendix_task.get("recommendedMaterials"))
    ])
    return recommended


def _reference_materials_for_fill(
    item: dict[str, Any],
    appendix_task: dict[str, Any],
    data: dict[str, Any],
    selected_ids: list[str],
) -> list[dict[str, Any]]:
    # 有来源规则时 context 同样不混入通用 matched/candidate 素材，
    # 只用规则命中与调用方显式指定的素材做 id → 素材详情解析。
    generic_pools = (
        []
        if effective_source_routing(item, appendix_task)
        else _object_items(item.get("matchedMaterials")) + _object_items(item.get("candidateMaterials"))
    )
    context = _dedupe_material_summaries(
        _object_items(data.get("referenceMaterials"))
        + generic_pools
        + _object_items(item.get("sourceRoutedMaterials"))
        + _object_items(appendix_task.get("recommendedMaterials"))
        + [
            material
            for task in _object_items(item.get("appendixTasks"))
            for material in _object_items(task.get("recommendedMaterials"))
        ]
    )
    by_id = {
        str(material.get("id") or "").strip(): material
        for material in context
        if str(material.get("id") or "").strip()
    }
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for material_id in selected_ids:
        if material_id in seen:
            continue
        seen.add(material_id)
        result.append(by_id.get(material_id) or {"id": material_id, "name": material_id})
    return result


def _allowed_technical_material_index(material_scope: dict[str, Any], turbine_model: dict[str, Any]) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    seen: set[str] = set()
    for scope in material_scope.get("readableScopes") or []:
        if not isinstance(scope, dict):
            continue
        folder_path = str(scope.get("path") or "").strip()
        if not folder_path:
            continue
        payload = _run_async(
            technical_material_store.raw_files(
                folder_path=folder_path,
                material_tier=str(scope.get("materialTier") or ""),
                turbine_model=turbine_model,
                recursive=True,
                page=1,
                page_size=1000,
            )
        )
        for raw in payload.get("items") or []:
            if not isinstance(raw, dict):
                continue
            material_id = str(raw.get("id") or "")
            if not material_id or material_id in seen:
                continue
            seen.add(material_id)
            items.append(
                {
                    "id": material_id,
                    "name": str(raw.get("name") or ""),
                    "folderPath": str(raw.get("folderPath") or ""),
                    "materialTier": str(raw.get("materialTier") or scope.get("materialTier") or ""),
                    "hasCleanedWord": bool(raw.get("hasCleanedWord")),
                    "cleanedFileName": str(raw.get("cleanedFileName") or ""),
                    "cleanStatus": str(raw.get("cleanStatus") or ""),
                    "turbineModelLabel": str(raw.get("turbineModelLabel") or ""),
                    "updatedAt": str(raw.get("updatedAt") or ""),
                }
            )
    return items


def _material_index_for_fill(
    project: dict[str, Any],
    plan: dict[str, Any],
    item: dict[str, Any],
    selected_ids: list[str],
    reference_materials: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    _ = project
    context = _dedupe_material_summaries(
        reference_materials
        + _object_items(item.get("sourceRoutedMaterials"))
        + _object_items(plan.get("materialIndex"))
        + _object_items(item.get("candidateMaterials"))
        + _object_items(item.get("matchedMaterials"))
        + [
            material
            for task in _object_items(item.get("appendixTasks"))
            for material in _object_items(task.get("recommendedMaterials"))
        ]
    )
    by_id = {
        _material_key(material): material
        for material in context
        if _material_key(material)
    }
    return _dedupe_material_summaries([
        by_id.get(material_id) or {"id": material_id, "name": material_id}
        for material_id in selected_ids
    ])


async def _downloadable_technical_fill_source_payload(material_id: str) -> tuple[dict[str, Any], str]:
    try:
        payload = await technical_material_store.raw_download_cleaned_content(material_id)
        return payload, "cleaned"
    except Exception:
        payload = await technical_material_store.raw_download_content(material_id)
    mime_type = str(payload.get("mimeType") or "")
    file_name = str(payload.get("fileName") or "").lower()
    allowed_ext = file_name.endswith((".docx", ".xlsx", ".xlsm", ".pdf"))
    allowed_mime = (
        "wordprocessingml" in mime_type
        or "spreadsheetml" in mime_type
        or "ms-excel" in mime_type
        or mime_type == "application/pdf"
    )
    if not allowed_ext and not allowed_mime:
        raise ValueError(f"素材 {material_id} 不是填写 Skill 可读取的 Word/Excel/PDF。")
    return payload, "raw"


_FILL_ORIGINAL_SUFFIXES = (".xlsx", ".xlsm")


async def _original_technical_fill_source_payload(material_id: str) -> dict[str, Any] | None:
    """取素材 Excel 原件的下载信息。

    清洗稿是 docx 文本化版本，丢掉了 sheet 结构：报价分项转写、功率曲线矩阵、
    参数表机型列定位都必须读原生 Excel。返回 None 表示该素材没有需要额外落地
    的 Excel 原件（本身是 Word/PDF，或取不到原件）。PDF 不在此列——其文本已由
    清洗稿承载，额外落地只会凭空增加 OCR 开销。
    """
    try:
        payload = await technical_material_store.raw_download_content(material_id)
    except Exception:
        return None
    file_name = str(payload.get("fileName") or "").lower()
    if not file_name.endswith(_FILL_ORIGINAL_SUFFIXES):
        return None
    return payload


def _material_may_have_excel_original(item: dict[str, Any]) -> bool:
    """按素材名判断是否值得回源取 Excel 原件，避免逐份多打一次 DB/MinIO。

    素材名是上传时的原始文件名（清洗稿名单独放 cleanedFileName），后缀已足够
    判定；名称缺失或后缀不认识时返回 True，交给下载分支探测。
    """
    for key in ("name", "title"):
        text = str(item.get(key) or "").strip().lower()
        if not text:
            continue
        if text.endswith(_FILL_ORIGINAL_SUFFIXES):
            return True
        if text.endswith((".docx", ".doc", ".pdf", ".txt", ".md")):
            return False
    return True


def _resolve_original_material_file(
    material_id: str,
    item: dict[str, Any],
    cache_dir: Path,
) -> Path | None:
    """把 Excel 原件落地到素材缓存，并在 item 上暴露 originalPath。

    已落地的就是原件（清洗稿缺失走 raw 分支）时直接复用，不重复下载。
    """
    current = Path(str(item.get("path") or ""))
    if current.suffix.lower() in _FILL_ORIGINAL_SUFFIXES and current.exists():
        item["originalPath"] = str(current)
        item["originalFileName"] = current.name
        return current
    if not _material_may_have_excel_original(item):
        return None
    awaitable = _original_technical_fill_source_payload(material_id)
    try:
        payload = _run_async(awaitable)
    except Exception:
        if hasattr(awaitable, "close"):
            awaitable.close()
        return None
    if not payload:
        return None
    file_name = safe_filename(str(payload.get("fileName") or ""), f"{material_id}原件")
    target_path = cache_dir / f"{material_id}-{file_name}"
    if not target_path.exists():
        try:
            minio_client.download_file(str(payload["bucket"]), str(payload["key"]), target_path)
        except Exception:
            return None
    item["originalPath"] = str(target_path)
    item["originalFileName"] = target_path.name
    return target_path


async def _downloadable_technical_word_payload(material_id: str) -> tuple[dict[str, Any], str]:
    try:
        payload = await technical_material_store.raw_download_cleaned_content(material_id)
        return payload, "cleaned"
    except Exception:
        payload = await technical_material_store.raw_download_content(material_id)
    mime_type = str(payload.get("mimeType") or "")
    file_name = str(payload.get("fileName") or "")
    if "wordprocessingml" not in mime_type and not file_name.lower().endswith(".docx"):
        raise ValueError(f"素材 {material_id} 没有可用于拼接的 Word 文件或清洗稿。")
    return payload, "raw"


def _atomic_write_text(path: Path, text: str) -> None:
    """写共享缓存文件：先写线程唯一的临时文件再原子改名。

    一键填写并发后，多个线程可能同时生成同一素材的 OCR sidecar；直接 write_text
    会让并发读者读到写了一半的文件。临时名带 pid/线程 id，互不覆盖。
    """
    temp_path = path.with_name(f"{path.name}.{os.getpid()}.{threading.get_ident()}.tmp")
    try:
        temp_path.write_text(text, encoding="utf-8")
        temp_path.replace(path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise


def _ensure_pdf_ocr_sidecar(pdf_path: Path) -> tuple[str, str]:
    """为证书类 PDF 素材生成 OCR 文本 sidecar（{stem}.ocr.txt，与 PDF 同目录缓存复用）。

    返回 (sidecar 路径, 状态)。OCR 未配置/失败时不抛出，仅返回空路径与原因，
    由 manifest 透出，填表侧对无 sidecar 的 PDF 显式报错。
    """
    sidecar_path = pdf_path.with_suffix(".ocr.txt")
    if sidecar_path.exists() and sidecar_path.stat().st_size > 0:
        return str(sidecar_path), "cached"
    try:
        text, _meta = _run_async(
            ocr_service.recognize_text_for_parse(
                file_name=pdf_path.name,
                content=pdf_path.read_bytes(),
                mime_type="application/pdf",
            )
        )
    except PeripheralError as exc:
        if exc.code == "OCR_CONFIG_REQUIRED":
            return "", "skipped: OCR 模型未启用或未配置"
        return "", f"failed: {exc.detail}"
    except Exception as exc:
        return "", f"failed: {exc}"
    if not str(text or "").strip():
        return "", "failed: OCR 识别结果为空"
    _atomic_write_text(sidecar_path, str(text))
    return str(sidecar_path), "generated"


def _ensure_pdf_ocr_sidecars_batch(pending: list[tuple[dict[str, Any], Path]]) -> int:
    """对本轮收集的待 OCR PDF 批量提交、统一等待并回填 sidecar 状态。

    与 _ensure_pdf_ocr_sidecar 的单份状态语义一致（generated / skipped / failed），
    区别在于全部任务一次性提交给 OCR worker 池并发执行，而不是逐份串行同步等待
    （R09-B07-01）。OCR 未配置/失败不抛出，仅写入各 item 的 ocrStatus。
    返回本轮新生成的 sidecar 数。
    """
    if not pending:
        return 0

    def _mark_all(status: str) -> int:
        for item, _pdf_path in pending:
            item["ocrStatus"] = status
        return 0

    try:
        results = _run_async(
            ocr_service.recognize_texts_for_parse_batch(
                files=[(pdf_path.name, pdf_path.read_bytes(), "application/pdf") for _, pdf_path in pending]
            )
        )
    except PeripheralError as exc:
        if exc.code == "OCR_CONFIG_REQUIRED":
            return _mark_all("skipped: OCR 模型未启用或未配置")
        return _mark_all(f"failed: {exc.detail}")
    except Exception as exc:
        return _mark_all(f"failed: {exc}")

    generated = 0
    for (item, pdf_path), result in zip(pending, results, strict=True):
        if isinstance(result, BaseException):
            if isinstance(result, PeripheralError) and result.code == "OCR_CONFIG_REQUIRED":
                item["ocrStatus"] = "skipped: OCR 模型未启用或未配置"
            elif isinstance(result, PeripheralError):
                item["ocrStatus"] = f"failed: {result.detail}"
            else:
                item["ocrStatus"] = f"failed: {result}"
            continue
        text, _meta = result
        if not str(text or "").strip():
            item["ocrStatus"] = "failed: OCR 识别结果为空"
            continue
        sidecar_path = pdf_path.with_suffix(".ocr.txt")
        # 批量等待期间其他任务可能已补齐同一 sidecar，缓存优先不覆写
        if not (sidecar_path.exists() and sidecar_path.stat().st_size > 0):
            _atomic_write_text(sidecar_path, str(text))
        item["ocrTextPath"] = str(sidecar_path)
        item["ocrStatus"] = "generated"
        generated += 1
    return generated


def _prepare_material_index_files(
    material_index: list[dict[str, Any]],
    work_dir: Path,
    *,
    cache_dir: Path | None = None,
    limit: int = 240,
    ocr_pdf: bool = False,
    ocr_budget: int = 12,
) -> list[dict[str, Any]]:
    prepared: list[dict[str, Any]] = []
    cache_dir = cache_dir or (work_dir / "material_index")
    # 两阶段：先逐份串行下载落地（缓存命中秒过），收集本轮需 OCR 的 PDF；
    # 再对收集列表批量提交、统一等待回填，避免逐份同步阻塞 OCR worker 池。
    pending_ocr: list[tuple[dict[str, Any], Path]] = []
    for material in material_index[:limit]:
        item = dict(material)
        material_id = str(item.get("id") or item.get("materialId") or "").strip()
        if not material_id:
            prepared.append(item)
            continue
        awaitable = _downloadable_technical_fill_source_payload(material_id)
        try:
            payload, source_kind = _run_async(awaitable)
        except Exception:
            if hasattr(awaitable, "close"):
                awaitable.close()
            prepared.append(item)
            continue
        file_name = safe_filename(
            str(item.get("cleanedFileName") or payload.get("fileName") or item.get("name") or f"{material_id}.docx"),
            f"{material_id}.docx",
        )
        target_path = cache_dir / f"{material_id}-{file_name}"
        if not target_path.exists():
            minio_client.download_file(str(payload["bucket"]), str(payload["key"]), target_path)
        item.update(
            {
                "path": str(target_path),
                "sourceKind": source_kind,
                "fileName": target_path.name,
            }
        )
        # 清洗稿是 docx 文本化版本；Excel 原件另行落地，供报价转写、曲线矩阵等
        # 依赖 sheet 结构的填表分支使用（见 _resolve_original_material_file）。
        _resolve_original_material_file(material_id, item, cache_dir)
        if ocr_pdf and target_path.suffix.lower() == ".pdf":
            has_cache = target_path.with_suffix(".ocr.txt").exists()
            if has_cache:
                sidecar_path, ocr_status = _ensure_pdf_ocr_sidecar(target_path)
                if sidecar_path:
                    item["ocrTextPath"] = sidecar_path
                item["ocrStatus"] = ocr_status
            elif len(pending_ocr) < ocr_budget:
                pending_ocr.append((item, target_path))
            else:
                item["ocrStatus"] = _OCR_BUDGET_SKIPPED_STATUS
        prepared.append(item)
    _ensure_pdf_ocr_sidecars_batch(pending_ocr)
    return prepared


# 单轮 OCR 配额跳过的状态文案。填表任务内部按轮循环补齐（见
# _prepare_fill_materials_with_ocr），一轮内仍受 ocr_budget 约束。
_OCR_BUDGET_SKIPPED_STATUS = "skipped: 本次运行 OCR 配额已用完，缓存后续运行补齐"
# 补齐循环兜底轮数：一轮三路最多各新 OCR ocr_budget 份，正常项目 2-3 轮内收敛。
_AI_FILL_OCR_PREP_MAX_ROUNDS = 10


def _prepare_fill_materials_with_ocr(
    material_index: list[dict[str, Any]],
    reference_materials: list[dict[str, Any]],
    recommended_materials: list[dict[str, Any]],
    work_dir: Path,
    *,
    cache_dir: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """在本次填表任务内部循环补齐 PDF OCR sidecar，使一次点击即可用全量素材。

    单轮准备仍受 ocr_budget 约束（排队与并发由 OcrTask worker 保证，配额只是
    单轮时间封顶）；sidecar 落共享缓存跨轮/跨缺口复用，已就绪的素材秒过。
    循环到没有"配额跳过"的 PDF 为止，之后填表只跑一次。
    防死循环：限轮数，且某轮就绪数零增长即停止（OCR 失败为 failed 状态，
    不计入跳过，不会导致空转）。
    """
    prev_ready = -1
    for _round in range(_AI_FILL_OCR_PREP_MAX_ROUNDS):
        material_index = _prepare_material_index_files(
            material_index,
            work_dir,
            cache_dir=cache_dir,
            ocr_pdf=True,
        )
        reference_materials = _prepare_material_index_files(
            reference_materials,
            work_dir,
            cache_dir=cache_dir,
            ocr_pdf=True,
        )
        recommended_materials = _prepare_material_index_files(
            recommended_materials,
            work_dir,
            cache_dir=cache_dir,
            ocr_pdf=True,
        )
        prepared_lists = (material_index, reference_materials, recommended_materials)
        deferred = sum(
            1
            for prepared in prepared_lists
            for item in prepared
            if str(item.get("ocrStatus") or "") == _OCR_BUDGET_SKIPPED_STATUS
        )
        ready = sum(
            1
            for prepared in prepared_lists
            for item in prepared
            if str(item.get("ocrTextPath") or "").strip()
        )
        if deferred == 0 or ready <= prev_ready:
            break
        prev_ready = ready
    return material_index, reference_materials, recommended_materials


def _prepare_word_blank_source(blank_source: dict[str, Any], work_dir: Path) -> dict[str, Any]:
    source = dict(blank_source)
    for key in ("docxPath", "path", "workspacePath"):
        candidate = Path(str(source.get(key) or ""))
        if candidate.exists() and candidate.suffix.lower() == ".docx":
            source["docxPath"] = str(candidate)
            return source

    material_id = str(source.get("materialId") or source.get("id") or "").strip()
    if not material_id:
        raise ValueError("待填写 Word 缺少 materialId/docxPath，无法准备模板。")
    payload, source_kind = _run_async(_downloadable_technical_word_payload(material_id))
    file_name = safe_filename(
        str(source.get("cleanedFileName") or payload.get("fileName") or source.get("title") or f"{material_id}.docx"),
        f"{material_id}.docx",
    )
    if not file_name.lower().endswith(".docx"):
        file_name = f"{Path(file_name).stem}.docx"
    target_path = work_dir / "blank_source" / f"{material_id}-{file_name}"
    if not target_path.exists():
        minio_client.download_file(str(payload["bucket"]), str(payload["key"]), target_path)
    source.update(
        {
            "docxPath": str(target_path),
            "path": str(target_path),
            "sourceKind": source_kind,
            "fileName": target_path.name,
        }
    )
    return source


# ---------- 待插入：整份素材嵌入的备料 ----------
#
# 正文占位符分两型，后缀是权威标记：`待填写` 填一个字段值，`待插入` 嵌一整份素材。
# 素材检索要联网查库，按「skill 只依据 manifest 工作」的既有约定放在后端：这里按项目
# 素材范围找同名素材、下载落地，写进 manifest.embedSources，脚本只管插入。

# 只识别待插入后缀，用于备料；占位符解析的权威实现在 filler 的 PLACEHOLDER_RE。
_EMBED_PLACEHOLDER_RE = re.compile(r"[\[【]\s*([^\]】\r\n]{1,80}?)\s*[,，、:：\s]*待插入\s*[\]】]")
# 与 filler 的 norm() 同款归一化：全半角括号、分隔符差异不该影响占位符与素材名的比对
_EMBED_NORM_RE = re.compile(r"[\s（）()、/\\:：；;，,。\-_—×*\[\]【】]+")
# 素材分层的特异性：同名素材优先取更专的一层，同层撞名才交人工
_EMBED_TIER_PRIORITY = {"project": 3, "customer": 2, "standard": 1}
_EMBED_UNSUPPORTED_SUFFIXES = {".xlsx", ".xls", ".xlsm"}


def _embed_norm(value: Any) -> str:
    return _EMBED_NORM_RE.sub("", str(value or "").replace("　", " ").strip().lower())


def _scan_embed_placeholders(docx_path: Path) -> list[str]:
    """扫描待填写 Word 里的待插入占位符，返回去重后的素材名。

    只看段落：待插入要整份嵌入，必须独占段落，表格单元格里的由 filler 标黄交人工，
    这里不必为它们备料。
    """
    from docx import Document

    labels: list[str] = []
    seen: set[str] = set()
    for paragraph in Document(str(docx_path)).paragraphs:
        for match in _EMBED_PLACEHOLDER_RE.finditer(paragraph.text or ""):
            label = str(match.group(1) or "").strip()
            key = _embed_norm(label)
            if not key or key in seen:
                continue
            seen.add(key)
            labels.append(label)
    return labels


def _pick_most_specific_material(materials: list[dict[str, Any]]) -> tuple[dict[str, Any], bool]:
    """同名素材取最专的一层（项目定制 > 客户定制 > 标准）；同层撞名视为分不开。"""

    def priority(material: dict[str, Any]) -> int:
        return _EMBED_TIER_PRIORITY.get(str(material.get("materialTier") or ""), 0)

    ranked = sorted(materials, key=priority, reverse=True)
    top = priority(ranked[0])
    return ranked[0], sum(1 for material in ranked if priority(material) == top) > 1


def _embed_sources_for_fill(
    project: dict[str, Any],
    blank_docx_path: Path,
    work_dir: Path,
) -> list[dict[str, Any]]:
    """为待填写 Word 里的每个待插入占位符备好可嵌入的 Word 素材。

    每条都带 status：只有 ready 才会被嵌入，其余由 filler 原地标黄并写明原因，
    不静默跳过——整份素材没进去却报成功，审核界面上看不出来。
    """
    labels = _scan_embed_placeholders(blank_docx_path)
    if not labels:
        return []
    material_scope = build_project_material_scope(project)
    candidates = _allowed_technical_material_index(material_scope, project_turbine_model(project))
    by_key: dict[str, list[dict[str, Any]]] = {}
    for material in candidates:
        key = _embed_norm(Path(str(material.get("name") or "")).stem)
        if key:
            by_key.setdefault(key, []).append(material)

    sources: list[dict[str, Any]] = []
    for label in labels:
        matches = by_key.get(_embed_norm(label)) or []
        if not matches:
            sources.append(
                {
                    "placeholder": label,
                    "status": "not_found",
                    "statusMessage": f"项目素材范围内未找到名为「{label}」的素材。",
                }
            )
            continue
        picked, ambiguous = _pick_most_specific_material(matches)
        material_id = str(picked.get("id") or "")
        name = str(picked.get("name") or "")
        entry = {
            "placeholder": label,
            "materialId": material_id,
            "name": name,
            "folderPath": str(picked.get("folderPath") or ""),
            "materialTier": str(picked.get("materialTier") or ""),
        }
        if ambiguous:
            sources.append(
                {
                    **entry,
                    "status": "ambiguous",
                    "statusMessage": f"同一层级存在多个名为「{label}」的素材，请人工指定。",
                    "candidateCount": len(matches),
                }
            )
            continue
        # 按素材原始后缀判断，不看能不能取到 docx：xlsx 若被 Wiki 预览转换过，
        # 取素材时会拿到自动转换的清洗稿，那份转换有损（合并单元格丢失、表头认错），
        # 悄悄嵌进投标材料就是静默降级。
        if Path(name).suffix.lower() in _EMBED_UNSUPPORTED_SUFFIXES:
            sources.append(
                {
                    **entry,
                    "status": "unsupported_format",
                    "statusMessage": f"「{name}」是 Excel 素材，请人工另存为 Word 后重新上传。",
                }
            )
            continue
        awaitable = _downloadable_technical_word_payload(material_id)
        try:
            payload, source_kind = _run_async(awaitable)
            file_name = safe_filename(
                str(picked.get("cleanedFileName") or payload.get("fileName") or name or f"{material_id}.docx"),
                f"{material_id}.docx",
            )
            if not file_name.lower().endswith(".docx"):
                file_name = f"{Path(file_name).stem}.docx"
            target_path = work_dir / "embed_sources" / f"{material_id}-{file_name}"
            if not target_path.exists():
                minio_client.download_file(str(payload["bucket"]), str(payload["key"]), target_path)
        except Exception as exc:  # noqa: BLE001 - 单份素材取不到不能中断整份文件的填写
            if hasattr(awaitable, "close"):
                awaitable.close()
            sources.append(
                {
                    **entry,
                    "status": "download_failed",
                    "statusMessage": f"素材「{name}」下载失败：{exc}",
                }
            )
            continue
        sources.append({**entry, "status": "ready", "sourceKind": source_kind, "docxPath": str(target_path)})
    return sources
