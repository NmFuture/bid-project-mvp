"""技术标项目事实表：素材事实提取（素材索引、docx/xlsx/自由文本启发式抽取）。

从 technical_gap_fact_table 拆出；对外仍经 app.services.technical_gap_fact_table 门面 re-export。
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

from app.core.config import settings
from app.services.bid_type import TECHNICAL_BID_TYPE
from app.services.file_utils import run_awaitable_sync
from app.services.identity import build_project_material_scope
from app.services.project_fact_materials import prepare_project_fact_material_files
from app.services.technical_material_store import technical_material_store
from app.services.turbine_models import project_turbine_model
from app.services.fact_table_common import (
    COMMON_PROJECT_FACT_LABELS,
    FACT_MATERIAL_SOURCE_PRIORITIES,
    canonical_fact_label,
)
from app.services.technical_fact_extract_parse import looks_like_table_field_label

logger = logging.getLogger(__name__)


def project_material_fact_fields(
    project: dict[str, Any],
    gap_state: dict[str, Any],
    *,
    excluded_paths: set[str] | None = None,
    specs: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    materials = project_fact_material_index(project, gap_state)
    if not materials:
        return []
    prepared = prepare_project_fact_materials(project, materials)
    # 延迟 import 避免循环：专项模块经门面复用本模块的 material_fact/clean_fact_text
    from app.services.technical_fact_special_extractors import (
        facts_from_certificate_materials,
        run_special_extractor,
        special_extractor_for_material,
    )

    facts: list[dict[str, Any]] = []
    cert_materials: list[tuple[dict[str, Any], Path]] = []
    for material in prepared:
        if not isinstance(material, dict):
            continue
        facts.extend(facts_from_material_name(material))
        path_text = str(material.get("path") or material.get("docx") or "").strip()
        if not path_text:
            continue
        path = Path(path_text)
        if not path.exists():
            continue
        if str(path.resolve()) in (excluded_paths or set()):
            continue
        kind = special_extractor_for_material(material)
        if kind == "certificate":
            # 证书按"型式认证 > 设计认证"成组处理
            cert_materials.append((material, path))
            continue
        if kind:
            special_facts = run_special_extractor(kind, path, material, project, specs=specs)
            if special_facts is not None:
                facts.extend(special_facts)
                continue
        suffix = path.suffix.lower()
        if suffix in {".docx", ".doc"}:
            facts.extend(facts_from_docx_material(path, material))
        elif suffix in {".xlsx", ".xlsm"}:
            facts.extend(facts_from_xlsx_material(path, material, project))
    facts.extend(facts_from_certificate_materials(cert_materials, project, specs=specs))
    # 延迟 import 避免循环：派生事实留在门面模块，门面在模块级 import 本模块
    from app.services.technical_gap_fact_table import derived_material_fact_fields

    facts.extend(derived_material_fact_fields(project, facts))
    return facts


# 素材索引口径：build 供规则抽取（只认项目定制素材的格式），
# curate 供 AI 补抽（机型标准配置类字段的来源只在标准文件目录下）
FACT_MATERIAL_USAGE_BUILD = "build"
FACT_MATERIAL_USAGE_CURATE = "curate"


def project_fact_material_index(
    project: dict[str, Any],
    gap_state: dict[str, Any],
    *,
    usage: str = FACT_MATERIAL_USAGE_BUILD,
) -> list[dict[str, Any]]:
    plan = gap_state.get("plan") if isinstance(gap_state.get("plan"), dict) else {}
    materials: list[dict[str, Any]] = []
    seen: set[str] = set()

    def material_key(item: dict[str, Any]) -> str:
        material_id = str(item.get("id") or item.get("materialId") or "").strip()
        if material_id:
            return f"id:{material_id}"
        path = str(item.get("path") or item.get("folderPath") or "").strip()
        name = str(item.get("name") or item.get("materialName") or item.get("fileName") or "").strip()
        return f"path:{path}/{name}" if path or name else ""

    def append_material(item: dict[str, Any]) -> None:
        key = material_key(item)
        if not key or key in seen:
            return
        seen.add(key)
        materials.append(item)

    for item in (plan.get("materialIndex") if isinstance(plan.get("materialIndex"), list) else []):
        if isinstance(item, dict):
            append_material(dict(item))

    if not materials and isinstance(plan.get("tasks"), list):
        for task in plan.get("tasks") or []:
            if not isinstance(task, dict):
                continue
            for raw in [*(task.get("candidateMaterials") or []), *(task.get("selectedMaterialRefs") or [])]:
                if not isinstance(raw, dict):
                    continue
                material_id = str(raw.get("id") or raw.get("materialId") or "").strip()
                if not material_id:
                    continue
                append_material(
                    {
                        "id": material_id,
                        "name": str(raw.get("name") or raw.get("materialName") or raw.get("fileName") or material_id),
                        "folderPath": str(raw.get("folderPath") or raw.get("path") or ""),
                        "materialTier": str(raw.get("materialTier") or raw.get("libraryScope") or ""),
                        "cleanedFileName": str(raw.get("cleanedFileName") or ""),
                        "hasCleanedWord": bool(raw.get("hasCleanedWord")),
                        "turbineModelLabel": str(raw.get("turbineModelLabel") or ""),
                    }
                )
    try:
        selected_model = project_turbine_model(project)

        def collect_scope_files(
            folder_path: str,
            material_tier: str,
            *,
            customer_name: str = "",
            project_id: str = "",
        ) -> None:
            payload = run_async_material_files(
                folder_path=folder_path,
                bid_type=TECHNICAL_BID_TYPE,
                material_tier=material_tier,
                customer_name=customer_name,
                project_id=project_id,
                turbine_model=selected_model,
                recursive=True,
                page=1,
                page_size=1000,
            )
            for raw in payload.get("items") or []:
                if not isinstance(raw, dict):
                    continue
                material_id = str(raw.get("id") or "")
                if not material_id:
                    continue
                append_material(
                    {
                        "id": material_id,
                        "name": str(raw.get("name") or ""),
                        "folderPath": str(raw.get("folderPath") or ""),
                        "materialTier": str(raw.get("materialTier") or material_tier),
                        "hasCleanedWord": bool(raw.get("hasCleanedWord")),
                        "cleanedFileName": str(raw.get("cleanedFileName") or ""),
                        "turbineModelLabel": str(raw.get("turbineModelLabel") or ""),
                        "size": int(raw.get("size") or 0),
                    }
                )

        # 无现成索引时，沿用项目默认目录扫描作为回退。
        if not materials:
            material_scope = build_project_material_scope(project)
            for scope in material_scope.get("readableScopes") or []:
                if not isinstance(scope, dict):
                    continue
                # build 只扫本项目「项目定制」目录（recursive 覆盖下一级子目录，
                # 相关项目素材由用户归置到该目录下）：规则抽取器只认这些素材的格式。
                # curate 三层都扫，标准文件目录已由 turbine_model 收敛到本机型。
                tier = str(scope.get("materialTier") or "")
                if usage == FACT_MATERIAL_USAGE_BUILD and tier != "project":
                    continue
                folder_path = str(scope.get("path") or "").strip()
                if not folder_path:
                    continue
                # 素材库按客户别名建目录（如「华能」而非规范名「华能集团」），拼全路径查不到。
                # 与 planner 同口径：客户/项目层路径截到标类根，改用 customer_name/project_id
                # 交给 customer_matches/project_matches 做别名匹配。
                if tier in {"customer", "project"}:
                    folder_path = "/".join([part for part in folder_path.split("/") if part][:2])
                collect_scope_files(
                    folder_path,
                    tier,
                    customer_name=str(scope.get("customerName") or "") if tier == "customer" else "",
                    project_id=str(scope.get("projectId") or "") if tier == "project" else "",
                )

        # 用户显式选择的目录始终叠加到计划索引，并按素材 ID 去重。
        custom_paths = gap_state.get("factMaterialPaths") if isinstance(gap_state.get("factMaterialPaths"), list) else []
        for raw_path in custom_paths:
            folder_path = str(raw_path or "").strip().strip("/")
            if folder_path and not folder_path.startswith(f"{TECHNICAL_BID_TYPE}/"):
                folder_path = f"{TECHNICAL_BID_TYPE}/{folder_path}"
            if folder_path:
                # 显式目录不按 tier 过滤，相关性由 material_is_fact_relevant 把守。
                collect_scope_files(folder_path, "")
    except Exception:
        logger.exception("项目事实素材索引查询失败，保留已收集素材继续构建")
    return [item for item in materials if material_is_fact_relevant(item, usage=usage)]


# 业主待填目标表格模板的文件名前缀（「待填写-附表X….docx」「待填写、待用印-….docx」等，
# 前缀取自业主下发的空白附表命名约定）：它们是要填的目标，不是取数素材，不进事实表素材体系
FILL_TEMPLATE_NAME_PREFIX = "待填写"


def material_is_fill_template(material: dict[str, Any]) -> bool:
    """按文件名前缀识别待填目标表格模板（不看 folderPath，避免按目录名猜内容）。"""
    name = str(material.get("name") or material.get("cleanedFileName") or "").strip()
    return name.startswith(FILL_TEMPLATE_NAME_PREFIX)


def material_is_fact_relevant(
    material: dict[str, Any],
    *,
    usage: str = FACT_MATERIAL_USAGE_BUILD,
) -> bool:
    if material_is_fill_template(material):
        return False
    tier = str(material.get("materialTier") or "").strip()
    if tier == "project":
        return True
    # curate 的范围已由 build_project_material_scope 按机型/客户收敛，无需再按关键词猜内容：
    # 白名单会挡掉机型标准配置类素材（如自动消防系统），而这类素材正是 Excel 未指定来源
    # （referenceFile 为「/」）的字段的唯一取值处。
    if usage == FACT_MATERIAL_USAGE_CURATE:
        return True
    text = " ".join(
        str(material.get(key) or "")
        for key in ("name", "cleanedFileName", "folderPath", "path")
    )
    return bool(
        re.search(
            r"参数|机型|功率曲线|风资源|发电量|报价|容量|安全|场址|载荷|工程量|技术承诺|投标关键数据|"
            r"弯矩|认证|承诺函|生产制造基地",
            text,
        )
    )


def project_fact_material_work_dir(project: dict[str, Any]) -> Path:
    """事实表素材物化目录：build 批量落地与 curate 按需拉取共用同一份缓存。"""
    project_id = str(project.get("id") or "project")
    return settings.documents_dir / project_id / "technical-workspace" / "gaps" / "fact_table_materials"


def prepare_project_fact_materials(project: dict[str, Any], materials: list[dict[str, Any]]) -> list[dict[str, Any]]:
    path_materials = [item for item in materials if item.get("path")]
    if path_materials and len(path_materials) == len(materials) and all(
        Path(str(item.get("path") or "")).exists() for item in path_materials
    ):
        return materials
    work_dir = project_fact_material_work_dir(project)
    work_dir.mkdir(parents=True, exist_ok=True)
    try:
        return prepare_project_fact_material_files(
            materials, work_dir, bid_type=TECHNICAL_BID_TYPE, limit=120
        )
    except Exception:
        return materials


def facts_from_material_name(material: dict[str, Any]) -> list[dict[str, Any]]:
    text = str(material.get("name") or material.get("cleanedFileName") or "")
    facts: list[dict[str, Any]] = []
    for pattern, label, unit in [
        (r"空气密度\s*([0-9]+(?:\.[0-9]+)?)", "空气密度", "kg/m3"),
        (r"湍流强度\s*([0-9]+(?:\.[0-9]+)?)", "湍流强度", ""),
        (r"风(?:剪切|切变)(?:指数)?\s*([0-9]+(?:\.[0-9]+)?)", "风剪切", ""),
    ]:
        match = re.search(pattern, text, flags=re.I)
        if match:
            facts.append(material_fact(label, match.group(1), material, unit=unit, confidence=0.9))
    return facts


def facts_from_docx_material(path: Path, material: dict[str, Any]) -> list[dict[str, Any]]:
    facts: list[dict[str, Any]] = []
    try:
        document = Document(str(path))
    except Exception:
        return facts
    text_parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for para_idx, paragraph in enumerate(document.paragraphs, start=1):
        text = clean_fact_text(paragraph.text)
        if not text or len(text) > 180:
            continue
        match = re.match(r"^([^:：]{2,40})[:：]\s*(.{1,100})$", text)
        if match:
            fact = material_fact_from_label_value(match.group(1), match.group(2), material, location=f"P{para_idx}", confidence=0.78)
            if fact:
                facts.append(fact)
    for table_idx, table in enumerate(document.tables, start=1):
        facts.extend(facts_from_guarantee_table(table, material, table_idx=table_idx))
        for row_idx, row in enumerate(table.rows, start=1):
            cells = [clean_fact_text(cell.text) for cell in row.cells]
            text_parts.append(" | ".join(cell for cell in cells if cell))
            facts.extend(facts_from_table_cells(cells, material, location=f"T{table_idx}/R{row_idx}"))
    facts.extend(facts_from_free_text("\n".join(text_parts), material))
    return facts


def facts_from_guarantee_table(table: Any, material: dict[str, Any], *, table_idx: int) -> list[dict[str, Any]]:
    if not getattr(table, "rows", None) or len(table.rows) < 2:
        return []
    header = " ".join(clean_fact_text(cell.text) for cell in table.rows[0].cells)
    if not ("年平均风速" in header and "保证年上网电量" in header and "满负荷小时" in header):
        return []
    facts: list[dict[str, Any]] = []
    for row_idx, row in enumerate(table.rows[1:], start=2):
        cells = [clean_fact_text(cell.text) for cell in row.cells]
        if len(cells) < 3:
            continue
        wind_speed = clean_fact_value("年平均风速", cells[0])
        energy = clean_fact_value("保证发电量", cells[1])
        hours = clean_fact_value("保证有效小时数", cells[2])
        if not (wind_speed and energy and hours):
            continue
        matrix_fact = material_fact(
            "__guaranteeMatrixRow",
            {"windSpeed": wind_speed, "energyMwh": energy, "hours": hours},
            material,
            location=f"T{table_idx}/R{row_idx}",
            confidence=0.82,
        )
        matrix_fact["internal"] = True
        facts.append(matrix_fact)
    return facts


def facts_from_xlsx_material(path: Path, material: dict[str, Any], project: dict[str, Any]) -> list[dict[str, Any]]:
    facts: list[dict[str, Any]] = []
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return facts
    project_model = str((project_turbine_model(project) or {}).get("model") or "")
    model_key = re.sub(r"(上置|下置|内置|外置|塔上|塔下)", "", project_model)
    for worksheet in workbook.worksheets:
        selected_col = xlsx_model_column(worksheet, model_key)
        for row_idx, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            cells = [clean_fact_text(cell) for cell in row]
            if selected_col is not None and selected_col < len(cells):
                for label_idx in (2, 1, 0):
                    if label_idx < len(cells):
                        fact = material_fact_from_label_value(
                            cells[label_idx],
                            cells[selected_col],
                            material,
                            unit=cells[3] if len(cells) > 3 else "",
                            location=f"{worksheet.title}!R{row_idx}",
                            confidence=0.82,
                        )
                        if fact:
                            facts.append(fact)
                            break
            facts.extend(facts_from_table_cells(cells, material, location=f"{worksheet.title}!R{row_idx}"))
            if len(facts) >= 800:
                return facts
    return facts


def xlsx_model_column(worksheet: Any, model_key: str) -> int | None:
    if not model_key:
        return None
    normalized_model = re.sub(r"\s+", "", model_key)
    for row in worksheet.iter_rows(min_row=1, max_row=min(12, worksheet.max_row), values_only=True):
        for index, value in enumerate(row):
            text = re.sub(r"\s+", "", str(value or ""))
            if normalized_model and normalized_model in text:
                return index
    return None


def facts_from_table_cells(cells: list[str], material: dict[str, Any], *, location: str) -> list[dict[str, Any]]:
    facts: list[dict[str, Any]] = []
    nonempty = [(index, value) for index, value in enumerate(cells) if value]
    if len(nonempty) < 2:
        return facts
    if len(cells) >= 4:
        fact = material_fact_from_label_value(cells[1], cells[3], material, unit=cells[2], location=location, confidence=0.88)
        if fact:
            facts.append(fact)
    for first, second in zip(nonempty, nonempty[1:]):
        fact = material_fact_from_label_value(first[1], second[1], material, location=location, confidence=0.76)
        if fact:
            facts.append(fact)
            break
    if len(cells) >= 3 and cells[0]:
        wind_fact = material_fact_from_label_value(cells[0], cells[2], material, unit=cells[1], location=location, confidence=0.84)
        if wind_fact:
            facts.append(wind_fact)
    return facts


def material_fact_from_label_value(
    label: Any,
    value: Any,
    material: dict[str, Any],
    *,
    unit: str = "",
    location: str = "",
    confidence: float = 0.78,
) -> dict[str, Any] | None:
    label_text = canonical_fact_label(label)
    value_text = clean_fact_value(label_text, value)
    if not label_text or not value_text:
        return None
    if label_text not in COMMON_PROJECT_FACT_LABELS and not looks_like_table_field_label(label_text):
        return None
    unit_text = clean_fact_unit(unit)
    raw_label = str(label or "")
    raw_value = str(value or "")
    if not unit_text:
        raw_context = f"{raw_label}{raw_value}"
        if label_text in {"轮毂高度", "叶轮直径"} and re.search(r"(?:m|米)", raw_context, flags=re.I):
            unit_text = "m"
        elif label_text in {"极端风速", "年平均风速"} and re.search(r"m/?s|米/秒", raw_context, flags=re.I):
            unit_text = "m/s"
        elif label_text == "空气密度" and re.search(r"kg/?m|kg/m3|kg/m³", raw_context, flags=re.I):
            unit_text = "kg/m3"
        elif label_text == "机组台数" and "台" in raw_value:
            unit_text = "台"
    if not unit_text:
        if label_text in {"轮毂高度", "叶轮直径"}:
            unit_text = "m"
        elif label_text in {"极端风速", "年平均风速"}:
            unit_text = "m/s"
        elif label_text == "空气密度":
            unit_text = "kg/m3"
        elif label_text == "机组台数":
            unit_text = "台"
    return material_fact(label_text, value_text, material, unit=unit_text, location=location, confidence=confidence)


def facts_from_free_text(text: str, material: dict[str, Any]) -> list[dict[str, Any]]:
    compact = clean_fact_text(text)
    facts: list[dict[str, Any]] = []
    patterns = [
        (r"(?:总装机容量|建设容量|标段规模|总容量)[^0-9]{0,12}([0-9]+(?:\.[0-9]+)?\s*(?:MW|万千瓦|kW)?)", "总装机容量", ""),
        (r"(?:机组台数|机组数量|风机台数|风机数量|安装)[^0-9]{0,12}([0-9]+)\s*台", "机组台数", "台"),
        (r"轮毂(?:中心)?高度[^0-9]{0,12}([0-9]+(?:\.[0-9]+)?\s*m?)", "轮毂高度", "m"),
        (r"(?:安全等级|适用等级|设计等级)[^A-Za-z0-9]{0,12}((?:IEC\s*)?[A-Z0-9][A-Z0-9/ .-]{0,20})", "安全等级", ""),
        (r"空气密度[^0-9]{0,12}([0-9]+(?:\.[0-9]+)?)", "空气密度", "kg/m3"),
        (r"湍流强度[^0-9]{0,12}([0-9]+(?:\.[0-9]+)?)", "湍流强度", ""),
        (r"(?:极端风速|极大风速|Ve50)[^0-9]{0,12}([0-9]+(?:\.[0-9]+)?\s*m/s?)", "极端风速", "m/s"),
        (r"年平均风速[^0-9]{0,12}([0-9]+(?:\.[0-9]+)?\s*m/s?)", "年平均风速", "m/s"),
    ]
    for pattern, label, unit in patterns:
        match = re.search(pattern, compact, flags=re.I)
        if match:
            value = clean_fact_value(label, match.group(1))
            if value:
                facts.append(material_fact(label, value, material, unit=unit, confidence=0.78))
    return facts


def clean_fact_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def clean_fact_unit(value: Any) -> str:
    text = re.sub(r"\s+", "", str(value or "")).strip()
    text = text.strip("：:；;，,、")
    if text in {"", "-", "/", "—", "NA", "N/A", "字段", "值", "年份", "参数内容", "结果", "说明", "备注", "机型", "型号"}:
        return ""
    text = text.replace("m³", "m3")
    text = re.sub(r"kg/?m3", "kg/m3", text, flags=re.I)
    text = re.sub(r"m/?s", "m/s", text, flags=re.I)
    return text


def clean_fact_value(label: str, value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"\s+", "", text)
    text = text.strip("：:；;，,、")
    if not text or len(text) > 120:
        return ""
    if any(token in text for token in ("待填写", "待人工", "未填写")):
        return ""
    if text in {"-", "/", "—", "无", "暂无", "值", "结果", "参数内容", "单位", "年份"}:
        return ""
    numeric_ranges = {
        "机组台数": (1, 1000),
        "轮毂高度": (40, 250),
        "叶轮直径": (50, 350),
        "空气密度": (0.7, 1.5),
        "湍流强度": (0, 1),
        "风剪切": (0, 1),
        "极端风速": (20, 100),
        "年平均风速": (2, 15),
    }
    numeric_noise = (
        "年份",
        "各年",
        "版本",
        "编制",
        "校核",
        "审核",
        "批准",
        "日期",
        "参数内容",
        "结果结果",
        "场址空气密度下",
    )
    if label in numeric_ranges and (
        any(token in text for token in numeric_noise)
        or re.search(r"\d{4}[-/年]\d{1,2}", text)
    ):
        return ""
    if label in {"总装机容量", "单机容量"}:
        match = re.search(r"([0-9]+(?:\.[0-9]+)?)(万千瓦|MW|kW)?", text, flags=re.I)
        if match:
            return f"{match.group(1)}{match.group(2) or ''}".strip()
    if label in {"机组台数"}:
        match = re.search(r"([0-9]+)", text)
        if not match:
            return ""
        number = float(match.group(1))
        low, high = numeric_ranges[label]
        return match.group(1) if low <= number <= high else ""
    if label in {"保证发电量", "保证有效小时数"}:
        if re.search(r"风电场|保证年上网电量|满负荷小时|字段|单位", text):
            return ""
        match = re.search(r"([0-9]+(?:\.[0-9]+)?)", text)
        if not match:
            return ""
        number = float(match.group(1))
        if label == "保证发电量" and not (1 <= number <= 10000000):
            return ""
        if label == "保证有效小时数" and not (1 <= number <= 8760):
            return ""
        return match.group(1)
    if label in {"极端风速", "年平均风速"}:
        match = re.search(r"([0-9]+(?:\.[0-9]+)?)", text)
        if match:
            number = float(match.group(1))
            low, high = numeric_ranges[label]
            if not (low <= number <= high):
                return ""
            return f"{match.group(1)}m/s" if re.search(r"m/?s|米/秒", text, flags=re.I) else match.group(1)
    if label in {"轮毂高度", "叶轮直径", "空气密度", "湍流强度", "风剪切"}:
        match = re.search(r"([0-9]+(?:\.[0-9]+)?)", text)
        if match:
            number = float(match.group(1))
            low, high = numeric_ranges[label]
            if not (low <= number <= high):
                return ""
            return match.group(1)
    if label in {"投标方案", "投标机型", "机组类型"}:
        if text in {"机型", "投标机型", "方案", "投标方案"}:
            return ""
        model_match = re.search(r"([A-Z]{1,6}\d+(?:\.\d+)?[-—]\d+(?:[-—]\d+)?)", text, flags=re.I)
        if model_match:
            return model_match.group(1).replace("—", "-")
    if label == "安全等级":
        text = re.sub(r"^IEC\s*", "IEC ", text, flags=re.I).strip()
    return text


def material_fact(
    label: str,
    value: Any,
    material: dict[str, Any],
    *,
    unit: str = "",
    location: str = "",
    confidence: float = 0.78,
) -> dict[str, Any]:
    tier = str(material.get("materialTier") or "").strip() or "standard"
    return {
        "label": canonical_fact_label(label),
        "value": value,
        "category": "素材库事实",
        "unit": clean_fact_unit(unit),
        "confidence": confidence,
        "sourcePriority": FACT_MATERIAL_SOURCE_PRIORITIES.get(tier, 50),
        "sourceRef": {
            "type": "materialFact",
            "materialId": str(material.get("id") or material.get("materialId") or ""),
            "materialTier": tier,
            "name": str(material.get("name") or material.get("fileName") or material.get("cleanedFileName") or ""),
            "folderPath": str(material.get("folderPath") or ""),
            "path": str(material.get("path") or ""),
            "location": location,
        },
    }


def run_async_material_files(**kwargs: Any) -> dict[str, Any]:
    kwargs.pop("bid_type", None)
    result = run_awaitable_sync(technical_material_store.raw_files(**kwargs))
    return result if isinstance(result, dict) else {}
