"""S2 TOC 工作区发布与路径重映射：staging→published 搬迁、manifest/产物路径改写、输入拷贝。

从 outline_generation 拆出；对外仍经 app.services.outline_generation 门面 re-export。
"""

from __future__ import annotations

import copy
import json
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document

from app.core.config import settings
from app.services.bid_type import require_bid_type
from app.services.file_utils import safe_filename
from app.services.parsing import IMAGE_SUFFIXES, _ocr_fallback_text
from app.services.template_store import is_valid_docx_file
from app.services.turbine_models import project_turbine_model
from app.services.workspace_artifacts import workspace_dir
from app.services.outline_result_loading import _is_business_bid, _load_json_dict


# 注入 S2 manifest 的事实表状态：有值即可用（confirmed）；未提取与不适用不注入。
# 历史值 extracted/pending_confirmation 一并保留，旧项目状态未重建时仍能注入。
MANIFEST_FACT_VALUE_STATUSES = {"confirmed", "extracted", "pending_confirmation"}


def project_facts_for_manifest(project: dict[str, Any]) -> dict[str, str]:
    """从 S3 项目事实表提取 label→value 映射，供 S2 manifest 注入。

    S2 通常早于 S3 事实表构建，gap_state 无事实表时返回空映射（调用方不写该键）。
    """
    gap_state = project.get("gap_state") if isinstance(project.get("gap_state"), dict) else {}
    fact_table = gap_state.get("projectFactTable") if isinstance(gap_state.get("projectFactTable"), dict) else {}
    facts: dict[str, str] = {}
    for field in fact_table.get("fields") or []:
        if not isinstance(field, dict):
            continue
        label = str(field.get("label") or "").strip()
        value = str(field.get("value") or "").strip()
        if label and value and str(field.get("status") or "") in MANIFEST_FACT_VALUE_STATUSES:
            facts.setdefault(label, value)
    return facts


def _prepare_toc_skill_workspace(
    *,
    project_id: str,
    project: dict[str, Any],
    parse_storage: dict[str, Any],
    tender_file_records: list[dict[str, Any]],
    template_file_records: list[dict[str, Any]],
) -> dict[str, Any]:
    bid_type = _outline_bid_type(str(project.get("bidType") or ""))
    project_dir = workspace_dir(project_id, bid_type)
    project_dir.mkdir(parents=True, exist_ok=True)

    published_work_dir = project_dir / "s2_toc_workdir"
    staging_work_dir = project_dir / "s2_toc_workdir.new"
    archive_root = project_dir / "s2_toc_workdir.runs"
    _archive_workspace_if_exists(staging_work_dir, archive_root, "stale")
    staging_work_dir.mkdir(parents=True, exist_ok=True)
    _remove_manifest_alias(project_dir)

    tender_inputs = _copy_tender_inputs(tender_file_records, staging_work_dir, parse_storage)
    template_path, attach_path = _copy_template_inputs(template_file_records, staging_work_dir, project_id)
    if template_path is None:
        raise ValueError("投标模板不存在，请先上传可读取的投标模板文件。")
    output_file = staging_work_dir / safe_filename(settings.s2_toc_output_file_name, "toc.json")
    manifest_path = staging_work_dir / "s2_input.json"
    manifest = {
        "projectId": project_id,
        "projectCode": str(project.get("projectCode") or project_id),
        "projectName": str(project.get("name") or project_id),
        "bidType": bid_type,
        "workDir": str(staging_work_dir),
        "tenderFiles": tender_inputs,
        "templateFile": str(template_path) if template_path else "",
        "attachFile": str(attach_path) if attach_path else "",
        "outputFile": str(output_file),
    }
    # 下游对接：项目机型与 S3 事实表值注入 manifest（须在 _trustedManifest 快照之前）。
    # S2 通常早于 S3 事实表构建，缺失时不写这两个键，不报错。
    turbine_model = project_turbine_model(project)
    if turbine_model:
        manifest["turbineModel"] = turbine_model
    project_facts = project_facts_for_manifest(project)
    if project_facts:
        manifest["projectFacts"] = project_facts
    if _is_business_bid(bid_type):
        manifest["evidenceFile"] = str(
            staging_work_dir / safe_filename(settings.s2_toc_evidence_file_name, "toc_evidence.json")
        )
    else:
        manifest["requireComposedOutline"] = True
    manifest_text = json.dumps(manifest, ensure_ascii=False, indent=2)
    manifest_path.write_text(manifest_text, encoding="utf-8")
    return {
        **manifest,
        "_trustedManifest": copy.deepcopy(manifest),
        "manifestPath": str(manifest_path),
        "canonicalManifestPath": str(manifest_path),
        "publishedWorkDir": str(published_work_dir),
        "stagingWorkDir": str(staging_work_dir),
        "archiveRoot": str(archive_root),
        "tenderFileCount": len(tender_inputs),
        "templateFileCount": 1 if template_path else 0,
        "hasAttachFile": bool(attach_path),
    }


def _publish_toc_skill_workspace(skill_workspace: dict[str, Any], toc_result: dict[str, Any]) -> dict[str, Any]:
    staging_work_dir = Path(str(skill_workspace.get("stagingWorkDir") or skill_workspace.get("workDir") or "")).expanduser()
    published_work_dir = Path(str(skill_workspace.get("publishedWorkDir") or skill_workspace.get("workDir") or "")).expanduser()
    archive_root = Path(str(skill_workspace.get("archiveRoot") or published_work_dir.with_name("s2_toc_workdir.runs"))).expanduser()
    if not staging_work_dir.exists():
        raise RuntimeError(f"S2 staging 工作目录不存在：{staging_work_dir}")

    replacements = {str(staging_work_dir): str(published_work_dir)}
    staging_manifest_path = staging_work_dir / "s2_input.json"
    current_manifest = _load_json_dict(staging_manifest_path)
    is_business_bid = _is_business_bid(skill_workspace.get("bidType"))
    trusted_manifest = skill_workspace.get("_trustedManifest")
    if not is_business_bid:
        if not isinstance(trusted_manifest, dict) or current_manifest != trusted_manifest:
            raise RuntimeError("技术标目录 manifest 在 Opencode 执行期间被修改。")
        manifest = copy.deepcopy(trusted_manifest)
    else:
        manifest = current_manifest
    staging_output_file = Path(
        str(manifest.get("outputFile") or staging_work_dir / safe_filename(settings.s2_toc_output_file_name, "toc.json"))
    ).expanduser()
    staging_evidence_file = (
        Path(str(manifest.get("evidenceFile") or "")).expanduser()
        if is_business_bid
        else None
    )
    if is_business_bid:
        manifest = _remap_workspace_paths(manifest, replacements)
    else:
        manifest["workDir"] = str(published_work_dir)
        manifest["outputFile"] = str(published_work_dir / staging_output_file.name)
        for key in ("templateFile", "attachFile"):
            if str(manifest.get(key) or ""):
                manifest[key] = _remap_workspace_paths(manifest[key], replacements)
        for tender_file in manifest.get("tenderFiles") or []:
            if not isinstance(tender_file, dict):
                continue
            for key in ("path", "originalPath"):
                if str(tender_file.get(key) or ""):
                    tender_file[key] = _remap_workspace_paths(tender_file[key], replacements)
    manifest_path = published_work_dir / "s2_input.json"
    output_file = Path(
        str(manifest.get("outputFile") or published_work_dir / safe_filename(settings.s2_toc_output_file_name, "toc.json"))
    ).expanduser()
    evidence_file = (
        Path(
            str(
                manifest.get("evidenceFile")
                or published_work_dir / safe_filename(settings.s2_toc_evidence_file_name, "toc_evidence.json")
            )
        ).expanduser()
        if is_business_bid
        else None
    )
    business_outline_file = published_work_dir / "outline.json"
    tender_map_inputs_file = published_work_dir / "tender_map_inputs.json"
    history_bid_outline_inputs_file = published_work_dir / "history_bid_outline_inputs.json"
    manifest["workDir"] = str(published_work_dir)
    manifest["outputFile"] = str(output_file)
    if evidence_file is not None:
        manifest["evidenceFile"] = str(evidence_file)
    else:
        manifest.pop("evidenceFile", None)
    staging_manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    if is_business_bid:
        _remap_json_file(
            staging_output_file,
            replacements,
            ({"outputFile": str(output_file), "evidenceFile": str(evidence_file)} if evidence_file is not None else None),
        )
    if staging_evidence_file is not None:
        _remap_json_file(staging_evidence_file, replacements)
    _remap_json_file(staging_work_dir / "outline.json", replacements)
    _remap_json_file(staging_work_dir / "tender_map_inputs.json", replacements)
    _remap_json_file(staging_work_dir / "history_bid_outline_inputs.json", replacements)
    compose_report_path = staging_work_dir / "outline_compose_report.json"
    if is_business_bid:
        _remap_json_file(compose_report_path, replacements)
    else:
        _remap_json_file(compose_report_path, {}, {"outputFile": str(output_file)})

    result = (
        _remap_workspace_paths(toc_result, replacements)
        if is_business_bid
        else copy.deepcopy(toc_result)
    )
    result["outputFile"] = str(output_file)
    if evidence_file is not None:
        result["evidenceFile"] = str(evidence_file)
    else:
        result.pop("evidenceFile", None)
    if (staging_work_dir / "outline.json").exists():
        result["businessOutlineFile"] = str(business_outline_file)
    if (staging_work_dir / "tender_map_inputs.json").exists():
        result["tenderMapInputsFile"] = str(tender_map_inputs_file)
    if (staging_work_dir / "history_bid_outline_inputs.json").exists():
        result["historyBidOutlineInputsFile"] = str(history_bid_outline_inputs_file)
    if isinstance(result.get("opencodeOutput"), dict):
        result["opencodeOutput"]["workDir"] = str(published_work_dir)
        result["opencodeOutput"]["manifestPath"] = str(manifest_path)
        result["opencodeOutput"]["canonicalManifestPath"] = str(manifest_path)
        result["opencodeOutput"]["tocJsonPath"] = str(output_file)
        if evidence_file is not None:
            result["opencodeOutput"]["evidencePath"] = str(evidence_file)
        else:
            result["opencodeOutput"].pop("evidencePath", None)
        if (staging_work_dir / "outline.json").exists():
            result["opencodeOutput"]["businessOutlinePath"] = str(business_outline_file)
        if (staging_work_dir / "tender_map_inputs.json").exists():
            result["opencodeOutput"]["tenderMapInputsPath"] = str(tender_map_inputs_file)
        if (staging_work_dir / "history_bid_outline_inputs.json").exists():
            result["opencodeOutput"]["historyBidOutlineInputsPath"] = str(history_bid_outline_inputs_file)

    previous_archive = ""
    try:
        previous_archive = _archive_workspace_if_exists(published_work_dir, archive_root, "previous")
        published_work_dir.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(staging_work_dir), str(published_work_dir))
    except Exception:
        previous_archive_path = Path(previous_archive) if previous_archive else None
        if previous_archive_path and previous_archive_path.exists() and not published_work_dir.exists():
            shutil.move(str(previous_archive_path), str(published_work_dir))
        raise

    publish_result = {
        "result": result,
        "workDir": str(published_work_dir),
        "stagingWorkDir": str(staging_work_dir),
        "archiveRoot": str(archive_root),
        "previousArchive": previous_archive,
        "manifestPath": str(manifest_path),
        "canonicalManifestPath": str(manifest_path),
        "outputFile": str(output_file),
    }
    if evidence_file is not None:
        publish_result["evidenceFile"] = str(evidence_file)
    return publish_result


def _archive_workspace_if_exists(target: Path, archive_root: Path, label: str) -> str:
    if not target.exists():
        return ""
    archive_root.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    archive_path = _unique_path(archive_root / f"{timestamp}-{label}-{target.name}")
    shutil.move(str(target), str(archive_path))
    return str(archive_path)


def _remove_manifest_alias(project_dir: Path) -> None:
    alias_path = project_dir / "s2.json"
    if alias_path.exists():
        alias_path.unlink()


def _remap_json_file(path: Path, replacements: dict[str, str], updates: dict[str, Any] | None = None) -> None:
    if not path.exists():
        return
    payload = _remap_workspace_paths(_load_json_dict(path), replacements)
    if updates:
        payload.update(updates)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _remap_workspace_paths(value: Any, replacements: dict[str, str]) -> Any:
    if isinstance(value, str):
        result = value
        for old, new in replacements.items():
            result = result.replace(old, new)
        return result
    if isinstance(value, list):
        return [_remap_workspace_paths(item, replacements) for item in value]
    if isinstance(value, dict):
        return {key: _remap_workspace_paths(item, replacements) for key, item in value.items()}
    return value



def _heading_style_for_line(line: str) -> str | None:
    text = str(line or "").strip()
    if re.match(r"^第[一二三四五六七八九十百千万零〇两0-9]+章", text):
        return "Heading 1"
    match = re.match(r"^(?P<number>\d+(?:\.\d+)*)(?:[\s　:：、.-]+)", text)
    if not match:
        return None
    level = min(match.group("number").count(".") + 1, 4)
    return f"Heading {level}"


def _write_text_docx(path: Path, title: str, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph(title)
    for line in str(text or "").splitlines():
        if line.strip():
            style = _heading_style_for_line(line)
            document.add_paragraph(line.strip(), style=style) if style else document.add_paragraph(line.strip())
    document.save(path)
    return path


def _copy_tender_inputs(file_records: list[dict[str, Any]], work_dir: Path, parse_storage: dict[str, Any]) -> list[dict[str, str]]:
    copied: list[dict[str, str]] = []
    for index, record in enumerate(file_records, start=1):
        source = Path(str(record.get("path") or "")).expanduser()
        if not source.exists() or source.suffix.lower() != ".docx":
            continue
        name = safe_filename(str(record.get("name") or source.name), f"tender-{index}.docx")
        destination = _unique_path(work_dir / name)
        shutil.copy2(source, destination)
        copied.append(
            {
                "id": str(record.get("id") or f"tender-{index}"),
                "name": name,
                "path": str(destination),
                "originalPath": str(source),
            }
        )
    if not copied:
        combined_text_path = Path(str(parse_storage.get("combinedTextPath") or "")).expanduser()
        if combined_text_path.exists():
            text = combined_text_path.read_text(encoding="utf-8", errors="replace")
            generated_path = _write_text_docx(work_dir / "tender-from-s1-text.docx", "招标文件解析文本", text)
            copied.append(
                {
                    "id": "TEN-S1-TEXT",
                    "name": "招标文件解析文本.docx",
                    "path": str(generated_path),
                    "originalPath": str(combined_text_path),
                }
            )
    return copied


def _copy_template_inputs(file_records: list[dict[str, Any]], work_dir: Path, project_id: str) -> tuple[Path | None, Path | None]:
    docx_records = [
        record
        for record in file_records
        if Path(str(record.get("path") or "")).expanduser().exists()
        and Path(str(record.get("path") or "")).suffix.lower() == ".docx"
    ]
    attach_record = next((record for record in docx_records if _looks_like_attachment_template(record)), None)
    template_record = next(
        (
            record
            for record in docx_records
            if record is not attach_record
        ),
        None,
    )
    if template_record is None and docx_records:
        template_record = docx_records[0]

    template_path = None
    if template_record is not None:
        template_path = _copy_single_template(template_record, work_dir / "template-main.docx")
    elif file_records:
        template_path = _copy_visual_template_input(project_id, file_records[0], work_dir / "template-main.docx")
    attach_path = (
        _copy_single_template(attach_record, work_dir / "template-attachment.docx")
        if attach_record is not None and attach_record is not template_record
        else None
    )
    return template_path, attach_path


def _looks_like_attachment_template(record: dict[str, Any]) -> bool:
    role = str(record.get("role") or record.get("templateRole") or record.get("type") or "").lower()
    if role in {"attachment", "attachments", "appendix", "appendices", "attach"}:
        return True
    name = str(record.get("name") or Path(str(record.get("path") or "")).name)
    return bool(re.search(r"(附表|附件|appendix|attachment|attach)", name, re.IGNORECASE))


def _copy_single_template(record: dict[str, Any], destination: Path) -> Path:
    source = Path(str(record.get("path") or "")).expanduser()
    if source.suffix.lower() == ".docx" and not is_valid_docx_file(source):
        source_label = "系统默认模板" if str(record.get("source") or "") == "system-default" else "投标模板"
        raise ValueError(f"{source_label}不是有效 DOCX 文件，请重新上传或更换默认模板。")
    resolved_destination = destination.with_suffix(source.suffix.lower() or ".docx")
    shutil.copy2(source, resolved_destination)
    return resolved_destination


def _copy_visual_template_input(project_id: str, record: dict[str, Any], destination: Path) -> Path | None:
    source = Path(str(record.get("path") or "")).expanduser()
    if not source.exists() or source.suffix.lower() not in {".pdf", *IMAGE_SUFFIXES}:
        return None
    text, meta = _ocr_fallback_text(project_id, record, source)
    if not text:
        message = meta.get("message") if isinstance(meta, dict) else ""
        raise ValueError(f"投标模板为图片或 PDF，但视觉模型未能读取：{message or '未知错误'}")
    return _write_text_docx(destination, str(record.get("name") or source.name), text)


def _unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem = path.stem
    suffix = path.suffix
    for index in range(2, 1000):
        candidate = path.with_name(f"{stem}-{index}{suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"无法为文件生成唯一路径：{path}")


def _outline_bid_type(value: str) -> str:
    return require_bid_type(
        value,
        error_message="目录工作区必须显式传入技术标或商务标。",
    )
