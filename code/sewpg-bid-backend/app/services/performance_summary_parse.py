"""业绩汇总表解析与字段推断（performance_package_service 拆分）。

职责：解析业绩汇总 .docx（汇总表选择、表头识别、字段 schema 推断）、
从行值提取核心字段与时间事实，以及业绩条目编辑后回写行值
（_sync_performance_item_row_values 族）。实现自
app.services.performance_package_service 纯搬迁，不改行为与签名；
门面 performance_package_service 仍 re-export 本模块符号。
"""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document

from app.services.peripheral import PeripheralError


PERFORMANCE_ITEM_EDITABLE_TEXT_FIELDS = {
    "projectName": "project_name",
    "customerName": "customer_name",
    "partnerName": "partner_name",
    "turbineModel": "turbine_model",
    "contractQuantity": "contract_quantity",
    "trialOperationQuantity": "trial_operation_quantity",
    "commissionedCapacityMw": "commissioned_capacity_mw",
    "deliveryOrOperationTime": "delivery_or_operation_time",
    "contactInfo": "contact_info",
}
PERFORMANCE_ITEM_EDITABLE_YEAR_FIELDS = {
    "contractYear": ("contract_year", "合同年"),
    "deliveryYear": ("delivery_year", "交货年"),
    "operationYear": ("operation_year", "投运年"),
}
PERFORMANCE_ITEM_ROW_VALUE_TEXT_FIELDS = {
    "projectName": ("项目名称", ("项目名称", "合同名称", "工程名称"), ()),
    "customerName": ("买方名称", ("买方名称", "业主", "客户", "采购方"), ("联系人", "电话")),
    "partnerName": ("项目合作方单位", ("项目合作方单位", "合作方单位", "合作方"), ()),
    "turbineModel": ("型号", ("型号和规格", "风机型号", "型号", "机型"), ("数量", "台数")),
    "contractQuantity": ("合同台数", ("合同台数", "台数", "数量"), ("试运行", "试运", "240", "型号", "机型")),
    "trialOperationQuantity": ("试运行台数", ("试运行台数", "试运行", "试运", "240小时"), ("型号", "机型")),
    "commissionedCapacityMw": ("投运容量(MW)", ("投运容量", "容量"), ("总容量", "装机容量")),
    "deliveryOrOperationTime": (
        "交货期/投运时间",
        (
            "交货期/投运时间",
            "交货/投运",
            "交付/投运",
            "投运时间",
            "投运日期",
            "并网时间",
            "运行时间",
            "交货期",
            "交付时间",
            "交货时间",
            "交付日期",
        ),
        (),
    ),
    "contactInfo": ("联系人及电话", ("联系人", "联系方式", "电话"), ()),
}
PERFORMANCE_ITEM_ROW_VALUE_YEAR_FIELDS = {
    "contractYear": ("合同时间", ("合同时间", "合同日期", "签约时间", "签订时间", "签订日期"), ()),
    "deliveryYear": (
        "交货时间",
        ("交货期", "交付时间", "交货时间", "交付日期"),
        ("交货期/投运时间", "交货/投运", "交付/投运"),
    ),
    "operationYear": (
        "投运时间",
        ("投运时间", "投运日期", "并网时间", "运行时间"),
        ("交货期/投运时间", "交货/投运", "交付/投运"),
    ),
}
PERFORMANCE_ITEM_ROW_VALUE_REPAIR_FIELDS = {"projectName", "turbineModel"}


def parse_performance_summary_docx(content: bytes, *, file_name: str = "performance.docx") -> dict[str, Any]:
    if not file_name.lower().endswith(".docx"):
        raise PeripheralError(400, "业绩汇总表解析仅支持 .docx 文件。", "PERFORMANCE_SUMMARY_DOCX_REQUIRED")
    try:
        doc = Document(BytesIO(content))
    except Exception as exc:
        raise PeripheralError(400, "业绩汇总表无法读取，请确认文件格式。", "PERFORMANCE_SUMMARY_DOCX_INVALID") from exc

    paragraphs = [_clean_text(paragraph.text) for paragraph in doc.paragraphs]
    category_name = _infer_category_name(paragraphs, file_name=file_name)
    summary = _infer_summary(paragraphs, category_name=category_name)
    scene = _infer_scene(category_name)
    power_rating = _infer_power_rating(category_name)

    selected_table = _select_summary_table(doc.tables)
    if selected_table is None:
        raise PeripheralError(400, "未在文件中识别到业绩汇总表。", "PERFORMANCE_SUMMARY_TABLE_NOT_FOUND")

    table_rows = _table_to_rows(selected_table)
    header_index = _find_header_row(table_rows)
    if header_index < 0:
        raise PeripheralError(400, "未识别到业绩汇总表表头。", "PERFORMANCE_SUMMARY_HEADER_NOT_FOUND")
    headers = _unique_headers(table_rows[header_index])
    field_schema = [_field_schema_item(header) for header in headers]
    detail_rows: list[dict[str, Any]] = []
    for index, values in enumerate(table_rows[header_index + 1 :], start=1):
        if not _is_data_row(values):
            continue
        row_values = {
            header: _normalize_empty(values[column_index] if column_index < len(values) else "")
            for column_index, header in enumerate(headers)
        }
        core = _core_values(row_values)
        detail_rows.append(
            {
                "rowIndex": index,
                "values": row_values,
                **core,
            }
        )

    if not detail_rows:
        raise PeripheralError(400, "业绩汇总表未识别到有效明细行。", "PERFORMANCE_SUMMARY_ROWS_EMPTY")

    return {
        "categoryName": category_name,
        "scene": scene,
        "powerRating": power_rating,
        "summary": summary,
        "fieldSchema": field_schema,
        "rows": detail_rows,
        "rowCount": len(detail_rows),
        "sourceFileName": file_name,
    }


def _select_summary_table(tables: Any) -> Any | None:
    best_table = None
    best_score = -1
    for table in tables:
        rows = _table_to_rows(table)
        if not rows:
            continue
        header_index = _find_header_row(rows)
        if header_index < 0:
            continue
        header = " ".join(rows[header_index])
        score = _header_score(rows[header_index]) * 10 + max(0, len(rows) - header_index - 1)
        if any(keyword in header for keyword in ("项目", "合同", "买方", "型号")):
            score += 10
        if score > best_score:
            best_score = score
            best_table = table
    return best_table


def _table_to_rows(table: Any) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [_clean_text(cell.text) for cell in row.cells]
        if any(cell for cell in cells):
            rows.append(cells)
    return rows


def _find_header_row(rows: list[list[str]]) -> int:
    best_index = -1
    best_score = 0
    for index, row in enumerate(rows[:8]):
        score = _header_score(row)
        non_empty = sum(1 for cell in row if cell)
        if non_empty >= 3 and score > best_score:
            best_index = index
            best_score = score
    return best_index if best_score >= 2 else -1


def _header_score(row: list[str]) -> int:
    joined = " ".join(row)
    score = 0
    for keyword in ("项目名称", "合同名称", "买方", "业主", "客户", "型号", "合同台数", "试运行", "投运容量", "联系人"):
        if keyword in joined:
            score += 1
    return score


def _infer_category_name(paragraphs: list[str], *, file_name: str) -> str:
    for paragraph in paragraphs[:6]:
        if paragraph and not paragraph.startswith("合同业绩台数"):
            return paragraph[:300]
    return Path(file_name).stem[:300]


def _infer_summary(paragraphs: list[str], *, category_name: str) -> str:
    for paragraph in paragraphs[:10]:
        if paragraph and paragraph != category_name and any(keyword in paragraph for keyword in ("合同业绩", "台数", "投运容量")):
            return paragraph
    return ""


def _infer_scene(text_value: str) -> str:
    text_value = str(text_value or "")
    if "海上" in text_value:
        return "海上"
    if "陆上" in text_value:
        return "陆上"
    return ""


def _infer_power_rating(text_value: str) -> str:
    match = re.search(r"(\d+(?:\.\d+)?)\s*MW\s*(?:及以上|以上)?", str(text_value or ""), flags=re.IGNORECASE)
    if not match:
        return ""
    suffix = "及以上" if "及以上" in match.group(0) or "以上" in match.group(0) else ""
    return f"{match.group(1)}MW{suffix}"


def _unique_headers(headers: list[str]) -> list[str]:
    result: list[str] = []
    counts: dict[str, int] = {}
    for index, header in enumerate(headers, start=1):
        value = _clean_text(header) or f"字段{index}"
        count = counts.get(value, 0) + 1
        counts[value] = count
        result.append(value if count == 1 else f"{value}_{count}")
    return result


def _field_schema_item(header: str) -> dict[str, str]:
    key = _field_key(header)
    return {"key": key, "label": header, "sourceHeader": header}


def _field_key(header: str) -> str:
    normalized = re.sub(r"[^0-9A-Za-z一-龥]+", "_", str(header or "").strip()).strip("_")
    return normalized or "field"


def _sync_performance_item_row_values(
    row_values: Any,
    *,
    existing: dict[str, Any],
    text_updates: dict[str, str],
    changed_text_fields: set[str],
    year_updates: dict[str, int | None],
) -> dict[str, Any]:
    original = dict(row_values or {})
    updated = dict(original)
    text_targets = {
        field: _row_value_update_key(
            original,
            aliases=PERFORMANCE_ITEM_ROW_VALUE_TEXT_FIELDS[field][1],
            excludes=PERFORMANCE_ITEM_ROW_VALUE_TEXT_FIELDS[field][2],
            existing_value=existing.get(PERFORMANCE_ITEM_EDITABLE_TEXT_FIELDS[field]),
        )
        for field in text_updates
    }
    year_targets = {
        field: _row_value_update_key(
            original,
            aliases=PERFORMANCE_ITEM_ROW_VALUE_YEAR_FIELDS[field][1],
            excludes=PERFORMANCE_ITEM_ROW_VALUE_YEAR_FIELDS[field][2],
            existing_value=existing.get(PERFORMANCE_ITEM_EDITABLE_YEAR_FIELDS[field][0]),
            year_value=True,
        )
        for field in year_updates
    }

    for field, value in text_updates.items():
        canonical = PERFORMANCE_ITEM_ROW_VALUE_TEXT_FIELDS[field][0]
        target = text_targets[field]
        should_repair = (
            field in PERFORMANCE_ITEM_ROW_VALUE_REPAIR_FIELDS
            and target is not None
            and _normalize_empty(updated.get(target)) != _normalize_empty(value)
        )
        if field not in changed_text_fields and not should_repair:
            continue
        if target:
            updated[target] = value
        elif field in changed_text_fields and value:
            updated[canonical] = value

    for field, year in year_updates.items():
        canonical = PERFORMANCE_ITEM_ROW_VALUE_YEAR_FIELDS[field][0]
        target = year_targets[field]
        if target:
            updated[target] = _replace_row_value_year(updated.get(target), year)
        elif year is not None:
            updated[canonical] = f"{year}年"
    return updated


def _row_value_update_key(
    row_values: dict[str, Any],
    *,
    aliases: tuple[str, ...],
    excludes: tuple[str, ...],
    existing_value: Any,
    year_value: bool = False,
) -> str | None:
    candidates: list[tuple[tuple[int, int, str], str, Any]] = []
    for raw_key, raw_value in row_values.items():
        key = str(raw_key)
        normalized_key = re.sub(r"\s+", "", key)
        if any(exclude in normalized_key for exclude in excludes):
            continue
        matched_indexes = [index for index, alias in enumerate(aliases) if alias in normalized_key]
        if not matched_indexes:
            continue
        alias_index = min(matched_indexes)
        exact_match = 0 if normalized_key == aliases[alias_index] else 1
        candidates.append(((alias_index, exact_match, normalized_key), key, raw_value))
    if not candidates:
        return None
    candidates.sort(key=lambda item: item[0])

    def matches_existing(value: Any) -> bool:
        if year_value:
            return existing_value is not None and _extract_year(str(value or "")) == existing_value
        normalized_existing = _normalize_empty(existing_value)
        return bool(normalized_existing) and _normalize_empty(value) == normalized_existing

    matching = [item for item in candidates if matches_existing(item[2])]
    primary = matching[0] if matching else candidates[0]
    return primary[1]


def _replace_row_value_year(value: Any, year: int | None) -> str:
    text_value = str(value or "").strip()
    year_pattern = r"(?:19\d{2}|20\d{2}|2100)"
    if year is None:
        return re.sub(rf"{year_pattern}\s*年?", "", text_value, count=1).strip()
    if re.search(year_pattern, text_value):
        return re.sub(year_pattern, str(year), text_value, count=1)
    return f"{year}年"


def _core_values(row_values: dict[str, str]) -> dict[str, str]:
    turbine_model = _first_value(row_values, ("型号和规格", "型号", "机型", "风机型号"))
    time_facts = _time_facts(row_values)
    return {
        "projectName": _first_value(row_values, ("项目名称", "合同名称", "工程名称")),
        "customerName": _first_value(row_values, ("买方名称", "业主", "客户", "采购方")),
        "partnerName": _first_value(row_values, ("项目合作方单位", "合作方单位", "合作方")),
        "turbineModel": turbine_model,
        "turbineModels": _split_turbine_models(turbine_model),
        "contractQuantity": _first_value(row_values, ("合同台数", "台数", "数量")),
        "trialOperationQuantity": _first_value(row_values, ("试运行台数", "240", "试运")),
        "commissionedCapacityMw": _first_value(row_values, ("投运容量", "容量")),
        "deliveryOrOperationTime": time_facts["deliveryOrOperationTimeRaw"],
        "contractYear": time_facts.get("contractYear"),
        "deliveryYear": time_facts.get("deliveryYear"),
        "operationYear": time_facts.get("operationYear"),
        "timeFacts": time_facts,
        "contactInfo": _first_value(row_values, ("联系人", "电话", "联系方式")),
    }


def _first_value(row_values: dict[str, str], keywords: tuple[str, ...]) -> str:
    for key, value in row_values.items():
        if any(keyword in key for keyword in keywords):
            return value
    return ""


def _split_turbine_models(value: str) -> list[str]:
    text_value = _normalize_empty(value)
    if not text_value:
        return []
    normalized = re.sub(r"[，,、/；;]+", " ", text_value)
    parts = re.split(r"\s+", normalized)
    result: list[str] = []
    seen: set[str] = set()
    for part in parts:
        candidate = part.strip("()（）[]【】")
        if not candidate:
            continue
        if not re.search(r"\d", candidate):
            continue
        if not re.search(r"[A-Za-z]|MW|mw|-", candidate):
            continue
        key = candidate.upper()
        if key in seen:
            continue
        seen.add(key)
        result.append(candidate)
    return result


def _time_facts(row_values: dict[str, str]) -> dict[str, Any]:
    contract_raw = _value_by_header(row_values, ("合同时间", "合同日期", "签约时间", "签订时间", "签订日期"))
    operation_raw = _value_by_header(row_values, ("投运时间", "投运日期", "并网时间", "运行时间"))
    delivery_raw = _value_by_header(row_values, ("交货期", "交付时间", "交货时间", "交付日期"))
    delivery_or_operation_raw = _value_by_header(row_values, ("交货期/投运时间", "交货/投运", "交付/投运"))
    combined_raw = delivery_or_operation_raw or operation_raw or delivery_raw
    contract_year = _extract_year(contract_raw)
    delivery_year = _extract_year(delivery_raw)
    operation_year = _extract_year(operation_raw)
    combined_year = _extract_year(combined_raw)
    if delivery_or_operation_raw:
        if delivery_year is None:
            delivery_year = combined_year
        if operation_year is None:
            operation_year = combined_year
    elif not operation_raw and delivery_raw and operation_year is None:
        operation_year = None
    return {
        "contractTimeRaw": contract_raw,
        "deliveryTimeRaw": delivery_raw,
        "operationTimeRaw": operation_raw,
        "deliveryOrOperationTimeRaw": combined_raw,
        "contractYear": contract_year,
        "deliveryYear": delivery_year,
        "operationYear": operation_year,
        "years": _unique_ints([contract_year, delivery_year, operation_year, combined_year]),
    }


def _value_by_header(row_values: dict[str, str], keywords: tuple[str, ...]) -> str:
    for key, value in row_values.items():
        normalized_key = re.sub(r"\s+", "", str(key or ""))
        if any(keyword in normalized_key for keyword in keywords):
            return value
    return ""


def _extract_year(value: str) -> int | None:
    text_value = _normalize_empty(value)
    if not text_value:
        return None
    match = re.search(r"(19\d{2}|20\d{2})", text_value)
    if not match:
        return None
    year = int(match.group(1))
    if 1990 <= year <= 2100:
        return year
    return None


def _unique_ints(values: list[int | None]) -> list[int]:
    result: list[int] = []
    seen: set[int] = set()
    for value in values:
        if value is None or value in seen:
            continue
        seen.add(value)
        result.append(value)
    return result


def _parse_year_filter(value: Any) -> int | None:
    year = _extract_year(str(value or ""))
    return year


def _parse_item_year(value: Any, *, field_label: str) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        raise PeripheralError(400, f"{field_label}无效，请输入 1990-2100 之间的年份。", "PERFORMANCE_ITEM_YEAR_INVALID")
    if isinstance(value, float) and value.is_integer():
        text_value = str(int(value))
    else:
        text_value = str(value).strip()
    if not text_value:
        return None
    year = _extract_year(text_value)
    if year is None:
        raise PeripheralError(400, f"{field_label}无效，请输入 1990-2100 之间的年份。", "PERFORMANCE_ITEM_YEAR_INVALID")
    return year


def _is_data_row(values: list[str]) -> bool:
    meaningful = [value for value in values if _normalize_empty(value)]
    if len(meaningful) < 2:
        return False
    joined = " ".join(meaningful)
    return not all(keyword in joined for keyword in ("序号", "型号", "项目"))


def _normalize_empty(value: Any) -> str:
    text_value = _clean_text(str(value or ""))
    if text_value in {"/", "-", "—", "无", "N/A", "NA"}:
        return ""
    return text_value


def _clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u3000", " ")).strip()
