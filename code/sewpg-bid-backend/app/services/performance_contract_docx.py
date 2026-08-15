"""业绩合同 docx 拆分/渲染与 OOXML 清洗（performance_package_service 拆分）。

职责：把整本合同 .docx 按标题/分页拆成单项目合同块、渲染单项合同 docx、
合同块与业绩行的匹配打分，以及字体/分页/DrawingML 等 OOXML 清洗与
soffice 规范化。实现自 app.services.performance_package_service 纯搬迁，
不改行为与签名；门面 performance_package_service 仍 re-export 本模块符号。
"""

from __future__ import annotations

import logging
import re
import shutil
import subprocess
import tempfile
import unicodedata
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from docx.shared import Pt

from app.services.peripheral import PeripheralError
from app.services.file_utils import safe_filename
from app.services.performance_summary_parse import _clean_text


logger = logging.getLogger(__name__)


CONTRACT_OUTPUT_EAST_ASIA_FONT = "宋体"
CONTRACT_OUTPUT_WESTERN_FONT = "Times New Roman"
CONTRACT_OUTPUT_SYMBOL_FONT = "Symbol"
RELATIONSHIP_NAMESPACE = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC_NAMESPACE = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
PIC_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/picture"
THEME_EAST_ASIA_SCRIPTS = {"Hans", "Hant", "Jpan", "Hang"}
OOXML_NAMESPACES = {
    "w": WORD_NAMESPACE,
    "mc": MC_NAMESPACE,
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": WP_NAMESPACE,
    "a": A_NAMESPACE,
    "pic": PIC_NAMESPACE,
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "wp14": "http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing",
    "w15": "http://schemas.microsoft.com/office/word/2012/wordml",
}

for _ooxml_prefix, _ooxml_namespace in OOXML_NAMESPACES.items():
    ET.register_namespace(_ooxml_prefix, _ooxml_namespace)


def split_performance_contract_docx(content: bytes, *, file_name: str = "contract.docx") -> list[dict[str, Any]]:
    if not file_name.lower().endswith(".docx"):
        return []
    try:
        source_doc = Document(BytesIO(content))
    except Exception as exc:
        raise PeripheralError(400, "合同附件无法读取，请确认文件格式。", "PERFORMANCE_CONTRACT_DOCX_INVALID") from exc

    body_children = list(source_doc.element.body.iterchildren())
    title_indexes = [
        index
        for index, child in enumerate(body_children)
        if _is_contract_title_block(child)
    ]
    if not title_indexes:
        title_indexes = _contract_title_indexes_from_page_breaks(body_children)
    if not title_indexes:
        title_indexes = [0] if body_children else []

    chunks: list[dict[str, Any]] = []
    for chunk_index, start in enumerate(title_indexes):
        end = title_indexes[chunk_index + 1] if chunk_index + 1 < len(title_indexes) else len(body_children)
        selected_blocks = [
            child
            for child in body_children[start:end]
            if _block_kind(child) != "sectPr"
        ]
        if not selected_blocks:
            continue
        title = _first_meaningful_block_text(selected_blocks) or f"{Path(file_name).stem}-{chunk_index + 1}"
        chunks.append(
            {
                "index": chunk_index + 1,
                "title": _dedupe_repeated_title(title)[:300],
                "blocks": selected_blocks,
                "sourceDoc": source_doc,
                "blockStart": int(start),
                "blockEnd": int(max(start, end - 1)),
            }
        )
    for chunk in chunks:
        content_bytes = render_contract_item_docx(chunk)
        chunk["content"] = content_bytes
        chunk["sizeBytes"] = len(content_bytes)
    return chunks


def ensure_contract_docx_file_name(file_name: str) -> None:
    """Reject legacy .doc contracts: only .docx files can flow through the split pipeline."""
    if not str(file_name or "").lower().endswith(".docx"):
        raise PeripheralError(
            400,
            f"合同附件 {file_name} 仅支持 .docx 格式，请先转换为 .docx 后再上传。",
            "PERFORMANCE_CONTRACT_DOC_NOT_SUPPORTED",
        )


def match_contract_chunks_to_items(chunks: list[dict[str, Any]], items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not chunks or not items:
        return []
    matches: list[dict[str, Any]] = []
    used_chunk_indexes: set[int] = set()
    for item_index, item in enumerate(items):
        best_chunk: dict[str, Any] | None = None
        best_score = 0
        for chunk_index, chunk in enumerate(chunks):
            if chunk_index in used_chunk_indexes:
                continue
            score = _contract_match_score(item, chunk)
            if score > best_score:
                best_score = score
                best_chunk = {**chunk, "_chunkIndex": chunk_index}
        if best_chunk is not None and best_score >= 58:
            used_chunk_indexes.add(int(best_chunk["_chunkIndex"]))
            matches.append({"item": item, "chunk": _without_private_keys(best_chunk), "confidence": best_score, "method": "project_name"})
            continue
        if item_index < len(chunks) and item_index not in used_chunk_indexes:
            used_chunk_indexes.add(item_index)
            matches.append({"item": item, "chunk": chunks[item_index], "confidence": 45, "method": "row_order"})
    return matches


def render_contract_item_docx(chunk: dict[str, Any], *, output_title: str = "") -> bytes:
    source_doc = chunk.get("sourceDoc")
    blocks = list(chunk.get("blocks") or [])
    if source_doc is None:
        content = chunk.get("content")
        if isinstance(content, bytes):
            return content
        return b""
    title = _clean_contract_output_title(output_title or str(chunk.get("title") or "项目合同"))
    content_blocks = _contract_content_blocks(blocks)
    return _docx_blocks_to_bytes(source_doc, content_blocks, title=title)


def _docx_blocks_to_bytes(source_doc: Any, blocks: list[Any], *, title: str = "") -> bytes:
    target_doc = Document()
    target_body = target_doc.element.body
    for child in list(target_body):
        target_body.remove(child)
    if title:
        title_paragraph = target_doc.add_paragraph()
        title_run = title_paragraph.add_run(title)
        title_run.bold = True
        title_run.font.size = Pt(12)
        title_format = title_paragraph.paragraph_format
        title_format.space_before = Pt(0)
        title_format.space_after = Pt(6)
        title_format.line_spacing = 1.5
        target_body.remove(title_paragraph._p)
        target_body.append(title_paragraph._p)
    for block in blocks:
        cloned = deepcopy(block)
        _copy_related_parts(source_doc.part, target_doc.part, cloned)
        _normalize_contract_block_format(cloned)
        _sanitize_contract_drawingml(cloned)
        _stabilize_contract_table_pagination(cloned)
        target_body.append(cloned)
    source_sect = source_doc.element.body.sectPr
    if source_sect is not None:
        target_body.append(deepcopy(source_sect))
    output = BytesIO()
    target_doc.save(output)
    sanitized = _sanitize_contract_docx_fonts(output.getvalue())
    return _normalize_contract_docx_with_soffice(sanitized)


def _normalize_contract_docx_with_soffice(content: bytes) -> bytes:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        logger.warning("Skip performance contract DOCX normalization: soffice/libreoffice not found.")
        return content
    try:
        with tempfile.TemporaryDirectory(prefix="performance-contract-docx-") as temp_root:
            root = Path(temp_root)
            input_dir = root / "input"
            output_dir = root / "output"
            profile_dir = root / "profile"
            input_dir.mkdir()
            output_dir.mkdir()
            profile_dir.mkdir()
            source_path = input_dir / "source.docx"
            source_path.write_bytes(content)
            completed = subprocess.run(
                [
                    executable,
                    "--headless",
                    f"-env:UserInstallation=file://{profile_dir.as_posix()}",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    str(output_dir),
                    str(source_path),
                ],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=90,
            )
            converted_path = output_dir / source_path.name
            if not converted_path.exists():
                candidates = list(output_dir.glob("*.docx"))
                converted_path = candidates[0] if candidates else converted_path
            if not converted_path.exists():
                logger.warning(
                    "Skip performance contract DOCX normalization: no output file. stdout=%s stderr=%s",
                    (completed.stdout or "").strip(),
                    (completed.stderr or "").strip(),
                )
                return content
            normalized = converted_path.read_bytes()
            return _sanitize_contract_docx_fonts(normalized) if normalized else content
    except FileNotFoundError:
        logger.warning("Skip performance contract DOCX normalization: soffice/libreoffice not found.")
    except subprocess.TimeoutExpired:
        logger.warning("Skip performance contract DOCX normalization: conversion timed out.")
    except subprocess.CalledProcessError as exc:
        logger.warning(
            "Skip performance contract DOCX normalization: conversion failed. stdout=%s stderr=%s",
            (exc.stdout or "").strip(),
            (exc.stderr or "").strip(),
        )
    except Exception as exc:  # pragma: no cover - normalization must not break import
        logger.warning("Skip performance contract DOCX normalization: %s", exc)
    return content


def _contract_content_blocks(blocks: list[Any]) -> list[Any]:
    result: list[Any] = []
    skipped_title = False
    had_title = False
    for block in blocks:
        if not skipped_title and _is_contract_title_block(block):
            skipped_title = True
            had_title = True
            continue
        result.append(block)
    if not had_title:
        result = _trim_leading_layout_blocks(result)
    return _trim_trailing_layout_blocks(result)


def _trim_leading_layout_blocks(blocks: list[Any]) -> list[Any]:
    start = 0
    while start < len(blocks) and _is_layout_only_paragraph(blocks[start]):
        start += 1
    return blocks[start:]


def _trim_trailing_layout_blocks(blocks: list[Any]) -> list[Any]:
    end = len(blocks)
    while end > 0 and _is_layout_only_paragraph(blocks[end - 1]):
        end -= 1
    return blocks[:end]


def _normalize_contract_block_format(block: Any) -> None:
    paragraph_nodes = [block] if _block_kind(block) == "p" else list(block.xpath('.//*[local-name()="p"]'))
    for paragraph in paragraph_nodes:
        if _is_layout_only_paragraph(paragraph):
            if _has_ooxml_ancestor(paragraph, "tc"):
                _normalize_layout_paragraph_spacing(paragraph)
            continue
        ppr = paragraph.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr")
        if ppr is None:
            ppr = _insert_ooxml_child(paragraph, "pPr", 0)
        spacing = ppr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing")
        if spacing is None:
            spacing = _append_ooxml_child(ppr, "spacing")
        spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}before", "0")
        spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}after", "0")
        spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}line", "360")
        spacing.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lineRule", "auto")
        rpr = ppr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
        if rpr is None:
            rpr = _append_ooxml_child(ppr, "rPr")
        _set_ooxml_rfonts(rpr)
    for run_properties in block.xpath('.//*[local-name()="rPr"]'):
        _set_ooxml_rfonts(run_properties)


def _sanitize_contract_drawingml(block: Any) -> None:
    for picture_index, inline in enumerate(block.xpath('.//*[local-name()="inline"]'), start=1):
        extent = next((node for node in inline.iter() if node.tag.rsplit("}", 1)[-1] == "extent"), None)
        extent_attrs = {attr: extent.attrib.get(attr) for attr in ("cx", "cy")} if extent is not None else {}
        for node in inline.iter():
            local_name = node.tag.rsplit("}", 1)[-1]
            if local_name in {"docPr", "cNvPr"} and "id" in node.attrib:
                node.set("id", str(picture_index))
            elif local_name == "ext" and extent_attrs.get("cx") and extent_attrs.get("cy"):
                node.set("cx", str(extent_attrs["cx"]))
                node.set("cy", str(extent_attrs["cy"]))
    for element in list(block.iter()):
        local_name = element.tag.rsplit("}", 1)[-1]
        if local_name == "effectExtent":
            for attr_name in ("l", "t", "r", "b"):
                element.set(attr_name, "0")
        elif local_name == "blip":
            for child in list(element):
                if child.tag.rsplit("}", 1)[-1] == "extLst":
                    element.remove(child)
        elif local_name == "spPr":
            for child in list(element):
                if child.tag.rsplit("}", 1)[-1] == "ln":
                    element.remove(child)


def _stabilize_contract_table_pagination(block: Any) -> None:
    if _block_kind(block) != "tbl":
        return
    rows = [child for child in block.iterchildren() if _block_kind(child) == "tr"]
    for index, row in enumerate(rows[:-1]):
        next_row = rows[index + 1]
        if not _is_contract_table_caption_row(row):
            continue
        if not _has_drawings(next_row):
            continue
        _set_row_cant_split(row)
        _set_row_cant_split(next_row)
        _set_row_keep_next(row)


def _is_contract_table_caption_row(row: Any) -> bool:
    text_value = _dedupe_repeated_title(_block_text(row))
    if not text_value or len(text_value) > 120:
        return False
    if _has_drawings(row):
        return False
    return any(keyword in text_value for keyword in ("页", "合同", "参数", "机型", "盖章", "首页"))


def _set_row_cant_split(row: Any) -> None:
    trpr = row.find(f"{{{WORD_NAMESPACE}}}trPr")
    if trpr is None:
        trpr = _insert_ooxml_child(row, "trPr", 0)
    if trpr.find(f"{{{WORD_NAMESPACE}}}cantSplit") is None:
        trpr.insert(0, trpr.makeelement(f"{{{WORD_NAMESPACE}}}cantSplit"))


def _set_row_keep_next(row: Any) -> None:
    for paragraph in row.xpath('.//*[local-name()="p"]'):
        ppr = paragraph.find(f"{{{WORD_NAMESPACE}}}pPr")
        if ppr is None:
            ppr = _insert_ooxml_child(paragraph, "pPr", 0)
        if ppr.find(f"{{{WORD_NAMESPACE}}}keepNext") is not None:
            continue
        keep_next = ppr.makeelement(f"{{{WORD_NAMESPACE}}}keepNext")
        insert_at = 0
        pstyle = ppr.find(f"{{{WORD_NAMESPACE}}}pStyle")
        if pstyle is not None:
            insert_at = list(ppr).index(pstyle) + 1
        ppr.insert(insert_at, keep_next)


def _normalize_layout_paragraph_spacing(paragraph: Any) -> None:
    ppr = paragraph.find(f"{{{WORD_NAMESPACE}}}pPr")
    if ppr is None:
        ppr = _insert_ooxml_child(paragraph, "pPr", 0)
    spacing = ppr.find(f"{{{WORD_NAMESPACE}}}spacing")
    if spacing is None:
        spacing = _append_ooxml_child(ppr, "spacing")
    spacing.set(f"{{{WORD_NAMESPACE}}}before", "0")
    spacing.set(f"{{{WORD_NAMESPACE}}}after", "0")
    spacing.set(f"{{{WORD_NAMESPACE}}}line", "360")
    spacing.set(f"{{{WORD_NAMESPACE}}}lineRule", "auto")


def _has_ooxml_ancestor(node: Any, local_name: str) -> bool:
    return any(_block_kind(ancestor) == local_name for ancestor in node.iterancestors())


def _set_ooxml_rfonts(run_properties: Any) -> None:
    rfonts = run_properties.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is None:
        rfonts = _insert_ooxml_child(run_properties, "rFonts", 0)
    rfonts.set(f"{{{WORD_NAMESPACE}}}ascii", CONTRACT_OUTPUT_WESTERN_FONT)
    rfonts.set(f"{{{WORD_NAMESPACE}}}hAnsi", CONTRACT_OUTPUT_WESTERN_FONT)
    rfonts.set(f"{{{WORD_NAMESPACE}}}eastAsia", CONTRACT_OUTPUT_EAST_ASIA_FONT)
    rfonts.set(f"{{{WORD_NAMESPACE}}}cs", CONTRACT_OUTPUT_WESTERN_FONT)
    for attr_name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "hint"):
        rfonts.attrib.pop(f"{{{WORD_NAMESPACE}}}{attr_name}", None)


def _sanitize_contract_docx_fonts(content: bytes) -> bytes:
    source = BytesIO(content)
    target = BytesIO()
    with ZipFile(source, "r") as src_zip, ZipFile(target, "w", compression=ZIP_DEFLATED) as dst_zip:
        for info in src_zip.infolist():
            data = src_zip.read(info.filename)
            if info.filename == "word/fontTable.xml":
                data = _contract_font_table_xml()
            elif info.filename == "word/theme/theme1.xml":
                data = _sanitize_theme_fonts_xml(data)
            elif info.filename == "word/settings.xml":
                data = _sanitize_contract_settings_xml(data)
            elif info.filename.startswith("word/") and info.filename.endswith(".xml"):
                data = _sanitize_word_fonts_xml(data)
            dst_zip.writestr(info, data)
    return target.getvalue()


def _sanitize_contract_settings_xml(data: bytes) -> bytes:
    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        return data
    for local_name in ("compat", "rsids"):
        for element in list(root):
            if element.tag == f"{{{WORD_NAMESPACE}}}{local_name}":
                root.remove(element)
    return _serialize_ooxml(root)


def _sanitize_word_fonts_xml(data: bytes) -> bytes:
    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        return data
    changed = False
    for element in root.iter():
        local_name = element.tag.rsplit("}", 1)[-1]
        if local_name == "rFonts":
            element.attrib[f"{{{WORD_NAMESPACE}}}ascii"] = CONTRACT_OUTPUT_WESTERN_FONT
            element.attrib[f"{{{WORD_NAMESPACE}}}hAnsi"] = CONTRACT_OUTPUT_WESTERN_FONT
            element.attrib[f"{{{WORD_NAMESPACE}}}eastAsia"] = CONTRACT_OUTPUT_EAST_ASIA_FONT
            element.attrib[f"{{{WORD_NAMESPACE}}}cs"] = CONTRACT_OUTPUT_WESTERN_FONT
            for attr_name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "hint"):
                element.attrib.pop(f"{{{WORD_NAMESPACE}}}{attr_name}", None)
            changed = True
    if not changed:
        return data
    return _serialize_ooxml(root)


def _sanitize_theme_fonts_xml(data: bytes) -> bytes:
    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        return data
    for tag_name in ("latin", "ea", "cs"):
        for element in root.iter(f"{{{A_NAMESPACE}}}{tag_name}"):
            element.set("typeface", CONTRACT_OUTPUT_EAST_ASIA_FONT if tag_name == "ea" else CONTRACT_OUTPUT_WESTERN_FONT)
    for element in root.iter(f"{{{A_NAMESPACE}}}font"):
        script = str(element.get("script") or "")
        element.set(
            "typeface",
            CONTRACT_OUTPUT_EAST_ASIA_FONT if script in THEME_EAST_ASIA_SCRIPTS else CONTRACT_OUTPUT_WESTERN_FONT,
        )
    return _serialize_ooxml(root)


def _serialize_ooxml(root: ET.Element) -> bytes:
    _sanitize_mc_ignorable(root)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _sanitize_mc_ignorable(root: ET.Element) -> None:
    ignorable_attr = f"{{{MC_NAMESPACE}}}Ignorable"
    value = str(root.attrib.get(ignorable_attr) or "").strip()
    if not value:
        return
    used_namespaces = {str(element.tag).split("}", 1)[0][1:] for element in root.iter() if str(element.tag).startswith("{")}
    kept = [
        prefix
        for prefix in value.split()
        if OOXML_NAMESPACES.get(prefix) in used_namespaces
    ]
    if kept:
        root.attrib[ignorable_attr] = " ".join(kept)
    else:
        root.attrib.pop(ignorable_attr, None)


def _contract_font_table_xml() -> bytes:
    return f"""<?xml version='1.0' encoding='UTF-8' standalone='yes'?>
<w:fonts xmlns:w="{WORD_NAMESPACE}">
  <w:font w:name="{CONTRACT_OUTPUT_EAST_ASIA_FONT}">
    <w:charset w:val="86"/>
    <w:family w:val="roman"/>
    <w:pitch w:val="variable"/>
  </w:font>
  <w:font w:name="{CONTRACT_OUTPUT_WESTERN_FONT}">
    <w:charset w:val="00"/>
    <w:family w:val="roman"/>
    <w:pitch w:val="variable"/>
  </w:font>
  <w:font w:name="{CONTRACT_OUTPUT_SYMBOL_FONT}">
    <w:charset w:val="02"/>
    <w:family w:val="auto"/>
    <w:pitch w:val="variable"/>
  </w:font>
</w:fonts>
""".encode("utf-8")


def _append_ooxml_child(parent: Any, local_name: str) -> Any:
    child = parent.makeelement(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{local_name}")
    parent.append(child)
    return child


def _insert_ooxml_child(parent: Any, local_name: str, index: int) -> Any:
    child = parent.makeelement(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{local_name}")
    parent.insert(index, child)
    return child


def _copy_related_parts(source_part: Any, target_part: Any, element: Any) -> None:
    for node in element.iter():
        for attr_name, attr_value in list(node.attrib.items()):
            if not attr_name.startswith(RELATIONSHIP_NAMESPACE):
                continue
            relation = source_part.rels.get(attr_value)
            if relation is None:
                continue
            if relation.is_external:
                new_rid = target_part.relate_to(relation.target_ref, relation.reltype, is_external=True)
            else:
                new_rid = target_part.relate_to(relation.target_part, relation.reltype)
            node.attrib[attr_name] = new_rid


def _clean_contract_output_title(value: str) -> str:
    text_value = _dedupe_repeated_title(_clean_text(value))
    text_value = re.sub(r"(采购)?合同$", "", text_value).strip(" ，,。；;:-_")
    if len(text_value) > 80:
        text_value = text_value[:80].rstrip(" ，,。；;:-_")
    return text_value or "项目合同"


def _is_contract_title_block(block: Any) -> bool:
    if _block_kind(block) != "p":
        return False
    text_value = _dedupe_repeated_title(_block_text(block))
    if len(text_value) < 8:
        return False
    if _has_drawings(block):
        return False
    if _is_empty_page_break_paragraph(block):
        return False
    if re.fullmatch(r"标段[一二三四五六七八九十\d]+[:：]?", text_value):
        return False
    if not any(keyword in text_value for keyword in ("合同", "项目", "工程", "风电", "设备", "采购", "供货", "EPC")):
        return False
    if any(keyword in text_value for keyword in ("合同首页", "签字盖章页", "技术数据页", "预验收证明", "运行证明", "并网证明")):
        return False
    return True


def _contract_title_indexes_from_page_breaks(blocks: list[Any]) -> list[int]:
    indexes: list[int] = []
    next_meaningful_starts_chunk = True
    for index, block in enumerate(blocks):
        if _block_kind(block) == "sectPr":
            continue
        text_value = _dedupe_repeated_title(_block_text(block))
        meaningful = bool(text_value) or _block_kind(block) == "tbl" or _has_drawings(block)
        if next_meaningful_starts_chunk and meaningful:
            indexes.append(index)
            next_meaningful_starts_chunk = False
        if _block_has_page_break(block):
            next_meaningful_starts_chunk = True
    return indexes


def _first_meaningful_block_text(blocks: list[Any]) -> str:
    for block in blocks:
        text_value = _dedupe_repeated_title(_block_text(block))
        if text_value:
            return text_value
    return ""


def _block_kind(block: Any) -> str:
    return str(getattr(block, "tag", "")).rsplit("}", 1)[-1]


def _block_text(block: Any) -> str:
    return _clean_text("".join(block.itertext()))


def _block_has_page_break(block: Any) -> bool:
    return bool(block.xpath('.//*[local-name()="br" and @*[local-name()="type"]="page"]'))


def _has_drawings(block: Any) -> bool:
    return bool(block.xpath('.//*[local-name()="drawing" or local-name()="pict"]'))


def _is_empty_page_break_paragraph(block: Any) -> bool:
    return _block_kind(block) == "p" and not _block_text(block) and _block_has_page_break(block)


def _is_layout_only_paragraph(block: Any) -> bool:
    return _block_kind(block) == "p" and not _block_text(block) and not _has_drawings(block)


def _dedupe_repeated_title(value: str) -> str:
    text_value = _clean_text(value)
    if len(text_value) < 8:
        return text_value
    for unit_length in range(4, max(5, len(text_value) // 2 + 1)):
        if len(text_value) % unit_length:
            continue
        unit = text_value[:unit_length]
        repeats = len(text_value) // unit_length
        if repeats >= 2 and unit * repeats == text_value:
            return unit
    half = len(text_value) // 2
    if half >= 8 and text_value[:half] == text_value[half:]:
        return text_value[:half]
    for end in range(8, min(len(text_value), 180)):
        unit = text_value[:end]
        if text_value.startswith(unit + unit):
            return unit
    return text_value


def _contract_match_score(item: dict[str, Any], chunk: dict[str, Any]) -> int:
    project_name = str(item.get("projectName") or "")
    chunk_title = str(chunk.get("title") or "")
    if not project_name or not chunk_title:
        return 0
    project_norm = _match_text(project_name)
    chunk_norm = _match_text(chunk_title)
    if not project_norm or not chunk_norm:
        return 0
    if project_norm in chunk_norm or chunk_norm in project_norm:
        return 96
    project_tokens = _match_tokens(project_name)
    chunk_tokens = _match_tokens(chunk_title)
    if not project_tokens:
        return 0
    overlap = project_tokens & chunk_tokens
    score = int(len(overlap) / max(1, len(project_tokens)) * 100)
    longest = _longest_common_substring_length(project_norm, chunk_norm)
    score = max(score, int(longest / max(1, len(project_norm)) * 100))
    ordered = _ordered_character_coverage(project_norm, chunk_norm)
    score = max(score, int(ordered * 100))
    model_values = item.get("turbineModels") or []
    if isinstance(model_values, list) and any(_match_text(model) and _match_text(model) in chunk_norm for model in model_values):
        score += 8
    customer = _match_text(item.get("customerName") or "")
    if customer and customer in chunk_norm:
        score += 8
    return max(0, min(100, score))


def _match_text(value: Any) -> str:
    text_value = unicodedata.normalize("NFKC", str(value or "")).lower()
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", text_value)


def _match_tokens(value: Any) -> set[str]:
    normalized = _match_text(value)
    if not normalized:
        return set()
    tokens = {match.group(0) for match in re.finditer(r"[a-z]+\d*(?:\.\d+)?|\d+(?:\.\d+)?mw|\d+(?:\.\d+)?万千瓦|\d+(?:\.\d+)?千瓦|\d+x\d+|\d+×\d+", normalized)}
    for width in (2, 3, 4):
        for start in range(0, max(0, len(normalized) - width + 1)):
            token = normalized[start : start + width]
            if len(token) == width:
                tokens.add(token)
    return tokens


def _ordered_character_coverage(left: str, right: str) -> float:
    if not left or not right:
        return 0.0
    right_index = 0
    matched = 0
    for char in left:
        found_at = right.find(char, right_index)
        if found_at < 0:
            continue
        matched += 1
        right_index = found_at + 1
    return matched / max(1, len(left))


def _longest_common_substring_length(left: str, right: str) -> int:
    if not left or not right:
        return 0
    previous = [0] * (len(right) + 1)
    best = 0
    for left_char in left:
        current = [0] * (len(right) + 1)
        for index, right_char in enumerate(right, start=1):
            if left_char == right_char:
                current[index] = previous[index - 1] + 1
                best = max(best, current[index])
        previous = current
    return best


def _without_private_keys(value: dict[str, Any]) -> dict[str, Any]:
    return {key: item for key, item in value.items() if not str(key).startswith("_")}


def _contract_item_file_name(row_index: Any, project_name: str, chunk_index: Any = 0) -> str:
    safe_project = safe_filename(_clean_text(project_name), "performance.docx")[:96].strip(" .-")
    if not safe_project:
        safe_project = f"项目{int(chunk_index or 0) or 1}"
    try:
        row_number = int(row_index or 0)
    except (TypeError, ValueError):
        row_number = 0
    prefix = f"{row_number:03d}-" if row_number > 0 else ""
    return f"{prefix}{safe_project}_合同.docx"
