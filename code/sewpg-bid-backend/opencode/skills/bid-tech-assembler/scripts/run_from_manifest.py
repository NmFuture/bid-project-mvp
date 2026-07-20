#!/usr/bin/env python3
"""Run bid-tech-assembler from a backend-prepared manifest."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent


def load_manifest(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise RuntimeError("manifest must be a JSON object")
    return data


def as_path(value: Any, *, required: bool = False, label: str = "path") -> Path | None:
    text = str(value or "").strip()
    if not text:
        if required:
            raise RuntimeError(f"{label} is required")
        return None
    path = Path(text)
    if required and not path.exists():
        raise RuntimeError(f"{label} does not exist: {path}")
    return path


def run_capture(command: list[str]) -> str:
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
    if result.returncode != 0:
        stderr = (result.stderr or "").strip()
        stdout = (result.stdout or "").strip()
        detail = "\n".join(part for part in (stdout, stderr) if part)
        raise RuntimeError(f"command failed ({result.returncode}): {' '.join(command)}\n{detail}")
    return result.stdout or ""


def safe_filename(value: str, fallback: str) -> str:
    text = re.sub(r"[\\/:*?\"<>|]+", "-", str(value or "").strip())
    text = re.sub(r"\s+", " ", text).strip(" .")
    return text or fallback


def create_simple_master(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("")
    doc.save(str(path))


def ensure_master_template(manifest: dict[str, Any], work_dir: Path) -> Path:
    template_dir = work_dir / "templates"
    target = template_dir / "技术投标母版模板.docx"
    if target.exists():
        return target

    sample = as_path(manifest.get("templateFile"), required=False)
    style = SKILL_DIR / "references" / "heading_style.json"
    if sample and sample.exists() and style.exists():
        try:
            run_capture(
                [
                    sys.executable,
                    str(SKILL_DIR / "tools" / "create_tech_master.py"),
                    "--sample",
                    str(sample),
                    "--style",
                    str(style),
                    "--out",
                    str(target),
                ]
            )
            return target
        except Exception:
            # The uploaded template can be malformed or too minimal in tests. Fall
            # back to a clean python-docx master; finalize.py reapplies styles.
            pass

    create_simple_master(target)
    return target


def summarize_plan(plan_path: Path) -> dict[str, Any]:
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    if not isinstance(plan, list):
        return {"total": 0, "byStatus": {}}
    counts = Counter(str(item.get("status") or "") for item in plan if isinstance(item, dict))
    used_paths: list[str] = []
    for item in plan:
        if not isinstance(item, dict):
            continue
        if str(item.get("status") or "") in {"MATCHED", "ADAPTED", "COVER"}:
            used_paths.extend(str(path) for path in item.get("paths") or [] if str(path).strip())
    return {
        "total": len(plan),
        "byStatus": dict(counts),
        "usedPathCount": len(set(used_paths)),
    }


def build_summary_and_warnings(
    plan: list[dict[str, Any]],
    merger_result: dict[str, Any],
    scan: dict[str, Any],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    counts = Counter(str(item.get("status") or "") for item in plan if isinstance(item, dict))
    used_paths = {
        str(path)
        for item in plan
        if isinstance(item, dict) and str(item.get("status") or "") in {"MATCHED", "ADAPTED", "COVER"}
        for path in item.get("paths") or []
        if str(path).strip()
    }
    placeholder_count = len(scan.get("placeholders") or [])
    empty_section_count = len(scan.get("empty_leaf_headings") or [])
    duplicate_heading_count = len(scan.get("dup_alerts") or [])
    abnormal_heading_count = sum(
        len(scan.get(key) or [])
        for key in ("ghost_chapters", "invalid_h1", "invalid_prefix")
    )
    format_risk_count = empty_section_count + duplicate_heading_count + abnormal_heading_count

    warnings: list[dict[str, Any]] = []
    for item in merger_result.get("warnings") or []:
        if not isinstance(item, dict):
            continue
        count = _safe_int(item.get("count"))
        if count <= 0:
            continue
        warnings.append(
            {
                "code": str(item.get("code") or "MATERIAL_MERGE_FAILED"),
                "message": str(item.get("message") or "素材处理失败，已跳过"),
                "count": count,
            }
        )

    unmatched_count = counts.get("UNMATCHED", 0)
    if unmatched_count:
        warnings.append(
            {
                "code": "DIRECTORY_UNMATCHED",
                "message": f"{unmatched_count} 个目录节点未匹配到素材，已保留标题并插入占位提示",
                "count": unmatched_count,
            }
        )
    if placeholder_count:
        warnings.append(
            {
                "code": "PLACEHOLDER_REMAINS",
                "message": f"组装稿中仍有 {placeholder_count} 个占位符，需后续补充或清洗",
                "count": placeholder_count,
            }
        )
    if format_risk_count:
        warnings.append(
            {
                "code": "FORMAT_RISK",
                "message": (
                    f"发现 {format_risk_count} 项格式风险"
                    f"（空章节 {empty_section_count}、重复标题 {duplicate_heading_count}、"
                    f"异常标题 {abnormal_heading_count}）"
                ),
                "count": format_risk_count,
            }
        )

    summary = {
        "total": len(plan),
        "byStatus": dict(counts),
        "usedPathCount": len(used_paths),
        "assembledCount": _safe_int(merger_result.get("merged_materials")),
        "unmatchedCount": unmatched_count,
        "needsReviewCount": counts.get("NEEDS_REVIEW", 0),
        "structuralCount": counts.get("STRUCTURAL", 0),
        "verification": {
            "placeholderCount": placeholder_count,
            "emptySectionCount": empty_section_count,
            "duplicateHeadingCount": duplicate_heading_count,
            "abnormalHeadingCount": abnormal_heading_count,
        },
        "warningCount": sum(item["count"] for item in warnings),
    }
    return summary, warnings


def _safe_int(value: Any, default: int = 0) -> int:
    if isinstance(value, bool):
        return default
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return default
    return parsed if parsed >= 0 else default


def run_from_manifest(manifest_path: Path) -> dict[str, Any]:
    manifest = load_manifest(manifest_path)
    work_dir = as_path(manifest.get("workDir"), required=False) or manifest_path.parent
    work_dir.mkdir(parents=True, exist_ok=True)

    toc_source = as_path(manifest.get("tocJsonPath") or manifest.get("tocPath"), required=True, label="tocJsonPath")
    wiki_dir = as_path(manifest.get("wikiDir"), required=True, label="wikiDir")
    material_library = as_path(manifest.get("materialLibraryDir"), required=True, label="materialLibraryDir")
    assert toc_source is not None
    assert wiki_dir is not None
    assert material_library is not None

    toc_json = work_dir / "toc_entries.json"
    params_path = as_path(manifest.get("projectParamsPath"), required=False) or (work_dir / "project_params.json")
    plan_path = work_dir / "assembly_plan.json"
    gap_plan_path = as_path(manifest.get("gapPlanPath"), required=False)
    merged_path = work_dir / "bid_merged.docx"
    merger_result_path = work_dir / "assembly_merge_result.json"
    verify_result_path = work_dir / "assembly_verify_result.json"

    project_params = manifest.get("projectParams")
    if isinstance(project_params, dict) and not params_path.exists():
        params_path.parent.mkdir(parents=True, exist_ok=True)
        params_path.write_text(json.dumps(project_params, ensure_ascii=False, indent=2), encoding="utf-8")

    run_capture([sys.executable, str(SCRIPT_DIR / "parse_toc.py"), str(toc_source), "--out", str(toc_json)])
    run_capture([sys.executable, str(SCRIPT_DIR / "init_params.py"), "--toc", str(toc_source), "--out", str(params_path)])
    run_capture(
        [
            sys.executable,
            str(SCRIPT_DIR / "build_assembly.py"),
            "--toc",
            str(toc_json),
            "--wiki",
            str(wiki_dir),
            "--params",
            str(params_path),
            "--out",
            str(plan_path),
        ]
        + (["--gap-plan", str(gap_plan_path)] if gap_plan_path and gap_plan_path.exists() else [])
    )

    template = ensure_master_template(manifest, work_dir)
    prep_dir = work_dir / "bid_prep"
    run_capture(
        [
            sys.executable,
            str(SCRIPT_DIR / "merger.py"),
            "--template",
            str(template),
            "--plan",
            str(plan_path),
            "--lib",
            str(material_library),
            "--params",
            str(params_path),
            "--prep-dir",
            str(prep_dir),
            "--out",
            str(merged_path),
            "--result",
            str(merger_result_path),
        ]
    )

    output_file = as_path(manifest.get("outputFile"), required=False)
    if output_file is None:
        params = json.loads(params_path.read_text(encoding="utf-8"))
        short = safe_filename(str(params.get("project_short") or params.get("project_name") or "项目"), "项目")
        stamp = datetime.now().strftime("%Y%m%d%H%M")
        output_file = work_dir / f"投标文件-正文_{short}_{stamp}.docx"

    run_capture(
        [
            sys.executable,
            str(SCRIPT_DIR / "finalize.py"),
            "--in",
            str(merged_path),
            "--params",
            str(params_path),
            "--style",
            str(SKILL_DIR / "references" / "heading_style.json"),
            "--out",
            str(output_file),
        ]
    )
    run_capture(
        [
            sys.executable,
            str(SCRIPT_DIR / "verify.py"),
            "--docx",
            str(output_file),
            "--plan",
            str(plan_path),
            "--params",
            str(params_path),
            "--result",
            str(verify_result_path),
        ]
    )

    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    merger_result = json.loads(merger_result_path.read_text(encoding="utf-8"))
    scan = json.loads(verify_result_path.read_text(encoding="utf-8"))
    summary, warnings = build_summary_and_warnings(plan, merger_result, scan)

    return {
        "schema_version": "bid-tech-assembly-v1",
        "workDir": str(work_dir),
        "tocJson": str(toc_json),
        "planFile": str(plan_path),
        "projectParamsFile": str(params_path),
        "gapPlanFile": str(gap_plan_path) if gap_plan_path else "",
        "outputFile": str(output_file),
        "assemblyReport": "",
        "needsReview": "",
        "summary": summary,
        "warnings": warnings,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--response", choices=("full", "summary"), default="summary")
    args = parser.parse_args()

    try:
        output = run_from_manifest(Path(args.manifest))
    except Exception as exc:
        print(f"run_from_manifest failed: {exc}", file=sys.stderr)
        raise SystemExit(1) from exc

    if args.response == "summary":
        payload = {
            "schema_version": output["schema_version"],
            "outputFile": output["outputFile"],
            "assemblyReport": output["assemblyReport"],
            "needsReview": output["needsReview"],
            "planFile": output["planFile"],
            "summary": output["summary"],
            "warnings": output["warnings"],
        }
        print(json.dumps(payload, ensure_ascii=False, indent=2))
    else:
        print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
