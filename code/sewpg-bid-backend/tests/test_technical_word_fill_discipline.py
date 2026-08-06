"""word-placeholder-filler 填写纪律单测。

正文填写只按事实表清单（待填写文件 + 原占位符位置）定位字段。旧的模糊匹配链路
及其防污染补丁已随链路一起删除，对应用例作废——现在的纪律由「定位不到就标黄」
保证，而不是靠给模糊匹配打补丁。
fixture 为脱敏合成数据（通用领域词面，无真实项目数据）。
"""
from __future__ import annotations

import importlib.util
import sys
import unittest
from pathlib import Path

_SRC = (
    Path(__file__).resolve().parents[1]
    / "opencode"
    / "skills"
    / "bid-tech-word-placeholder-filler"
    / "scripts"
    / "run_from_manifest.py"
)
_SPEC = importlib.util.spec_from_file_location("tech_word_filler_under_test", _SRC)
filler = importlib.util.module_from_spec(_SPEC)
assert _SPEC.loader is not None
sys.modules["tech_word_filler_under_test"] = filler
_SPEC.loader.exec_module(filler)


def _spec_fact(label: str, value: str, placeholders: list[str], targets: list[str]) -> dict:
    return {
        "label": label,
        "value": value,
        "reviewLabel": label,
        "sourceKind": "manifest",
        "confidence": 0.97,
        "sourcePriority": 90,
        "specPlaceholders": [filler.placeholder_key(item) for item in placeholders],
        "specTargets": [filler.file_key(item) for item in targets],
    }


class NormalizationTests(unittest.TestCase):
    def test_norm_strips_punctuation(self) -> None:
        # 修复前字符类在 `\\]` 处提前闭合，norm() 实际不剥离任何标点
        self.assertEqual(filler.norm("第1段（底）塔节底部直径（m）"), "第1段底塔节底部直径m")

    def test_placeholder_key_bridges_bracket_and_comma_variants(self) -> None:
        keys = {
            filler.placeholder_key("[安全等级，待填写]"),
            filler.placeholder_key("【安全等级, 待填写】"),
            filler.placeholder_key("[待填写：安全等级]"),
        }
        self.assertEqual(keys, {"安全等级"})

    def test_split_spec_cell_handles_newline_and_semicolons(self) -> None:
        self.assertEqual(filler.split_spec_cell("[甲，待填写]；[乙，待填写]\n[丙，待填写]"), ["[甲，待填写]", "[乙，待填写]", "[丙，待填写]"])

    def test_file_key_ignores_path_and_extension(self) -> None:
        self.assertEqual(filler.file_key("华能/待填写-塔筒设计方案.docx"), filler.file_key("待填写-塔筒设计方案"))


class SpecLocateTests(unittest.TestCase):
    def _index(self, facts: list[dict], blank: str = "待填写-塔筒设计方案.docx") -> filler.SpecIndex:
        manifest = {"projectFactTable": {"fields": [
            {"placeholder": "；".join(f"[{key}，待填写]" for key in fact["specPlaceholders"])} for fact in facts
        ]}}
        return filler.build_spec_index(manifest, facts, {filler.file_key(blank)})

    def test_unique_placeholder_fills_without_guessing(self) -> None:
        fact = _spec_fact("场址要求安全等级", "S级", ["[场址安全等级，待填写]"], ["华能/待填写-塔筒设计方案.docx"])
        index = self._index([fact])
        selected, _alts, status = filler.spec_locate(index, {"label": "场址安全等级"}, "任意上下文")
        self.assertEqual(status, "unique")
        self.assertEqual(selected["value"], "S级")

    def test_table_column_header_separates_adjacent_fields(self) -> None:
        # 同列相邻字段只差「底部/顶部」，靠行标签 + 列头的完整包含区分
        facts = [
            _spec_fact("第1段（底）塔节底部直径（m）", "4.5", ["[技术方案，待填写]"], ["待填写-塔筒设计方案.docx"]),
            _spec_fact("第1段（底）塔节顶部直径（m）", "3.8", ["[技术方案，待填写]"], ["待填写-塔筒设计方案.docx"]),
            _spec_fact("第2段塔节底部直径（m）", "3.8", ["[技术方案，待填写]"], ["待填写-塔筒设计方案.docx"]),
        ]
        index = self._index(facts)
        context = "T1R2C2 / 第1段（底） / [技术方案，待填写] / 底部直径（m）"
        selected, _alts, status = filler.spec_locate(index, {"label": "技术方案"}, context)
        self.assertEqual(status, "disambiguated")
        self.assertEqual(selected["label"], "第1段（底）塔节底部直径（m）")

    def test_indistinguishable_candidates_go_manual(self) -> None:
        # 上下文没有区分信息时宁空勿错，交人工并进诊断表
        facts = [
            _spec_fact("场址要求安全等级", "S级", ["[安全等级，待填写]"], ["待填写-塔筒设计方案.docx"]),
            _spec_fact("机型认证安全等级", "IEC IA", ["[安全等级，待填写]"], ["待填写-塔筒设计方案.docx"]),
        ]
        index = self._index(facts)
        selected, alternatives, status = filler.spec_locate(index, {"label": "安全等级"}, "本项目安全等级为")
        self.assertEqual(status, "ambiguous")
        self.assertIsNone(selected)
        self.assertEqual(len(alternatives), 2)

    def test_placeholder_outside_spec_goes_manual(self) -> None:
        fact = _spec_fact("塔筒段数", "5", ["[塔筒段数，待填写]"], ["待填写-塔筒设计方案.docx"])
        index = self._index([fact])
        selected, _alts, status = filler.spec_locate(index, {"label": "服务承诺"}, "售后服务响应时间")
        self.assertEqual(status, "not_in_spec")
        self.assertIsNone(selected)

    def test_placeholder_equal_to_field_name_hits_directly(self) -> None:
        # 素材库多数占位符已拆细成字段名本身，这条是零歧义的确定性路径
        fact = _spec_fact("单台机组功率曲线保证率（%）", "97", ["[投标方案，待填写]"], ["待填写-塔筒设计方案.docx"])
        index = self._index([fact])
        selected, _alts, status = filler.spec_locate(index, {"label": "单台机组功率曲线保证率（%）"}, "任意上下文")
        self.assertEqual(status, "field_name")
        self.assertEqual(selected["value"], "97")

    def test_derived_facts_never_hijack_generic_placeholders(self) -> None:
        # 派生事实（无清单元数据）不得进字段名索引：label 恰好等于泛占位符文字时会劫持全文
        hijacker = {"label": "投标方案", "value": "EW10.0-220", "sourceKind": "manifest", "confidence": 0.9}
        index = filler.build_spec_index({}, [hijacker], {filler.file_key("待填写-塔筒设计方案.docx")})
        self.assertFalse(index.enabled)

    def test_target_file_is_soft_filter(self) -> None:
        # 清单文件名与实际待填写文件对不上时不做排除，否则整份文件静默填不进任何字段
        fact = _spec_fact("塔筒段数（段）", "5", ["[技术方案，待填写]"], ["别的文件.docx"])
        index = self._index([fact], blank="待填写-塔筒设计方案.docx")
        selected, _alts, status = filler.spec_locate(index, {"label": "技术方案"}, "上下文")
        self.assertEqual(status, "unique")
        self.assertEqual(selected["value"], "5")

    def test_same_placeholder_shared_by_synonym_fields_fills_when_values_agree(self) -> None:
        # 清单里同一个占位符被多个同义字段共用（年等效满负荷小时数 / 等效上网小时数 …），
        # 取值一致时没有歧义，照填
        facts = [
            _spec_fact("年等效满负荷小时数（保证值，h）", "2836", ["[投标方案，待填写]"], ["待填写-塔筒设计方案.docx"]),
            _spec_fact("等效上网小时数（保证值，h）", "2836", ["[投标方案，待填写]"], ["待填写-塔筒设计方案.docx"]),
        ]
        index = self._index(facts)
        selected, _alts, status = filler.spec_locate(index, {"label": "年等效满负荷小时数（保证值，h）"}, "上下文")
        self.assertEqual(status, "field_name")
        self.assertEqual(selected["value"], "2836")

    def test_same_placeholder_shared_by_conflicting_fields_goes_manual(self) -> None:
        # 取值不一致时不得静默取第一个：上下文分不开就交人工
        facts = [
            _spec_fact("年等效满负荷小时数（保证值，h）", "2836", ["[投标方案，待填写]"], ["待填写-塔筒设计方案.docx"]),
            _spec_fact("年等效满负荷小时数（考核值，h）", "2650", ["[投标方案，待填写]"], ["待填写-塔筒设计方案.docx"]),
        ]
        facts[1]["reviewLabel"] = "年等效满负荷小时数（保证值，h）"  # 别名撞上另一个字段的正式名
        index = self._index(facts)
        selected, alternatives, status = filler.spec_locate(index, {"label": "年等效满负荷小时数（保证值，h）"}, "无区分信息")
        self.assertEqual(status, "ambiguous")
        self.assertIsNone(selected)
        self.assertEqual(len(alternatives), 2)

    def test_composite_placeholder_is_reported_not_filled(self) -> None:
        # 一格填多个字段（`[投标机型、台数，待填写]`）不自动填：分隔方式是产品决策
        facts = [_spec_fact("投标机型", "EW10.0-220", ["[机型，待填写]"], ["待填写-塔筒设计方案.docx"])]
        index = self._index(facts)
        self.assertEqual(filler.composite_field_names(index, "投标机型、台数"), ["投标机型"])
        # 字段名自带括号逗号时不得被切碎
        self.assertEqual(filler.composite_field_names(index, "折减系数（考核值，%）"), [])

    def test_missing_fact_table_disables_the_index(self) -> None:
        index = filler.build_spec_index({}, [], {"任意文件"})
        self.assertFalse(index.enabled)
        _selected, _alts, status = filler.spec_locate(index, {"label": "任意占位符"}, "上下文")
        self.assertEqual(status, "skipped")


class DerivedComputedFactsTests(unittest.TestCase):
    """事实之间的拼接：删素材抽取时保留下来，喂给关键数据表的跨字段语义校验。"""

    def _facts(self, *pairs: tuple[str, str]) -> list[dict]:
        return [
            {"label": label, "value": value, "sourceKind": "manifest", "confidence": 0.97, "sourcePriority": 90}
            for label, value in pairs
        ]

    def _exact(self, facts: list[dict], label: str) -> list[str]:
        # 不用 first_fact_value：它带包含匹配，查「容量」会捞到「单机容量」
        return [str(fact["value"]) for fact in facts if fact["label"] == label]

    def test_capacity_is_never_computed(self) -> None:
        # 「容量 = 台数 × 单机容量」原本就恒不执行：查「容量」必然先命中「单机容量」，
        # 分支前提（容量为空）与它自己的输入（单机容量有值）互斥
        facts = self._facts(("台数", "60"), ("单机容量", "10"))
        filler.derive_computed_facts(facts)
        self.assertEqual(self._exact(facts, "容量"), [])

    def test_scheme_is_composed_from_model_and_hub_height(self) -> None:
        facts = self._facts(("投标机型", "EW10.0-220"), ("轮毂高度", "125"), ("台数", "60"))
        filler.derive_computed_facts(facts)
        self.assertEqual(self._exact(facts, "投标方案"), ["60台EW10.0-220-125"])
        self.assertEqual(self._exact(facts, "方案"), ["60*EW10.0-220-125"])

    def test_no_material_reading_involved(self) -> None:
        # 只做算术与拼接：缺机型/轮毂高度时不产出方案，也不去任何文档里找
        facts = self._facts(("台数", "60"))
        filler.derive_computed_facts(facts)
        self.assertEqual(self._exact(facts, "投标方案"), [])


class SpeclessManifestTests(unittest.TestCase):
    """清单元数据缺失时脚本直接失败，不再退回模糊匹配填出可疑值。"""

    def test_run_from_manifest_refuses_specless_fact_table(self) -> None:
        import json
        import tempfile

        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            blank = work / "待填写-样例.docx"
            document = Document()
            document.add_paragraph("投标机型[投标机型，待填写]")
            document.save(str(blank))
            manifest_path = work / "word_fill_input.json"
            manifest_path.write_text(
                json.dumps(
                    {
                        "blankSource": {"docxPath": str(blank), "title": blank.name},
                        "outputFile": str(work / "out.docx"),
                        # 事实表有值但没有清单第 2/3 列 —— 正是过去静默走旧链路的入口
                        "projectFactTable": {
                            "status": "confirmed",
                            "fields": [{"label": "投标机型", "value": "EW10.0-220", "status": "confirmed"}],
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            with self.assertRaises(RuntimeError) as ctx:
                filler.run_from_manifest(manifest_path)
            self.assertIn("重新上传", str(ctx.exception))
            self.assertFalse((work / "out.docx").exists())


if __name__ == "__main__":
    unittest.main()
