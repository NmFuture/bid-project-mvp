from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from toc_skill_helpers import (
    OUTLINE_SCRIPT_DIR,
    TocSkillScriptTestBase,
    json_load,
    load_outline_script,
    write_decision_context_fixture,
)


class TocSkillScriptTests(TocSkillScriptTestBase):

    def test_bid_outline_template_headings_pages_complete_template_structure(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest, manifest_path = write_decision_context_fixture(
                root,
                template_count=50,
                heading_count=0,
            )

            first = outline_runner.dispatch_command(
                "template-headings", manifest, manifest_path, ["--cursor", "0", "--page-size", "20"]
            )
            second = outline_runner.dispatch_command(
                "template-headings",
                manifest,
                manifest_path,
                ["--cursor", first["next_cursor"], "--page-size", "20"],
            )
            third = outline_runner.dispatch_command(
                "template-headings",
                manifest,
                manifest_path,
                ["--cursor", second["next_cursor"], "--page-size", "20"],
            )

        self.assertEqual(len(first["items"]), 20)
        self.assertEqual(len(second["items"]), 20)
        self.assertEqual(len(third["items"]), 10)
        self.assertEqual(first["item_count"], 50)
        self.assertFalse(first["complete"])
        self.assertTrue(third["complete"])
        self.assertEqual(third["next_cursor"], "")
        self.assertEqual(
            [item["target_id"] for page in (first, second, third) for item in page["items"]],
            [f"TPL-{index:04d}" for index in range(1, 51)],
        )



    def test_bid_outline_prepare_builds_ordered_tender_review_chunks(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"

            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)

            tender_doc = Document()
            tender_doc.add_paragraph("1. 总则", style="Heading 1")
            tender_doc.add_paragraph("投标人应提供总体技术方案。")
            table = tender_doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "项目"
            table.cell(0, 1).text = "要求"
            table.cell(0, 0).merge(table.cell(0, 1))
            table.cell(1, 0).text = "机型"
            table.cell(1, 1).text = "投标人填写"
            tender_doc.add_paragraph("2. 专题方案", style="Heading 1")
            tender_doc.add_paragraph("投标人须编制场址安全适应性报告。")
            tender_doc.save(tender)

            manifest = {
                "projectId": "PRJ-TEST",
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

            result = outline_runner.write_template_structure(manifest, manifest_path)
            chunks = json_load(root / "tender_review_chunks.json")
            state = json_load(root / "tender_review_state.json")

        self.assertEqual(result["tenderReviewChunkCount"], len(chunks["chunks"]))
        self.assertEqual(chunks["schema_version"], "tender-review-chunks.v1")
        self.assertEqual(chunks["source_block_count"], 5)
        self.assertEqual(
            [block["type"] for chunk in chunks["chunks"] for block in chunk["blocks"]],
            ["paragraph", "paragraph", "table", "paragraph", "paragraph"],
        )
        table_block = next(
            block for chunk in chunks["chunks"] for block in chunk["blocks"] if block["type"] == "table"
        )
        self.assertEqual(table_block["rows"][0]["cells"], ["项目 要求", "项目 要求"])
        self.assertEqual(state["schema_version"], "tender-review-state.v1")
        self.assertEqual(state["reviewed_chunk_count"], 0)
        self.assertEqual(state["pending_chunk_count"], len(chunks["chunks"]))



    def test_bid_outline_review_workspace_uses_lightweight_docx_xml_parser(self) -> None:
        source = (OUTLINE_SCRIPT_DIR / "review_workflow.py").read_text(encoding="utf-8")

        self.assertNotIn("from docx", source)
        self.assertNotIn("Document(str(path))", source)



    def test_bid_outline_headings_prefers_toc_and_skips_body_titles_without_marking_reviewed(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")
        review_workflow = load_outline_script("review_workflow")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"

            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)

            tender_doc = Document()
            if "TOC 1" not in [style.name for style in tender_doc.styles]:
                tender_doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
            tender_doc.add_paragraph("第1章 总体要求 ........ 1", style="TOC 1")
            tender_doc.add_paragraph("第1章 总体要求", style="Heading 1")
            tender_doc.add_paragraph("投标人应提供总体技术方案。")
            tender_doc.add_paragraph("2. 专题方案")
            tender_doc.add_paragraph("附表A.1 技术参数表")
            table = tender_doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "参数"
            table.cell(0, 1).text = "要求"
            table.cell(1, 0).text = "额定功率"
            table.cell(1, 1).text = "投标人填写"
            tender_doc.save(tender)

            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)

            result = outline_runner.dispatch_command("headings", manifest, manifest_path, [])
            mapped_section = outline_runner.dispatch_command(
                "section",
                manifest,
                manifest_path,
                [result["files"][0]["items"][0]["section_id"]],
            )
            status = review_workflow.review_status(root)

        self.assertEqual(result["schema_version"], "tender-headings.v1")
        self.assertEqual(
            [(item["kind"], item["text"]) for item in result["files"][0]["items"]],
            [
                ("toc", "第1章 总体要求 ........ 1"),
            ],
        )
        self.assertEqual(result["files"][0]["source"], "toc")
        self.assertEqual(result["files"][0]["items"][0]["section_id"], "TEN-1:S0001")
        self.assertIn(
            "投标人应提供总体技术方案。",
            [item["text"] for item in mapped_section["records"]],
        )
        self.assertTrue(result["complete"])
        self.assertEqual(result["appendix_count"], 1)
        self.assertNotIn("appendices", result)
        self.assertEqual(status["reviewed_chunk_count"], 0)
        self.assertEqual(status["pending_chunk_count"], status["chunk_count"])



    def test_bid_outline_headings_pages_body_titles_when_toc_is_missing(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"

            template_doc = Document()
            template_doc.add_paragraph("Template", style="Heading 1")
            template_doc.save(template)

            tender_doc = Document()
            tender_doc.add_paragraph("Section one", style="Heading 1")
            tender_doc.add_paragraph("Section two", style="Heading 1")
            tender_doc.add_paragraph("Section three", style="Heading 1")
            tender_doc.save(tender)

            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)

            first = outline_runner.dispatch_command(
                "headings", manifest, manifest_path, ["--page-size", "2"]
            )
            second = outline_runner.dispatch_command(
                "headings",
                manifest,
                manifest_path,
                ["--cursor", first["next_cursor"], "--page-size", "2"],
            )

        self.assertEqual(first["files"][0]["source"], "body_headings")
        self.assertEqual(
            [item["text"] for item in first["files"][0]["items"]],
            ["Section one", "Section two"],
        )
        self.assertFalse(first["complete"])
        self.assertEqual(first["next_cursor"], "2")
        self.assertEqual(
            [item["text"] for item in second["files"][0]["items"]],
            ["Section three"],
        )
        self.assertTrue(second["complete"])
        self.assertEqual(second["next_cursor"], "")



    def test_bid_outline_headings_exposes_stable_sections_and_section_reads_continuous_body(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"

            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("1 总体技术要求", style="Heading 1")
            tender_doc.add_paragraph("投标人应提交总体技术方案和项目组织方案。")
            tender_doc.add_paragraph("1.1 专题报告", style="Heading 2")
            tender_doc.add_paragraph("投标人应提交场址安全适应性专题报告。")
            tender_doc.add_paragraph("2 供货范围", style="Heading 1")
            tender_doc.add_paragraph("投标人应提供完整供货清单。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": tender.name, "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)

            headings = outline_runner.dispatch_command("headings", manifest, manifest_path, [])
            section_id = headings["files"][0]["items"][0]["section_id"]
            first_page = outline_runner.dispatch_command(
                "section",
                manifest,
                manifest_path,
                [section_id, "--cursor", "0", "--max-chars", "20"],
            )
            second_page = outline_runner.dispatch_command(
                "section",
                manifest,
                manifest_path,
                [section_id, "--cursor", first_page["next_cursor"], "--max-chars", "200"],
            )
            review_headings = outline_runner.dispatch_command(
                "headings",
                manifest,
                manifest_path,
                ["--review", "--cursor", "0", "--page-size", "1"],
            )

        self.assertEqual(section_id, "TEN-1:S0001")
        self.assertEqual(first_page["section"]["title"], "1 总体技术要求")
        combined = first_page["records"] + second_page["records"]
        self.assertIn("投标人应提交场址安全适应性专题报告。", [item["text"] for item in combined])
        self.assertNotIn("2 供货范围", [item["text"] for item in combined])
        self.assertTrue(second_page["complete"])
        self.assertTrue(review_headings["review"])
        self.assertGreater(review_headings["returned_heading_count"], 0)



    def test_bid_outline_search_locates_full_text_without_marking_evidence_as_read(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")
        review_workflow = load_outline_script("review_workflow")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("1 安全要求", style="Heading 1")
            tender_doc.add_paragraph("投标人必须提交海上运输安全专项方案。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": tender.name, "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            headings = outline_runner.dispatch_command("headings", manifest, manifest_path, [])

            result = outline_runner.dispatch_command(
                "search", manifest, manifest_path, ["运输安全", "--max-results", "10"]
            )
            evidence_id = result["results"][0]["evidence_id"]
            with self.assertRaisesRegex(SystemExit, "尚未通过受控阅读"):
                review_workflow.resolve_tender_basis(root, evidence_id)
            outline_runner.dispatch_command(
                "section",
                manifest,
                manifest_path,
                [headings["files"][0]["items"][0]["section_id"]],
            )
            basis = review_workflow.resolve_tender_basis(root, evidence_id)

        self.assertEqual(result["results"][0]["section_id"], "TEN-1:S0001")
        self.assertEqual(basis["evidence_id"], evidence_id)
        self.assertEqual(basis["file_id"], "TEN-1")
        self.assertEqual(basis["search_text"], "投标人必须提交海上运输安全专项方案。")



    def test_bid_outline_table_evidence_uses_meaningful_cell_text_for_location(self) -> None:
        review_workflow = load_outline_script("review_workflow")

        self.assertEqual(
            review_workflow.evidence_search_text(
                {"cells": ["2", "抗低温", "√"], "text": "2 | 抗低温 | √"}
            ),
            "抗低温",
        )
        self.assertEqual(
            review_workflow.evidence_search_text(
                {
                    "type": "table",
                    "rows": [
                        {"cells": ["序号", "货物名称", "品牌或制造商名称"]},
                        {"cells": ["1", "主控系统", "自主可控品牌"]},
                    ],
                    "text": "序号 | 货物名称 | 品牌或制造商名称 | 1 | 主控系统 | 自主可控品牌",
                }
            ),
            "自主可控品牌",
        )



    def test_bid_outline_tender_search_validation_uses_controlled_table_evidence(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")
        review_workflow = load_outline_script("review_workflow")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("1 环境适应性", style="Heading 1")
            table = tender_doc.add_table(rows=2, cols=3)
            for column, value in enumerate(["序号", "要求", "响应"]):
                table.cell(0, column).text = value
            for column, value in enumerate(["2", "抗低温", "√"]):
                table.cell(1, column).text = value
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": tender.name, "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            table_result = review_workflow.read_table(root, "TEN-1:T0001", start=1, end=2)
            evidence_id = table_result["rows"][1]["evidence_id"]
            basis = review_workflow.resolve_tender_basis(root, evidence_id)
            nodes = [{"tender_basis": basis, "children": []}]

            access_path = root / "tender_evidence_access.json"
            access = json_load(access_path)
            access["evidence_ids"] = []
            access["events"] = []
            access_path.write_text(json.dumps(access, ensure_ascii=False), encoding="utf-8")

            outline_runner.validate_tender_search_texts(nodes, manifest, work_dir=root)
            nodes[0]["tender_basis"] = {**basis, "search_text": "被篡改的定位文本"}
            with self.assertRaisesRegex(SystemExit, "受控证据不一致"):
                outline_runner.validate_tender_search_texts(nodes, manifest, work_dir=root)



    def test_bid_outline_headings_pages_toc_items_with_same_cursor_contract(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"

            template_doc = Document()
            template_doc.add_paragraph("Template", style="Heading 1")
            template_doc.save(template)

            tender_doc = Document()
            if "TOC 1" not in [style.name for style in tender_doc.styles]:
                tender_doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
            for index in range(1, 6):
                tender_doc.add_paragraph(f"Section {index} ........ {index}", style="TOC 1")
            tender_doc.save(tender)

            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)

            first = outline_runner.dispatch_command(
                "headings", manifest, manifest_path, ["--page-size", "2"]
            )
            self.assertEqual(first["returned_heading_count"], 2)
            self.assertEqual(first["next_cursor"], "2")
            self.assertFalse(first["complete"])

            second = outline_runner.dispatch_command(
                "headings",
                manifest,
                manifest_path,
                ["--cursor", first["next_cursor"], "--page-size", "2"],
            )
            self.assertEqual(second["returned_heading_count"], 2)
            self.assertEqual(second["next_cursor"], "4")
            self.assertFalse(second["complete"])

            third = outline_runner.dispatch_command(
                "headings",
                manifest,
                manifest_path,
                ["--cursor", second["next_cursor"], "--page-size", "2"],
            )
            self.assertEqual(third["returned_heading_count"], 1)
            self.assertEqual(third["next_cursor"], "")
            self.assertTrue(third["complete"])



    def test_bid_outline_headings_requires_full_review_when_no_structure_exists(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("本项目位于沿海区域。")
            tender_doc.add_paragraph("投标人应提交完整技术响应文件。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": tender.name, "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
                "requireComposedOutline": True,
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)

            first = outline_runner.dispatch_command("headings", manifest, manifest_path, [])
            with self.assertRaisesRegex(SystemExit, "全文审阅"):
                outline_runner.dispatch_command("decision-next", manifest, manifest_path, [])
            batch = outline_runner.dispatch_command("next-batch", manifest, manifest_path, [])
            outline_runner.dispatch_command(
                "review-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "chunk_ids": batch["chunk_ids"],
                            "review_summary": "已逐段审阅无结构招标文件。",
                            "requirements": [],
                        },
                        ensure_ascii=False,
                    )
                ],
            )
            second = outline_runner.dispatch_command("headings", manifest, manifest_path, [])

        self.assertEqual(first["heading_count"], 0)
        self.assertTrue(first["requires_full_review"])
        self.assertFalse(first["complete"])
        self.assertGreater(first["full_review_pending_chunk_count"], 0)
        self.assertTrue(second["complete"])
        self.assertEqual(second["full_review_pending_chunk_count"], 0)



    def test_bid_outline_table_navigation_reports_continuation_and_truncation(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")
        review_workflow = load_outline_script("review_workflow")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("附表A.1 参数表")
            table = tender_doc.add_table(rows=30, cols=1)
            for row_index in range(30):
                table.cell(row_index, 0).text = "超长参数" * 80 if row_index == 0 else f"参数{row_index + 1}"
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)

            first_page = review_workflow.read_table(root, "TEN-1:T0001", start=1, end=24, max_chars=2_400)
            truncated_row = review_workflow.read_table(root, "TEN-1:T0001", start=1, end=1, max_chars=80)
            state = json_load(root / "tender_review_state.json")

        self.assertEqual(first_page["table"]["row_count"], 30)
        self.assertEqual(first_page["returned_range"], {"start": 1, "end": 24})
        self.assertTrue(first_page["has_more"])
        self.assertEqual(first_page["next_range"], "25-30")
        self.assertEqual(truncated_row["truncated_rows"], [1])
        table_chunk = next(item for item in state["chunks"] if item["chunk_id"].endswith("C0002"))
        self.assertEqual(table_chunk["table_read_ranges"], [{"start": 1, "end": 24}])



    def test_bid_outline_batch_review_keeps_business_decisions_agentic(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")
        review_workflow = load_outline_script("review_workflow")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("1. 总则", style="Heading 1")
            tender_doc.add_paragraph("投标人应提供总体技术方案。")
            table = tender_doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "项目"
            table.cell(0, 1).text = "要求"
            table.cell(1, 0).text = "报告"
            table.cell(1, 1).text = "投标人填写"
            tender_doc.add_paragraph("2. 专题", style="Heading 1")
            tender_doc.add_paragraph("投标人应提供独立专题报告。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)

            batch = review_workflow.next_review_batch(root, max_chunks=2, max_chars=30_000)
            chunk_ids = batch["chunk_ids"]
            repeated_batch = review_workflow.next_review_batch(root, max_chunks=8, max_chars=60_000)
            all_chunk_ids = [chunk["chunk_id"] for chunk in json_load(root / "tender_review_chunks.json")["chunks"]]
            next_chunk_id = next(chunk_id for chunk_id in all_chunk_ids if chunk_id not in chunk_ids)
            table_id = next(
                block["table_id"]
                for chunk in batch["chunks"]
                for block in chunk["blocks"]
                if block["type"] == "table"
            )
            first_paragraph_id = next(
                block["evidence_id"]
                for chunk in batch["chunks"]
                for block in chunk["blocks"]
                if block["type"] == "paragraph" and "投标人应提供" in block["text"]
            )

            with self.assertRaisesRegex(SystemExit, "当前受控批次完全一致"):
                review_workflow.submit_batch_review(
                    root,
                    [*chunk_ids, next_chunk_id],
                    {"review_summary": "试图扩展当前批次。", "requirements": []},
                )
            with self.assertRaisesRegex(SystemExit, "table chunk must be fully read"):
                review_workflow.submit_batch_review(
                    root,
                    chunk_ids,
                    {"review_summary": "表格尚未读完。", "requirements": []},
                )

            tables = review_workflow.read_tables(
                root,
                [table_id],
                start=1,
                end=24,
                max_chars=8_000,
            )
            progress_batch = review_workflow.next_review_batch(root, max_chunks=8, max_chars=60_000)
            table_progress = next(
                block
                for chunk in progress_batch["chunks"]
                for block in chunk["blocks"]
                if block["type"] == "table"
            )
            with self.assertRaisesRegex(SystemExit, "单一目录节点"):
                review_workflow.submit_batch_review(
                    root,
                    chunk_ids,
                    {
                        "review_summary": "错误地把多个承接节点写在一个 target_node。",
                        "requirements": [
                            {
                                "evidence_ids": [first_paragraph_id],
                                "obligation": "投标人应提供总体技术方案",
                                "disposition": "map_existing",
                                "target_node": "第1章/第2章",
                            }
                        ],
                    },
                )
            first_submitted = review_workflow.submit_batch_review(
                root,
                chunk_ids,
                {
                    "review_summary": "逐项审阅本批段落和表格。",
                    "requirements": [
                        {
                            "evidence_ids": [first_paragraph_id],
                            "obligation": "投标人应提供总体技术方案",
                            "disposition": "map_existing",
                            "target_node": "第1章",
                        }
                    ],
                },
            )
            second_batch = review_workflow.next_review_batch(root, max_chunks=8, max_chars=30_000)
            second_paragraph_id = next(
                block["evidence_id"]
                for chunk in second_batch["chunks"]
                for block in chunk["blocks"]
                if block["type"] == "paragraph" and "独立专题报告" in block["text"]
            )
            second_submitted = review_workflow.submit_batch_review(
                root,
                second_batch["chunk_ids"],
                {
                    "review_summary": "逐项审阅第二批段落。",
                    "requirements": [
                        {
                            "evidence_ids": [second_paragraph_id],
                            "obligation": "投标人应提供独立专题报告",
                            "disposition": "suggest_add",
                            "target_node": "第1章",
                            "proposed_title": "独立专题报告",
                            "reason": "招标明确要求独立报告，模板无语义等价节点。",
                        },
                    ],
                },
            )
            status = review_workflow.review_status(root)

        self.assertEqual(len(batch["chunks"]), 2)
        self.assertEqual(repeated_batch["chunk_ids"], chunk_ids)
        self.assertEqual(tables["table_count"], 1)
        self.assertEqual(tables["tables"][0]["table"]["table_id"], table_id)
        self.assertTrue(table_progress["fully_read"])
        self.assertEqual(table_progress["read_ranges"], [{"start": 1, "end": 2}])
        self.assertEqual(table_progress["truncated_rows"], [])
        self.assertEqual(first_submitted["reviewed_batch_chunk_count"], len(chunk_ids))
        self.assertEqual(first_submitted["added_requirement_count"], 1)
        self.assertEqual(second_submitted["added_requirement_count"], 1)
        self.assertEqual(status["pending_chunk_count"], 0)
        self.assertEqual(status["unfinished_table_count"], 0)



    def test_bid_outline_cli_accepts_navigation_options_after_manifest(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("投标人应提供总体技术方案。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            chunk = outline_runner.dispatch_command("next", manifest, manifest_path, [])["chunk"]
            evidence_id = chunk["blocks"][0]["evidence_id"]

            with patch.object(
                sys,
                "argv",
                [
                    "run_from_manifest.py",
                    "window",
                    str(manifest_path),
                    evidence_id,
                    "--before",
                    "0",
                    "--after",
                    "0",
                ],
            ), patch("builtins.print") as print_mock:
                result = outline_runner.main()

        self.assertEqual(result, 0)
        payload = json.loads(print_mock.call_args.args[0])
        self.assertEqual(payload["center"], evidence_id)
        self.assertEqual(len(payload["blocks"]), 1)
