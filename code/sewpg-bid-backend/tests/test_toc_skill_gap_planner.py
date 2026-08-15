from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch
from docx import Document

from toc_skill_helpers import (
    TocSkillScriptTestBase,
    json_load,
    load_gap_planner_script,
)


class TocSkillScriptTests(TocSkillScriptTestBase):

    def test_bid_gap_planner_routes_fill_template_material_to_ai_fill(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            toc_path.write_text(
                json.dumps(
                    {
                        "items": [
                            {"number": "第1章", "title": "标前概述", "level": 1},
                            {"number": "1.1", "title": "技术评分标准索引表", "level": 2},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            parse_path.write_text(json.dumps({"structured": {}}, ensure_ascii=False), encoding="utf-8")
            manifest = {
                "projectId": "PRJ-TEST",
                "projectName": "测试项目",
                "bidType": "技术标",
                "tocJsonPath": str(toc_path),
                "parseResultPath": str(parse_path),
                "materialScope": {"paths": ["技术标/客户素材/华能集团"]},
                "projectTurbineModel": {"model": "EW10.0-220上置"},
                "materialIndex": [
                    {
                        "id": "RAW-0453",
                        "name": "待填写-技术评分标准索引表.docx",
                        "folderPath": "技术标/客户素材/华能集团/技术标-标前概述",
                        "materialTier": "customer",
                        "hasCleanedWord": True,
                        "cleanedFileName": "待填写-技术评分标准索引表.docx",
                        "requiresFill": True,
                        "placeholderCount": 3,
                        "placeholderLabels": ["投标机型", "投标方案", "章节索引"],
                    }
                ],
            }

            plan = gap_runner.build_gap_plan(manifest)

        item = next(entry for entry in plan["items"] if entry["number"] == "1.1")
        self.assertEqual(item["decision"], "fill_required")
        self.assertEqual(item["status"], "needs_input")
        self.assertEqual(item["usage"], "section_fill")
        self.assertEqual(item["matchedMaterials"], [])
        self.assertEqual(item["fillTasks"][0]["skill"], "bid-tech-word-placeholder-filler")
        self.assertEqual(item["fillTasks"][0]["blankSource"]["id"], "RAW-0453")
        self.assertEqual(item["fillTasks"][0]["blankSource"]["placeholderCount"], 3)
        self.assertEqual(item["candidateMaterials"][0]["id"], "RAW-0453")



    def test_bid_gap_planner_requires_one_result_per_confirmed_toc_item(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            toc_items = [
                {"number": "第1章", "title": "标前概述", "level": 1},
                {"number": "1.1", "title": "技术评分标准索引表", "level": 2},
                {"number": "1.2", "title": "与华能集团签署的战略合作协议", "level": 2},
                {"number": "1.3", "title": "缺失专题", "level": 2},
            ]
            toc_path.write_text(json.dumps({"items": toc_items}, ensure_ascii=False), encoding="utf-8")
            parse_path.write_text(json.dumps({"structured": {}}, ensure_ascii=False), encoding="utf-8")
            manifest = {
                "projectId": "PRJ-TEST",
                "projectName": "测试项目",
                "bidType": "技术标",
                "tocJsonPath": str(toc_path),
                "parseResultPath": str(parse_path),
                "materialScope": {"paths": ["技术标/客户素材/华能集团"]},
                "projectTurbineModel": {"model": "EW10.0-220上置"},
                "materialIndex": [
                    {
                        "id": "RAW-0453",
                        "name": "待填写-技术评分标准索引表.docx",
                        "folderPath": "技术标/客户素材/华能集团/技术标-标前概述",
                        "materialTier": "customer",
                        "requiresFill": True,
                        "placeholderCount": 2,
                    },
                    {
                        "id": "RAW-0454",
                        "name": "固定-与华能集团签署的战略合作协议.docx",
                        "folderPath": "技术标/客户素材/华能集团/技术标-标前概述",
                        "materialTier": "customer",
                    },
                ],
            }

            plan = gap_runner.build_gap_plan(manifest)

        self.assertEqual(plan["summary"]["totalTocItems"], len(toc_items))
        self.assertEqual(len(plan["items"]), len(toc_items))
        self.assertEqual([item["number"] for item in plan["items"]], [item["number"] for item in toc_items])
        self.assertEqual(plan["integrity"]["coverageStatus"], "passed")
        self.assertEqual(plan["items"][1]["decision"], "fill_required")
        self.assertEqual(plan["items"][2]["decision"], "ready")
        self.assertEqual(plan["items"][3]["decision"], "material_required")



    def test_bid_gap_planner_carries_parse_fields_to_appendix_tasks(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            appendix_doc = root / "APPX-A1.docx"
            Document().save(appendix_doc)
            toc_path.write_text(
                json.dumps({"items": [{"number": "附表A.1", "title": "附表A.1 投标机型总方案信息表", "level": 2}]}, ensure_ascii=False),
                encoding="utf-8",
            )
            parse_path.write_text(
                json.dumps(
                    {
                        "items": [
                            *[
                                {
                                    "id": f"REQ-NOISE-{idx}",
                                    "type": "项目基础信息",
                                    "title": "招标人",
                                    "value": f"噪声字段{idx}",
                                    "sourceFile": "招标文件.docx",
                                }
                                for idx in range(180)
                            ],
                            {
                                "id": "REQ-SCALE",
                                "type": "项目基础信息",
                                "title": "标段规模",
                                "value": "600MW",
                                "sourceFile": "招标文件.docx",
                            }
                        ],
                        "structured": {
                            "appendices": [
                                {
                                    "id": "APPX-A1",
                                    "title": "附表A.1 投标机型总方案信息表",
                                    "docxPath": str(appendix_doc),
                                }
                            ]
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            plan = gap_runner.build_gap_plan(
                {
                    "projectId": "PRJ-TEST",
                    "projectName": "测试项目",
                    "tocJsonPath": str(toc_path),
                    "parseResultPath": str(parse_path),
                    "materialIndex": [],
                }
            )

        task = plan["items"][0]["appendixTasks"][0]
        self.assertEqual(task["availableParseFields"][0]["id"], "REQ-SCALE")
        self.assertEqual(task["availableParseFields"][0]["label"], "标段规模")
        self.assertGreater(len(task["availableParseFields"]), 160)



    def test_bid_gap_planner_summary_stdout_exposes_coverage_counts(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            manifest_path = root / "s4_gap_input.json"
            output_path = root / "gap_plan.json"
            toc_path.write_text(
                json.dumps(
                    {
                        "items": [
                            {"number": "第1章", "title": "标前概述", "level": 1},
                            {"number": "1.1", "title": "技术评分标准索引表", "level": 2},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            parse_path.write_text(json.dumps({"structured": {}}, ensure_ascii=False), encoding="utf-8")
            manifest_path.write_text(
                json.dumps(
                    {
                        "projectId": "PRJ-TEST",
                        "projectName": "测试项目",
                        "bidType": "技术标",
                        "tocJsonPath": str(toc_path),
                        "parseResultPath": str(parse_path),
                        "materialScope": {"paths": ["技术标/客户素材/华能集团"]},
                        "projectTurbineModel": {"model": "EW10.0-220上置"},
                        "materialIndex": [],
                        "outputFile": str(output_path),
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            with patch.object(sys, "argv", ["run_from_manifest.py", "--manifest", str(manifest_path)]), \
                patch("builtins.print") as mocked_print:
                gap_runner.main()

            response = json.loads(mocked_print.call_args.args[0])

        self.assertEqual(response["tocItemCount"], 2)
        self.assertEqual(response["itemCount"], 2)
        self.assertEqual(response["coverageStatus"], "passed")



    def test_bid_gap_planner_uses_project_commitment_chapter_word_for_chapter_four_children(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            toc_path.write_text(
                json.dumps(
                    {
                        "items": [
                            {"number": "第4章", "title": "项目技术承诺函", "level": 1},
                            {"number": "4.1", "title": "发电小时数承诺函", "level": 2},
                            {"number": "4.2", "title": "投标机组可利用率承诺", "level": 2},
                            {"number": "4.3", "title": "投标机组功率曲线保证率承诺", "level": 2},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            parse_path.write_text(json.dumps({"structured": {}}, ensure_ascii=False), encoding="utf-8")
            manifest = {
                "projectId": "PRJ-TEST",
                "projectName": "测试项目",
                "bidType": "技术标",
                "tocJsonPath": str(toc_path),
                "parseResultPath": str(parse_path),
                "materialScope": {"paths": ["技术标/项目素材/MAT-HN-CHIFENG-001"]},
                "projectTurbineModel": {"model": "EW10.0-220上置"},
                "materialIndex": [
                    {
                        "id": "RAW-0472",
                        "name": "定制-项目技术承诺函.docx",
                        "folderPath": "技术标/项目素材/MAT-HN-CHIFENG-001/技术标-项目技术承诺函",
                        "materialTier": "project",
                        "hasCleanedWord": True,
                        "cleanedFileName": "定制-项目技术承诺函.docx",
                        "requiresFill": True,
                        "placeholderCount": 1,
                        "placeholderLabels": ["项目承诺函"],
                    }
                ],
            }

            plan = gap_runner.build_gap_plan(manifest)

        parent = next(item for item in plan["items"] if item["number"] == "第4章")
        self.assertEqual(parent["decision"], "fill_required")
        self.assertEqual(parent["status"], "needs_input")
        self.assertEqual(parent["usage"], "chapter_fill")
        self.assertEqual(parent["coverageRole"], "chapter_master")
        self.assertEqual(parent["fillTasks"][0]["blankSource"]["id"], "RAW-0472")
        for number in ["4.1", "4.2", "4.3"]:
            child = next(item for item in plan["items"] if item["number"] == number)
            self.assertEqual(child["decision"], "fill_required")
            self.assertEqual(child["status"], "needs_input")
            self.assertEqual(child["usage"], "covered_by_parent")
            self.assertEqual(child["coverageRole"], "covered_by_parent")
            self.assertEqual(child["coveredByParent"], parent["id"])
            self.assertEqual(child["matchedMaterials"], [])
            self.assertEqual(child["fillTasks"], [])



    def test_bid_gap_planner_uses_standard_delivery_acceptance_word_for_chapter_six_children(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            toc_path.write_text(
                json.dumps(
                    {
                        "items": [
                            {"number": "第6章", "title": "产品交付、考核及验收", "level": 1},
                            {"number": "6.1", "title": "技术资料和交付进度", "level": 2},
                            {"number": "6.2", "title": "试验、检验和监造", "level": 2},
                            {"number": "6.3", "title": "设备安装、调试与试运行", "level": 2},
                            {"number": "6.4", "title": "考核指标", "level": 2},
                            {"number": "6.5", "title": "项目验收", "level": 2},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            parse_path.write_text(json.dumps({"structured": {}}, ensure_ascii=False), encoding="utf-8")
            manifest = {
                "projectId": "PRJ-TEST",
                "projectName": "测试项目",
                "bidType": "技术标",
                "tocJsonPath": str(toc_path),
                "parseResultPath": str(parse_path),
                "materialScope": {"paths": ["技术标/通用素材"]},
                "projectTurbineModel": {"model": "EW10.0-220上置"},
                "materialIndex": [
                    {
                        "id": "RAW-0429",
                        "name": "待填写-产品交付、考核及验收.docx",
                        "folderPath": "技术标/通用素材/技术标-产品交付、考核及验收",
                        "materialTier": "standard",
                        "hasCleanedWord": True,
                        "cleanedFileName": "待填写-产品交付、考核及验收.docx",
                        "requiresFill": True,
                        "placeholderCount": 1,
                        "placeholderLabels": ["招标要求"],
                    }
                ],
            }

            plan = gap_runner.build_gap_plan(manifest)

        parent = next(item for item in plan["items"] if item["number"] == "第6章")
        self.assertEqual(parent["decision"], "fill_required")
        self.assertEqual(parent["status"], "needs_input")
        self.assertEqual(parent["usage"], "chapter_fill")
        self.assertEqual(parent["coverageRole"], "chapter_master")
        self.assertEqual(parent["fillTasks"][0]["blankSource"]["id"], "RAW-0429")
        for number in ["6.1", "6.2", "6.3", "6.4", "6.5"]:
            child = next(item for item in plan["items"] if item["number"] == number)
            self.assertEqual(child["decision"], "fill_required")
            self.assertEqual(child["status"], "needs_input")
            self.assertEqual(child["usage"], "covered_by_parent")
            self.assertEqual(child["coverageRole"], "covered_by_parent")
            self.assertEqual(child["coveredByParent"], parent["id"])
            self.assertEqual(child["matchedMaterials"], [])
            self.assertEqual(child["fillTasks"], [])



    def test_bid_gap_planner_generically_uses_parent_chapter_word_for_any_chapter_children(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            toc_path.write_text(
                json.dumps(
                    {
                        "items": [
                            {"number": "第8章", "title": "运输安装与验收方案", "level": 1},
                            {"number": "8.1", "title": "运输组织方案", "level": 2},
                            {"number": "8.2", "title": "安装调试计划", "level": 2},
                            {"number": "8.3", "title": "验收交付安排", "level": 2},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            parse_path.write_text(json.dumps({"structured": {}}, ensure_ascii=False), encoding="utf-8")
            manifest = {
                "projectId": "PRJ-TEST",
                "projectName": "测试项目",
                "bidType": "技术标",
                "tocJsonPath": str(toc_path),
                "parseResultPath": str(parse_path),
                "materialScope": {"paths": ["技术标/项目素材/MAT-TEST"]},
                "projectTurbineModel": {"model": "EW10.0-220上置"},
                "materialIndex": [
                    {
                        "id": "RAW-0800",
                        "name": "定制-运输安装与验收方案.docx",
                        "folderPath": "技术标/项目素材/MAT-TEST/技术标-运输安装与验收方案",
                        "materialTier": "project",
                        "hasCleanedWord": True,
                        "cleanedFileName": "定制-运输安装与验收方案.docx",
                    },
                    {
                        "id": "RAW-0801",
                        "name": "固定-运输组织方案.docx",
                        "folderPath": "技术标/项目素材/MAT-TEST/技术标-运输安装与验收方案/子节",
                        "materialTier": "project",
                    },
                ],
            }

            plan = gap_runner.build_gap_plan(manifest)

        parent = next(item for item in plan["items"] if item["number"] == "第8章")
        self.assertEqual(parent["decision"], "ready")
        self.assertEqual(parent["status"], "matched")
        self.assertEqual(parent["usage"], "chapter_master")
        self.assertEqual(parent["coverageRole"], "chapter_master")
        self.assertEqual(parent["matchedMaterials"][0]["id"], "RAW-0800")
        # 金标反评 R3b：整章覆盖是默认值不是锁——8.1 自身存在剥修饰同名素材
        # （固定-运输组织方案.docx），保留子节自主匹配，不被父章吞并。
        own_child = next(item for item in plan["items"] if item["number"] == "8.1")
        self.assertEqual(own_child["decision"], "ready")
        self.assertEqual(own_child["status"], "matched")
        self.assertEqual(own_child["matchedMaterials"][0]["id"], "RAW-0801")
        self.assertEqual(own_child["fillTasks"], [])
        for number in ["8.2", "8.3"]:
            child = next(item for item in plan["items"] if item["number"] == number)
            self.assertEqual(child["decision"], "ready")
            self.assertEqual(child["status"], "matched")
            self.assertEqual(child["usage"], "covered_by_parent")
            self.assertEqual(child["coverageRole"], "covered_by_parent")
            self.assertEqual(child["coveredByParent"], parent["id"])
            self.assertEqual(child["matchedMaterials"], [])
            self.assertEqual(child["fillTasks"], [])



    def test_bid_gap_planner_uses_generic_chapter_title_match_for_wind_resource_chapter(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            toc_path.write_text(
                json.dumps(
                    {
                        "items": [
                            {"number": "第3章", "title": "风资源评估与机位排布方案", "level": 1},
                            {"number": "3.1", "title": "项目概况", "level": 2},
                            {"number": "3.2", "title": "风资源分析", "level": 2},
                            {"number": "3.3", "title": "机组选型", "level": 2},
                            {"number": "3.4", "title": "方案及发电量结果", "level": 2},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            parse_path.write_text(json.dumps({"structured": {}}, ensure_ascii=False), encoding="utf-8")
            manifest = {
                "projectId": "PRJ-TEST",
                "projectName": "测试项目",
                "bidType": "技术标",
                "tocJsonPath": str(toc_path),
                "parseResultPath": str(parse_path),
                "materialScope": {"paths": ["技术标/项目素材/MAT-HN-CHIFENG-001"]},
                "projectTurbineModel": {"model": "EW10.0-220上置"},
                "materialIndex": [
                    {
                        "id": "RAW-0471",
                        "name": "定制-项目风资源评估与机组选型排布及发电量计算.docx",
                        "folderPath": "技术标/项目素材/MAT-HN-CHIFENG-001/技术标-专题方案要求",
                        "materialTier": "project",
                    },
                    {
                        "id": "RAW-0473",
                        "name": "定制-风资源评估与机位排布方案.docx",
                        "folderPath": "技术标/项目素材/MAT-HN-CHIFENG-001/技术标-风资源评估与机位排布方案",
                        "materialTier": "project",
                        "hasCleanedWord": True,
                        "cleanedFileName": "定制-风资源评估与机位排布方案.docx",
                    },
                    {
                        "id": "RAW-0478",
                        "name": "风资源评估报告.docx",
                        "folderPath": "技术标/项目素材/MAT-HN-CHIFENG-001/风资源评估报告",
                        "materialTier": "project",
                    },
                ],
            }

            plan = gap_runner.build_gap_plan(manifest)

        parent = next(item for item in plan["items"] if item["number"] == "第3章")
        self.assertEqual(parent["decision"], "ready")
        self.assertEqual(parent["status"], "matched")
        self.assertEqual(parent["coverageRole"], "chapter_master")
        self.assertEqual(parent["matchedMaterials"][0]["id"], "RAW-0473")
        self.assertNotIn("RAW-0478", [item["id"] for item in parent["matchedMaterials"]])
        for number in ["3.1", "3.2", "3.3", "3.4"]:
            child = next(item for item in plan["items"] if item["number"] == number)
            self.assertEqual(child["coverageRole"], "covered_by_parent")
            self.assertEqual(child["coveredByParent"], parent["id"])
            self.assertEqual(child["matchedMaterials"], [])



    def test_bid_gap_planner_filters_toc_refs_by_allowed_material_index(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            toc_path.write_text(
                json.dumps(
                    {
                        "items": [
                            {
                                "number": "1.2",
                                "title": "与华能集团签署的战略合作协议",
                                "level": 2,
                                "material_refs": [
                                    {
                                        "id": "RAW-OUTSIDE",
                                        "docx": "技术标/客户素材/其他客户/战略合作协议.docx",
                                    }
                                ],
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            parse_path.write_text(json.dumps({"structured": {}}, ensure_ascii=False), encoding="utf-8")
            manifest = {
                "projectId": "PRJ-TEST",
                "projectName": "测试项目",
                "bidType": "技术标",
                "tocJsonPath": str(toc_path),
                "parseResultPath": str(parse_path),
                "materialScope": {"paths": ["技术标/客户素材/华能集团"]},
                "projectTurbineModel": {"model": "EW10.0-220上置"},
                "materialIndex": [
                    {
                        "id": "RAW-ALLOWED",
                        "name": "固定-其他资料.docx",
                        "folderPath": "技术标/客户素材/华能集团/其他",
                        "materialTier": "customer",
                    }
                ],
            }

            plan = gap_runner.build_gap_plan(manifest)

        item = plan["items"][0]
        self.assertEqual(item["decision"], "material_required")
        self.assertEqual(item["status"], "missing")
        self.assertEqual(item["matchedMaterials"], [])



    def test_bid_gap_planner_routes_appendix_sources_by_customer_matrix(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            output_path = root / "gap_plan.json"
            toc_path.write_text(
                json.dumps(
                    {
                        "items": [
                            {"number": "附表C.1", "title": "附表C.1 总体技术参数与规格", "level": 2},
                            {"number": "附表D.3", "title": "附表D.3 功率曲线", "level": 2},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            blank = root / "附表C.1 总体技术参数与规格.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=4)
            table.cell(0, 0).text = "序号"
            table.cell(0, 1).text = "参数名称"
            table.cell(0, 2).text = "投标人响应值"
            table.cell(0, 3).text = "单位"
            doc.save(blank)
            range_blank = root / "附表D.3 功率曲线.docx"
            doc.save(range_blank)
            parse_path.write_text(
                json.dumps(
                    {
                        "structured": {
                            "appendices": [
                                {
                                    "id": "APPX-C1",
                                    "title": "附表C.1 总体技术参数与规格",
                                    "docxPath": str(blank),
                                },
                                {
                                    "id": "APPX-D3",
                                    "title": "附表D.3 功率曲线",
                                    "docxPath": str(range_blank),
                                }
                            ]
                        }
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            manifest_path = root / "manifest.json"
            manifest_path.write_text(
                json.dumps(
                    {
                        "projectId": "PRJ-HN",
                        "projectName": "华能项目",
                        "bidType": "技术标",
                        "customerName": "华能",
                        "tocJsonPath": str(toc_path),
                        "parseResultPath": str(parse_path),
                        "materialScope": {"paths": ["技术标/标准文件", "技术标/项目定制"]},
                        "appendixSourceMatrix": {
                            "rows": [
                                {
                                    "id": "Sheet1!R35",
                                    "customer": "华能",
                                    "tableTitle": "附表C.1 总体技术参数与规格",
                                    "projectSources": [],
                                    "standardSources": ["机型参数表"],
                                    "otherSources": [],
                                },
                                {
                                    "id": "Sheet1!R40",
                                    "customer": "华能",
                                    "tableTitle": "附表D.1-D.6",
                                    "projectSources": ["功率曲线"],
                                    "standardSources": [],
                                    "otherSources": [],
                                }
                            ]
                        },
                        "materialIndex": [
                            {
                                "id": "RAW-WIND",
                                "name": "风资源评估报告.docx",
                                "folderPath": "技术标/项目定制/风资源评估报告",
                                "materialTier": "project",
                            },
                            {
                                "id": "RAW-PARAM",
                                "name": "X2平台机型投标参数_20250106.xlsx",
                                "folderPath": "技术标/标准文件/机型参数表",
                                "materialTier": "standard",
                            },
                            {
                                "id": "RAW-POWER",
                                "name": "项目功率曲线.xlsx",
                                "folderPath": "技术标/项目定制/功率曲线",
                                "materialTier": "project",
                            },
                        ],
                        "outputFile": str(output_path),
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = gap_runner.build_gap_plan(json_load(manifest_path))

            c_item = result["items"][0]
            c_task = c_item["appendixTasks"][0]
            self.assertEqual(c_item["usage"], "appendix_fill")
            self.assertEqual(c_task["sourceRouting"]["source"], "appendix_source_matrix")
            self.assertEqual(c_task["sourceRouting"]["standardSources"], ["机型参数表"])
            self.assertEqual(c_task["recommendedMaterials"][0]["id"], "RAW-PARAM")
            self.assertIn("standard 来源规定命中", c_task["recommendedMaterials"][0]["matchReason"])

            d_item = result["items"][1]
            d_task = d_item["appendixTasks"][0]
            self.assertEqual(d_task["sourceRouting"]["ruleId"], "Sheet1!R40")
            self.assertEqual(d_task["sourceRouting"]["projectSources"], ["功率曲线"])
            self.assertEqual(d_task["recommendedMaterials"][0]["id"], "RAW-POWER")
            self.assertIn("project 来源规定命中", d_task["recommendedMaterials"][0]["matchReason"])



    def test_bid_gap_planner_routes_tender_rule_to_all_project_tender_documents(self) -> None:
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            blank = root / "附表B.6 技术服务响应表.docx"
            Document().save(blank)
            toc_path.write_text(
                json.dumps({"items": [{"number": "附表B.6", "title": "附表B.6 技术服务响应表", "level": 2}]}, ensure_ascii=False),
                encoding="utf-8",
            )
            parse_path.write_text(
                json.dumps(
                    {
                        "documents": [
                            {"id": "TEN-1", "name": "技术规范.pdf", "status": "completed"},
                            {"id": "TEN-2", "name": "招标附图.docx", "status": "completed"},
                            {"id": "TEN-3", "name": "损坏文件.pdf", "status": "failed"},
                        ],
                        "structured": {
                            "appendices": [
                                {"id": "APPX-B6", "title": "附表B.6 技术服务响应表", "docxPath": str(blank)}
                            ]
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            plan = gap_runner.build_gap_plan(
                {
                    "projectId": "PRJ-TENDER",
                    "projectName": "招标文件填写项目",
                    "tocJsonPath": str(toc_path),
                    "parseResultPath": str(parse_path),
                    "materialIndex": [],
                    "appendixSourceMatrix": {
                        "rows": [
                            {
                                "id": "Sheet1!R8",
                                "tableTitle": "附表B.6 技术服务响应表",
                                "projectSources": [],
                                "standardSources": [],
                                "otherSources": ["响应招标文件填写"],
                            }
                        ]
                    },
                }
            )

        routing = plan["items"][0]["appendixTasks"][0]["sourceRouting"]
        self.assertEqual(routing["status"], "tender_parse_fields")
        self.assertTrue(routing["useTenderParseFields"])
        self.assertEqual(routing["tenderDocumentStatus"], "available")
        self.assertEqual(routing["tenderDocumentCount"], 2)
        self.assertEqual([item["name"] for item in routing["tenderDocuments"]], ["技术规范.pdf", "招标附图.docx"])



    def test_bid_gap_planner_parent_rule_covers_sub_numbered_appendix(self) -> None:
        """规则只写父级编号（附表F.2）时应覆盖子编号附表（F.2.1/F.2.2）。"""
        gap_runner = load_gap_planner_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_path = root / "toc.json"
            parse_path = root / "parse.json"
            output_path = root / "gap_plan.json"
            toc_path.write_text(
                json.dumps(
                    {
                        "items": [
                            {"number": "附表F.2.1", "title": "附表F.2.1 投标机组设计认证", "level": 2},
                            {"number": "附表G.2.1", "title": "附表G.2.1 场址载荷仿真关键计算方法", "level": 2},
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            blank = root / "附表F.2.1 投标机组设计认证.docx"
            Document().save(blank)
            other_blank = root / "附表G.2.1 场址载荷仿真关键计算方法.docx"
            Document().save(other_blank)
            parse_path.write_text(
                json.dumps(
                    {
                        "structured": {
                            "appendices": [
                                {"id": "APPX-F21", "title": "附表F.2.1 投标机组设计认证", "docxPath": str(blank)},
                                {
                                    "id": "APPX-G21",
                                    "title": "附表G.2.1 场址载荷仿真关键计算方法",
                                    "docxPath": str(other_blank),
                                },
                            ]
                        }
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            manifest_path = root / "manifest.json"
            manifest_path.write_text(
                json.dumps(
                    {
                        "projectId": "PRJ-HN",
                        "projectName": "华能项目",
                        "bidType": "技术标",
                        "customerName": "华能",
                        "tocJsonPath": str(toc_path),
                        "parseResultPath": str(parse_path),
                        "materialScope": {"paths": ["技术标/标准文件", "技术标/项目定制"]},
                        "appendixSourceMatrix": {
                            "rows": [
                                {
                                    "id": "Sheet1!R33",
                                    "customer": "华能",
                                    "tableTitle": "附表F.2 投标机型整机认证",
                                    "projectSources": [],
                                    "standardSources": ["认证证书"],
                                    "otherSources": [],
                                }
                            ]
                        },
                        "materialIndex": [
                            {
                                "id": "RAW-CERT",
                                "name": "EW10.0-220上置型式认证证书.pdf",
                                "folderPath": "技术标/标准文件/EW10.0-220上置/认证证书",
                                "materialTier": "standard",
                            }
                        ],
                        "outputFile": str(output_path),
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = gap_runner.build_gap_plan(json_load(manifest_path))

            f_item = result["items"][0]
            f_task = f_item["appendixTasks"][0]
            self.assertEqual(f_task["sourceRouting"]["source"], "appendix_source_matrix")
            self.assertEqual(f_task["sourceRouting"]["ruleId"], "Sheet1!R33")
            self.assertEqual(f_task["sourceRouting"]["standardSources"], ["认证证书"])
            self.assertEqual(f_task["recommendedMaterials"][0]["id"], "RAW-CERT")
            # 前缀不同的子编号（G.2.1）不应被 F.2 规则覆盖
            g_task = result["items"][1]["appendixTasks"][0]
            self.assertNotEqual(
                (g_task.get("sourceRouting") or {}).get("source"), "appendix_source_matrix"
            )
