from __future__ import annotations

import json
import tempfile
from pathlib import Path
from docx import Document

from toc_skill_helpers import (
    TocSkillScriptTestBase,
    complete_outline_review,
    json_load,
    load_outline_script,
)


class TocSkillScriptTests(TocSkillScriptTestBase):

    def test_bid_outline_finalize_allows_three_levels_and_rejects_fourth_level(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "toc.json"
            manifest_path = root / "s2_input.json"
            manifest = {"workDir": str(root), "outputFile": str(output)}
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            payload = {
                "schema_version": "technical-outline.v1",
                "nodes": [
                    {
                        "number": "1",
                        "title": "总体技术方案",
                        "suggestion_action": "必要",
                        "suggestion_reason": "",
                        "children": [
                            {
                                "number": "1.1",
                                "title": "机组设计",
                                "suggestion_action": "必要",
                                "suggestion_reason": "",
                                "children": [
                                    {
                                        "number": "1.1.1",
                                        "title": "关键部件设计",
                                        "suggestion_action": "必要",
                                        "suggestion_reason": "",
                                        "children": [],
                                    }
                                ],
                            }
                        ],
                    }
                ],
            }
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

            result = outline_runner.finalize_manifest(manifest, manifest_path)
            self.assertEqual(result["summary"]["total_nodes"], 3)

            payload["nodes"][0]["children"][0]["children"][0]["children"].append(
                {
                    "number": "1.1.1.1",
                    "title": "叶片设计参数",
                    "suggestion_action": "必要",
                    "suggestion_reason": "",
                    "children": [],
                }
            )
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

            with self.assertRaisesRegex(SystemExit, "技术标目录最多三级"):
                outline_runner.finalize_manifest(manifest, manifest_path)



    def test_bid_outline_finalize_allows_selective_review_and_realizes_dynamic_additions(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")
        review_workflow = load_outline_script("review_workflow")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            output = root / "toc.json"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("投标人应提供项目场址安全适应性报告。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(output),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            payload = {
                "schema_version": "technical-outline.v1",
                "nodes": [
                    {
                        "number": "第1章",
                        "title": "技术方案",
                        "suggestion_action": "必要",
                        "suggestion_reason": "",
                        "children": [],
                    }
                ],
            }
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

            selective_result = outline_runner.finalize_manifest(manifest, manifest_path)
            self.assertEqual(selective_result["summary"]["reviewCoverage"], 0.0)

            chunk = review_workflow.next_review_chunk(root)["chunk"]
            evidence_id = chunk["blocks"][0]["evidence_id"]
            review_workflow.submit_chunk_review(
                root,
                chunk["chunk_id"],
                {
                    "review_summary": "发现一项独立专项报告要求。",
                    "requirements": [
                        {
                            "evidence_ids": [evidence_id],
                            "obligation": "投标人应提供项目场址安全适应性报告",
                            "disposition": "suggest_add",
                            "target_node": "第1章",
                            "proposed_title": "项目场址安全适应性报告",
                            "reason": "招标明确要求独立报告，模板没有等价节点。",
                        }
                    ],
                },
            )
            with self.assertRaisesRegex(SystemExit, "未落实到最终目录"):
                outline_runner.finalize_manifest(manifest, manifest_path)

            payload["nodes"][0]["children"].append(
                {
                    "number": "1.1",
                    "title": "项目场址安全适应性报告",
                    "suggestion_action": "建议增加",
                    "suggestion_reason": "招标明确要求独立报告，模板没有等价节点。",
                    "tender_basis": {
                        "file_id": "TEN-1",
                        "search_text": "投标人应提供项目场址安全适应性报告",
                    },
                    "children": [],
                }
            )
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            result = outline_runner.finalize_manifest(manifest, manifest_path)

        self.assertEqual(result["summary"]["reviewCoverage"], 1.0)
        self.assertEqual(result["summary"]["requirementCount"], 1)
        self.assertEqual(result["summary"]["unfinishedTableCount"], 0)



    def test_bid_outline_finalize_does_not_require_reading_appendix_table_contents(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            output = root / "toc.json"
            manifest_path = root / "s2_input.json"

            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("附表A.1 技术参数表")
            table = tender_doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "参数"
            table.cell(0, 1).text = "要求"
            table.cell(1, 0).text = "额定功率"
            table.cell(1, 1).text = "投标人填写"
            tender_doc.save(tender)

            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(output),
                "requireComposedOutline": True,
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            headings = outline_runner.dispatch_command("headings", manifest, manifest_path, [])
            self.assertTrue(headings["complete"])
            batch = outline_runner.dispatch_command("decision-next", manifest, manifest_path, [])
            appendix_heading = outline_runner.dispatch_command(
                "search", manifest, manifest_path, ["附表A.1"]
            )
            outline_runner.dispatch_command(
                "read",
                manifest,
                manifest_path,
                [appendix_heading["results"][0]["evidence_id"]],
            )
            outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": batch["batch_token"],
                            "items": [
                                {
                                    "target_id": item["target_id"],
                                    "decision": "retain",
                                    "reason": "历史模板专家经验保留。",
                                }
                                for item in batch["items"]
                            ],
                            "additions": [],
                        },
                        ensure_ascii=False,
                    )
                ],
            )
            appendix_batch = outline_runner.dispatch_command(
                "appendix-next", manifest, manifest_path, []
            )
            outline_runner.dispatch_command(
                "appendix-decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": appendix_batch["batch_token"],
                            "items": [
                                {
                                    "appendix_id": "APP-0001",
                                    "decision": "include",
                                    "node_id": "ADD-APPENDIX-A1",
                                    "parent_id": "ADD-APPENDIX",
                                    "reason": "招标文件包含独立填写表格。",
                                }
                            ],
                            "root_addition": {
                                "node_id": "ADD-APPENDIX",
                                "parent_id": None,
                                "number": "第2章",
                                "title": "技术附表",
                                "reason": "招标文件包含独立附表。",
                            },
                        },
                        ensure_ascii=False,
                    )
                ],
            )
            outline_runner.dispatch_command(
                "read",
                manifest,
                manifest_path,
                [appendix_batch["items"][0]["evidence_id"]],
            )
            complete_outline_review(outline_runner, manifest, manifest_path)
            outline_runner.dispatch_command("decisions", manifest, manifest_path, [])
            outline_runner.compose_manifest(manifest, manifest_path)

            result = outline_runner.finalize_manifest(manifest, manifest_path)

        self.assertEqual(result["summary"]["workflowStage"], "finalized")
        self.assertEqual(result["summary"]["reviewCoverage"], 0.0)
        self.assertEqual(result["summary"]["unfinishedTableCount"], 1)



    def test_bid_outline_finalize_rejects_nested_technical_appendix(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "toc.json"
            manifest_path = root / "s2_input.json"
            manifest = {"workDir": str(root), "outputFile": str(output)}
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            output.write_text(
                json.dumps(
                    {
                        "schema_version": "technical-outline.v1",
                        "nodes": [
                            {
                                "number": "第6章",
                                "title": "技术方案",
                                "suggestion_action": "必要",
                                "suggestion_reason": "",
                                "children": [
                                    {
                                        "number": "6.6",
                                        "title": "技术附表",
                                        "suggestion_action": "建议增加",
                                        "suggestion_reason": "招标包含技术附表。",
                                        "children": [],
                                    }
                                ],
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            with self.assertRaisesRegex(SystemExit, "技术附表必须是唯一的最后一个根节点"):
                outline_runner.finalize_manifest(manifest, manifest_path)



    def test_bid_outline_finalize_rejects_missing_included_appendix(self) -> None:
        def remove_included_addition(state: dict, _items: list[dict]) -> None:
            state["additions"] = [
                addition
                for addition in state["additions"]
                if addition.get("node_id") != "ADD-APPENDIX-1"
            ]

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            outline_runner, manifest, manifest_path = self._prepare_finalized_appendix_outline(
                root,
                ["include", "exclude"],
                mutate_state=remove_included_addition,
            )

            with self.assertRaisesRegex(SystemExit, "APP-0001"):
                outline_runner.finalize_manifest(manifest, manifest_path)



    def test_bid_outline_finalize_rejects_excluded_appendix_in_output(self) -> None:
        def exclude_included_addition(state: dict, items: list[dict]) -> None:
            appendix_id = items[0]["appendix_id"]
            state["appendix_decisions"][appendix_id] = {
                "appendix_id": appendix_id,
                "decision": "exclude",
                "reason": "逐项核验后排除。",
            }

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            outline_runner, manifest, manifest_path = self._prepare_finalized_appendix_outline(
                root,
                ["include", "exclude"],
                mutate_state=exclude_included_addition,
            )

            with self.assertRaisesRegex(SystemExit, "APP-0001"):
                outline_runner.finalize_manifest(manifest, manifest_path)



    def test_bid_outline_finalize_accepts_exact_ai_appendix_decisions(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            outline_runner, manifest, manifest_path = self._prepare_finalized_appendix_outline(
                root,
                ["include", "exclude"],
            )

            result = outline_runner.finalize_manifest(manifest, manifest_path)

        self.assertEqual(result["summary"]["workflowStage"], "finalized")



    def test_bid_outline_finalize_allows_missing_appendices_and_free_number_title_split(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "toc.json"
            manifest_path = root / "s2_input.json"
            manifest = {"workDir": str(root), "outputFile": str(output)}
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            (root / "template_structure.json").write_text(
                json.dumps(
                    {
                        "schema_version": "template-structure.v1",
                        "items": [{"number": "第1章", "title": "技术方案", "level": 1}],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            (root / "tender_appendix_inventory.json").write_text(
                json.dumps(
                    {
                        "schema_version": "tender-appendix-inventory.v1",
                        "items": [
                            {
                                "number": "附表A.1",
                                "title": "机型信息表",
                                "following_table_count": 1,
                            },
                            {
                                "number": "附表A.2",
                                "title": "特殊方案表",
                                "following_table_count": 1,
                            },
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            output.write_text(
                json.dumps(
                    {
                        "schema_version": "technical-outline.v1",
                        "nodes": [
                            {
                                "number": "第1章",
                                "title": "技术方案",
                                "suggestion_action": "必要",
                                "suggestion_reason": "",
                                "children": [],
                            },
                            {
                                "number": "第2章",
                                "title": "技术附表",
                                "suggestion_action": "建议增加",
                                "suggestion_reason": "招标文件新增独立附表章节。",
                                "children": [
                                    {
                                        "number": "A.1",
                                        "title": "附表A.1 机型信息表",
                                        "suggestion_action": "建议增加",
                                        "suggestion_reason": "招标文件新增独立表格。",
                                        "children": [],
                                    }
                                ],
                            },
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = outline_runner.finalize_manifest(manifest, manifest_path)
            payload = json_load(output)
            payload["nodes"] = payload["nodes"][:1]
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            result_without_appendix = outline_runner.finalize_manifest(manifest, manifest_path)

        self.assertEqual(result["summary"]["workflowStage"], "finalized")
        self.assertEqual(result_without_appendix["summary"]["workflowStage"], "finalized")



    def test_bid_outline_generator_runner_only_validates_inputs_and_writes_no_final_toc(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            output = root / "toc.json"
            manifest_path = root / "s2_input.json"

            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)

            tender_doc = Document()
            tender_doc.add_paragraph("投标人应提供实施方案。")
            tender_doc.save(tender)

            manifest = {
                "projectId": "PRJ-TEST",
                "projectName": "测试项目",
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(output),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
            result = outline_runner.run_manifest(manifest, manifest_path)

            self.assertEqual(result["summary"]["schema_version"], "technical-outline-inputs.v1")
            self.assertEqual(result["summary"]["templateFile"], str(template))
            self.assertEqual(result["summary"]["tenderFileCount"], 1)
            self.assertEqual(result["summary"]["outputFile"], str(output))
            self.assertFalse(output.exists())
            self.assertFalse((root / "agent_review_input.json").exists())



    def test_bid_outline_generator_finalize_validates_existing_final_outputs(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")
        review_workflow = load_outline_script("review_workflow")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            output = root / "toc.json"
            manifest_path = root / "s2_input.json"

            template_doc = Document()
            template_doc.add_paragraph("第1章 总体技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("投标人应提供总体技术方案")
            tender_doc.save(tender)
            manifest = {
                "projectId": "PRJ-TEST",
                "projectName": "test",
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(output),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            chunk = review_workflow.next_review_chunk(root)["chunk"]
            evidence_id = chunk["blocks"][0]["evidence_id"]
            review_workflow.submit_chunk_review(
                root,
                chunk["chunk_id"],
                {
                    "review_summary": "总体技术方案要求由模板根节点承接。",
                    "requirements": [
                        {
                            "evidence_ids": [evidence_id],
                            "obligation": "投标人应提供总体技术方案",
                            "disposition": "map_existing",
                            "target_node": "第1章",
                        }
                    ],
                },
            )
            output.write_text(
                json.dumps(
                    {
                        "schema_version": "technical-outline.v1",
                        "nodes": [
                            {
                                "number": "第1章",
                                "title": "总体技术方案",
                                "suggestion_action": "必要",
                                "suggestion_reason": "",
                                "tender_basis": {
                                    "file_id": "TEN-1",
                                    "search_text": "投标人应提供总体技术方案",
                                },
                                "children": [
                                    {
                                        "number": "1.1",
                                        "title": "机组选型",
                                        "suggestion_action": "建议增加",
                                        "suggestion_reason": "招标文件要求单独编制机组选型方案。",
                                        "children": [],
                                    }
                                ],
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            result = outline_runner.finalize_manifest(manifest, manifest_path)

        self.assertEqual(result["schema_version"], "technical-outline.v1")
        self.assertEqual(result["summary"]["workflowStage"], "finalized")
        self.assertEqual(result["summary"]["total_nodes"], 2)
        self.assertEqual(result["summary"]["action_counts"], {"必要": 1, "建议增加": 1})
        self.assertNotIn("evidenceFile", result)



    def test_bid_outline_finalize_rejects_unlocatable_tender_basis(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")
        review_workflow = load_outline_script("review_workflow")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            output = root / "toc.json"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 总体技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("投标人应提供真实、完整的总体技术方案。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "outputFile": str(output),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            chunk = review_workflow.next_review_chunk(root)["chunk"]
            evidence_id = chunk["blocks"][0]["evidence_id"]
            review_workflow.submit_chunk_review(
                root,
                chunk["chunk_id"],
                {
                    "review_summary": "总体技术方案要求由模板根节点承接。",
                    "requirements": [
                        {
                            "evidence_ids": [evidence_id],
                            "obligation": "投标人应提供真实、完整的总体技术方案",
                            "disposition": "map_existing",
                            "target_node": "第1章",
                        }
                    ],
                },
            )
            payload = {
                "schema_version": "technical-outline.v1",
                "nodes": [
                    {
                        "number": "第1章",
                        "title": "总体技术方案",
                        "suggestion_action": "必要",
                        "suggestion_reason": "",
                        "tender_basis": {"file_id": "TEN-1", "search_text": "原文不存在的编造依据"},
                        "children": [],
                    }
                ],
            }
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

            with self.assertRaisesRegex(SystemExit, "search_text 无法在 tenderFile 定位"):
                outline_runner.finalize_manifest(manifest, manifest_path)

            payload["nodes"][0]["tender_basis"]["search_text"] = "投标人应提供真实、完整的总体技术方案"
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            result = outline_runner.finalize_manifest(manifest, manifest_path)

        self.assertEqual(result["summary"]["total_nodes"], 1)



    def test_bid_outline_finalize_rejects_necessary_appendix_absent_from_template(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "toc.json"
            manifest_path = root / "s2_input.json"
            manifest = {"workDir": str(root), "outputFile": str(output)}
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            (root / "template_structure.json").write_text(
                json.dumps(
                    {
                        "schema_version": "template-structure.v1",
                        "items": [{"number": "第1章", "title": "技术方案", "level": 1}],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            appendix_inventory = {
                "schema_version": "tender-appendix-inventory.v1",
                "items": [
                    {
                        "file_id": "TEN-1",
                        "file_name": "tender.docx",
                        "number": "附表A.1",
                        "title": "投标机型总方案信息表",
                        "raw_text": "附表A.1 投标机型总方案信息表",
                        "paragraph_index": 20,
                        "following_table_count": 0,
                        "following_text_count": 0,
                    },
                    {
                        "file_id": "TEN-1",
                        "file_name": "tender.docx",
                        "number": "附表B.1",
                        "title": "另一张实际表单",
                        "raw_text": "附表B.1 另一张实际表单",
                        "paragraph_index": 30,
                        "following_table_count": 1,
                        "following_text_count": 0,
                    },
                ],
            }
            (root / "tender_appendix_inventory.json").write_text(
                json.dumps(appendix_inventory, ensure_ascii=False),
                encoding="utf-8",
            )
            output.write_text(
                json.dumps(
                    {
                        "schema_version": "technical-outline.v1",
                        "nodes": [
                            {
                                "number": "第1章",
                                "title": "技术方案",
                                "suggestion_action": "必要",
                                "suggestion_reason": "",
                                "children": [],
                            },
                            {
                                "number": "第2章",
                                "title": "技术附表",
                                "suggestion_action": "必要",
                                "suggestion_reason": "",
                                "children": [
                                    {
                                        "number": "附表A.1",
                                        "title": "投标机型总方案信息表",
                                        "suggestion_action": "必要",
                                        "suggestion_reason": "",
                                        "children": [],
                                    }
                                ],
                            },
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            with self.assertRaisesRegex(SystemExit, "模板目录不存在.*建议增加"):
                outline_runner.finalize_manifest(manifest, manifest_path)

            payload = json_load(output)
            appendix = payload["nodes"][-1]
            appendix["suggestion_action"] = "建议增加"
            appendix["suggestion_reason"] = "招标文件新增技术附表。"
            appendix["children"][0]["suggestion_action"] = "建议增加"
            appendix["children"][0]["suggestion_reason"] = "招标文件新增独立表单。"
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

            with self.assertRaisesRegex(SystemExit, "没有独立填写表格"):
                outline_runner.finalize_manifest(manifest, manifest_path)

            appendix_inventory["items"][0]["following_table_count"] = 1
            (root / "tender_appendix_inventory.json").write_text(
                json.dumps(appendix_inventory, ensure_ascii=False),
                encoding="utf-8",
            )
            appendix["children"].append(
                {
                    "number": "附表B.1",
                    "title": "另一张实际表单",
                    "suggestion_action": "建议增加",
                    "suggestion_reason": "招标文件新增独立表单。",
                    "children": [],
                }
            )
            output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            result = outline_runner.finalize_manifest(manifest, manifest_path)

        self.assertEqual(result["summary"]["action_counts"], {"必要": 1, "建议增加": 3})
