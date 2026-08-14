"""正文标题层级编号契约：一级「第N章」，二级及以下「N.M」「N.M.K」。

编号一律由 chapter_no_flat 推导，与素材内部标题的父前缀同源；确认目录里的
chapter_no（模板原样号，可能是「1、」这类段内序号）只作阶段间素材关联键。
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.document_processing.technical_document.assembly.merger import merge
from app.document_processing.technical_document.assembly.parse_toc import display_chapter_no
from app.document_processing.technical_document.assembly.numbering_fixer import (
    _NUMBERED_HEADING_TEXT_RE,
    strip_handwritten_numbering_in_body,
)
from app.services.tech_assembly import _tech_format_sections_from_toc_items


COMMITMENT_LETTER_CLAUSES = [
    "2.1上述承诺电量是基于测风塔轮毂高度处代表年平均风速计算。",
    "2.2 上述承诺电量是基于卖方提供的由EW10.0-220-125机型评估得出，轮毂高度125m，"
    "如果发生机型、轮毂高度变更，卖方须重新评估承诺电量。",
    "2.3 上述承诺电量基于买方提供的风资源信息表。若机位排布、海拔高度等风资源发生变更，"
    "将基于新的信息输入重新进行电量核算并编制发电量承诺表。",
    "2.4上述承诺电量计量点为本项目的送出线路关口表。",
    "2.5上述承诺电量不包括由于非卖方机组原因造成的发电量损失（如电网限电、现场阻工、"
    "居民点过近造成限功率运行或停机、业主方计划停机、电网和升压站可靠性损失、线损和厂用电等）。",
    "2.6上述承诺电量不包括由于不可抗力因素导致的发电量损失（如极端天气、地震、雷击、冰冻、洪水等自然灾害）。",
    "2.7若风场周边未来新建风电场对本风场造成尾流影响，双方应协商并评估周边风电场影响。",
    "2.8对于上述承诺电量的考核须建立在考核期内风机设备由卖方运维的前提下，"
    "并且招标方需要保证现场满足按时维护的条件。",
    "2.9发电量保证期为整机质保期，全场所有风机预验收完成证书签署之日起进入发电量保证考核周期，"
    "自发电量保证考核期第一日起，针对发电量保证值每1年考核一次，在合同约定的质保期结束后统一结算。",
]


def _structural_entry(level: int, chapter_no: str, flat: str, title: str) -> dict:
    return {
        "toc_idx": 0,
        "level": level,
        "chapter_no": chapter_no,
        "chapter_no_flat": flat,
        "title": title,
        "raw_text": f"{chapter_no} {title}",
        "tag": "normal",
        "is_preface": False,
        "is_appendix": False,
        "paths": [],
        "shifts": [],
        "attach_modes": [],
        "status": "STRUCTURAL",
        "note": "",
    }


class DisplayChapterNoTest(unittest.TestCase):
    def test_level_one_becomes_chapter_form(self) -> None:
        self.assertEqual(display_chapter_no({"chapter_no": "1", "chapter_no_flat": "1"}), "第1章")
        self.assertEqual(display_chapter_no({"chapter_no": "第5章", "chapter_no_flat": "5"}), "第5章")

    def test_deeper_levels_use_dotted_flat_number(self) -> None:
        self.assertEqual(display_chapter_no({"chapter_no": "1、", "chapter_no_flat": "1.1"}), "1.1")
        self.assertEqual(display_chapter_no({"chapter_no": "2、", "chapter_no_flat": "1.2"}), "1.2")
        self.assertEqual(display_chapter_no({"chapter_no": "1、", "chapter_no_flat": "1.1.1"}), "1.1.1")

    def test_template_ordinal_never_reaches_the_body(self) -> None:
        """模板抄来的「N、」是段内序号，正文标题里不能出现。"""
        for flat, expected in (("1", "第1章"), ("1.3", "1.3"), ("1.3.1", "1.3.1")):
            self.assertEqual(display_chapter_no({"chapter_no": "3、", "chapter_no_flat": flat}), expected)

    def test_preface_and_appendix_keep_their_own_form(self) -> None:
        self.assertEqual(display_chapter_no({"chapter_no": "前言", "chapter_no_flat": ""}), "前言")
        self.assertEqual(display_chapter_no({"chapter_no": "附表A", "chapter_no_flat": ""}), "附表A")
        self.assertEqual(display_chapter_no({"chapter_no": "附", "chapter_no_flat": ""}), "附")

    def test_accepts_toc_item_key_spelling(self) -> None:
        self.assertEqual(display_chapter_no({"number": "1、", "chapter_no_flat": "2.4"}), "2.4")
        self.assertEqual(display_chapter_no({"number": "附表A"}), "附表A")


class NumberedHeadingGuardTest(unittest.TestCase):
    def test_chapter_form_counts_as_text_numbered(self) -> None:
        """「第N章 xxx」也是已写入文本编号，必须被 Word 自动编号抑制覆盖到。"""
        self.assertTrue(_NUMBERED_HEADING_TEXT_RE.match("第1章 标前概述"))
        self.assertTrue(_NUMBERED_HEADING_TEXT_RE.match("第 12 章  投标技术方案"))
        self.assertTrue(_NUMBERED_HEADING_TEXT_RE.match("1.3.1 华能汕头勒门"))

    def test_unnumbered_and_appendix_headings_are_untouched(self) -> None:
        self.assertIsNone(_NUMBERED_HEADING_TEXT_RE.match("标前概述"))
        self.assertIsNone(_NUMBERED_HEADING_TEXT_RE.match("附表A 参数表"))


class VerifyGuardsTest(unittest.TestCase):
    """verify 的两条启发式原本假设「第N章」只会是素材漏进标题的幽灵编号。"""

    def test_arabic_chapter_heading_is_a_valid_h1(self) -> None:
        from app.document_processing.technical_document.assembly.verify import ALLOWED_H1_PATTERNS

        for text in ("第1章 标前概述", "第7章 技术附表", "第十章 附录"):
            self.assertTrue(any(pattern.search(text) for pattern in ALLOWED_H1_PATTERNS), text)

    def test_seventh_chapter_is_not_a_ghost(self) -> None:
        """项目可以有 7 个以上章节；开头的第N章是正式章号，不是幽灵。"""
        from app.document_processing.technical_document.assembly.verify import GHOST_CHAPTER

        match = GHOST_CHAPTER.search("第7章  技术附表")
        self.assertIsNotNone(match)
        self.assertEqual(match.start(), 0, "开头匹配才会被 gm.start() > 0 放行")

        leaked = GHOST_CHAPTER.search("1.2.3 关于第9章的说明")
        self.assertIsNotNone(leaked)
        self.assertGreater(leaked.start(), 0, "标题中间的章节引用仍要被判为幽灵")


class FormatCleanerOutlineTest(unittest.TestCase):
    def test_cleaner_outline_uses_the_same_numbering_as_the_body(self) -> None:
        """cleaner 按「编号+标题」匹配正文，两边编号必须同源，否则全部落成未匹配。"""
        sections = _tech_format_sections_from_toc_items(
            [
                {"level": 1, "number": "1", "chapter_no_flat": "1", "title": "标前概述"},
                {"level": 2, "number": "1、", "chapter_no_flat": "1.1", "title": "技术评分标准索引表"},
                {"level": 3, "number": "1、", "chapter_no_flat": "1.1.1", "title": "华能汕头勒门"},
                {"level": 1, "number": "附表A", "chapter_no_flat": "", "title": "参数表"},
            ]
        )
        chapter = sections[0]
        self.assertEqual(chapter["number"], "第1章")
        self.assertEqual(chapter["children"][0]["number"], "1.1")
        self.assertEqual(chapter["children"][0]["children"][0]["number"], "1.1.1")
        self.assertEqual(sections[1]["number"], "附表A")


class MergerHeadingTextTest(unittest.TestCase):
    def test_merged_body_headings_carry_hierarchical_numbers(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            template = work / "母版.docx"
            Document().save(str(template))
            plan = [
                _structural_entry(1, "1", "1", "标前概述"),
                _structural_entry(2, "1、", "1.1", "技术评分标准索引表"),
                _structural_entry(2, "2、", "1.2", "战略合作协议"),
                _structural_entry(3, "1、", "1.2.1", "海上示范应用"),
                _structural_entry(1, "2", "2", "技术标准"),
            ]
            output = work / "merged.docx"
            merge(template, plan, work, {}, work / "prep", output)

            headings = [
                (paragraph.style.name, paragraph.text)
                for paragraph in Document(str(output)).paragraphs
                if paragraph.style and paragraph.style.name.startswith("Heading") and paragraph.text.strip()
            ]
            self.assertEqual(
                headings,
                [
                    ("Heading 1", "第1章  标前概述"),
                    ("Heading 2", "1.1  技术评分标准索引表"),
                    ("Heading 2", "1.2  战略合作协议"),
                    ("Heading 3", "1.2.1  海上示范应用"),
                    ("Heading 1", "第2章  技术标准"),
                ],
            )


class BodyHandwrittenNumberingTest(unittest.TestCase):
    """正文条款编号（承诺函 2.1~2.9）有法律引用意义，不能被伪标题编号擦除吃掉。

    条款文本取自素材《发电小时数承诺函（承诺保证值）》，同时覆盖「编号后有空格」
    与「编号后无空格」两种写法——历史上前者被擦、后者保留，同一份文件被擦掉一半。
    """

    def test_commitment_letter_clause_numbering_survives(self) -> None:
        doc = Document()
        paragraphs = [doc.add_paragraph(text) for text in COMMITMENT_LETTER_CLAUSES]

        self.assertEqual(strip_handwritten_numbering_in_body(doc), 0)
        self.assertEqual([p.text for p in paragraphs], COMMITMENT_LETTER_CLAUSES)

    def test_pseudo_heading_numbering_still_stripped(self) -> None:
        doc = Document()
        paragraphs = [
            doc.add_paragraph("7.10 符合招标公告及招标文件要求的业绩情况"),
            doc.add_paragraph("5.1 项目概况"),
        ]

        self.assertEqual(strip_handwritten_numbering_in_body(doc), 2)
        self.assertEqual(
            [p.text for p in paragraphs],
            ["符合招标公告及招标文件要求的业绩情况", "项目概况"],
        )


if __name__ == "__main__":
    unittest.main()
