"""媒体旁路的保真契约：图片必须逐字节原样回到成稿，一个都不能少、不能改。"""

from __future__ import annotations

import hashlib
import struct
import zipfile
import zlib
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Cm

from app.document_processing.technical_document.assembly.docx_package import (
    read_parts,
    rewrite_parts,
)
from app.document_processing.technical_document.assembly.media_vault import (
    STUB_MAGIC,
    MediaVault,
    restore_media,
    strip_media,
)
from app.document_processing.technical_document.assembly.preprocess import (
    sanitize_dangling_media_refs,
)

_VML_NS = "urn:schemas-microsoft-com:vml"
_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _png(path: Path, *, color: bytes) -> Path:
    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    width = height = 8
    raw = b"".join(b"\x00" + color * width for _ in range(height))
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    return path


def _docx_with_images(path: Path, images: list[Path]) -> Path:
    doc = Document()
    doc.add_paragraph("正文段落")
    for image in images:
        doc.add_paragraph().add_run().add_picture(str(image), width=Cm(3), height=Cm(2))
    doc.save(str(path))
    return path


def _media_map(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {
            info.filename: archive.read(info)
            for info in archive.infolist()
            if info.filename.startswith("word/media/")
        }


def _part_map(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {info.filename: archive.read(info) for info in archive.infolist()}


@pytest.fixture()
def sample(tmp_path: Path) -> Path:
    red = _png(tmp_path / "red.png", color=b"\xff\x00\x00")
    blue = _png(tmp_path / "blue.png", color=b"\x00\x00\xff")
    return _docx_with_images(tmp_path / "sample.docx", [red, blue])


def test_round_trip_restores_every_image_byte_for_byte(sample: Path, tmp_path: Path) -> None:
    before = _media_map(sample)
    assert before, "样例文档应当带图，否则这条契约测了个寂寞"

    vault = MediaVault()
    light = tmp_path / "light.docx"
    strip_media(sample, light, vault)
    restored = tmp_path / "restored.docx"
    restore_media(light, restored, vault)

    assert _media_map(restored) == before
    # 非媒体部件同样不能被动过
    assert {k: v for k, v in _part_map(restored).items() if not k.startswith("word/media/")} == {
        k: v for k, v in _part_map(sample).items() if not k.startswith("word/media/")
    }


def test_stripped_document_carries_no_image_bytes(sample: Path, tmp_path: Path) -> None:
    """轻量稿必须真的把图片字节留在原件里，否则整个方案没有意义。"""
    vault = MediaVault()
    light = tmp_path / "light.docx"
    strip_media(sample, light, vault)

    assert light.stat().st_size < sample.stat().st_size
    for name, payload in _media_map(light).items():
        assert payload.startswith(STUB_MAGIC), name
    assert len(vault) == len(_media_map(sample))
    assert Document(str(light)).paragraphs[0].text == "正文段落"


def test_identical_images_share_one_vault_entry(tmp_path: Path) -> None:
    """同图去重语义必须与原本按 sha1 去重一致，不能因为旁路而在成稿里存出多份。"""
    red = _png(tmp_path / "red.png", color=b"\xff\x00\x00")
    source = _docx_with_images(tmp_path / "twins.docx", [red, red])

    vault = MediaVault()
    strip_media(source, tmp_path / "light.docx", vault)

    stubs = set(_media_map(tmp_path / "light.docx").values())
    assert len(vault) == 1
    assert len(stubs) == 1


def test_strip_is_idempotent(sample: Path, tmp_path: Path) -> None:
    """重复剥离不能把占位符自身当原件登记，否则"还原"出来的会是占位符。"""
    vault = MediaVault()
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    strip_media(sample, first, vault)
    strip_media(first, second, vault)

    assert len(vault) == len(_media_map(sample))
    restored = tmp_path / "restored.docx"
    restore_media(second, restored, vault)
    assert _media_map(restored) == _media_map(sample)


def test_restore_fails_loudly_when_source_is_gone(sample: Path, tmp_path: Path) -> None:
    """索引对不上时必须报错，不能静默产出缺图的标书。"""
    vault = MediaVault()
    light = tmp_path / "light.docx"
    strip_media(sample, light, vault)
    sample.unlink()

    with pytest.raises(RuntimeError, match="媒体还原失败"):
        restore_media(light, tmp_path / "restored.docx", vault)


def test_rewrite_parts_touches_only_the_named_part(sample: Path) -> None:
    """外科式改写：换一个 XML 部件，其余条目（尤其图片）必须原样不动。"""
    before = _part_map(sample)
    settings_name = "word/settings.xml"
    assert settings_name in before

    replaced = rewrite_parts(sample, {settings_name: b"<w:settings/>"})

    after = _part_map(sample)
    assert replaced == 1
    assert after[settings_name] == b"<w:settings/>"
    assert {k: v for k, v in after.items() if k != settings_name} == {
        k: v for k, v in before.items() if k != settings_name
    }


def test_read_parts_returns_only_requested_names(sample: Path) -> None:
    parts = read_parts(sample, {"word/settings.xml", "word/does-not-exist.xml"})
    assert set(parts) == {"word/settings.xml"}


def test_sanitize_drops_imagedata_without_relationship_but_keeps_text(tmp_path: Path) -> None:
    """WPS 转档常留下连 r:id 都没有的空 v:imagedata，docxcompose 会因此整份素材失败。

    摘掉这类元素不能连带丢掉同一个图形里的文字。
    """
    doc = Document()
    paragraph = doc.add_paragraph()
    pict = paragraph._p.makeelement(f"{{{_WORD_NS}}}pict", {})
    shape = pict.makeelement(f"{{{_VML_NS}}}shape", {})
    imagedata = shape.makeelement(f"{{{_VML_NS}}}imagedata", {})
    textbox = shape.makeelement(f"{{{_VML_NS}}}textbox", {})
    shape.append(imagedata)
    shape.append(textbox)
    pict.append(shape)
    paragraph._p.append(pict)

    stats = sanitize_dangling_media_refs(doc)

    assert stats["vml_dropped"] == 1
    assert doc.element.find(f".//{{{_VML_NS}}}imagedata") is None
    assert doc.element.find(f".//{{{_VML_NS}}}textbox") is not None


def test_sanitize_keeps_valid_image_references(sample: Path) -> None:
    doc = Document(str(sample))
    digest_before = hashlib.sha1(sample.read_bytes()).hexdigest()

    stats = sanitize_dangling_media_refs(doc)

    assert stats == {
        "vml_repaired": 0,
        "vml_dropped": 0,
        "blip_dropped": 0,
        "diagram_dropped": 0,
    }
    assert hashlib.sha1(sample.read_bytes()).hexdigest() == digest_before
