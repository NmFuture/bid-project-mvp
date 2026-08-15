from __future__ import annotations

import json
import tempfile
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from toc_skill_helpers import (
    TocSkillScriptTestBase,
    json_load,
    load_material_cleaner_script,
    load_wiki_script,
    load_word_filler_script,
)


class TocSkillScriptTests(TocSkillScriptTestBase):

    def test_bid_wiki_builder_writes_full_blueprint_and_returns_small_summary(self) -> None:
        wiki_runner = load_wiki_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest_path = root / "wiki_manifest.json"
            output_file = root / "wiki_blueprint.json"
            manifest = {
                "bidType": "技术标",
                "rootTitle": "技术标Wiki（自动生成）",
                "workDir": str(root),
                "outputFile": str(output_file),
                "stats": {"fileCount": 1},
                "tiers": [
                    {
                        "name": "标准文件",
                        "tier": "standard",
                        "path": "技术标/标准文件",
                        "fileCount": 1,
                        "folders": [
                            {
                                "name": "EW5.0",
                                "path": "技术标/标准文件/EW5.0",
                                "fileCount": 1,
                                "files": [
                                    {
                                        "id": "M1",
                                        "name": "总体方案.docx",
                                        "path": "技术标/标准文件/EW5.0/总体方案.docx",
                                        "ext": "docx",
                                    }
                                ],
                            }
                        ],
                    }
                ],
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

            response = wiki_runner.run_manifest(manifest, manifest_path)

            self.assertEqual(response["outputFile"], str(output_file))
            self.assertIn("nodeTitles", response)
            self.assertLess(len(json.dumps(response, ensure_ascii=False)), 1000)
            blueprint = json_load(output_file)
            self.assertEqual(blueprint["rootTitle"], "技术标Wiki（自动生成）")
            self.assertEqual(
                [node["title"] for node in blueprint["nodes"]],
                ["标准文件"],
            )
            self.assertEqual(blueprint["nodes"][0]["children"][0]["children"][0]["title"], "总体方案.docx")



    def test_bid_word_placeholder_filler_fills_key_data_table_from_spec_columns(self) -> None:
        """清单驱动端到端：占位符即字段名直接命中，关键数据表语义校验全通过。

        原先这一区有 5 个用例走的是「事实表没有清单第 2/3 列 → 上下文规则 + 全库模糊
        匹配 + 从素材抽事实」的旧链路。该链路已删除，缺清单元数据现在直接报错，对应契约
        由 tests/test_technical_word_fill_discipline.py 的 SpecLocateTests（逐级定位）
        与 SpeclessManifestTests（缺清单即失败）覆盖。
        """
        word_filler = load_word_filler_script("run_from_manifest")

        blank_name = "待填写-投标关键数据一览表.docx"

        def _field(label: str, value: str, placeholder: str) -> dict:
            return {
                "label": label,
                "value": value,
                "status": "confirmed",
                "placeholder": placeholder,
                "targetFile": f"客户定制/华能/{blank_name}",
            }

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            template = tmp_path / blank_name
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.cell(0, 0).text = "项目名称"
            table.cell(0, 1).text = "[项目名称，待填写]"
            for row in (
                ["轮毂高度（m）", "[轮毂高度，待填写]"],
                ["台数", "[机组台数，待填写]"],
                ["容量（MW）", "[总装机容量，待填写]"],
                ["有效小时数（h）", "[保证有效小时数，待填写]"],
                ["单台机组功率曲线保证率（%）", "[功率曲线保证率，待填写]"],
            ):
                cells = table.add_row().cells
                cells[0].text, cells[1].text = row
            doc.save(template)

            manifest_path = tmp_path / "manifest.json"
            output = tmp_path / "filled.docx"
            manifest_path.write_text(
                json.dumps(
                    {
                        "schemaVersion": "bid-tech-word-placeholder-fill-v1",
                        "projectName": "测试项目",
                        "blankSource": {
                            "id": "RAW-KEY-DATA",
                            "title": blank_name,
                            "docxPath": str(template),
                            "placeholderCount": 6,
                        },
                        "projectFactTable": {
                            "status": "confirmed",
                            "fields": [
                                _field("项目名称", "测试项目", "[项目名称，待填写]"),
                                _field("轮毂高度", "125", "[轮毂高度，待填写]"),
                                _field("机组台数", "60", "[机组台数，待填写]"),
                                _field("总装机容量", "600", "[总装机容量，待填写]"),
                                _field("保证有效小时数", "2836", "[保证有效小时数，待填写]"),
                                _field("功率曲线保证率", "97%", "[功率曲线保证率，待填写]"),
                            ],
                        },
                        "outputFile": str(output),
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = word_filler.run_from_manifest(manifest_path)

            self.assertEqual(result["unfilledFields"], [])
            self.assertTrue(result["fillReport"]["specDriven"])
            # 六处占位符全部走「占位符即字段名」的确定性路径，没有一处靠上下文消歧
            self.assertEqual(result["fillReport"]["specFieldNameHitCount"], 6)
            rows = Document(str(output)).tables[0].rows
            self.assertEqual(rows[0].cells[1].text, "测试项目")
            self.assertEqual(rows[1].cells[1].text, "125")
            self.assertEqual(rows[2].cells[1].text, "60")
            self.assertEqual(rows[3].cells[1].text, "600")
            self.assertEqual(rows[4].cells[1].text, "2836")
            self.assertEqual(rows[5].cells[1].text, "97%")
            # 关键数据表语义校验仍然生效：删素材链路没有把这层跨字段复核一起带走
            self.assertEqual(result["fillReport"]["semanticCheckCount"], 6)
            self.assertEqual(result["fillReport"]["semanticFailedCount"], 0)
            self.assertEqual(result["fillReport"]["semanticValidationRate"], 1.0)



    def test_material_cleaner_does_not_fabricate_heading_levels(self) -> None:
        """1.9 回归：清洗不按"带编号、文字较短"把正文猜成标题。"""
        word_cleaner = load_material_cleaner_script("word_cleaner")

        doc = Document()
        body = doc.add_paragraph("1、载荷仿真分析能力")  # Normal 样式 + 手写编号
        heading = doc.add_paragraph("2、自主研发控制算法", style="Heading 1")

        word_cleaner._strip_numbered_heading_prefixes(doc)

        # 正文不升格、文本保持原样
        self.assertEqual(body.style.name, "Normal")
        self.assertEqual(body.text, "1、载荷仿真分析能力")
        self.assertIsNone(body._p.find(qn("w:pPr")))
        # 真 Heading 仍剥手写前缀
        self.assertEqual(heading.text, "自主研发控制算法")



    def test_material_cleaner_preserves_basedon_heading_level(self) -> None:
        """自定义样式 basedOn Heading 6 时按真实层级清理，不按文本编号推断。"""
        word_cleaner = load_material_cleaner_script("word_cleaner")

        doc = Document()
        custom = doc.styles.add_style("标题6-标书", WD_STYLE_TYPE.PARAGRAPH)
        custom.base_style = doc.styles["Heading 6"]
        heading = doc.add_paragraph("1.7 自定义技术路线", style=custom)

        normalized = word_cleaner._strip_numbered_heading_prefixes(doc)

        self.assertEqual(normalized, 1)
        self.assertEqual(heading.text, "自定义技术路线")
        self.assertEqual(heading.style.name, "Heading 6")
        outline = heading._p.find(qn("w:pPr")).find(qn("w:outlineLvl"))
        self.assertEqual(outline.get(qn("w:val")), "5")



    def test_material_cleaner_suppresses_inherited_numbering_on_direct_outline_heading(self) -> None:
        """direct-outline 标题应局部抑制样式链编号，且不影响共用样式的正文。"""
        word_cleaner = load_material_cleaner_script("word_cleaner")

        doc = Document()
        numbered_base = self._make_hidden_numbered_custom_style(
            doc,
            "自动编号基准",
            "Normal",
            "7",
        )
        shared = doc.styles.add_style("共享编号正文", WD_STYLE_TYPE.PARAGRAPH)
        shared.base_style = numbered_base

        heading = doc.add_paragraph("总体技术路线", style=shared)
        heading_p_pr = heading._p.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "2")
        heading_p_pr.append(outline)
        body_item = doc.add_paragraph("正文列表项", style=shared)

        normalized = word_cleaner._strip_numbered_heading_prefixes(doc)

        self.assertEqual(normalized, 1)
        heading_num_pr = heading._p.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertIsNotNone(heading_num_pr)
        self.assertEqual(heading_num_pr.find(qn("w:numId")).get(qn("w:val")), "0")
        style_num_pr = numbered_base.element.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertEqual(style_num_pr.find(qn("w:numId")).get(qn("w:val")), "7")
        body_p_pr = body_item._p.find(qn("w:pPr"))
        self.assertIsNotNone(body_p_pr)
        self.assertIsNone(body_p_pr.find(qn("w:numPr")))
