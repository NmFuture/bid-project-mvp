from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from app.core.config import settings
from app.main import app
from app.services.onlyoffice_documents import document_path
from app.services.store import store
from app.services.technical_score_index_state import score_index_state

PLACEHOLDER = "[待人工补充：章节索引]"


def _state_for_tests(project_id: str) -> dict:
    return score_index_state(store.get_project_runtime_state(project_id))


def _document_state_for_tests(project_id: str) -> dict:
    return dict(store.get_project_runtime_state(project_id)["document_state"])


def _write_draft_with_index_table(path: Path) -> None:
    """造一份带评分索引表的成稿，形状与格式清洗产物一致。"""
    doc = Document()
    doc.add_heading("5 投标技术方案", level=1)
    doc.add_paragraph("总述正文。")
    doc.add_heading("5.1 投标总体方案概述", level=2)
    doc.add_paragraph("概述正文。")
    doc.add_heading("5.3.3 叶片设计", level=3)
    doc.add_paragraph("叶片正文。")

    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "序号"
    table.cell(0, 1).text = "评审因素"
    table.cell(0, 2).text = "章节索引"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "风轮系统先进性及可靠性"
    table.cell(1, 2).text = PLACEHOLDER
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


class TechnicalScoreIndexFlowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        base = Path(self.temp_dir.name)
        settings.uploads_dir = base / "uploads"
        settings.documents_dir = base / "documents"
        settings.parsed_dir = base / "parsed"
        settings.ensure_dirs()

        store.reset_for_tests()
        self.client = TestClient(app, base_url="http://127.0.0.1:8000")
        login = self.client.post("/api/auth/login", json={"email": "admin@sewpg.com", "password": "123456"})
        self.assertEqual(login.status_code, 200)
        self.headers = {"Authorization": f"Bearer {login.json()['token']}"}

        created = self.client.post(
            "/api/technical/projects",
            json={"name": "重新生成索引项目", "customerName": "测试业主"},
        )
        created.raise_for_status()
        self.project_id = created.json()["id"]

    def tearDown(self) -> None:
        self.client.close()
        self.temp_dir.cleanup()

    def test_status_defaults_to_idle_before_any_run(self) -> None:
        response = self.client.get(f"/api/technical/projects/{self.project_id}/score-index", headers=self.headers)

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["status"], "idle")
        self.assertEqual(payload["percentage"], 0)
        self.assertIsNone(payload["indexProgress"])

    def test_run_returns_running_state_immediately_and_records_audit(self) -> None:
        with patch("app.services.technical_score_index_flow._schedule_score_index_job"):
            response = self.client.post(
                f"/api/technical/projects/{self.project_id}/score-index/run",
                headers=self.headers,
            )

        self.assertEqual(response.status_code, 202)
        payload = response.json()
        self.assertEqual(payload["status"], "running")
        self.assertEqual(payload["percentage"], 3)
        self.assertEqual(payload["events"][0]["step"], "bootstrap")

        audit = self.client.get("/api/technical/audit", headers=self.headers)
        self.assertEqual(audit.status_code, 200)
        logs = [item for item in audit.json()["items"] if item["action"] == "开始重新生成章节索引"]
        self.assertGreaterEqual(len(logs), 1)

    def test_second_run_while_running_does_not_start_another_job(self) -> None:
        with patch("app.services.technical_score_index_flow._schedule_score_index_job") as schedule:
            self.client.post(f"/api/technical/projects/{self.project_id}/score-index/run", headers=self.headers)
            second = self.client.post(
                f"/api/technical/projects/{self.project_id}/score-index/run",
                headers=self.headers,
            )

        self.assertEqual(second.status_code, 202)
        self.assertEqual(second.json()["message"], "章节索引正在生成中，请稍候。")
        self.assertEqual(schedule.call_count, 1)

    def test_run_is_rejected_while_body_generation_is_running(self) -> None:
        project = store.require_project_for_update(self.project_id)
        project["fill_state"] = {"status": "running", "percentage": 30}
        store.persist_project_state(project)

        response = self.client.post(
            f"/api/technical/projects/{self.project_id}/score-index/run",
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 409)
        self.assertIn("正文正在重新生成", response.json()["detail"])

    def test_body_regeneration_is_rejected_while_score_index_is_running(self) -> None:
        """反向互斥：两条链路都以成稿为最终产物，并行会互相顶掉。"""
        with patch("app.services.technical_score_index_flow._schedule_score_index_job"):
            self.client.post(f"/api/technical/projects/{self.project_id}/score-index/run", headers=self.headers)

        response = self.client.post(
            f"/api/technical/projects/{self.project_id}/fill-generation/run",
            headers=self.headers,
        )

        self.assertEqual(response.status_code, 409)
        self.assertIn("章节索引正在重新生成", response.json()["detail"])

    def test_job_writes_quantified_progress_then_completed_state(self) -> None:
        from app.services import technical_score_index_flow as flow

        def fake_regenerate(project_id, progress_callback=None):
            progress_callback("calling_score_index_xref", {})
            progress_callback(
                "score_index_xref_probed",
                {"tableFound": True, "rowCount": 26, "pendingRowCount": 26},
            )
            progress_callback("score_index_xref_mapping_requested", {"pendingRowCount": 26})
            progress_callback(
                "score_index_xref_completed",
                {"summary": {"rowCount": 26, "filledRowCount": 26, "linkedCount": 78}},
            )
            return {
                "status": "completed",
                "applied": True,
                "runDurationSec": 42,
                "summary": {
                    "rowCount": 26,
                    "filledRowCount": 26,
                    "linkedCount": 78,
                    "unresolvedCount": 0,
                    "pageNumbersResolved": False,
                },
                "warnings": [],
            }

        with patch.object(flow, "regenerate_score_index_xref_for_project", fake_regenerate):
            flow.run_score_index_job(self.project_id, {}, None)

        state = _state_for_tests(self.project_id)
        self.assertEqual(state["status"], "completed")
        self.assertEqual(state["percentage"], 100)
        self.assertEqual(state["runDurationSec"], 42)
        self.assertEqual(state["indexProgress"]["done"], 26)
        self.assertEqual(state["indexProgress"]["total"], 26)
        self.assertTrue(state["output"]["applied"])
        self.assertFalse(state["output"]["pageNumbersResolved"])
        steps = [event["step"] for event in state["events"]]
        self.assertIn("score_index_probed", steps)
        self.assertIn("score_index_built", steps)

    def test_job_reports_skip_without_touching_the_document(self) -> None:
        from app.services import technical_score_index_flow as flow

        def fake_regenerate(project_id, progress_callback=None):
            progress_callback(
                "score_index_xref_probed",
                {"tableFound": False, "rowCount": 0, "pendingRowCount": 0},
            )
            return {"status": "skipped", "applied": False, "runDurationSec": 3, "summary": {}, "warnings": []}

        with patch.object(flow, "regenerate_score_index_xref_for_project", fake_regenerate):
            flow.run_score_index_job(self.project_id, {}, None)

        state = _state_for_tests(self.project_id)
        self.assertEqual(state["status"], "completed")
        self.assertFalse(state["output"]["applied"])
        self.assertIn("未找到技术评分标准索引表", state["summary"])

    def test_job_failure_is_surfaced_and_says_document_untouched(self) -> None:
        from app.services import technical_score_index_flow as flow

        def boom(project_id, progress_callback=None):
            raise ValueError("当前项目还没有可用的技术标成稿，请先生成正文。")

        with patch.object(flow, "regenerate_score_index_xref_for_project", boom):
            flow.run_score_index_job(self.project_id, {}, None)

        state = _state_for_tests(self.project_id)
        self.assertEqual(state["status"], "failed")
        self.assertIn("请先生成正文", state["summary"])
        self.assertEqual(state["events"][-1]["level"], "error")

    def test_real_regenerate_rewrites_draft_and_invalidates_onlyoffice_cache(self) -> None:
        """真跑一遍重新生成索引：成稿被换掉，且 OnlyOffice 的 documentKey 必须跟着变。

        OnlyOffice 按 documentKey 认文档，换了盘上的 docx 而 key 不变，
        编辑器会继续显示缓存里的旧稿——用户点了按钮却看不到变化。
        这条不变量没有任何机制强制执行，只能靠测试钉住。
        """
        from app.services import tech_assembly

        draft = document_path(self.project_id)
        _write_draft_with_index_table(draft)
        before = _document_state_for_tests(self.project_id)
        before_key = before["onlyoffice"]["documentKey"]
        before_bytes = draft.read_bytes()

        def fake_skill(brief_path: Path, mapping_path: Path):
            payload = json.loads(brief_path.read_text(encoding="utf-8"))
            factor = payload["rows"][0]["factor"]
            mapping_path.write_text(
                json.dumps({factor: ["5.3.3"]}, ensure_ascii=False), encoding="utf-8"
            )
            return {"mappingFile": str(mapping_path), "factorCount": 1}

        original = tech_assembly.run_technical_score_index_xref_skill
        tech_assembly.run_technical_score_index_xref_skill = fake_skill
        try:
            result = tech_assembly.regenerate_score_index_xref_for_project(self.project_id)
        finally:
            tech_assembly.run_technical_score_index_xref_skill = original

        self.assertEqual(result["status"], "completed")
        self.assertTrue(result["applied"])
        self.assertEqual(result["summary"]["filledRowCount"], 1)

        # 成稿确实被换掉了
        self.assertNotEqual(draft.read_bytes(), before_bytes)
        with ZipFile(draft) as archive:
            self.assertIn("5.3.3 叶片设计", archive.read("word/document.xml").decode("utf-8"))

        # 换了文件就必须让编辑器重新加载
        after = _document_state_for_tests(self.project_id)
        self.assertGreater(int(after["version"]), int(before["version"]))
        self.assertNotEqual(after["onlyoffice"]["documentKey"], before_key)
        self.assertTrue(after["lastSavedAt"])

    def test_real_regenerate_without_draft_fails_with_actionable_message(self) -> None:
        """还没生成正文就点按钮：要给出能照做的提示，而不是抛底层异常。"""
        from app.services import tech_assembly

        with self.assertRaises(ValueError) as ctx:
            tech_assembly.regenerate_score_index_xref_for_project(self.project_id)
        self.assertIn("请先生成正文", str(ctx.exception))

    def test_real_regenerate_leaves_draft_untouched_when_no_index_table(self) -> None:
        """成稿里没有评分索引表：如实报未改动，不能动 documentKey 白刷编辑器。"""
        from app.services import tech_assembly

        draft = document_path(self.project_id)
        doc = Document()
        doc.add_heading("5 投标技术方案", level=1)
        doc.add_paragraph("正文里没有评分索引表。")
        draft.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(draft))
        before = _document_state_for_tests(self.project_id)
        before_bytes = draft.read_bytes()

        result = tech_assembly.regenerate_score_index_xref_for_project(self.project_id)

        self.assertEqual(result["status"], "skipped")
        self.assertFalse(result["applied"])
        self.assertEqual(draft.read_bytes(), before_bytes)
        after = _document_state_for_tests(self.project_id)
        self.assertEqual(int(after["version"]), int(before["version"]))
        self.assertEqual(after["onlyoffice"]["documentKey"], before["onlyoffice"]["documentKey"])

    def test_running_percentage_never_goes_backwards(self) -> None:
        from app.services import technical_score_index_flow as flow

        flow._write_state(self.project_id, status="running", percentage=90)
        flow._write_state(self.project_id, percentage=25)

        self.assertEqual(_state_for_tests(self.project_id)["percentage"], 90)


if __name__ == "__main__":
    unittest.main()
