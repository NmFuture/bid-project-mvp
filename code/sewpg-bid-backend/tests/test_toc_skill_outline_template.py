from __future__ import annotations

import io
import json
import sys
import tempfile
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from toc_skill_helpers import (
    TocSkillScriptTestBase,
    json_load,
    load_outline_script,
    write_decision_context_fixture,
)


class TocSkillScriptTests(TocSkillScriptTestBase):

    def test_bid_outline_appendix_output_limit_rolls_back_state_and_can_retry(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _, manifest_path = write_decision_context_fixture(root, heading_count=0)

            decision_stdout = io.StringIO()
            with patch.object(
                sys,
                "argv",
                [
                    "run_from_manifest.py",
                    "decision-next",
                    str(manifest_path),
                ],
            ), redirect_stdout(decision_stdout):
                outline_runner.main()
            decision_batch = json.loads(decision_stdout.getvalue())

            decision_payload = json.dumps(
                {
                    "batch_token": decision_batch["batch_token"],
                    "items": [
                        {
                            "target_id": decision_batch["items"][0]["target_id"],
                            "decision": "retain",
                        }
                    ],
                    "additions": [],
                }
            )
            with patch.object(
                sys,
                "argv",
                [
                    "run_from_manifest.py",
                    "decision-batch",
                    str(manifest_path),
                    decision_payload,
                ],
            ), redirect_stdout(io.StringIO()):
                outline_runner.main()

            long_title = "超长技术响应附表" * 3000
            (root / "tender_appendix_inventory.json").write_text(
                json.dumps(
                    {
                        "schema_version": "tender-appendix-inventory.v1",
                        "items": [
                            {
                                "file_id": "TEN-1",
                                "number": f"附表A.{index}",
                                "title": long_title,
                                "following_table_count": 1,
                            }
                            for index in range(1, 3)
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            state_path = root / outline_runner.decision_workflow.STATE_FILE_NAME
            state_before = state_path.read_bytes()

            stdout = io.StringIO()
            with patch.object(
                sys,
                "argv",
                [
                    "run_from_manifest.py",
                    "appendix-next",
                    str(manifest_path),
                    "--max-items",
                    "2",
                ],
            ), redirect_stdout(stdout), self.assertRaisesRegex(
                SystemExit, r"command=appendix-next, actual_chars=\d+"
            ) as raised:
                outline_runner.main()

            self.assertEqual(stdout.getvalue(), "")
            self.assertEqual(state_path.read_bytes(), state_before)
            self.assertIn("--max-items", str(raised.exception))

            retry_stdout = io.StringIO()
            with patch.object(
                sys,
                "argv",
                [
                    "run_from_manifest.py",
                    "appendix-next",
                    str(manifest_path),
                    "--max-items",
                    "1",
                ],
            ), redirect_stdout(retry_stdout):
                outline_runner.main()
            retry = json.loads(retry_stdout.getvalue())

        self.assertEqual(len(retry["items"]), 1)
        self.assertEqual(retry["remaining_count"], 2)



    def test_bid_outline_template_structure_supplements_anchored_body_level_three(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            template = Path(tmp) / "template.docx"
            doc = Document()
            if "TOC 1" not in [style.name for style in doc.styles]:
                doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
            if "TOC 2" not in [style.name for style in doc.styles]:
                doc.styles.add_style("TOC 2", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("第1章 总体技术方案 ........ 1", style="TOC 1")
            doc.add_paragraph("1.1 机组选型 ........ 2", style="TOC 2")
            doc.add_paragraph("1.2 供货范围 ........ 3", style="TOC 2")
            doc.add_paragraph("第1章 总体技术方案", style="Heading 1")
            doc.add_paragraph("1.1 机组选型", style="Heading 2")
            doc.add_paragraph("1.1.1 关键部件选型", style="Heading 3")
            doc.add_paragraph("1.1.1.1 叶片设计参数", style="Heading 4")
            doc.add_paragraph("1.2 供货范围", style="Heading 2")
            doc.add_paragraph("1.2.1 主机供货范围", style="Heading 3")
            doc.add_paragraph("第9章 正文噪声", style="Heading 1")
            doc.add_paragraph("9.1 非模板章节", style="Heading 2")
            doc.add_paragraph("9.1.1 不应补充", style="Heading 3")
            doc.save(template)

            result = outline_runner.extract_template_structure(template)

        self.assertEqual(result["source"], "automatic_toc")
        self.assertEqual(
            [(item["number"], item["title"], item["level"]) for item in result["items"]],
            [
                ("第1章", "总体技术方案", 1),
                ("1.1", "机组选型", 2),
                ("1.1.1", "关键部件选型", 3),
                ("1.2", "供货范围", 2),
                ("1.2.1", "主机供货范围", 3),
            ],
        )



    def test_bid_outline_template_structure_numbers_blank_body_level_three_in_document_order(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            template = Path(tmp) / "template.docx"
            doc = Document()
            for style_name in ("TOC 1", "TOC 2"):
                if style_name not in [style.name for style in doc.styles]:
                    doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("\u7b2c1\u7ae0 Overview ........ 1", style="TOC 1")
            doc.add_paragraph("5.18 Digital Wind Farm ........ 2", style="TOC 2")
            doc.add_paragraph("\u7b2c1\u7ae0 Overview", style="Heading 1")
            doc.add_paragraph("5.18 Digital Wind Farm", style="Heading 2")
            doc.add_paragraph("SCADA System", style="Heading 3")
            doc.add_paragraph("Wind Farm Control", style="Heading 3")
            doc.save(template)

            result = outline_runner.extract_template_structure(template)

        self.assertEqual(
            [
                (item["number"], item["title"])
                for item in result["items"]
                if item["level"] == 3
            ],
            [
                ("5.18.1", "SCADA System"),
                ("5.18.2", "Wind Farm Control"),
            ],
        )



    def test_bid_outline_template_structure_deduplicates_existing_level_three(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            template = Path(tmp) / "template.docx"
            doc = Document()
            for style_name in ("TOC 1", "TOC 2", "TOC 3"):
                if style_name not in [style.name for style in doc.styles]:
                    doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("第1章 总体技术方案 ........ 1", style="TOC 1")
            doc.add_paragraph("1.1 机组选型 ........ 2", style="TOC 2")
            doc.add_paragraph("1.1.1 关键部件选型 ........ 3", style="TOC 3")
            doc.add_paragraph("总体技术方案", style="Heading 1")
            doc.add_paragraph("机组选型", style="Heading 2")
            doc.add_paragraph("关键部件选型", style="Heading 3")
            doc.save(template)

            result = outline_runner.extract_template_structure(template)

        self.assertEqual(
            [(item["number"], item["title"], item["level"]) for item in result["items"]],
            [
                ("第1章", "总体技术方案", 1),
                ("1.1", "机组选型", 2),
                ("1.1.1", "关键部件选型", 3),
            ],
        )



    def test_bid_outline_template_structure_preserves_same_title_with_different_numbers(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            template = Path(tmp) / "template.docx"
            doc = Document()
            for style_name in ("TOC 1", "TOC 2"):
                if style_name not in [style.name for style in doc.styles]:
                    doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("第1章 总体技术方案 ........ 1", style="TOC 1")
            doc.add_paragraph("1.1 认证情况 ........ 2", style="TOC 2")
            doc.add_paragraph("第1章 总体技术方案", style="Heading 1")
            doc.add_paragraph("1.1 认证情况", style="Heading 2")
            doc.add_paragraph("1.1.1 认证未完成或存在待解决项", style="Heading 3")
            doc.add_paragraph("1.1.2 认证未完成或存在待解决项", style="Heading 3")
            doc.save(template)

            result = outline_runner.extract_template_structure(template)

        self.assertEqual(
            [item["number"] for item in result["items"] if item["level"] == 3],
            ["1.1.1", "1.1.2"],
        )



    def test_bid_outline_template_structure_interleaves_supplemented_level_three(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            template = Path(tmp) / "template.docx"
            doc = Document()
            for style_name in ("TOC 1", "TOC 2", "TOC 3", "TOC 4"):
                if style_name not in [style.name for style in doc.styles]:
                    doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("第1章 总体技术方案 ........ 1", style="TOC 1")
            doc.add_paragraph("1.1 机组选型 ........ 2", style="TOC 2")
            doc.add_paragraph("1.1.1 叶片专题 ........ 3", style="TOC 3")
            doc.add_paragraph("1.1.1.1 叶片参数 ........ 4", style="TOC 4")
            doc.add_paragraph("1.1.3 齿轮箱专题 ........ 5", style="TOC 3")
            doc.add_paragraph("第1章 总体技术方案", style="Heading 1")
            doc.add_paragraph("1.1 机组选型", style="Heading 2")
            doc.add_paragraph("1.1.1 叶片专题", style="Heading 3")
            doc.add_paragraph("1.1.1.1 叶片参数", style="Heading 4")
            doc.add_paragraph("1.1.2 变桨系统专题", style="Heading 3")
            doc.add_paragraph("1.1.3 齿轮箱专题", style="Heading 3")
            doc.save(template)

            result = outline_runner.extract_template_structure(template)

        self.assertEqual(
            [(item["number"], item["level"]) for item in result["items"]],
            [
                ("第1章", 1),
                ("1.1", 2),
                ("1.1.1", 3),
                ("1.1.1.1", 4),
                ("1.1.2", 3),
                ("1.1.3", 3),
            ],
        )



    def test_bid_outline_template_command_writes_tender_appendix_heading_inventory(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"

            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)

            tender_doc = Document()
            tender_doc.add_paragraph("附表H.1 工程进度表 177")
            tender_doc.add_paragraph("附表H.2 交货进度表 178")
            tender_doc.add_paragraph("技术附表I 技术条款偏差表 199")
            tender_doc.add_paragraph("技术附表H 进度表")
            tender_doc.add_paragraph("附表H.1 工程进度表")
            tender_doc.add_table(rows=1, cols=2)
            tender_doc.add_paragraph("附表H.2 交货进度表")
            tender_doc.add_table(rows=2, cols=2)
            tender_doc.add_paragraph("技术附表I 技术条款偏差表")
            tender_doc.add_table(rows=1, cols=3)
            tender_doc.add_paragraph("附表F.5-2 认证未完成或存在待解决项")
            tender_doc.add_table(rows=1, cols=2)
            tender_doc.add_paragraph("投标人应逐项填写，不得遗漏。")
            tender_doc.save(tender)

            manifest = {
                "projectId": "PRJ-TEST",
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

            result = outline_runner.write_template_structure(manifest, manifest_path)
            structure = json_load(root / "template_structure.json")
            inventory = json_load(root / "tender_appendix_inventory.json")

        self.assertEqual(result["inputFingerprint"], structure["input_fingerprint"])
        self.assertRegex(result["inputFingerprint"], r"^[0-9a-f]{64}$")
        self.assertEqual(result["tenderAppendixItemCount"], 5)
        self.assertEqual(result["tenderAppendixInventoryFile"], str(root / "tender_appendix_inventory.json"))
        self.assertEqual(inventory["schema_version"], "tender-appendix-inventory.v1")
        self.assertEqual(
            [(item["number"], item["title"], item["file_id"]) for item in inventory["items"]],
            [
                ("技术附表H", "进度表", "TEN-1"),
                ("附表H.1", "工程进度表", "TEN-1"),
                ("附表H.2", "交货进度表", "TEN-1"),
                ("技术附表I", "技术条款偏差表", "TEN-1"),
                ("附表F.5-2", "认证未完成或存在待解决项", "TEN-1"),
            ],
        )
        self.assertEqual([item["following_table_count"] for item in inventory["items"]], [0, 1, 1, 1, 1])



    def test_bid_outline_template_structure_falls_back_to_toc_page_then_body_headings(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            toc_page = root / "toc-page.docx"
            toc_doc = Document()
            toc_doc.add_paragraph("目录")
            toc_doc.add_paragraph("第1章 总体技术方案 ........ 1")
            toc_doc.add_paragraph("1.1 机组选型 ........ 2")
            toc_doc.save(toc_page)

            body = root / "body.docx"
            body_doc = Document()
            body_doc.add_paragraph("第1章 总体技术方案", style="Heading 1")
            body_doc.add_paragraph("1.1 机组选型", style="Heading 2")
            body_doc.save(body)

            toc_result = outline_runner.extract_template_structure(toc_page)
            body_result = outline_runner.extract_template_structure(body)

        self.assertEqual(toc_result["source"], "toc_page")
        self.assertEqual(body_result["source"], "body_headings")
        self.assertEqual([item["level"] for item in body_result["items"]], [1, 2])
