from __future__ import annotations

import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document


BACKEND_ROOT = Path(__file__).resolve().parents[1]
SKILL_DIR = BACKEND_ROOT / "opencode" / "skills" / "bid-tech-score-index-xref"
RUNNER_PATH = SKILL_DIR / "scripts" / "run_from_manifest.py"


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


runner = _load("score_index_xref_runner_test", RUNNER_PATH)
xref = runner.xref_module


PLACEHOLDER = "[待人工补充：章节索引]"


def _build_docx(
    path: Path,
    *,
    index_entries: list[str] | None,
    with_table: bool = True,
    with_ordinal_headings: bool = False,
) -> None:
    """造一份「文本编号标题 + 评分索引表」的成稿，形状与格式清洗产物一致。"""
    doc = Document()
    doc.add_heading("5 投标技术方案", level=1)
    doc.add_paragraph("总述正文。")
    doc.add_heading("5.1 投标总体方案概述", level=2)
    doc.add_paragraph("概述正文。")
    doc.add_heading("5.3.3 叶片设计", level=3)
    doc.add_paragraph("叶片正文。")
    doc.add_heading("附表2 发电量计算表", level=1)
    doc.add_paragraph("附表正文。")
    if with_ordinal_headings:
        # 素材自带的段内列表序号，正文里会大量重复
        doc.add_heading("1、 与华能集团签署的战略合作协议", level=2)
        doc.add_paragraph("协议正文。")
        doc.add_heading("1、 华能汕头勒门海上机组国产化示范", level=2)
        doc.add_paragraph("示范正文。")

    if with_table:
        rows = 1 + len(index_entries or [""])
        table = doc.add_table(rows=rows, cols=3)
        table.cell(0, 0).text = "序号"
        table.cell(0, 1).text = "评审因素"
        table.cell(0, 2).text = "章节索引"
        for offset, entry in enumerate(index_entries or [""], start=1):
            table.cell(offset, 0).text = str(offset)
            table.cell(offset, 1).text = "风轮系统先进性及可靠性"
            table.cell(offset, 2).text = entry
    doc.save(str(path))


def _document_xml(path: Path) -> str:
    with ZipFile(path) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def _write_manifest(path: Path, **fields) -> Path:
    path.write_text(json.dumps(fields, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


class ScoreIndexXrefTest(unittest.TestCase):
    def test_builds_hyperlink_and_pageref_for_filled_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=["5.1 投标总体方案概述", "5.3.3 叶片设计"])
            manifest = _write_manifest(
                work / "manifest.json",
                inputFile=str(source),
                outputFile=str(output),
            )

            result = runner.run_manifest(manifest)

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["summary"]["linkedCount"], 2)
            self.assertEqual(result["summary"]["unresolvedCount"], 0)
            self.assertEqual(result["summary"]["pagerefCount"], 2)
            self.assertTrue(result["summary"]["bookmarksResolvable"])
            self.assertTrue(output.exists())

            xml = _document_xml(output)
            self.assertIn("PAGEREF _Xref_", xml)
            self.assertIn('w:anchor="_Xref_0001"', xml)
            # 未算页码时留占位符，并明确提示需要 F9
            codes = {item["code"] for item in result["warnings"]}
            self.assertIn("xref_page_number_pending", codes)

    def test_matches_bare_chapter_number_after_format_cleaning(self) -> None:
        """格式清洗后标题是「文本编号 + Heading 样式」，映射只写章节号也要能命中。"""
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=[""])
            mapping = work / "mapping.json"
            mapping.write_text(
                json.dumps({"风轮系统先进性及可靠性": ["5.1", "5.3.3", "附表2"]}, ensure_ascii=False),
                encoding="utf-8",
            )
            manifest = _write_manifest(
                work / "manifest.json",
                inputFile=str(source),
                outputFile=str(output),
                mappingFile=str(mapping),
            )

            result = runner.run_manifest(manifest)

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["summary"]["filledRowCount"], 1)
            self.assertEqual(result["summary"]["linkedCount"], 3)
            self.assertEqual(result["summary"]["unresolvedCount"], 0)

    def test_skips_when_no_score_index_table(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=None, with_table=False)
            manifest = _write_manifest(
                work / "manifest.json",
                inputFile=str(source),
                outputFile=str(output),
            )

            result = runner.run_manifest(manifest)

            self.assertEqual(result["status"], "skipped")
            self.assertEqual(result["outputFile"], "")
            self.assertFalse(output.exists())
            self.assertEqual([item["code"] for item in result["warnings"]], ["index_table_not_found"])

    def test_unresolved_entry_is_reported_and_left_untouched(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=["9.9 并不存在的章节"])
            manifest = _write_manifest(
                work / "manifest.json",
                inputFile=str(source),
                outputFile=str(output),
            )

            result = runner.run_manifest(manifest)

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["summary"]["linkedCount"], 0)
            self.assertEqual(result["summary"]["unresolvedCount"], 1)
            codes = {item["code"] for item in result["warnings"]}
            self.assertIn("xref_unresolved_entry", codes)
            texts = [p.text for row in Document(str(output)).tables[0].rows for p in row.cells[2].paragraphs]
            self.assertIn("9.9 并不存在的章节", texts)

    def test_rerun_is_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            first = work / "一次.docx"
            second = work / "二次.docx"
            _build_docx(source, index_entries=["5.1 投标总体方案概述"])
            runner.run_manifest(
                _write_manifest(work / "m1.json", inputFile=str(source), outputFile=str(first))
            )
            result = runner.run_manifest(
                _write_manifest(work / "m2.json", inputFile=str(first), outputFile=str(second))
            )

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["summary"]["linkedCount"], 1)
            self.assertEqual(result["summary"]["pagerefCount"], 1)
            # 复用已有书签，不新增
            self.assertEqual(result["summary"]["bookmarksCreated"], 0)
            cell_text = Document(str(second)).tables[0].rows[1].cells[2].text
            self.assertEqual(cell_text.count("，P"), 1)

    def test_mapping_with_unknown_section_never_writes_a_wrong_number(self) -> None:
        """映射里有正文不存在的章节号：不写错号，退回无映射重建并显式报出。"""
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=[PLACEHOLDER])
            mapping = work / "mapping.json"
            mapping.write_text(json.dumps({"风轮系统先进性及可靠性": ["8.8"]}, ensure_ascii=False), encoding="utf-8")

            result = runner.run_manifest(
                _write_manifest(
                    work / "manifest.json",
                    inputFile=str(source),
                    outputFile=str(output),
                    mappingFile=str(mapping),
                )
            )

            self.assertEqual(result["status"], "completed")
            codes = {item["code"] for item in result["warnings"]}
            self.assertIn("mapping_unresolved", codes)
            self.assertEqual(result["summary"]["filledRowCount"], 0)
            self.assertEqual(result["summary"]["linkedCount"], 0)
            cell = Document(str(output)).tables[0].rows[1].cells[2].text
            self.assertIn("待人工补充", cell)

    def test_placeholder_column_is_pending_not_unresolved(self) -> None:
        """S3 留下的 `[待人工补充：章节索引]` 是待填标记，不是定位不到的章节。"""
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=[PLACEHOLDER, PLACEHOLDER])

            result = runner.run_manifest(
                _write_manifest(work / "manifest.json", inputFile=str(source), outputFile=str(output))
            )

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["summary"]["placeholderCount"], 2)
            self.assertEqual(result["summary"]["unresolvedCount"], 0)
            codes = {item["code"] for item in result["warnings"]}
            self.assertIn("xref_index_column_pending", codes)
            self.assertNotIn("xref_unresolved_entry", codes)

    def test_mapping_fills_over_placeholder_without_overwrite(self) -> None:
        """占位符视同空格子，不开 overwrite 也要能填进去。"""
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=[PLACEHOLDER])
            mapping = work / "mapping.json"
            mapping.write_text(
                json.dumps({"风轮系统先进性及可靠性": ["5.1", "5.3.3"]}, ensure_ascii=False), encoding="utf-8"
            )

            result = runner.run_manifest(
                _write_manifest(
                    work / "manifest.json",
                    inputFile=str(source),
                    outputFile=str(output),
                    mappingFile=str(mapping),
                )
            )

            self.assertEqual(result["summary"]["filledRowCount"], 1)
            self.assertEqual(result["summary"]["linkedCount"], 2)
            self.assertEqual(result["summary"]["placeholderCount"], 0)
            cell = Document(str(output)).tables[0].rows[1].cells[2].text
            self.assertNotIn("待人工补充", cell)

    def test_list_ordinal_is_not_treated_as_chapter_number(self) -> None:
        """"1、xxx" 是段内列表序号；当成章节号会让映射写 "1" 时静默链到错章节。"""
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            _build_docx(source, index_entries=[PLACEHOLDER], with_ordinal_headings=True)
            doc, _, _, _, _, headings = xref.prepare(
                str(source), xref.DEFAULT_INDEX_HEADERS, xref.DEFAULT_FACTOR_HEADERS
            )

            ordinal = [h for h in headings if h.title.startswith("与华能") or h.title.startswith("华能汕头")]
            self.assertEqual(len(ordinal), 0, "列表序号标题不应被拆出 number")
            # 附表类标题不以数字开头，本就没有 number，靠标题前缀匹配
            numbers = [h.number for h in headings if h.number]
            self.assertEqual(sorted(numbers), ["5", "5.1", "5.3.3"])
            # 撞号的编号不进 by_num，映射引用它时显式失败而不是链到第一个
            by_num, _, _ = xref.build_lookups(headings)
            self.assertNotIn("1", by_num)

    def test_ambiguous_chapter_number_fails_instead_of_linking_wrong_section(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=[PLACEHOLDER], with_ordinal_headings=True)
            mapping = work / "mapping.json"
            mapping.write_text(json.dumps({"风轮系统先进性及可靠性": ["1"]}, ensure_ascii=False), encoding="utf-8")

            result = runner.run_manifest(
                _write_manifest(
                    work / "manifest.json",
                    inputFile=str(source),
                    outputFile=str(output),
                    mappingFile=str(mapping),
                )
            )

            # 退回无映射重建：成品仍在，但明确报出映射不可用与该列待填
            self.assertEqual(result["status"], "completed")
            codes = {item["code"] for item in result["warnings"]}
            self.assertIn("mapping_unresolved", codes)
            self.assertIn("xref_index_column_pending", codes)
            self.assertEqual(result["summary"]["linkedCount"], 0)
            self.assertTrue(output.exists())

    def test_inspect_mode_reports_pending_rows_and_writes_brief(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            brief = work / "brief.json"
            _build_docx(source, index_entries=[PLACEHOLDER, "5.1 投标总体方案概述"])

            result = runner.run_manifest(
                _write_manifest(
                    work / "manifest.json",
                    mode="inspect",
                    inputFile=str(source),
                    outDir=str(work),
                    briefFile=str(brief),
                )
            )

            self.assertEqual(result["status"], "completed")
            self.assertTrue(result["tableFound"])
            self.assertEqual(result["rowCount"], 2)
            self.assertEqual(result["pendingRowCount"], 1)
            payload = json.loads(brief.read_text(encoding="utf-8"))
            self.assertEqual(payload["rows"][0]["factor"], "风轮系统先进性及可靠性")
            self.assertTrue(payload["rows"][0]["pending"])
            self.assertFalse(payload["rows"][1]["pending"])
            self.assertIn({"number": "5.1", "title": "投标总体方案概述", "level": 2}, payload["headings"])

    def test_inspect_mode_reports_missing_table_without_raising(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            _build_docx(source, index_entries=None, with_table=False)

            result = runner.run_manifest(
                _write_manifest(work / "manifest.json", mode="inspect", inputFile=str(source), outDir=str(work))
            )

            self.assertEqual(result["status"], "skipped")
            self.assertFalse(result["tableFound"])

    def test_rejects_output_equal_to_input(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            _build_docx(source, index_entries=["5.1 投标总体方案概述"])
            with self.assertRaises(ValueError):
                runner.run_manifest(
                    _write_manifest(work / "manifest.json", inputFile=str(source), outputFile=str(source))
                )


class ScoreIndexXrefPipelineTest(unittest.TestCase):
    def test_pipeline_step_falls_back_to_cleaner_output_on_failure(self) -> None:
        from app.services import tech_assembly

        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            missing = work / "不存在.docx"
            result = tech_assembly._run_tech_score_index_xref_step(
                input_path=missing,
                output_path=work / "out.docx",
                work_dir=work,
            )

            self.assertEqual(result["status"], "failed")
            self.assertEqual(result["outputFile"], "")
            self.assertEqual([item["code"] for item in result["warnings"]], ["score_index_xref_failed"])

    def test_pipeline_step_reports_skip_without_output(self) -> None:
        from app.services import tech_assembly

        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            _build_docx(source, index_entries=None, with_table=False)
            result = tech_assembly._run_tech_score_index_xref_step(
                input_path=source,
                output_path=work / "成稿_xref.docx",
                work_dir=work,
            )

            self.assertEqual(result["status"], "skipped")
            self.assertEqual(result["outputFile"], "")

    def test_pipeline_asks_agent_for_mapping_when_column_is_pending(self) -> None:
        from app.services import tech_assembly

        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            _build_docx(source, index_entries=[PLACEHOLDER])
            briefs: list[Path] = []

            def fake_skill(brief_path: Path, mapping_path: Path):
                briefs.append(brief_path)
                payload = json.loads(brief_path.read_text(encoding="utf-8"))
                factor = payload["rows"][0]["factor"]
                mapping_path.write_text(
                    json.dumps({factor: ["5.1", "5.3.3"]}, ensure_ascii=False), encoding="utf-8"
                )
                return {"mappingFile": str(mapping_path), "factorCount": 1}

            stages: list[str] = []
            original = tech_assembly.run_technical_score_index_xref_skill
            tech_assembly.run_technical_score_index_xref_skill = fake_skill
            try:
                result = tech_assembly._run_tech_score_index_xref_step(
                    input_path=source,
                    output_path=work / "成稿_xref.docx",
                    work_dir=work,
                    progress_callback=lambda stage, _meta: stages.append(stage),
                )
            finally:
                tech_assembly.run_technical_score_index_xref_skill = original

            self.assertEqual(len(briefs), 1)
            self.assertIn("score_index_xref_mapping_requested", stages)
            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["mappingSource"], "opencode")
            self.assertEqual(result["summary"]["linkedCount"], 2)
            self.assertEqual(result["summary"]["placeholderCount"], 0)

    def test_pipeline_keeps_going_when_agent_unavailable(self) -> None:
        from app.services import tech_assembly

        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=[PLACEHOLDER])

            def boom(brief_path: Path, mapping_path: Path):
                raise RuntimeError("opencode 不可用")

            original = tech_assembly.run_technical_score_index_xref_skill
            tech_assembly.run_technical_score_index_xref_skill = boom
            try:
                result = tech_assembly._run_tech_score_index_xref_step(
                    input_path=source,
                    output_path=output,
                    work_dir=work,
                )
            finally:
                tech_assembly.run_technical_score_index_xref_skill = original

            # agent 挂了不阻断出稿，但必须把「这一列没填」说出来
            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["mappingSource"], "")
            codes = {item["code"] for item in result["warnings"]}
            self.assertIn("score_index_xref_mapping_unavailable", codes)
            self.assertIn("xref_index_column_pending", codes)
            self.assertTrue(output.exists())

    def test_prefilled_column_is_rebuilt_against_the_current_document(self) -> None:
        """素材模板自带的章节索引指向上一版章节号，组装正文时必须按当前文档重判并覆盖。

        旧行为是「列非空就跳过 agent」，素材里的旧号会原样留到成稿里链到错章节。
        """
        from app.services import tech_assembly

        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=["5.1 投标总体方案概述"])
            pending_counts: list[int] = []

            def fake_skill(brief_path: Path, mapping_path: Path):
                payload = json.loads(brief_path.read_text(encoding="utf-8"))
                pending_counts.append(int(payload["pendingRowCount"]))
                factor = payload["rows"][0]["factor"]
                mapping_path.write_text(
                    json.dumps({factor: ["5.3.3"]}, ensure_ascii=False), encoding="utf-8"
                )
                return {"mappingFile": str(mapping_path), "factorCount": 1}

            original = tech_assembly.run_technical_score_index_xref_skill
            tech_assembly.run_technical_score_index_xref_skill = fake_skill
            try:
                result = tech_assembly._run_tech_score_index_xref_step(
                    input_path=source,
                    output_path=output,
                    work_dir=work,
                )
            finally:
                tech_assembly.run_technical_score_index_xref_skill = original

            # 所有行都被重新标成待判断，agent 才会重做已经填过的这一行
            self.assertEqual(pending_counts, [1])
            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["summary"]["filledRowCount"], 1)
            # 覆盖生效：旧的 5.1 被换成按当前文档判断出来的 5.3.3
            xml = _document_xml(output)
            self.assertIn("5.3.3 叶片设计", xml)
            self.assertNotIn("5.1 投标总体方案概述", xml.split("<w:tbl>")[-1])

    def test_stale_mapping_from_a_previous_run_is_not_reused(self) -> None:
        """工作目录里留着上一轮的映射时也要重判，否则重新生成索引会原样复现旧结果。"""
        from app.services import tech_assembly

        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=[PLACEHOLDER])
            stale = work / "tech_score_index_xref_mapping.json"
            stale.write_text(
                json.dumps({"风轮系统先进性及可靠性": ["5.1"]}, ensure_ascii=False), encoding="utf-8"
            )

            def fake_skill(brief_path: Path, mapping_path: Path):
                payload = json.loads(brief_path.read_text(encoding="utf-8"))
                factor = payload["rows"][0]["factor"]
                mapping_path.write_text(
                    json.dumps({factor: ["5.3.3"]}, ensure_ascii=False), encoding="utf-8"
                )
                return {"mappingFile": str(mapping_path), "factorCount": 1}

            original = tech_assembly.run_technical_score_index_xref_skill
            tech_assembly.run_technical_score_index_xref_skill = fake_skill
            try:
                tech_assembly._run_tech_score_index_xref_step(
                    input_path=source,
                    output_path=output,
                    work_dir=work,
                )
            finally:
                tech_assembly.run_technical_score_index_xref_skill = original

            xml = _document_xml(output)
            self.assertIn("5.3.3 叶片设计", xml)
            self.assertNotIn("5.1 投标总体方案概述", xml.split("<w:tbl>")[-1])

    def test_second_run_redecides_over_previously_generated_links(self) -> None:
        """「重新生成索引」的真实场景：列里已经是上一轮生成的超链接 + PAGEREF 域，仍要重判并换掉。"""
        from app.services import tech_assembly

        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            first_out = work / "成稿_xref1.docx"
            second_out = work / "成稿_xref2.docx"
            _build_docx(source, index_entries=[PLACEHOLDER])
            answers = [["5.1"], ["5.3.3"]]

            def fake_skill(brief_path: Path, mapping_path: Path):
                payload = json.loads(brief_path.read_text(encoding="utf-8"))
                factor = payload["rows"][0]["factor"]
                mapping_path.write_text(
                    json.dumps({factor: answers.pop(0)}, ensure_ascii=False), encoding="utf-8"
                )
                return {"mappingFile": str(mapping_path), "factorCount": 1}

            original = tech_assembly.run_technical_score_index_xref_skill
            tech_assembly.run_technical_score_index_xref_skill = fake_skill
            try:
                first = tech_assembly._run_tech_score_index_xref_step(
                    input_path=source, output_path=first_out, work_dir=work,
                )
                # 第一轮产物就是第二轮的输入，等同于共创导出页上点「重新生成索引」
                second = tech_assembly._run_tech_score_index_xref_step(
                    input_path=first_out, output_path=second_out, work_dir=work,
                )
            finally:
                tech_assembly.run_technical_score_index_xref_skill = original

            self.assertEqual(first["summary"]["linkedCount"], 1)
            self.assertEqual(second["status"], "completed")
            self.assertEqual(second["summary"]["filledRowCount"], 1)
            self.assertEqual(second["summary"]["linkedCount"], 1)
            # 两次判断都被采纳，且第二轮把第一轮的结果换掉而不是并存
            first_xml = _document_xml(first_out)
            second_xml = _document_xml(second_out)
            self.assertIn("5.1 投标总体方案概述", first_xml)
            self.assertIn("5.3.3 叶片设计", second_xml)
            second_table = second_xml.split("<w:tbl>")[-1]
            self.assertNotIn("5.1 投标总体方案概述", second_table)
            # 旧的页码尾巴没有被累加成「，P1，P1」
            self.assertEqual(second_table.count("PAGEREF"), 1)

    def test_existing_index_is_kept_when_agent_returns_nothing(self) -> None:
        """覆盖只作用在拿到映射的行；agent 不可用时原有索引不能被清空。"""
        from app.services import tech_assembly

        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            output = work / "成稿_xref.docx"
            _build_docx(source, index_entries=["5.1 投标总体方案概述"])

            def boom(brief_path: Path, mapping_path: Path):
                raise RuntimeError("opencode 不可用")

            original = tech_assembly.run_technical_score_index_xref_skill
            tech_assembly.run_technical_score_index_xref_skill = boom
            try:
                result = tech_assembly._run_tech_score_index_xref_step(
                    input_path=source,
                    output_path=output,
                    work_dir=work,
                )
            finally:
                tech_assembly.run_technical_score_index_xref_skill = original

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["summary"]["linkedCount"], 1)
            self.assertIn("5.1 投标总体方案概述", _document_xml(output))

    def test_probe_stage_reports_row_counts_for_progress(self) -> None:
        """进度条的量化指标来自体检回调，缺了它前端只能显示百分比。"""
        from app.services import tech_assembly

        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            source = work / "成稿.docx"
            _build_docx(source, index_entries=[PLACEHOLDER, PLACEHOLDER])
            probed: list[dict] = []

            def record(stage: str, meta=None):
                if stage == "score_index_xref_probed":
                    probed.append(dict(meta or {}))

            original = tech_assembly.run_technical_score_index_xref_skill
            tech_assembly.run_technical_score_index_xref_skill = lambda *_args: {}
            try:
                tech_assembly._run_tech_score_index_xref_step(
                    input_path=source,
                    output_path=work / "成稿_xref.docx",
                    work_dir=work,
                    progress_callback=record,
                )
            finally:
                tech_assembly.run_technical_score_index_xref_skill = original

            self.assertEqual(len(probed), 1)
            self.assertTrue(probed[0]["tableFound"])
            self.assertEqual(probed[0]["rowCount"], 2)
            self.assertEqual(probed[0]["pendingRowCount"], 2)


if __name__ == "__main__":
    unittest.main()
