"""附表 AI 填写取值来源优先级（L1 事实表 > L2 规则/人工素材 > L3 其他）单测。

覆盖：
1. prepare：事实表字段命中时 targetField 带 preferredRoute/preferredValue
   （含概念同义词命中），不匹配或同分歧义时不带（保守，宁缺勿滥）。
2. materials tier 标注：人工指定/招标文件/规则命中素材 = 2，索引/推荐 = 3。
3. apply：无文件路由（factTable/parseFields/projectTurbineModel）不再要求
   excerpt，改值一致性校验——值一致无 excerpt 也过，不一致降级 manual；
   preferredValue 偏离（素材路由填了与事实表矛盾的值）降级 manual；
   素材路由 excerpt 逐字校验行为不变。
"""
from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

_SRC = (
    Path(__file__).resolve().parents[1]
    / "opencode"
    / "skills"
    / "bid-tech-table-filler"
    / "scripts"
    / "run_from_manifest.py"
)
_SPEC = importlib.util.spec_from_file_location("tech_table_filler_priority_under_test", _SRC)
filler = importlib.util.module_from_spec(_SPEC)
assert _SPEC.loader is not None
sys.modules["tech_table_filler_priority_under_test"] = filler
_SPEC.loader.exec_module(filler)


def _build_docx(path: Path, heading: str, rows: list[list[str]]) -> Path:
    doc = Document()
    doc.add_paragraph(heading)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            table.rows[ri].cells[ci].text = value
    doc.save(str(path))
    return path


def _build_text_docx(path: Path, paragraphs: list[str]) -> Path:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(str(path))
    return path


# 参数表：编号 + 项目 + 主要项目 + 技术参数与规格（响应列）+ 计量单位 + 备注
_PARAM_ROWS = [
    ["编号", "项目", "主要项目", "技术参数与规格", "计量单位", "备注"],
    ["1", "机型总体参数", "投标机型", "", "", ""],
    ["2", "机型总体参数", "单机容量", "", "MW", ""],
    ["3", "机型总体参数", "叶轮直径", "", "m", ""],
]

_MATERIAL_PARAGRAPHS = [
    "机型参数说明",
    "投标机型为 EW6.25-220 机组。",
    "单机容量 7000 kW，叶轮直径 220 m。",
]


class PriorityTestBase(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.base = Path(self._tmp.name)
        self.blank = _build_docx(self.base / "blank.docx", "附表C.1 机型总体参数与规格", _PARAM_ROWS)
        self.material = _build_text_docx(self.base / "material.docx", _MATERIAL_PARAGRAPHS)
        self.tender = _build_text_docx(self.base / "tender.docx", ["招标文件正文。"])
        self.indexed = _build_text_docx(self.base / "indexed.docx", ["索引素材。"])
        self.output = self.base / "out" / "filled.docx"

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def _write_manifest(self, **overrides) -> Path:
        manifest = {
            "blankSource": {"docxPath": str(self.blank), "title": "附表C.1 机型总体参数与规格", "id": "APPX-TEST"},
            "outputFile": str(self.output),
            "referenceMaterials": [{"id": "RAW-1", "name": "机型参数材料.docx", "path": str(self.material)}],
            "projectFactTable": {
                "status": "confirmed",
                "fields": [
                    {"label": "单机容量", "value": "6.25", "unit": "MW", "status": "confirmed"},
                ],
            },
            "parseFields": [{"label": "叶轮直径", "value": "220", "unit": "m"}],
            "projectTurbineModel": {"model": "EW6.25-220", "ratedPowerKw": 6250},
        }
        manifest.update(overrides)
        path = self.base / "manifest.json"
        path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        return path

    def _write_plan(self, fills: list[dict]) -> Path:
        plan = {"schemaVersion": "bid-tech-table-fill-plan-v1", "fills": fills}
        path = self.base / "fill_plan.json"
        path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        return path

    def _fill(self, target_field_id: str, field: str, value: str, excerpt: str = "", **extra) -> dict:
        fill = {
            "targetFieldId": target_field_id,
            "field": field,
            "action": "fill",
            "value": value,
            "unit": "",
            "confidence": 0.9,
            "evidence": {"sourceRoute": "referenceMaterial", "sourcePath": str(self.material), "excerpt": excerpt},
            "reason": "测试取值",
        }
        fill.update(extra)
        return fill


class PreparePreferredBindingTests(PriorityTestBase):
    def _brief(self, manifest_path: Path) -> dict:
        filler.run_prepare(manifest_path)
        return json.loads((self.base / "fill_brief.json").read_text(encoding="utf-8"))

    def test_fact_table_hit_binds_preferred_value(self) -> None:
        brief = self._brief(self._write_manifest())
        by_field = {f["field"]: f for f in brief["targetFields"]}
        capacity = by_field["单机容量"]
        self.assertEqual(capacity["preferredRoute"], "factTable")
        self.assertEqual(capacity["preferredValue"], "6.25")
        self.assertEqual(capacity["preferredUnit"], "MW")
        self.assertEqual(capacity["preferredLabel"], "单机容量")

    def test_concept_synonym_hit_binds(self) -> None:
        # 事实表标签「额定功率」与字段「单机容量」同属 rated_power 概念 → 绑定
        manifest_path = self._write_manifest(
            projectFactTable={
                "status": "confirmed",
                "fields": [{"label": "额定功率", "value": "6.25", "unit": "MW", "status": "confirmed"}],
            }
        )
        brief = self._brief(manifest_path)
        by_field = {f["field"]: f for f in brief["targetFields"]}
        self.assertEqual(by_field["单机容量"].get("preferredValue"), "6.25")

    def test_no_match_no_binding(self) -> None:
        brief = self._brief(self._write_manifest())
        by_field = {f["field"]: f for f in brief["targetFields"]}
        # 事实表里没有投标机型/叶轮直径对应字段 → 不绑定（保守）
        for name in ("投标机型", "叶轮直径"):
            self.assertNotIn("preferredRoute", by_field[name])

    def test_ambiguous_tie_no_binding(self) -> None:
        # 同概念两个不同事实同分 → 宁缺勿滥不绑定
        manifest_path = self._write_manifest(
            projectFactTable={
                "status": "confirmed",
                "fields": [
                    {"label": "额定功率", "value": "6.25", "unit": "MW", "status": "confirmed"},
                    {"label": "单机功率", "value": "7.0", "unit": "MW", "status": "extracted"},
                ],
            }
        )
        brief = self._brief(manifest_path)
        by_field = {f["field"]: f for f in brief["targetFields"]}
        self.assertNotIn("preferredRoute", by_field["单机容量"])

    def test_pending_fact_not_bound(self) -> None:
        # pending_confirmation 的事实不进 brief，也不参与绑定
        manifest_path = self._write_manifest(
            projectFactTable={
                "status": "confirmed",
                "fields": [{"label": "单机容量", "value": "6.25", "unit": "MW", "status": "pending_confirmation"}],
            }
        )
        brief = self._brief(manifest_path)
        by_field = {f["field"]: f for f in brief["targetFields"]}
        self.assertNotIn("preferredRoute", by_field["单机容量"])

    def test_qualifier_field_does_not_poison_concept_match(self) -> None:
        # 回归（PRJ-0001 实测）：「单机功率曲线考核阈值」「招标单机容量（出口端，MW）」
        # 这类限定字段的子串命中曾在 rated_power 里制造伪歧义，导致「机组功率」
        # 与「单机容量」该绑的绑不上。概念命中改为别名精确等价后应唯一绑定。
        manifest_path = self._write_manifest(
            projectFactTable={
                "status": "confirmed",
                "fields": [
                    {"label": "单机容量", "value": "10", "unit": "MW", "status": "confirmed"},
                    {"label": "招标单机容量（出口端，MW）", "value": "10", "status": "confirmed"},
                    {"label": "单机功率曲线考核阈值", "value": "97", "status": "confirmed"},
                ],
            }
        )
        brief = self._brief(manifest_path)
        by_field = {f["field"]: f for f in brief["targetFields"]}
        self.assertEqual(by_field["单机容量"].get("preferredValue"), "10")
        self.assertEqual(by_field["单机容量"].get("preferredLabel"), "单机容量")

    def test_parenthetical_label_variant_binds(self) -> None:
        # 「额定功率（MW）」这类带单位注释的事实表标签，去括号后与别名精确等价 → 绑定
        manifest_path = self._write_manifest(
            projectFactTable={
                "status": "confirmed",
                "fields": [{"label": "额定功率（MW）", "value": "6.25", "status": "confirmed"}],
            }
        )
        brief = self._brief(manifest_path)
        by_field = {f["field"]: f for f in brief["targetFields"]}
        self.assertEqual(by_field["单机容量"].get("preferredValue"), "6.25")


class PrepareMaterialTierTests(PriorityTestBase):
    def test_tier_annotations(self) -> None:
        manifest_path = self._write_manifest(
            tenderDocuments=[{"id": "T-1", "name": "招标文件.docx", "path": str(self.tender)}],
            materialIndex=[{"id": "IDX-1", "name": "索引素材.docx", "path": str(self.indexed)}],
            recommendedMaterials=[
                {"id": "REC-1", "name": "推荐素材.docx", "path": str(self.indexed)},
                {
                    "id": "REC-2",
                    "name": "规则命中素材.docx",
                    "path": str(self.indexed),
                    "sourceRouting": {"source": "appendix_source_matrix", "ruleId": "Sheet1!R40"},
                },
            ],
        )
        filler.run_prepare(manifest_path)
        brief = json.loads((self.base / "fill_brief.json").read_text(encoding="utf-8"))
        tiers = {(m["route"], m["id"]): m["tier"] for m in brief["materials"]}
        # 人工最终指定 / 规则要求给入的招标文件 / 规则命中推荐素材 = L2
        self.assertEqual(tiers[("referenceMaterial", "RAW-1")], 2)
        self.assertEqual(tiers[("tenderDocument", "T-1")], 2)
        self.assertEqual(tiers[("recommendedMaterials", "REC-2")], 2)
        # 素材索引 / 普通推荐 = L3
        self.assertEqual(tiers[("materialIndex", "IDX-1")], 3)
        self.assertEqual(tiers[("recommendedMaterials", "REC-1")], 3)


class ApplyFilelessRouteTests(PriorityTestBase):
    def test_fact_table_route_consistent_without_excerpt(self) -> None:
        manifest_path = self._write_manifest()
        fill = self._fill(
            "C1-R02", "单机容量", "6250",
            unit="kW",
            evidence={"sourceRoute": "factTable"},
        )
        self._write_plan([fill])
        result = filler.run_apply(manifest_path)
        decision = result["mapping"]["decisions"][1]
        self.assertEqual(decision["action"], "fill")
        self.assertEqual(decision["value"], "6.25")
        row = Document(str(self.output)).tables[0].rows[2]
        self.assertEqual(row.cells[3].text.strip(), "6.25")

    def test_fact_table_route_inconsistent_downgrades(self) -> None:
        manifest_path = self._write_manifest()
        fill = self._fill("C1-R02", "单机容量", "7.0", unit="MW", evidence={"sourceRoute": "factTable"})
        self._write_plan([fill])
        result = filler.run_apply(manifest_path)
        decision = result["mapping"]["decisions"][1]
        self.assertEqual(decision["action"], "manual")
        self.assertIn("与事实表值不一致", decision["reason"])

    def test_fact_table_route_unknown_field_downgrades(self) -> None:
        # 事实表没有「叶轮直径」字段，无法核对 → 降级而不是放行
        manifest_path = self._write_manifest()
        fill = self._fill("C1-R03", "叶轮直径", "220", unit="m", evidence={"sourceRoute": "factTable"})
        self._write_plan([fill])
        result = filler.run_apply(manifest_path)
        decision = result["mapping"]["decisions"][2]
        self.assertEqual(decision["action"], "manual")
        self.assertIn("事实表中找不到", decision["reason"])

    def test_parse_fields_route_value_consistency(self) -> None:
        manifest_path = self._write_manifest()
        ok = self._fill("C1-R03", "叶轮直径", "220", unit="m", evidence={"sourceRoute": "parseFields"})
        self._write_plan([ok])
        result = filler.run_apply(manifest_path)
        self.assertEqual(result["mapping"]["decisions"][2]["action"], "fill")

    def test_parse_fields_route_inconsistent_downgrades(self) -> None:
        manifest_path = self._write_manifest()
        bad = self._fill("C1-R03", "叶轮直径", "250", unit="m", evidence={"sourceRoute": "parseFields"})
        self._write_plan([bad])
        result = filler.run_apply(manifest_path)
        decision = result["mapping"]["decisions"][2]
        self.assertEqual(decision["action"], "manual")
        self.assertIn("与解析字段值不一致", decision["reason"])

    def test_turbine_model_route_value_consistency(self) -> None:
        manifest_path = self._write_manifest()
        ok = self._fill("C1-R01", "投标机型", "EW6.25-220", evidence={"sourceRoute": "projectTurbineModel"})
        self._write_plan([ok])
        result = filler.run_apply(manifest_path)
        self.assertEqual(result["mapping"]["decisions"][0]["action"], "fill")

    def test_turbine_model_route_inconsistent_downgrades(self) -> None:
        manifest_path = self._write_manifest()
        bad = self._fill("C1-R01", "投标机型", "EW5.0-190", evidence={"sourceRoute": "projectTurbineModel"})
        self._write_plan([bad])
        result = filler.run_apply(manifest_path)
        decision = result["mapping"]["decisions"][0]
        self.assertEqual(decision["action"], "manual")
        self.assertIn("与投标机型信息不一致", decision["reason"])


class ApplyPreferredDeviationTests(PriorityTestBase):
    def test_material_route_deviating_from_fact_table_downgrades(self) -> None:
        # 素材路由（excerpt 合法命中）但值与事实表 6.25 MW 矛盾 → 事实表优先降级
        manifest_path = self._write_manifest()
        fill = self._fill("C1-R02", "单机容量", "7000", unit="kW", excerpt="单机容量 7000 kW")
        self._write_plan([fill])
        result = filler.run_apply(manifest_path)
        decision = result["mapping"]["decisions"][1]
        self.assertEqual(decision["action"], "manual")
        self.assertIn("事实表优先", decision["reason"])
        self.assertIn("6.25", decision["reason"])

    def test_material_route_consistent_with_fact_table_fills(self) -> None:
        # 素材路由取值与事实表一致（kW→MW 换算后相等）→ 正常填写
        manifest_path = self._write_manifest()
        fill = self._fill("C1-R02", "单机容量", "6250", unit="kW", excerpt="单机容量 7000 kW")
        self._write_plan([fill])
        result = filler.run_apply(manifest_path)
        decision = result["mapping"]["decisions"][1]
        self.assertEqual(decision["action"], "fill")
        self.assertEqual(decision["value"], "6.25")

    def test_material_route_excerpt_check_unchanged(self) -> None:
        # 无事实表绑定的字段：excerpt 逐字校验行为不变，命中不了照样降级
        manifest_path = self._write_manifest()
        fill = self._fill("C1-R03", "叶轮直径", "220", excerpt="叶轮直径 999 m")
        self._write_plan([fill])
        result = filler.run_apply(manifest_path)
        decision = result["mapping"]["decisions"][2]
        self.assertEqual(decision["action"], "manual")
        self.assertIn("证据未命中", decision["reason"])


if __name__ == "__main__":
    unittest.main()
