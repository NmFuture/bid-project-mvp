"""table-filler 原生 Excel 源选择单测（B 报价类 / D 曲线类 0 填写攻坚）。

金标反评定位：素材缓存里落地的是清洗后的 docx 文本稿，sheet 结构已丢失，
凡依赖 openpyxl 读 sheet 的分支（报价分项转写、功率曲线矩阵、参数表机型列）
拿到的 kind 都是 docx，条件不成立直接跳过，整簇附表 0 填写。
本组测试覆盖新增的原件优先路径：originalPath 在时 Source 按 xlsx 走、
清洗 docx 仍保留在 cleaned_docx_path 供整表移植回退、原件缺失/非 Excel 不改行为。
"""
from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

_SRC = (
    Path(__file__).resolve().parents[1]
    / "opencode"
    / "skills"
    / "bid-tech-table-filler"
    / "scripts"
    / "run_from_manifest.py"
)
_SPEC = importlib.util.spec_from_file_location("tech_table_filler_original_under_test", _SRC)
filler = importlib.util.module_from_spec(_SPEC)
assert _SPEC.loader is not None
sys.modules["tech_table_filler_original_under_test"] = filler
_SPEC.loader.exec_module(filler)


def _write_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("清洗后的文本稿")
    doc.save(str(path))
    return path


def _write_xlsx(path: Path, sheet_name: str = "B 设备的分项报价") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws["A1"] = "序号"
    wb.save(str(path))
    return path


class AddSourceFromMaterialOriginalTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.base = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def _add(self, material: dict) -> list:
        sources: list = []
        filler.add_source_from_material(
            sources,
            material,
            self.base,
            priority=70,
            route="test",
        )
        return sources

    def test_original_xlsx_wins_over_cleaned_docx(self) -> None:
        cleaned = _write_docx(self.base / "RAW-0136-报价文件.docx")
        original = _write_xlsx(self.base / "RAW-0136-报价文件.xlsx")
        sources = self._add(
            {
                "id": "RAW-0136",
                "name": "报价文件.xlsx",
                "path": str(cleaned),
                "originalPath": str(original),
            }
        )
        self.assertEqual(len(sources), 1)
        source = sources[0]
        # 依赖 openpyxl 的分支据 kind=xlsx 与 path 后缀判定，两者都必须指向原件
        self.assertEqual(source.kind, "xlsx")
        self.assertEqual(source.path, original.resolve())
        # 清洗稿不丢：整表移植等按 Word 表格工作的分支仍可回退使用
        self.assertEqual(source.cleaned_docx_path, cleaned.resolve())

    def test_quote_source_matched_after_original_landed(self) -> None:
        cleaned = _write_docx(self.base / "RAW-0136-无价格-…报价文件.docx")
        original = _write_xlsx(self.base / "RAW-0136-无价格-…报价文件.xlsx")
        sources = self._add(
            {
                "id": "RAW-0136",
                "name": "无价格-标段一：…报价文件.xlsx",
                "path": str(cleaned),
                "originalPath": str(original),
            }
        )
        # 报价转写器要求 kind=xlsx 且名含"报价文件"，此前恒为 None 导致 B 簇全空
        self.assertIsNotNone(filler.quote_xlsx_source(sources))

    def test_missing_original_keeps_cleaned_docx_behaviour(self) -> None:
        cleaned = _write_docx(self.base / "RAW-0161-风资源评估报告.docx")
        sources = self._add({"id": "RAW-0161", "name": "风资源评估报告.docx", "path": str(cleaned)})
        source = sources[0]
        self.assertEqual(source.kind, "docx")
        self.assertEqual(source.path, cleaned.resolve())
        self.assertIsNone(source.cleaned_docx_path)

    def test_non_excel_original_path_ignored(self) -> None:
        cleaned = _write_docx(self.base / "RAW-0152-载荷安全性评估报告.docx")
        pdf = self.base / "RAW-0152-载荷安全性评估报告.pdf"
        pdf.write_bytes(b"%PDF-1.4 fake")
        sources = self._add(
            {
                "id": "RAW-0152",
                "name": "载荷安全性评估报告.pdf",
                "path": str(cleaned),
                "originalPath": str(pdf),
            }
        )
        source = sources[0]
        # PDF 文本已由清洗稿承载，不改读原件，避免凭空要求 OCR sidecar
        self.assertEqual(source.kind, "docx")
        self.assertEqual(source.path, cleaned.resolve())

    def test_original_path_not_on_disk_falls_back(self) -> None:
        cleaned = _write_docx(self.base / "RAW-0132-功率曲线.docx")
        sources = self._add(
            {
                "id": "RAW-0132",
                "name": "功率曲线.xlsx",
                "path": str(cleaned),
                "originalPath": str(self.base / "缺失.xlsx"),
            }
        )
        self.assertEqual(sources[0].kind, "docx")


class SourceDocxPathTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.base = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_docx_source_uses_itself(self) -> None:
        path = _write_docx(self.base / "a.docx")
        source = filler.Source(name="a", path=path, kind="docx", priority=70, route="test")
        self.assertEqual(filler.source_docx_path(source), path)

    def test_xlsx_source_falls_back_to_cleaned_docx(self) -> None:
        cleaned = _write_docx(self.base / "b.docx")
        original = _write_xlsx(self.base / "b.xlsx")
        source = filler.Source(
            name="b",
            path=original,
            kind="xlsx",
            priority=70,
            route="test",
            cleaned_docx_path=cleaned,
        )
        self.assertEqual(filler.source_docx_path(source), cleaned)

    def test_xlsx_source_without_cleaned_docx_is_none(self) -> None:
        original = _write_xlsx(self.base / "c.xlsx")
        source = filler.Source(name="c", path=original, kind="xlsx", priority=70, route="test")
        self.assertIsNone(filler.source_docx_path(source))


if __name__ == "__main__":
    unittest.main()
