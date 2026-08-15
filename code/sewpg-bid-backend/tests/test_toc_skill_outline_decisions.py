from __future__ import annotations

import json
import tempfile
from pathlib import Path
from docx import Document

from toc_skill_helpers import (
    TocSkillScriptTestBase,
    complete_outline_review,
    json_load,
    load_outline_script,
    submit_outline_changes,
    write_decision_context_fixture,
)


class TocSkillScriptTests(TocSkillScriptTestBase):

    def test_bid_outline_compose_requires_current_explicit_decisions(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest_path = root / "s2_input.json"
            manifest = {"workDir": str(root), "outputFile": str(root / "toc.json")}
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            structure_path = root / "template_structure.json"
            structure_path.write_text(
                json.dumps(
                    {
                        "schema_version": "template-structure.v1",
                        "items": [{"number": "第1章", "title": "技术方案", "level": 1}],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            with self.assertRaisesRegex(SystemExit, "必须先执行 s2outline decisions"):
                outline_runner.compose_manifest(manifest, manifest_path)

            with self.assertRaisesRegex(SystemExit, "input_fingerprint.*required"):
                outline_runner.submit_outline_decisions(
                    manifest,
                    manifest_path,
                    {"schema_version": "technical-outline-decisions.v1", "changes": []},
                )

            structure = json_load(structure_path)
            fingerprint = outline_runner.outline_composer.annotate_template_structure(structure)[
                "input_fingerprint"
            ]
            with self.assertRaisesRegex(SystemExit, "does not match"):
                outline_runner.submit_outline_decisions(
                    manifest,
                    manifest_path,
                    {
                        "schema_version": "technical-outline-decisions.v1",
                        "input_fingerprint": "stale-template-fingerprint",
                        "template_decisions": [],
                        "changes": [],
                    },
                )

            outline_runner.submit_outline_decisions(
                manifest,
                manifest_path,
                {
                    "schema_version": "technical-outline-decisions.v1",
                    "input_fingerprint": fingerprint,
                    "template_decisions": [
                        {"target_id": "TPL-0001", "decision": "retain"}
                    ],
                    "changes": [],
                },
            )
            result = outline_runner.compose_manifest(manifest, manifest_path)

        self.assertEqual(result["summary"]["total_nodes"], 1)



    def test_bid_outline_decisions_require_an_explicit_choice_for_every_template_node(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest_path = root / "s2_input.json"
            manifest = {"workDir": str(root), "outputFile": str(root / "toc.json")}
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            structure = {
                "schema_version": "template-structure.v1",
                "items": [
                    {"number": "Chapter 1", "title": "Overview", "level": 1},
                    {"number": "1.1", "title": "Legacy topic", "level": 2},
                ],
            }
            (root / "template_structure.json").write_text(
                json.dumps(structure), encoding="utf-8"
            )
            annotated = outline_runner.outline_composer.annotate_template_structure(structure)
            incomplete = {
                "schema_version": "technical-outline-decisions.v1",
                "input_fingerprint": annotated["input_fingerprint"],
                "template_decisions": [
                    {"target_id": "TPL-0001", "decision": "retain"},
                ],
                "changes": [],
            }

            with self.assertRaisesRegex(SystemExit, "TPL-0002"):
                outline_runner.submit_outline_decisions(manifest, manifest_path, incomplete)

            complete = {
                **incomplete,
                "template_decisions": [
                    {"target_id": "TPL-0001", "decision": "retain"},
                    {
                        "target_id": "TPL-0002",
                        "decision": "suggest_delete",
                        "reason": "The tender explicitly excludes this scope.",
                    },
                ],
            }
            outline_runner.submit_outline_decisions(manifest, manifest_path, complete)
            outline_runner.compose_manifest(manifest, manifest_path)
            outline = json_load(root / "toc.json")

        self.assertEqual(outline["nodes"][0]["suggestion_action"], "必要")
        self.assertEqual(
            outline["nodes"][0]["children"][0]["suggestion_action"],
            "建议删除",
        )



    def test_bid_outline_decision_next_returns_only_current_template_chapter(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest, manifest_path = write_decision_context_fixture(
                root,
                template_count=50,
                template_title="超长中文技术方案章节" * 30,
            )

            batch = outline_runner.dispatch_command(
                "decision-next", manifest, manifest_path, []
            )

        compact = json.dumps(batch, ensure_ascii=False, separators=(",", ":"))
        self.assertLess(len(compact), 45000)
        self.assertEqual(len(batch["items"]), 1)
        self.assertEqual(batch["remaining_count"], 50)
        self.assertNotIn("comparison_context", batch)
        self.assertEqual(batch["decision_level"], 2)
        # 该夹具的章只有章根（无二级节点），属稀疏章：切换为连续通读构建纪律
        self.assertEqual(batch["authoring_mode"], "sparse_chapter")
        self.assertEqual(
            batch["decision_steps"],
            [
                "本章模板二级节点稀疏，需要从招标原文构建子目录：先通读完整招标目录，再自主圈定与本章主题对应的少数上级章节",
                "对圈定章节用 section --max-chars 30000 连续通读原文，不逐段跳读；search 全会话最多 2 次，仅用于跨章定位",
                "从已读原文自主提炼响应单元并确定标题与粒度，一次性提交全部新增；每个新增 reason + evidence_id",
                "章根与既有二级节点仍按 retain / suggest_delete 表态",
            ],
        )
        self.assertEqual(
            batch["submission_contract"],
            {
                "required_fields": ["batch_token", "items", "additions"],
                "items_must_match_batch": True,
                "additions_must_be_explicit": True,
                "decision_covers_subtree": "对二级节点的判断适用于其下全部三级节点",
                "addition_levels": "章节会话只能在当前一级章下新增二级节点，parent_id 必须引用当前章根；不新增一级章或三级节点",
            },
        )



    def test_bid_outline_decision_next_uses_one_complete_root_chapter(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest, manifest_path = write_decision_context_fixture(root, heading_count=0)
            (root / "template_structure.json").write_text(
                json.dumps(
                    {
                        "schema_version": "template-structure.v1",
                        "items": [
                            {"number": "第1章", "title": "总体方案", "level": 1},
                            {"number": "1.1", "title": "设计依据", "level": 2},
                            {"number": "1.1.1", "title": "采用标准", "level": 3},
                            {"number": "第2章", "title": "供货方案", "level": 1},
                            {"number": "2.1", "title": "供货范围", "level": 2},
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            first = outline_runner.dispatch_command(
                "decision-next", manifest, manifest_path, []
            )
            outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": first["batch_token"],
                            "items": [
                                {"target_id": item["target_id"], "decision": "retain"}
                                for item in first["items"]
                            ],
                            "additions": [],
                        },
                        ensure_ascii=False,
                    )
                ],
            )
            second = outline_runner.dispatch_command(
                "decision-next", manifest, manifest_path, []
            )

        self.assertEqual(first["chapter_id"], "TPL-0001")
        # TPL-0003 是三级节点，跟随二级父节点 TPL-0002，不进入决策批次。
        self.assertEqual([item["target_id"] for item in first["items"]], ["TPL-0001", "TPL-0002"])
        self.assertEqual(second["chapter_id"], "TPL-0004")
        self.assertEqual([item["target_id"] for item in second["items"]], ["TPL-0004", "TPL-0005"])



    def test_bid_outline_decision_next_keeps_one_batch_per_chapter_at_level_two(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest, manifest_path = write_decision_context_fixture(root, heading_count=0)
            items = [{"number": "第1章", "title": "专题方案", "level": 1}]
            for section in range(1, 4):
                items.append(
                    {"number": f"1.{section}", "title": f"专题{section}", "level": 2}
                )
                items.extend(
                    {
                        "number": f"1.{section}.{child}",
                        "title": f"专题{section}响应内容{child}",
                        "level": 3,
                    }
                    for child in range(1, 21)
                )
            (root / "template_structure.json").write_text(
                json.dumps(
                    {"schema_version": "template-structure.v1", "items": items},
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            chapter = outline_runner.dispatch_command(
                "decision-next", manifest, manifest_path, []
            )
            outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": chapter["batch_token"],
                            "items": [
                                {
                                    "target_id": item["target_id"],
                                    "decision": "retain",
                                    "reason": "保留完整专题方案章节。",
                                }
                                for item in chapter["items"]
                            ],
                            "additions": [],
                        },
                        ensure_ascii=False,
                    )
                ],
            )
            after = outline_runner.dispatch_command(
                "decision-next", manifest, manifest_path, []
            )

        # 60 个三级节点跟随各自的二级父节点，整章只剩章根 + 3 个二级节点一批决完。
        self.assertEqual(chapter["chapter_id"], "TPL-0001")
        self.assertEqual(
            [item["number"] for item in chapter["items"]],
            ["第1章", "1.1", "1.2", "1.3"],
        )
        self.assertTrue(after["complete"])



    def test_bid_outline_decision_next_uses_actual_first_page_size_for_large_chapter(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest, manifest_path = write_decision_context_fixture(
                root,
                heading_count=160,
                heading_text="招标技术要求与独立响应事项" * 4,
            )
            items = [{"number": "1", "title": "总体技术方案", "level": 1}]
            items.extend(
                {
                    "number": f"1.{index}",
                    "title": "需要结合招标目录逐项判断的专业技术专题" * 6,
                    "level": 2,
                }
                for index in range(1, 30)
            )
            (root / "template_structure.json").write_text(
                json.dumps({"schema_version": "template-structure.v1", "items": items}, ensure_ascii=False),
                encoding="utf-8",
            )

            batch = outline_runner.dispatch_command("decision-next", manifest, manifest_path, [])

        compact = json.dumps(batch, ensure_ascii=False, separators=(",", ":"))
        self.assertEqual(len(batch["items"]), 30)
        self.assertNotIn("comparison_context", batch)
        self.assertLess(len(compact), 45000)



    def test_bid_outline_decision_batch_requires_explicit_additions(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest, manifest_path = write_decision_context_fixture(root, heading_count=0)
            batch = outline_runner.dispatch_command(
                "decision-next", manifest, manifest_path, []
            )

            with self.assertRaisesRegex(SystemExit, "additions.*required"):
                outline_runner.dispatch_command(
                    "decision-batch",
                    manifest,
                    manifest_path,
                    [
                        json.dumps(
                            {
                                "batch_token": batch["batch_token"],
                                "items": [
                                    {"target_id": item["target_id"], "decision": "retain"}
                                    for item in batch["items"]
                                ],
                            }
                        )
                    ],
                )



    def test_bid_outline_decision_batch_requires_reason_or_read_tender_evidence(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("1 专项方案", style="Heading 1")
            tender_doc.add_paragraph("投标人必须提交海上运输安全专项方案。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": tender.name, "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            headings = outline_runner.dispatch_command("headings", manifest, manifest_path, [])
            batch = outline_runner.dispatch_command("decision-next", manifest, manifest_path, [])

            self.assertNotIn("comparison_context", batch)
            with self.assertRaisesRegex(SystemExit, "evidence_id.*reason"):
                outline_runner.dispatch_command(
                    "decision-batch",
                    manifest,
                    manifest_path,
                    [
                        json.dumps(
                            {
                                "batch_token": batch["batch_token"],
                                "items": [
                                    {"target_id": batch["items"][0]["target_id"], "decision": "retain"}
                                ],
                                "additions": [],
                            },
                            ensure_ascii=False,
                        )
                    ],
                )

            search = outline_runner.dispatch_command(
                "search", manifest, manifest_path, ["运输安全"]
            )
            evidence_id = search["results"][0]["evidence_id"]
            payload = {
                "batch_token": batch["batch_token"],
                "items": [
                    {
                        "target_id": batch["items"][0]["target_id"],
                        "decision": "retain",
                        "evidence_id": evidence_id,
                        "reason": "招标要求本章承接整体技术方案。",
                    }
                ],
                "additions": [
                    {
                        "node_id": "AI-0001",
                        "parent_id": batch["items"][0]["target_id"],
                        "number": "1.1",
                        "title": "海上运输安全专项方案",
                        "reason": "招标文件明确要求独立提交。",
                        "evidence_id": evidence_id,
                    }
                ],
            }
            with self.assertRaisesRegex(SystemExit, "尚未通过受控阅读"):
                outline_runner.dispatch_command(
                    "decision-batch",
                    manifest,
                    manifest_path,
                    [json.dumps(payload, ensure_ascii=False)],
                )
            outline_runner.dispatch_command(
                "section",
                manifest,
                manifest_path,
                [headings["files"][0]["items"][0]["section_id"]],
            )
            result = outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [json.dumps(payload, ensure_ascii=False)],
            )
            state = json_load(root / "outline_decision_state.json")

        self.assertEqual(result["decided_count"], 1)
        # evidence_id 与 reason 是前端展示契约：有依据就能跳转原文，reason 同时作为说明保留。
        retained = state["template_decisions"][batch["items"][0]["target_id"]]
        self.assertEqual(retained["tender_basis"]["evidence_id"], evidence_id)
        self.assertEqual(retained["reason"], "招标要求本章承接整体技术方案。")
        self.assertEqual(state["additions"][0]["tender_basis"]["evidence_id"], evidence_id)



    def test_bid_outline_controlled_decision_batches_build_final_decisions(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest_path = root / "s2_input.json"
            manifest = {"workDir": str(root), "outputFile": str(root / "toc.json")}
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            (root / "template_structure.json").write_text(
                json.dumps(
                    {
                        "schema_version": "template-structure.v1",
                        "items": [
                            {"number": "Chapter 1", "title": "Overview", "level": 1},
                            {"number": "1.1", "title": "Design", "level": 2},
                            {"number": "1.1.1", "title": "Legacy case", "level": 3},
                        ],
                    }
                ),
                encoding="utf-8",
            )

            first = outline_runner.dispatch_command("decision-next", manifest, manifest_path, [])
            first_result = outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": first["batch_token"],
                            "items": [
                                {"target_id": item["target_id"], "decision": "retain"}
                                for item in first["items"]
                            ],
                            "additions": [],
                        }
                    )
                ],
            )
            complete_outline_review(outline_runner, manifest, manifest_path)
            finalized = outline_runner.dispatch_command(
                "decisions", manifest, manifest_path, []
            )

        # 一级章 + 二级节点各一次决策；三级 Legacy case 跟随二级父节点。
        self.assertEqual(first_result["decided_count"], 2)
        self.assertEqual(finalized["templateDecisionCount"], 2)
        self.assertEqual(finalized["remainingTemplateDecisionCount"], 0)



    def test_bid_outline_level_two_decision_cascades_to_level_three_subtree(self) -> None:
        outline_composer = load_outline_script("outline_composer")
        structure = {
            "schema_version": "template-structure.v1",
            "items": [
                {"number": "第1章", "title": "总体方案", "level": 1},
                {"number": "1.1", "title": "设计依据", "level": 2},
                {"number": "1.1.1", "title": "采用标准", "level": 3},
                {"number": "1.1.2", "title": "设计输入", "level": 3},
                {"number": "1.2", "title": "重复小节", "level": 2},
                {"number": "1.2.1", "title": "重复细项", "level": 3},
                {"number": "第2章", "title": "其他分册内容", "level": 1},
            ],
        }
        basis = {"evidence_id": "TEN-1:B000123", "file_id": "TEN-1", "search_text": "设计依据"}
        decisions = {
            "schema_version": "technical-outline-decisions.v1",
            "input_fingerprint": outline_composer.annotate_template_structure(structure)[
                "input_fingerprint"
            ],
            "template_decisions": [
                {"target_id": "TPL-0001", "decision": "retain", "reason": "承接整体技术论证。"},
                {"target_id": "TPL-0002", "decision": "retain", "tender_basis": basis},
                {"target_id": "TPL-0005", "decision": "suggest_delete", "reason": "与 1.1 语义重复。"},
                {"target_id": "TPL-0007", "decision": "suggest_delete", "reason": "属于商务分册。"},
            ],
            "changes": [],
        }

        outline, _ = outline_composer.build_composition(structure, decisions)

        chapter, other = outline["nodes"]
        design, duplicate = chapter["children"]
        # 保留的二级节点：整个三级子树跟随标签与招标依据。
        self.assertEqual([child["title"] for child in design["children"]], ["采用标准", "设计输入"])
        for child in design["children"]:
            self.assertEqual(child["suggestion_action"], "必要")
            self.assertEqual(child["tender_basis"], basis)
        # 建议删除的二级节点：子树全部建议删除，并复用父节点理由。
        self.assertEqual(duplicate["suggestion_action"], "建议删除")
        self.assertEqual(
            [
                (child["suggestion_action"], child["suggestion_reason"])
                for child in duplicate["children"]
            ],
            [("建议删除", "与 1.1 语义重复。")],
        )
        # 整章删除同样成立，且一级章自身可以独立决策。
        self.assertEqual(other["suggestion_action"], "建议删除")
        self.assertEqual(other["suggestion_reason"], "属于商务分册。")



    def test_bid_outline_additions_are_limited_to_chapter_and_section_levels(self) -> None:
        decision_workflow = load_outline_script("run_from_manifest").decision_workflow
        structure = {
            "schema_version": "template-structure.v1",
            "items": [
                {"number": "第1章", "title": "总体方案", "level": 1},
                {"number": "1.1", "title": "设计依据", "level": 2},
                {"number": "1.1.1", "title": "采用标准", "level": 3},
            ],
        }

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            batch = decision_workflow.next_decision_batch(root, structure)
            self.assertEqual(
                [item["target_id"] for item in batch["items"]], ["TPL-0001", "TPL-0002"]
            )

            def submit(additions: list[dict]) -> dict:
                return decision_workflow.submit_decision_batch(
                    root,
                    structure,
                    {
                        "batch_token": batch["batch_token"],
                        "items": [
                            {"target_id": item["target_id"], "decision": "retain"}
                            for item in batch["items"]
                        ],
                        "additions": additions,
                    },
                )

            for parent_id in ("TPL-0002", "TPL-0003"):
                with self.subTest(parent_id=parent_id), self.assertRaisesRegex(
                    SystemExit, "只能新增一级章或二级节点"
                ):
                    submit(
                        [
                            {
                                "node_id": "ADD-0001",
                                "parent_id": parent_id,
                                "number": "1.1.2",
                                "title": "专项细项",
                                "reason": "招标要求补充。",
                            }
                        ]
                    )

            result = submit(
                [
                    {
                        "node_id": "ADD-CHAPTER",
                        "parent_id": None,
                        "number": "第2章",
                        "title": "专项试验方案",
                        "reason": "招标要求独立成章。",
                    },
                    {
                        "node_id": "ADD-SECTION",
                        "parent_id": "TPL-0001",
                        "number": "1.2",
                        "title": "海上运输安全专项方案",
                        "reason": "招标要求独立提交。",
                    },
                ]
            )

        self.assertEqual(result["addition_count"], 2)



    def test_bid_outline_chapter_decisions_are_isolated_and_merge_completely(self) -> None:
        decision_workflow = load_outline_script("run_from_manifest").decision_workflow
        structure = {
            "schema_version": "template-structure.v1",
            "items": [
                {"number": "1", "title": "总体方案", "level": 1},
                {"number": "1.1", "title": "实施组织", "level": 2},
                {"number": "2", "title": "质量安全", "level": 1},
                {"number": "2.1", "title": "质量保证", "level": 2},
            ],
        }

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            chapter_catalog = decision_workflow.decision_chapters(structure)
            chapter_dirs = {
                item["chapter_id"]: root / item["chapter_id"]
                for item in chapter_catalog["chapters"]
            }
            chapter_titles = {}
            chapter_parent_ids = {}
            for chapter_id, chapter_dir in chapter_dirs.items():
                chapter_dir.mkdir()
                while True:
                    batch = decision_workflow.next_decision_batch(
                        chapter_dir,
                        structure,
                        chapter_id=chapter_id,
                    )
                    if batch["complete"]:
                        break
                    chapter_titles[chapter_id] = tuple(item["title"] for item in batch["items"])
                    chapter_parent_ids[chapter_id] = batch["items"][0]["target_id"]
                    other_parent_id = next(
                        item["template_id"]
                        for item in decision_workflow._annotated_items(structure)[1]
                        if item["template_id"] not in {
                            batch_item["target_id"] for batch_item in batch["items"]
                        }
                    )
                    with self.assertRaisesRegex(SystemExit, "outside the active chapter"):
                        decision_workflow.submit_decision_batch(
                            chapter_dir,
                            structure,
                            {
                                "batch_token": batch["batch_token"],
                                "items": [
                                    {
                                        "target_id": item["target_id"],
                                        "decision": "retain",
                                        "reason": "历史模板专家经验仍适用。",
                                    }
                                    for item in batch["items"]
                                ],
                                "additions": [
                                    {
                                        "node_id": "ADD-0001",
                                        "parent_id": other_parent_id,
                                        "number": "1.9",
                                        "title": "跨章新增",
                                        "reason": "不允许跨章挂载。",
                                    }
                                ],
                            },
                            chapter_id=chapter_id,
                        )
                    with self.assertRaisesRegex(SystemExit, "must stay under the active chapter"):
                        decision_workflow.submit_decision_batch(
                            chapter_dir,
                            structure,
                            {
                                "batch_token": batch["batch_token"],
                                "items": [
                                    {
                                        "target_id": item["target_id"],
                                        "decision": "retain",
                                        "reason": "历史模板专家经验仍适用。",
                                    }
                                    for item in batch["items"]
                                ],
                                "additions": [
                                    {
                                        "node_id": "ADD-ROOT",
                                        "parent_id": None,
                                        "number": "9",
                                        "title": "跨章一级新增",
                                        "reason": "章节会话不得创建无归属的一级章。",
                                    }
                                ],
                            },
                            chapter_id=chapter_id,
                        )
                    decision_workflow.submit_decision_batch(
                        chapter_dir,
                        structure,
                        {
                            "batch_token": batch["batch_token"],
                            "items": [
                                {
                                    "target_id": item["target_id"],
                                    "decision": "retain",
                                    "reason": "历史模板专家经验仍适用。",
                                }
                                for item in batch["items"]
                            ],
                            "additions": [
                                {
                                    "node_id": "ADD-0001",
                                    "parent_id": batch["items"][0]["target_id"],
                                    "number": "1.9",
                                    "title": "本章专项响应",
                                    "reason": "本章需要独立表达。",
                                }
                            ],
                        },
                        chapter_id=chapter_id,
                    )

            merged = decision_workflow.merge_chapter_decisions(
                root,
                structure,
                chapter_dirs,
            )
            state = json_load(root / "outline_decision_state.json")

        self.assertEqual(chapter_catalog["chapter_count"], 2)
        self.assertEqual(set(chapter_titles.values()), {
            ("总体方案", "实施组织"),
            ("质量安全", "质量保证"),
        })
        self.assertTrue(merged["complete"])
        self.assertEqual(merged["decided_count"], 4)
        self.assertEqual(merged["addition_count"], 2)
        self.assertEqual(len(state["template_decisions"]), 4)
        self.assertEqual(len({item["node_id"] for item in state["additions"]}), 2)
        for addition in state["additions"]:
            owner = state["addition_chapters"][addition["node_id"]]
            self.assertEqual(addition["parent_id"], chapter_parent_ids[owner])



    def test_bid_outline_sparse_chapter_unit_switches_to_contiguous_reading_mode(self) -> None:
        decision_workflow = load_outline_script("run_from_manifest").decision_workflow
        structure = {
            "schema_version": "template-structure.v1",
            "items": [
                {"number": "1", "title": "总体方案", "level": 1},
                {"number": "1.1", "title": "实施组织", "level": 2},
                {"number": "1.2", "title": "进度计划", "level": 2},
                {"number": "2", "title": "技术规范响应", "level": 1},
            ],
        }

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            normal = decision_workflow.next_decision_batch(root, structure)
            self.assertNotIn("authoring_mode", normal)
            self.assertFalse(any("连续通读" in step for step in normal["decision_steps"]))
            decision_workflow.submit_decision_batch(
                root,
                structure,
                {
                    "batch_token": normal["batch_token"],
                    "items": [
                        {
                            "target_id": item["target_id"],
                            "decision": "retain",
                            "reason": "历史模板专家经验仍适用。",
                        }
                        for item in normal["items"]
                    ],
                    "additions": [],
                },
            )
            sparse = decision_workflow.next_decision_batch(root, structure)
            replay = decision_workflow.next_decision_batch(root, structure)

        # 有二级节点的章走默认判定流程；空章切换为"圈定上级章节 + 大页连续通读"纪律
        self.assertEqual(len(sparse["items"]), 1)
        self.assertEqual(sparse["authoring_mode"], "sparse_chapter")
        self.assertTrue(any("连续通读" in step for step in sparse["decision_steps"]))
        self.assertTrue(any("最多 2 次" in step for step in sparse["decision_steps"]))
        self.assertTrue(any("--max-chars 30000" in step for step in sparse["decision_steps"]))
        # 未提交前重复 decision-next 返回同一活动批次，模式标记保持
        self.assertEqual(replay["batch_token"], sparse["batch_token"])
        self.assertEqual(replay["authoring_mode"], "sparse_chapter")



    def test_bid_outline_decisions_reject_add_or_move_under_collapsed_parent(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest_path = root / "s2_input.json"
            manifest = {"workDir": str(root), "outputFile": str(root / "toc.json")}
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            structure = {
                "schema_version": "template-structure.v1",
                "items": [
                    {"number": "第1章", "title": "技术方案", "level": 1},
                    {"number": "1.1", "title": "待收敛节点", "level": 2},
                    {"number": "1.2", "title": "保留节点", "level": 2},
                    {"number": "1.2.1", "title": "待移动专题", "level": 3},
                ],
            }
            structure_path = root / "template_structure.json"
            structure_path.write_text(json.dumps(structure, ensure_ascii=False), encoding="utf-8")
            invalid_changes = [
                [
                    {
                        "operation": "collapse",
                        "target_id": "TPL-0002",
                        "reason": "由根节点统一承载。",
                    },
                    {
                        "operation": "add",
                        "node_id": "ADD-0001",
                        "parent_id": "TPL-0002",
                        "number": "1.1.1",
                        "title": "新增专题",
                        "suggestion_action": "建议增加",
                        "suggestion_reason": "招标要求独立提交。",
                    },
                ],
                [
                    {
                        "operation": "collapse",
                        "target_id": "TPL-0002",
                        "reason": "由根节点统一承载。",
                    },
                    {
                        "operation": "update",
                        "target_id": "TPL-0004",
                        "parent_id": "TPL-0002",
                        "reason": "调整专题归属。",
                    },
                ],
            ]

            for changes in invalid_changes:
                with self.subTest(operation=changes[-1]["operation"]), self.assertRaisesRegex(
                    SystemExit, "collapsed"
                ):
                    submit_outline_changes(outline_runner, manifest, manifest_path, changes)



    def test_bid_outline_strict_compose_requires_controlled_decisions_and_binds_state(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            output = root / "toc.json"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("第1章 招标技术要求", style="Heading 1")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": tender.name, "path": str(tender)}],
                "outputFile": str(output),
                "requireComposedOutline": True,
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            with self.assertRaisesRegex(SystemExit, "完整读取招标目录"):
                outline_runner.compose_manifest(manifest, manifest_path)
            headings = outline_runner.dispatch_command("headings", manifest, manifest_path, [])
            section = outline_runner.dispatch_command(
                "section",
                manifest,
                manifest_path,
                [headings["files"][0]["items"][0]["section_id"]],
            )
            evidence_id = section["records"][0]["evidence_id"]
            structure = json_load(root / "template_structure.json")
            outline_runner.submit_outline_decisions(
                manifest,
                manifest_path,
                {
                    "schema_version": "technical-outline-decisions.v1",
                    "input_fingerprint": structure["input_fingerprint"],
                    "template_decisions": [
                        {"target_id": "TPL-0001", "decision": "retain"}
                    ],
                    "changes": [],
                },
            )
            with self.assertRaisesRegex(SystemExit, "受控决策"):
                outline_runner.compose_manifest(manifest, manifest_path)

            decision_batch = outline_runner.dispatch_command(
                "decision-next", manifest, manifest_path, []
            )
            outline_runner.dispatch_command(
                "read", manifest, manifest_path, [evidence_id]
            )
            outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": decision_batch["batch_token"],
                            "items": [
                                {
                                    "target_id": "TPL-0001",
                                    "decision": "retain",
                                    "evidence_id": evidence_id,
                                }
                            ],
                            "additions": [],
                        },
                        ensure_ascii=False,
                    )
                ],
            )
            outline_runner.dispatch_command(
                "read", manifest, manifest_path, [evidence_id]
            )
            complete_outline_review(outline_runner, manifest, manifest_path)
            outline_runner.dispatch_command("decisions", manifest, manifest_path, [])
            outline_runner.compose_manifest(manifest, manifest_path)
            report = json_load(root / "outline_compose_report.json")

            state_path = root / "outline_decision_state.json"
            state = json_load(state_path)
            state["tampered"] = True
            state_path.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
            with self.assertRaisesRegex(SystemExit, "decisionStateDigest"):
                outline_runner.finalize_manifest(manifest, manifest_path)

        self.assertTrue(report["tenderInputsDigest"])
        self.assertTrue(report["headingsStateDigest"])
        self.assertTrue(report["decisionStateDigest"])
