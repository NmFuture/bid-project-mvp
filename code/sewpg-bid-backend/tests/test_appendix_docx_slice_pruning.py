from __future__ import annotations

import struct
import zipfile
import zlib
from pathlib import Path

from docx import Document
from docx.shared import Cm

from app.document_processing.technical_document.assembly.create_tech_master import prune_unreferenced_media
from app.services.parsing import _slice_appendix_from_source


def _write_png(path: Path, color: bytes) -> Path:
    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    width = height = 8
    raw = b"".join(b"\x00" + color * width for _ in range(height))
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    return path


def _source_docx(tmp_path: Path) -> Path:
    image = _write_png(tmp_path / "outside.png", b"\xff\x00\x00")
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("附表前正文")
    document.add_paragraph().add_run().add_picture(str(image), width=Cm(2))
    document.add_paragraph("附表A.1 投标机型总方案信息表")
    document.save(source)
    return source


def _media_entries(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return [name for name in archive.namelist() if name.startswith("word/media/")]


def test_slice_appendix_drops_media_not_referenced_by_kept_body(tmp_path: Path) -> None:
    source = _source_docx(tmp_path)
    target = tmp_path / "appendix.docx"

    assert _slice_appendix_from_source(source, target, 2, 2, prune_unused_parts=True) is True

    assert _media_entries(target) == []
    assert [paragraph.text for paragraph in Document(target).paragraphs] == ["附表A.1 投标机型总方案信息表"]


def test_slice_appendix_keeps_media_referenced_by_kept_body(tmp_path: Path) -> None:
    source = _source_docx(tmp_path)
    target = tmp_path / "appendix-with-image.docx"

    assert _slice_appendix_from_source(source, target, 1, 2, prune_unused_parts=True) is True

    assert len(_media_entries(target)) == 1


def test_slice_appendix_does_not_prune_shared_business_path_by_default(tmp_path: Path) -> None:
    source = _source_docx(tmp_path)
    target = tmp_path / "business-appendix.docx"

    assert _slice_appendix_from_source(source, target, 2, 2) is True

    assert len(_media_entries(target)) == 1


def test_prune_keeps_relationship_marked_as_external(tmp_path: Path) -> None:
    source = tmp_path / "external-link.docx"
    Document().save(source)
    with zipfile.ZipFile(source, "r") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    relationship = (
        b'<Relationship Id="rIdExternal" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        b'Target="mailto:bid@example.com" TargetMode="External"/>'
    )
    parts["word/_rels/document.xml.rels"] = parts["word/_rels/document.xml.rels"].replace(
        b"</Relationships>", relationship + b"</Relationships>",
    )
    hyperlink = b'<w:hyperlink r:id="rIdExternal"><w:r><w:t>Email</w:t></w:r></w:hyperlink>'
    parts["word/document.xml"] = parts["word/document.xml"].replace(
        b"</w:body>", hyperlink + b"</w:body>",
    )
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)

    prune_unreferenced_media(source)

    with zipfile.ZipFile(source, "r") as archive:
        rels = archive.read("word/_rels/document.xml.rels")
    assert b'Id="rIdExternal"' in rels
    assert b'TargetMode="External"' in rels
