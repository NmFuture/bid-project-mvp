from __future__ import annotations

import io
import json
import sys
import tempfile
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch
from docx import Document

from toc_skill_helpers import (
    BACKEND_ROOT,
    OUTLINE_SCRIPT_DIR,
    TocSkillScriptTestBase,
    complete_outline_review,
    json_load,
    load_outline_script,
    write_decision_context_fixture,
)


class TocSkillScriptTests(TocSkillScriptTestBase):

    def test_bid_outline_skill_is_compact_generic_and_autonomous(self) -> None:
        skill_path = OUTLINE_SCRIPT_DIR.parent / "SKILL.md"
        content = skill_path.read_text(encoding="utf-8")

        self.assertLessEqual(len(content.splitlines()), 200)
        for case_specific_text in (
            "PRJ-0119",
            "5.8.1",
            "附表 G.2.3",
            "附表G.2.3",
            "华能",
            "上海电气",
            "projectId",
            "projectName",
            "projectCode",
        ):
            self.assertNotIn(case_specific_text, content)

        for principle in (
            "历史模板提供成熟投标经验",
            "当前招标文件决定本项目约束",
            "决策只到二级",
            "跟随其二级父节点",
            "父节点保留则整个子树保留",
            "父节点建议删除则整个子树建议删除",
            "一至三级目录",
            "s2outline prepare",
            "s2outline template-headings",
            "s2outline headings",
            "s2outline search",
            "s2outline section",
            "next_cursor",
            "requires_full_review=true",
            "s2outline next-batch",
            "s2outline read",
            "s2outline window",
            "s2outline table",
            "s2outline review-batch",
            "s2outline decision-next",
            "s2outline decision-batch",
            "一个完整决策单元",
            "一个一级章的章根加它下面的全部二级节点",
            "不做章节复核",
            "retain",
            "suggest_delete",
            "招标目录没有同名标题，不等于该节点应删除",
            "独立编制、提交、评审或评分",
            "`additions` 即使为空也必须写 `[]`",
            "s2outline appendix-next",
            "s2outline appendix-decision-batch",
            "--max-items 40",
            "严格保持返回顺序",
            "`source_status=missing` 必须 `exclude`",
            "root_addition",
            "`present`",
            "`missing`",
            "只选 `include` 或 `exclude`",
            "只做一次全局复核",
            "s2outline review-corrections",
            "s2outline review-complete",
            "s2outline decisions",
            "s2outline compose",
            "不要自行写 `manifest.outputFile`",
            "s2outline finalize",
            "timeout=300000",
            "不要检查脚本或包装器",
            "不要执行同功能的 `template`",
            "不要直接读取 `template_structure.json`",
            "内容有投标表达价值，不等于必须独立成章",
            "宽泛父节点不当然覆盖",
            "企业通用能力介绍与本项目专项响应",
            "不得改用 `reason` 规避校验",
            "从招标侧检查遗漏",
            "从模板侧检查不适用、重复或可合并节点",
            "一个短关键词或短语",
            "不能把多个无关关键词拼成一次查询",
            "疑似独立成果必须逐项读原文",
            "附表只覆盖表格填写",
            "不当然覆盖正文方案、说明、报告或承诺",
            "逐项重扫完整招标目录",
            "完整掌握模板结构",
            "先识别响应单元，再比较目录节点",
            "语义等价且粒度相当",
            "单独表达能够让评审人更清楚地看到",
            "父章节能够容纳内容，不等于目录已经覆盖",
            "仅有营销属性不是删除理由",
            "招标目录只用于定位，不能据标题判定覆盖",
            "把它的整个三级子树当成一个整体",
            "不新增三级节点",
            "并行章节会话中，`parent_id` 必须引用当前章根",
            "给了 `evidence_id`，前端就能点击跳转招标原文",
            "会把每个二级决策下沉到它的三级子树",
        ):
            self.assertIn(principle, content)

        self.assertNotIn('"operation":"collapse"', content)
        self.assertNotIn("由父节点统一承载", content)
        self.assertNotIn("模板是主骨架", content)
        self.assertNotIn("仅因招标未提及不能建议删除", content)
        self.assertNotIn("确认没有不适用证据后选择 `retain`", content)
        # 已删除的程序性门禁不应再出现在指令里。
        self.assertNotIn("新的受控正文阅读", content)
        self.assertNotIn("超过 50 个节点的超大章", content)
        self.assertNotIn("--max-items 50", content)

        for redundant_contract in (
            "required_status",
            "review_required",
            "review_note",
            "source_refs",
            "itemId",
            "toc_evidence.json",
            "当前环境不提供 `read` 工具",
            "用 `bash` 调用 `python-docx`",
            "全文审阅是完成条件",
            "不得跳过任何分块",
            "未读完的表格",
        ):
            self.assertNotIn(redundant_contract, content)



    def test_bid_outline_docker_command_exposes_review_navigation(self) -> None:
        # harness-09：wrapper 改由 entrypoint 按 skills/commands.json 运行时生成，
        # 这里直接校验注册表条目；生成行为由 test_skill_command_registration.py 覆盖。
        registry_path = BACKEND_ROOT / "opencode" / "skills" / "commands.json"
        registry = json.loads(registry_path.read_text(encoding="utf-8"))
        entries = {entry["name"]: entry for entry in registry["commands"]}

        commands = (
            "prepare|template|template-headings|headings|search|section|next-batch|read|window|table|tables|"
            "review-batch|decision-next|decision-batch|decision-reopen|review-corrections|"
            "appendix-next|appendix-decision-batch|appendix-predecision-next|"
            "appendix-predecision-batch|review-complete|decisions|compose|"
            "validate|status|finalize"
        )
        entry = entries["s2outline"]
        self.assertEqual("|".join(entry["subcommands"]), commands)
        self.assertEqual(entry["usage"], f"s2outline [{commands}] <manifest> [...]")
        self.assertEqual(entry["prefix_args"], ["--require-compose"])
        self.assertEqual(
            entry["script"], "bid-tech-outline-generator/scripts/run_from_manifest.py"
        )



    def test_bid_outline_agent_navigation_outputs_stay_below_hard_limit(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest, manifest_path = write_decision_context_fixture(
                root,
                heading_count=300,
                heading_text="招标技术要求",
            )
            (root / "tender_appendix_inventory.json").write_text(
                json.dumps(
                    {
                        "schema_version": "tender-appendix-inventory.v1",
                        "items": [
                            {
                                "file_id": "TEN-1",
                                "number": f"附表A.{index}",
                                "title": f"技术响应表{index}",
                                "following_table_count": 1,
                            }
                            for index in range(1, 65)
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            chunks = json_load(root / "tender_review_chunks.json")
            chunks["input_fingerprint"] = "test-navigation-input"
            (root / "tender_review_chunks.json").write_text(
                json.dumps(chunks, ensure_ascii=False), encoding="utf-8"
            )
            files_by_id, _ = outline_runner.review_workflow._collect_heading_files(chunks)
            inventory = json_load(root / "tender_appendix_inventory.json")
            headings_state = json_load(root / "tender_headings_state.json")
            headings_state["input_fingerprint"] = "test-navigation-input"
            headings_state["headings_catalog_digest"] = (
                outline_runner.review_workflow._heading_catalog_digest(
                    files_by_id, inventory["items"]
                )
            )
            (root / "tender_headings_state.json").write_text(
                json.dumps(headings_state, ensure_ascii=False), encoding="utf-8"
            )

            stdout_sizes: dict[str, int] = {}

            def invoke(command: str, *args: str) -> dict:
                stdout = io.StringIO()
                with patch.object(
                    sys,
                    "argv",
                    ["run_from_manifest.py", command, str(manifest_path), *args],
                ), redirect_stdout(stdout):
                    outline_runner.main()
                output = stdout.getvalue()
                stdout_sizes[command] = len(output)
                self.assertLess(stdout_sizes[command], 45000, stdout_sizes)
                self.assertNotIn("output truncated", output.lower())
                return json.loads(output)

            invoke("headings", "--page-size", "80")
            decision_batch = invoke("decision-next")
            self.assertNotIn("comparison_context", decision_batch)
            outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": decision_batch["batch_token"],
                            "items": [
                                {
                                    "target_id": decision_batch["items"][0]["target_id"],
                                    "decision": "retain",
                                    "reason": "历史模板专家经验保留。",
                                }
                            ],
                            "additions": [],
                        }
                    )
                ],
            )
            appendix_batch = invoke("appendix-next", "--max-items", "20")

        self.assertEqual(len(appendix_batch["items"]), 20)
        self.assertEqual(
            set(stdout_sizes),
            {"headings", "decision-next", "appendix-next"},
        )



    def test_bid_outline_headings_page_auto_shrinks_below_output_limit(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _, manifest_path = write_decision_context_fixture(
                root,
                heading_count=2,
                heading_text="超长目录标题" * 3400,
            )
            chunks = json_load(root / "tender_review_chunks.json")
            chunks["input_fingerprint"] = "headings-rollback-input"
            (root / "tender_review_chunks.json").write_text(
                json.dumps(chunks, ensure_ascii=False), encoding="utf-8"
            )
            files_by_id, _ = outline_runner.review_workflow._collect_heading_files(chunks)
            state_path = root / "tender_headings_state.json"
            headings_state = json_load(state_path)
            headings_state["input_fingerprint"] = "headings-rollback-input"
            headings_state["headings_catalog_digest"] = (
                outline_runner.review_workflow._heading_catalog_digest(files_by_id, [])
            )
            state_path.write_text(
                json.dumps(headings_state, ensure_ascii=False), encoding="utf-8"
            )
            # 两条 2 万字符级超长标题 + page-size 2：脚本按字符预算自动收缩为单条返回，
            # 不再触发 45000 字符硬限回滚，按 next_cursor 连续读完。
            first_stdout = io.StringIO()
            with patch.object(
                sys,
                "argv",
                [
                    "run_from_manifest.py",
                    "headings",
                    str(manifest_path),
                    "--page-size",
                    "2",
                ],
            ), redirect_stdout(first_stdout):
                outline_runner.main()
            first = json.loads(first_stdout.getvalue())
            self.assertEqual(first["cursor"], "0")
            self.assertEqual(first["next_cursor"], "1")
            self.assertEqual(first["returned_heading_count"], 1)

            second_stdout = io.StringIO()
            with patch.object(
                sys,
                "argv",
                [
                    "run_from_manifest.py",
                    "headings",
                    str(manifest_path),
                    "--cursor",
                    "1",
                    "--page-size",
                    "2",
                ],
            ), redirect_stdout(second_stdout):
                outline_runner.main()
            second = json.loads(second_stdout.getvalue())

        self.assertEqual(second["cursor"], "1")
        self.assertEqual(second["next_cursor"], "")
        self.assertTrue(second["complete"])
        self.assertEqual(second["returned_heading_count"], 1)



    def test_bid_outline_section_default_budget_stays_below_char_limit_on_chinese_text(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            # 100 段密集中文正文：默认 --max-chars 12000 在字符预算下应整页承载，
            # 每页真实携带 1.2 万字，不再被字节预算压缩成更小的页。
            tender_doc = Document()
            tender_doc.add_paragraph("1 技术要求", style="Heading 1")
            for index in range(100):
                tender_doc.add_paragraph(f"第{index}条 " + "招标技术条款正文" * 60)
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": tender.name, "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)

            headings = outline_runner.dispatch_command("headings", manifest, manifest_path, [])
            section_id = headings["files"][0]["items"][0]["section_id"]
            with self.assertRaisesRegex(SystemExit, r"items\[\]\.section_id"):
                outline_runner.dispatch_command(
                    "section", manifest, manifest_path, ["TEN-1:B000001"]
                )

            cursor = 0
            rounds = 0
            seen_records = 0
            while True:
                stdout = io.StringIO()
                with patch.object(
                    sys,
                    "argv",
                    [
                        "run_from_manifest.py",
                        "section",
                        str(manifest_path),
                        section_id,
                        "--cursor",
                        str(cursor),
                    ],
                ), redirect_stdout(stdout):
                    outline_runner.main()
                output = stdout.getvalue()
                self.assertLess(len(output), 45000)
                page = json.loads(output)
                seen_records += len(page["records"])
                rounds += 1
                if page["complete"]:
                    break
                cursor = int(page["next_cursor"])
                self.assertLess(rounds, 60)

        self.assertGreaterEqual(seen_records, 100)
        # 约 4.9 万字正文按每页 1.2 万字（--max-chars 默认值）读完，页数不应超过 6；
        # 若回退为字节预算，每页只能装约 6600 字，页数会膨胀到 8 页以上。
        self.assertLessEqual(rounds, 6)



    def test_bid_outline_navigation_output_limit_counts_linux_lf_boundary(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            manifest_path = Path(tmp) / "s2_input.json"
            manifest_path.write_text("{}", encoding="utf-8")

            below_limit = {"payload": "x" * 44984}
            stdout = io.StringIO()
            with patch.object(
                sys,
                "argv",
                ["run_from_manifest.py", "headings", str(manifest_path)],
            ), patch.object(
                outline_runner, "dispatch_command", return_value=below_limit
            ), redirect_stdout(stdout):
                outline_runner.main()
            self.assertEqual(len(stdout.getvalue()), 44999)
            self.assertTrue(stdout.getvalue().endswith("\n"))
            self.assertFalse(stdout.getvalue().endswith("\r\n"))

            at_limit = {"payload": "x" * 44985}
            decision_state_path = (
                manifest_path.parent / outline_runner.decision_workflow.STATE_FILE_NAME
            )

            def write_new_state(*_args, **_kwargs):
                decision_state_path.write_bytes(b"new navigation state")
                return at_limit

            for command in ("decision-next",):
                with self.subTest(command=command), patch.object(
                    sys,
                    "argv",
                    ["run_from_manifest.py", command, str(manifest_path)],
                ), patch.object(
                    outline_runner, "dispatch_command", side_effect=write_new_state
                ), redirect_stdout(io.StringIO()), self.assertRaisesRegex(
                    SystemExit, rf"command={command}, actual_chars=45000"
                ):
                    outline_runner.main()
                self.assertFalse(decision_state_path.exists())



    def test_bid_outline_global_review_is_required_and_invalidated_by_reopen(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest, manifest_path = write_decision_context_fixture(root, heading_count=0)
            batch = outline_runner.dispatch_command("decision-next", manifest, manifest_path, [])
            outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": batch["batch_token"],
                            "items": [{"target_id": batch["items"][0]["target_id"], "decision": "retain"}],
                            "additions": [],
                        }
                    )
                ],
            )

            with self.assertRaisesRegex(SystemExit, "全局复核"):
                outline_runner.dispatch_command("decisions", manifest, manifest_path, [])
            reviewed = outline_runner.dispatch_command(
                "review-complete",
                manifest,
                manifest_path,
                [json.dumps({"review_summary": "已对照招标目录完成全局复核。", "issues": []}, ensure_ascii=False)],
            )
            outline_runner.dispatch_command("decisions", manifest, manifest_path, [])
            outline_runner.dispatch_command(
                "decision-reopen", manifest, manifest_path, [batch["chapter_id"]]
            )
            reopened = outline_runner.dispatch_command("decision-next", manifest, manifest_path, [])
            outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": reopened["batch_token"],
                            "items": [{"target_id": reopened["items"][0]["target_id"], "decision": "retain"}],
                            "additions": [],
                        }
                    )
                ],
            )
            with self.assertRaisesRegex(SystemExit, "全局复核"):
                outline_runner.dispatch_command("decisions", manifest, manifest_path, [])

        self.assertTrue(reviewed["review_digest"])



    def test_bid_outline_global_review_completes_without_forced_extra_read(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("1 技术要求", style="Heading 1")
            tender_doc.add_paragraph("投标人应提交完整技术方案。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": tender.name, "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            headings = outline_runner.dispatch_command("headings", manifest, manifest_path, [])
            section_id = headings["files"][0]["items"][0]["section_id"]
            section = outline_runner.dispatch_command(
                "section", manifest, manifest_path, [section_id]
            )
            evidence_id = section["records"][1]["evidence_id"]
            batch = outline_runner.dispatch_command("decision-next", manifest, manifest_path, [])
            outline_runner.dispatch_command(
                "read", manifest, manifest_path, [evidence_id]
            )
            outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": batch["batch_token"],
                            "items": [
                                {
                                    "target_id": batch["items"][0]["target_id"],
                                    "decision": "retain",
                                    "evidence_id": evidence_id,
                                }
                            ],
                            "additions": [],
                        },
                        ensure_ascii=False,
                    )
                ],
            )
            payload = json.dumps(
                {"review_summary": "已从招标侧完成全局查漏。", "issues": []},
                ensure_ascii=False,
            )
            # 查漏由 opencode 自主安排，不再用阅读次数基线当门禁；只有摘要缺失才拒绝。
            with self.assertRaisesRegex(SystemExit, "review_summary is required"):
                outline_runner.dispatch_command(
                    "review-complete",
                    manifest,
                    manifest_path,
                    [json.dumps({"review_summary": "", "issues": []})],
                )
            completed = outline_runner.dispatch_command(
                "review-complete", manifest, manifest_path, [payload]
            )

        self.assertTrue(completed["review_complete"])



    def test_bid_outline_global_review_corrections_apply_omissions_without_reopening_chapter(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("1 专项要求", style="Heading 1")
            tender_doc.add_paragraph("投标人应独立提交项目专项承诺。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": tender.name, "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            headings = outline_runner.dispatch_command("headings", manifest, manifest_path, [])
            section = outline_runner.dispatch_command(
                "section",
                manifest,
                manifest_path,
                [headings["files"][0]["items"][0]["section_id"]],
            )
            evidence_id = section["records"][1]["evidence_id"]
            batch = outline_runner.dispatch_command("decision-next", manifest, manifest_path, [])
            target_id = batch["items"][0]["target_id"]
            outline_runner.dispatch_command("read", manifest, manifest_path, [evidence_id])
            outline_runner.dispatch_command(
                "decision-batch",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "batch_token": batch["batch_token"],
                            "items": [
                                {
                                    "target_id": target_id,
                                    "decision": "retain",
                                    "evidence_id": evidence_id,
                                }
                            ],
                            "additions": [],
                        }
                    )
                ],
            )
            outline_runner.dispatch_command("read", manifest, manifest_path, [evidence_id])

            corrected = outline_runner.dispatch_command(
                "review-corrections",
                manifest,
                manifest_path,
                [
                    json.dumps(
                        {
                            "items": [
                                {
                                    "target_id": target_id,
                                    "decision": "suggest_delete",
                                    "reason": "全局复核发现该节点与新增专项重复。",
                                }
                            ],
                            "additions": [
                                {
                                    "node_id": "ADD-REVIEW-1",
                                    "parent_id": target_id,
                                    "number": "1.1",
                                    "title": "项目专项承诺",
                                    "reason": "全局复核发现招标要求独立提交。",
                                    "evidence_id": evidence_id,
                                }
                            ],
                        },
                        ensure_ascii=False,
                    )
                ],
            )
            with self.assertRaisesRegex(SystemExit, "全局复核"):
                outline_runner.dispatch_command("decisions", manifest, manifest_path, [])
            outline_runner.dispatch_command("read", manifest, manifest_path, [evidence_id])
            complete_outline_review(outline_runner, manifest, manifest_path)
            decisions = outline_runner.dispatch_command(
                "decisions", manifest, manifest_path, []
            )
            decision_payload = json_load(Path(decisions["decisionsFile"]))

        self.assertEqual(corrected["corrected_item_count"], 1)
        self.assertEqual(corrected["added_count"], 1)
        self.assertEqual(decision_payload["changes"][0]["node_id"], "ADD-REVIEW-1")
        self.assertEqual(
            decision_payload["template_decisions"][0],
            {
                "target_id": target_id,
                "decision": "suggest_delete",
                "reason": "全局复核发现该节点与新增专项重复。",
            },
        )



    def test_bid_outline_review_navigation_resumes_from_first_pending_chunk(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")
        review_workflow = load_outline_script("review_workflow")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("1. 总则", style="Heading 1")
            tender_doc.add_paragraph("投标人应提供总体技术方案。")
            tender_doc.add_paragraph("2. 专题", style="Heading 1")
            tender_doc.add_paragraph("投标人应提供场址安全适应性报告。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)

            first = review_workflow.next_review_chunk(root)
            review_workflow.submit_chunk_review(
                root,
                first["chunk"]["chunk_id"],
                {"review_summary": "已审阅总则要求。", "requirements": []},
            )
            second = review_workflow.next_review_chunk(root)
            state = review_workflow.review_status(root)

        self.assertNotEqual(first["chunk"]["chunk_id"], second["chunk"]["chunk_id"])
        self.assertEqual(first["remaining_chunk_count"], 2)
        self.assertEqual(second["remaining_chunk_count"], 1)
        self.assertEqual(state["reviewed_chunk_count"], 1)
        self.assertEqual(state["pending_chunk_count"], 1)



    def test_bid_outline_chunk_review_validates_dynamic_requirement_dispositions(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")
        review_workflow = load_outline_script("review_workflow")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("5.8 项目风机各子系统专题", style="Heading 1")
            tender_doc.add_paragraph("投标人应提供叶片专题。")
            tender_doc.add_paragraph("投标人应提供场址安全适应性报告。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            outline_runner.write_template_structure(manifest, manifest_path)
            chunk = review_workflow.next_review_chunk(root)["chunk"]
            paragraph_ids = [block["evidence_id"] for block in chunk["blocks"] if block["type"] == "paragraph"]

            with self.assertRaisesRegex(SystemExit, "target_node"):
                review_workflow.submit_chunk_review(
                    root,
                    chunk["chunk_id"],
                    {
                        "review_summary": "已识别子系统专题。",
                        "requirements": [
                            {
                                "evidence_ids": [paragraph_ids[1]],
                                "obligation": "投标人应提供叶片专题",
                                "disposition": "map_existing",
                            }
                        ],
                    },
                )
            with self.assertRaisesRegex(SystemExit, "不属于当前分块"):
                review_workflow.submit_chunk_review(
                    root,
                    chunk["chunk_id"],
                    {
                        "review_summary": "证据校验。",
                        "requirements": [
                            {
                                "evidence_ids": ["TEN-1:B999999"],
                                "obligation": "不存在的义务",
                                "disposition": "reference_only",
                                "reason": "仅作参考",
                            }
                        ],
                    },
                )

            result = review_workflow.submit_chunk_review(
                root,
                chunk["chunk_id"],
                {
                    "review_summary": "已逐项判断两个独立响应要求。",
                    "requirements": [
                        {
                            "evidence_ids": [paragraph_ids[1]],
                            "obligation": "投标人应提供叶片专题",
                            "disposition": "map_existing",
                            "target_node": "5.8",
                        },
                        {
                            "evidence_ids": [paragraph_ids[2]],
                            "obligation": "投标人应提供场址安全适应性报告",
                            "disposition": "suggest_add",
                            "target_node": "5.7",
                            "proposed_title": "项目场址安全适应性报告",
                            "reason": "招标明确要求独立报告，模板没有语义等价节点。",
                        },
                    ],
                },
            )
            ledger = json_load(root / "requirement_ledger.json")

        self.assertEqual(result["added_requirement_count"], 2)
        self.assertEqual(ledger["requirement_count"], 2)
        self.assertEqual(
            [item["disposition"] for item in ledger["requirements"]],
            ["map_existing", "suggest_add"],
        )
        self.assertEqual(ledger["requirements"][1]["proposed_title"], "项目场址安全适应性报告")



    def test_bid_outline_agentic_commands_drive_review_without_raw_file_reads(self) -> None:
        outline_runner = load_outline_script("run_from_manifest")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            tender = root / "tender.docx"
            manifest_path = root / "s2_input.json"
            template_doc = Document()
            template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
            template_doc.save(template)
            tender_doc = Document()
            tender_doc.add_paragraph("1. 总则", style="Heading 1")
            tender_doc.add_paragraph("投标人应提供总体技术方案。")
            tender_doc.save(tender)
            manifest = {
                "workDir": str(root),
                "templateFile": str(template),
                "tenderFiles": [{"id": "TEN-1", "name": "tender.docx", "path": str(tender)}],
                "outputFile": str(root / "toc.json"),
            }
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

            prepared = outline_runner.dispatch_command("prepare", manifest, manifest_path, [])
            next_payload = outline_runner.dispatch_command("next", manifest, manifest_path, [])
            chunk = next_payload["chunk"]
            evidence_id = chunk["blocks"][1]["evidence_id"]
            read_payload = outline_runner.dispatch_command("read", manifest, manifest_path, [evidence_id])
            window_payload = outline_runner.dispatch_command(
                "window",
                manifest,
                manifest_path,
                [evidence_id, "--before", "1", "--after", "1"],
            )
            submitted = outline_runner.dispatch_command(
                "review-chunk",
                manifest,
                manifest_path,
                [
                    chunk["chunk_id"],
                    json.dumps(
                        {
                            "review_summary": "总体技术方案由模板节点承接。",
                            "requirements": [
                                {
                                    "evidence_ids": [evidence_id],
                                    "obligation": "投标人应提供总体技术方案",
                                    "disposition": "map_existing",
                                    "target_node": "第1章",
                                }
                            ],
                        },
                        ensure_ascii=False,
                    ),
                ],
            )
            status = outline_runner.dispatch_command("status", manifest, manifest_path, [])

        self.assertEqual(prepared["tenderReviewChunkCount"], 1)
        self.assertEqual(read_payload["record"]["evidence_id"], evidence_id)
        self.assertEqual(len(window_payload["blocks"]), 2)
        self.assertEqual(submitted["status"], "reviewed")
        self.assertEqual(status["review_coverage"], 1.0)
        self.assertEqual(status["requirement_count"], 1)
