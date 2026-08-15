from __future__ import annotations

import re
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from toc_skill_helpers import (
    TocSkillScriptTestBase,
    load_assembler_script,
)


class TocSkillScriptTests(TocSkillScriptTestBase):

    def test_bid_assembler_parse_toc_accepts_current_s2_json(self) -> None:
        parse_toc = load_assembler_script("parse_toc")

        with tempfile.TemporaryDirectory() as tmp:
            toc_path = Path(tmp) / "投标文件-总目录.json"
            toc_path.write_text(
                """{
  "schema_version": "bid-toc-json-v1",
  "document_title": "测试项目投标文件总目录",
  "items": [
    {"order": 1, "level": 1, "number": "1", "title": "项目概况", "annotation": "保留"},
    {"order": 2, "level": 2, "number": "1.1", "title": "项目背景", "annotation": "适配"},
    {"order": 3, "level": 1, "number": "附表", "title": "", "annotation": "保留"}
  ]
}""",
                encoding="utf-8",
            )

            entries = parse_toc.parse_toc_json(toc_path)

        self.assertEqual(entries[0]["chapter_no_flat"], "1")
        self.assertEqual(entries[0]["title"], "项目概况")
        self.assertEqual(entries[1]["tag"], "适配")
        self.assertEqual(entries[2]["title"], "附表")



    def test_bid_assembler_inject_prefix_collapses_heading_level_gaps(self) -> None:
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        doc.add_paragraph("上海电气优势简介", style="Heading 2")
        doc.add_paragraph("基本情况", style="Heading 3")
        doc.add_paragraph("集团概况", style="Heading 4")
        doc.add_paragraph("载荷仿真分析能力", style="Heading 1")
        doc.add_paragraph("测试验证技术", style="Heading 6")

        stats = numbering_fixer.inject_prefix_to_headings(
            doc,
            "1.9",
            toc_title="上海电气优势简介",
            skip_first_if_match=True,
        )
        headings = [para.text.strip().replace("  ", " ") for para in doc.paragraphs if para.text.strip()]

        self.assertTrue(stats["skipped_first"])
        self.assertEqual(
            headings,
            [
                "1.9.1 基本情况",
                "1.9.1.1 集团概况",
                "1.9.2 载荷仿真分析能力",
                "1.9.2.1 测试验证技术",
            ],
        )
        self.assertFalse(any(".0." in item or item.endswith(".0") for item in headings))



    def test_bid_assembler_inject_prefix_preserves_style_and_updates_navigation(self) -> None:
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        doc.add_paragraph("上海电气优势简介", style="Heading 2")
        stale = doc.add_paragraph("载荷仿真分析能力", style="Heading 1")
        p_pr = stale._p.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "0")
        p_pr.append(outline)

        numbering_fixer.inject_prefix_to_headings(
            doc,
            "1.9",
            toc_title="上海电气优势简介",
            skip_first_if_match=True,
        )

        remaining = [para for para in doc.paragraphs if "载荷仿真分析能力" in para.text][0]
        self.assertEqual(remaining.style.name, "Heading 1")
        p_pr = remaining._p.find(qn("w:pPr"))
        self.assertIsNotNone(p_pr)
        self.assertEqual(p_pr.find(qn("w:outlineLvl")).get(qn("w:val")), "2")



    def test_bid_assembler_demotes_material_headings_to_body(self) -> None:
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        doc.add_paragraph("上海电气优势简介", style="Heading 2")
        doc.add_paragraph("基本情况", style="Heading 3")
        stale = doc.add_paragraph("载荷仿真分析能力", style="Heading 1")
        p_pr = stale._p.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "0")
        p_pr.append(outline)

        stats = numbering_fixer.demote_headings_to_body(
            doc,
            toc_title="上海电气优势简介",
            remove_first_if_match=True,
        )

        self.assertEqual(stats["removed"], 1)
        self.assertEqual(stats["demoted"], 2)
        self.assertEqual([para.text for para in doc.paragraphs], ["基本情况", "载荷仿真分析能力"])
        self.assertFalse(any((para.style.name or "").startswith("Heading") for para in doc.paragraphs))
        self.assertTrue(all(para._p.find(qn("w:pPr")).find(qn("w:outlineLvl")) is None for para in doc.paragraphs))



    def test_bid_assembler_demotes_direct_outline_without_body_style(self) -> None:
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        styled = doc.add_paragraph("发电小时数承诺函", style="Heading 2")
        direct_only = doc.add_paragraph("发电量保证矩阵表")
        direct_p_pr = direct_only._p.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "0")
        direct_p_pr.append(outline)

        styles_el = doc.styles._element
        for style in list(styles_el.findall(qn("w:style"))):
            if style.get(qn("w:styleId")) == "Normal":
                styles_el.remove(style)

        stats = numbering_fixer.demote_headings_to_body(doc)

        self.assertEqual(stats["demoted"], 2)
        for para in (styled, direct_only):
            p_pr = para._p.find(qn("w:pPr"))
            self.assertIsNotNone(p_pr)
            self.assertIsNone(p_pr.find(qn("w:pStyle")))
            self.assertIsNone(p_pr.find(qn("w:outlineLvl")))



    def test_bid_assembler_keeps_s2_child_headings_from_material(self) -> None:
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        doc.add_paragraph("项目技术承诺函", style="Heading 1")
        child = doc.add_paragraph("发电小时数承诺函", style="Heading 2")
        extra = doc.add_paragraph("发电量保证矩阵表")
        p_pr = extra._p.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "0")
        p_pr.append(outline)

        stats = numbering_fixer.demote_headings_to_body(
            doc,
            toc_title="项目技术承诺函",
            remove_first_if_match=True,
            keep_heading_map={
                "发电小时数承诺函": {
                    "chapter_no": "4.1",
                    "title": "发电小时数承诺函",
                    "level": 2,
                }
            },
        )

        self.assertEqual(stats["removed"], 1)
        self.assertEqual(stats["kept"], 1)
        self.assertEqual(stats["demoted"], 1)
        self.assertEqual(child.text.strip().replace("  ", " "), "4.1 发电小时数承诺函")
        self.assertEqual(child.style.name, "Heading 2")
        extra_p_pr = extra._p.find(qn("w:pPr"))
        self.assertIsNone(extra_p_pr.find(qn("w:pStyle")))
        self.assertIsNone(extra_p_pr.find(qn("w:outlineLvl")))



    def test_bid_assembler_remaps_material_headings_to_navigation(self) -> None:
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        doc.add_paragraph("设备运行和维护专题", style="Heading 2")
        project_flow = doc.add_paragraph("项目流程", style="Heading 1")
        list_item = doc.add_paragraph("（3）机组调试")
        static = doc.add_paragraph()
        static.add_run("静态调试").bold = True
        table_heading = doc.add_paragraph("表C.1 总体技术参数与规格", style="Heading 2")

        stats = numbering_fixer.remap_material_headings_to_navigation(
            doc,
            toc_title="设备运行和维护专题",
            remove_first_if_match=True,
            parent_level=2,
        )

        self.assertEqual(stats["removed"], 1)
        self.assertEqual(stats["remapped"], 1)
        self.assertEqual(stats["bold_subheadings"], 0)
        self.assertEqual(stats["demoted"], 1)
        self.assertEqual(project_flow.style.name, "Heading 1")
        self.assertEqual(
            project_flow._p.find(qn("w:pPr")).find(qn("w:outlineLvl")).get(qn("w:val")),
            "2",
        )
        self.assertFalse((list_item.style.name or "").startswith("Heading"))
        self.assertFalse((static.style.name or "").startswith("Heading"))
        self.assertFalse((table_heading.style.name or "").startswith("Heading"))



    def test_bid_assembler_strips_heading_style_numbering(self) -> None:
        numbering_fixer = load_assembler_script("numbering_fixer")

        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "numbered-heading.docx"
            doc = Document()
            heading_style = doc.styles["Heading 2"]
            p_pr = heading_style.element.get_or_add_pPr()
            num_pr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "1")
            num_id = OxmlElement("w:numId")
            num_id.set(qn("w:val"), "1")
            num_pr.append(ilvl)
            num_pr.append(num_id)
            p_pr.append(num_pr)
            doc.add_paragraph("1.7 投标方案优势说明", style="Heading 2")

            self.assertEqual(numbering_fixer.strip_numPr_from_heading_styles(doc), 1)
            doc.save(docx_path)

            with zipfile.ZipFile(docx_path) as zf:
                styles_xml = zf.read("word/styles.xml").decode("utf-8")

        match = re.search(r'(<w:style[^>]+w:styleId="Heading2"[^>]*>.*?</w:style>)', styles_xml)
        self.assertIsNotNone(match)
        self.assertNotIn("<w:numPr>", match.group(1))



    def test_bid_assembler_remap_preserves_numid_zero_suppression(self) -> None:
        """1.7 回归：注入层级时不得删除抑制隐藏自动编号的 numId=0。"""
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        self._make_hidden_numbered_custom_style(doc, "标题6-标书", "Heading 6", "7")
        para = doc.add_paragraph("总体技术路线", style="标题6-标书")
        self._add_paragraph_numpr(para, "0")  # 源文档用 numId=0 抑制样式隐藏编号

        stats = numbering_fixer.remap_material_headings_to_navigation(doc, parent_level=2)

        self.assertEqual(stats["remapped"], 1)
        num_pr = para._p.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertIsNotNone(num_pr)
        self.assertEqual(num_pr.find(qn("w:numId")).get(qn("w:val")), "0")



    def test_bid_assembler_remap_strips_active_paragraph_numbering(self) -> None:
        """注入层级时真正生效的段落自动编号（numId>0）仍要剥掉，避免双编号。"""
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        para = doc.add_paragraph("总体技术路线", style="Heading 6")
        self._add_paragraph_numpr(para, "7")

        stats = numbering_fixer.remap_material_headings_to_navigation(doc, parent_level=2)

        self.assertEqual(stats["remapped"], 1)
        num_pr = para._p.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertIsNone(num_pr)



    def test_bid_assembler_strips_basedon_chain_heading_style_numbering(self) -> None:
        """1.7 回归：basedOn 链指向 Heading 的自定义样式也要剥样式级自动编号。"""
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        custom = self._make_hidden_numbered_custom_style(doc, "标题6-标书", "Heading 6", "7")

        self.assertEqual(numbering_fixer.strip_numPr_from_heading_styles(doc), 1)
        self.assertIsNone(custom.element.find(qn("w:pPr")).find(qn("w:numPr")))



    def test_bid_assembler_enforce_invariant_clears_residual_numbering(self) -> None:
        """不变量：已写入文本编号的 Heading 不得再有有效 Word 自动编号。"""
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        custom = self._make_hidden_numbered_custom_style(doc, "标题6-标书", "Heading 6", "7")
        # 场景1：段落 numId>0 → 改为段落级 numId=0
        active = doc.add_paragraph("1.7.3.1 总体技术路线", style="Heading 3")
        self._add_paragraph_numpr(active, "5")
        # 场景2：段落 numId=0 抑制仍在 → 不动
        suppressed = doc.add_paragraph("1.7.3.2 关键技术路线", style="标题6-标书")
        self._add_paragraph_numpr(suppressed, "0")
        # 场景3：样式链编号且无段落抑制 → 只给当前段落加 numId=0
        inherited = doc.add_paragraph("1.7.3.3 其他技术路线", style="标题6-标书")

        fixed = numbering_fixer.enforce_no_auto_numbering_on_numbered_headings(doc)

        self.assertEqual(fixed, 2)
        active_num_pr = active._p.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertEqual(active_num_pr.find(qn("w:numId")).get(qn("w:val")), "0")
        suppressed_num_pr = suppressed._p.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertIsNotNone(suppressed_num_pr)
        self.assertEqual(suppressed_num_pr.find(qn("w:numId")).get(qn("w:val")), "0")
        inherited_num_pr = inherited._p.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertEqual(inherited_num_pr.find(qn("w:numId")).get(qn("w:val")), "0")
        style_num_pr = custom.element.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertEqual(style_num_pr.find(qn("w:numId")).get(qn("w:val")), "7")



    def test_bid_assembler_invariant_keeps_shared_body_list_style_numbering(self) -> None:
        """标题局部抑制自动编号时，不得破坏同样式的正文列表。"""
        numbering_fixer = load_assembler_script("numbering_fixer")

        doc = Document()
        shared = self._make_hidden_numbered_custom_style(
            doc,
            "共享正文列表",
            "Normal",
            "7",
        )
        heading = doc.add_paragraph("1.7.4 列表样式标题", style=shared)
        heading_p_pr = heading._p.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "2")
        heading_p_pr.append(outline)
        body_item = doc.add_paragraph("正文列表项", style=shared)

        fixed = numbering_fixer.enforce_no_auto_numbering_on_numbered_headings(doc)

        self.assertEqual(fixed, 1)
        heading_num_pr = heading._p.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertEqual(heading_num_pr.find(qn("w:numId")).get(qn("w:val")), "0")
        style_num_pr = shared.element.find(qn("w:pPr")).find(qn("w:numPr"))
        self.assertEqual(style_num_pr.find(qn("w:numId")).get(qn("w:val")), "7")
        body_p_pr = body_item._p.find(qn("w:pPr"))
        self.assertIsNotNone(body_p_pr)
        self.assertIsNone(body_p_pr.find(qn("w:numPr")))



    def test_bid_assembler_merges_oversize_material_instead_of_placeholder(self) -> None:
        merger = load_assembler_script("merger")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            source = root / "lib" / "投标机型业绩情况.docx"
            out = root / "out.docx"
            prep = root / "prep"

            master = Document()
            master.add_paragraph("")
            master.save(template)

            source.parent.mkdir(parents=True)
            doc = Document()
            doc.add_paragraph("投标机型业绩情况", style="Heading 1")
            doc.add_paragraph("这里是业绩正文")
            doc.save(source)

            plan = [
                {
                    "status": "MATCHED",
                    "level": 2,
                    "title": "投标机型业绩情况",
                    "chapter_no": "1.8",
                    "chapter_no_flat": "1.8",
                    "paths": [source.name],
                }
            ]
            original_stat = Path.stat

            class FakeStat:
                def __init__(self, wrapped):
                    self._wrapped = wrapped
                    self.st_size = 234 * 1024 * 1024

                def __getattr__(self, name):
                    return getattr(self._wrapped, name)

            def fake_stat(path, *args, **kwargs):
                value = original_stat(path, *args, **kwargs)
                if Path(path).resolve() == source.resolve():
                    return FakeStat(value)
                return value

            with patch.object(Path, "stat", fake_stat):
                stats = merger.merge(template, plan, source.parent, {}, prep, out)

            result = Document(str(out))
            text = "\n".join(para.text for para in result.paragraphs)

        self.assertEqual(stats["merged_materials"], 1)
        self.assertNotIn("大素材跳过", text)
        self.assertIn("这里是业绩正文", text)



    def test_bid_assembler_merger_remaps_material_headings_in_navigation(self) -> None:
        merger = load_assembler_script("merger")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            source = root / "lib" / "上海电气优势简介.docx"
            out = root / "out.docx"
            prep = root / "prep"

            master = Document()
            master.add_paragraph("")
            master.save(template)

            source.parent.mkdir(parents=True)
            doc = Document()
            doc.add_paragraph("上海电气优势简介", style="Heading 2")
            doc.add_paragraph("基本情况", style="Heading 3")
            stale = doc.add_paragraph("载荷仿真分析能力", style="Heading 1")
            p_pr = stale._p.get_or_add_pPr()
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), "0")
            p_pr.append(outline)
            doc.add_paragraph("这里是正文")
            doc.save(source)

            plan = [
                {
                    "status": "MATCHED",
                    "level": 2,
                    "title": "上海电气优势简介",
                    "chapter_no": "1.9",
                    "chapter_no_flat": "1.9",
                    "paths": [source.name],
                }
            ]
            merger.merge(template, plan, source.parent, {}, prep, out)

            result = Document(str(out))
            headings = [
                para.text.strip().replace("  ", " ")
                for para in result.paragraphs
                if (para.style.name or "").startswith("Heading") and para.text.strip()
            ]
            text = "\n".join(para.text for para in result.paragraphs)

        self.assertEqual(
            headings,
            ["1.9 上海电气优势简介", "1.9.1 基本情况", "1.9.2 载荷仿真分析能力"],
        )
        self.assertIn("基本情况", text)
        self.assertIn("载荷仿真分析能力", text)



    def test_bid_assembler_merger_keeps_matching_child_heading_in_place(self) -> None:
        merger = load_assembler_script("merger")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / "template.docx"
            source = root / "lib" / "项目技术承诺函.docx"
            out = root / "out.docx"
            prep = root / "prep"

            master = Document()
            master.add_paragraph("")
            master.save(template)

            source.parent.mkdir(parents=True)
            doc = Document()
            doc.add_paragraph("项目技术承诺函", style="Heading 1")
            doc.add_paragraph("发电小时数承诺函", style="Heading 2")
            doc.add_paragraph("这里是 4.1 正文")
            doc.save(source)

            plan = [
                {
                    "status": "MATCHED",
                    "level": 1,
                    "title": "项目技术承诺函",
                    "chapter_no": "第四章",
                    "chapter_no_flat": "4",
                    "paths": [source.name],
                },
                {
                    "status": "UNMATCHED",
                    "level": 2,
                    "title": "发电小时数承诺函",
                    "chapter_no": "4.1",
                    "chapter_no_flat": "4.1",
                    "paths": [],
                },
            ]
            merger.merge(template, plan, source.parent, {}, prep, out)

            result = Document(str(out))
            headings = [
                para.text.strip().replace("  ", " ")
                for para in result.paragraphs
                if (para.style.name or "").startswith("Heading") and para.text.strip()
            ]
            text = "\n".join(para.text for para in result.paragraphs)

        self.assertEqual(headings, ["第四章 项目技术承诺函", "4.1 发电小时数承诺函"])
        self.assertIn("这里是 4.1 正文", text)
        self.assertNotIn("[缺失：发电小时数承诺函", text)



    def test_bid_assembler_verify_uses_direct_outline_level_priority(self) -> None:
        verify = load_assembler_script("verify")

        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "stale-outline.docx"
            doc = Document()
            stale = doc.add_paragraph("1.9.4  载荷仿真分析能力", style="Heading 3")
            p_pr = stale._p.get_or_add_pPr()
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), "0")
            p_pr.append(outline)
            doc.save(docx_path)

            scan = verify.scan_docx(docx_path)

        self.assertEqual(scan["heading_counts"], {"Heading 1": 1})
        self.assertIn("1.9.4  载荷仿真分析能力", scan["invalid_h1"])
