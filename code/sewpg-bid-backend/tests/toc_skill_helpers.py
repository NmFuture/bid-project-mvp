"""test_toc_skill_* 共享夹具：skill 脚本目录定位、脚本加载器、评审/决策夹具与基类。

由 test_toc_skill_scripts.py 拆分而来，只搬不改。
"""
from __future__ import annotations

import importlib.util
import json
import sys
import unittest
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BACKEND_ROOT = Path(__file__).resolve().parents[1]


def skill_script_dir(skill_name: str) -> Path:
    for root_name in ("skills", "skill"):
        candidate = BACKEND_ROOT / "opencode" / root_name / skill_name / "scripts"
        if candidate.exists():
            return candidate
    return BACKEND_ROOT / "opencode" / "skills" / skill_name / "scripts"


ASSEMBLER_SCRIPT_DIR = skill_script_dir("bid-tech-assembler")


OUTLINE_SCRIPT_DIR = skill_script_dir("bid-tech-outline-generator")


WIKI_SCRIPT_DIR = skill_script_dir("bid-tech-wiki-material-builder")


GAP_PLANNER_SCRIPT_DIR = skill_script_dir("bid-tech-gap-planner")


WORD_FILLER_SCRIPT_DIR = skill_script_dir("bid-tech-word-placeholder-filler")


MATERIAL_CLEANER_SCRIPT_DIR = skill_script_dir("bid-material-format-cleaner")


def load_assembler_script(name: str):
    spec = importlib.util.spec_from_file_location(name, ASSEMBLER_SCRIPT_DIR / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def load_outline_script(name: str):
    module_name = f"outline_{name}"
    spec = importlib.util.spec_from_file_location(module_name, OUTLINE_SCRIPT_DIR / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


def load_wiki_script(name: str):
    module_name = f"wiki_{name}"
    spec = importlib.util.spec_from_file_location(module_name, WIKI_SCRIPT_DIR / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


def load_gap_planner_script(name: str):
    module_name = f"gap_planner_{name}"
    spec = importlib.util.spec_from_file_location(module_name, GAP_PLANNER_SCRIPT_DIR / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


def load_word_filler_script(name: str):
    module_name = f"word_filler_{name}"
    spec = importlib.util.spec_from_file_location(module_name, WORD_FILLER_SCRIPT_DIR / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


def load_material_cleaner_script(name: str):
    module_name = f"material_cleaner_{name}"
    spec = importlib.util.spec_from_file_location(module_name, MATERIAL_CLEANER_SCRIPT_DIR / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


def json_load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def complete_outline_review(outline_runner, manifest: dict, manifest_path: Path) -> dict:
    return outline_runner.dispatch_command(
        "review-complete",
        manifest,
        manifest_path,
        [json.dumps({"review_summary": "已完成全局复核。", "issues": []}, ensure_ascii=False)],
    )


def write_decision_context_fixture(
    root: Path,
    *,
    heading_count: int = 80,
    heading_text: str | None = None,
    template_count: int = 1,
    template_title: str = "技术方案",
) -> tuple[dict, Path]:
    manifest_path = root / "s2_input.json"
    manifest = {"workDir": str(root), "outputFile": str(root / "toc.json")}
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    (root / "template_structure.json").write_text(
        json.dumps(
            {
                "schema_version": "template-structure.v1",
                "items": [
                    {
                        "number": f"第{index}章",
                        "title": f"{template_title}{index}",
                        "level": 1,
                    }
                    for index in range(1, template_count + 1)
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    blocks = [
        {
            "type": "paragraph",
            "evidence_id": f"TEN-1:B{index:06d}",
            "body_index": index,
            "text": heading_text or f"{index} 招标技术要求 " + ("详细条款" * 70),
            "toc_level": 1,
            "heading_level": 0,
            "structural_title_level": 1,
        }
        for index in range(1, heading_count + 1)
    ]
    (root / "tender_review_chunks.json").write_text(
        json.dumps(
            {
                "schema_version": "tender-review-chunks.v1",
                "chunks": [
                    {
                        "file_id": "TEN-1",
                        "file_name": "招标文件.docx",
                        "blocks": blocks,
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (root / "tender_headings_state.json").write_text(
        json.dumps(
            {
                "schema_version": "tender-headings-state.v1",
                "next_cursor": 0,
                "complete": True,
            }
        ),
        encoding="utf-8",
    )
    return manifest, manifest_path


def submit_outline_changes(outline_runner, manifest: dict, manifest_path: Path, changes: list[dict]):
    structure = json_load(Path(manifest["workDir"]) / "template_structure.json")
    annotated = outline_runner.outline_composer.annotate_template_structure(structure)
    fingerprint = annotated["input_fingerprint"]
    deleted = {
        str(change.get("target_id") or ""): change
        for change in changes
        if change.get("operation") == "suggest_delete"
    }
    decision_max_level = outline_runner.outline_composer.DECISION_MAX_LEVEL
    template_decisions = []
    for item in annotated["items"]:
        if int(item.get("level") or 1) > decision_max_level:
            continue
        target_id = item["template_id"]
        change = deleted.get(target_id)
        if change:
            decision = {
                "target_id": target_id,
                "decision": "suggest_delete",
                "reason": change["reason"],
            }
            if "tender_basis" in change:
                decision["tender_basis"] = change["tender_basis"]
        else:
            decision = {"target_id": target_id, "decision": "retain"}
        template_decisions.append(decision)
    return outline_runner.submit_outline_decisions(
        manifest,
        manifest_path,
        {
            "schema_version": "technical-outline-decisions.v1",
            "input_fingerprint": fingerprint,
            "template_decisions": template_decisions,
            "changes": changes,
        },
    )


if __name__ == "__main__":
    unittest.main()


class TocSkillScriptTestBase(unittest.TestCase):

    def _prepare_finalized_appendix_outline(
        self,
        root: Path,
        decisions: list[str],
        *,
        mutate_state=None,
    ) -> tuple[object, dict, Path]:
        outline_runner = load_outline_script("run_from_manifest")
        template = root / "template.docx"
        tender = root / "tender.docx"
        output = root / "toc.json"
        manifest_path = root / "s2_input.json"

        template_doc = Document()
        template_doc.add_paragraph("第1章 技术方案", style="Heading 1")
        template_doc.save(template)
        tender_doc = Document()
        for number, title in (("附表A.1", "技术参数表"), ("附表A.2", "供货范围表")):
            tender_doc.add_paragraph(f"{number} {title}")
            table = tender_doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "参数"
            table.cell(0, 1).text = "要求"
            table.cell(1, 0).text = "示例"
            table.cell(1, 1).text = "投标人填写"
        tender_doc.save(tender)
        manifest = {
            "workDir": str(root),
            "templateFile": str(template),
            "tenderFiles": [{"id": "TEN-1", "name": tender.name, "path": str(tender)}],
            "outputFile": str(output),
            "requireComposedOutline": True,
        }
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        outline_runner.write_template_structure(manifest, manifest_path)
        outline_runner.dispatch_command("headings", manifest, manifest_path, [])
        batch = outline_runner.dispatch_command("decision-next", manifest, manifest_path, [])
        appendix_heading = outline_runner.dispatch_command(
            "search", manifest, manifest_path, ["附表A.1"]
        )
        outline_runner.dispatch_command(
            "read",
            manifest,
            manifest_path,
            [appendix_heading["results"][0]["evidence_id"]],
        )
        outline_runner.dispatch_command(
            "decision-batch",
            manifest,
            manifest_path,
            [
                json.dumps(
                    {
                        "batch_token": batch["batch_token"],
                        "items": [
                            {
                                "target_id": item["target_id"],
                                "decision": "retain",
                                "reason": "历史模板专家经验保留。",
                            }
                            for item in batch["items"]
                        ],
                        "additions": [],
                    },
                    ensure_ascii=False,
                )
            ],
        )
        appendix_batch = outline_runner.dispatch_command(
            "appendix-next", manifest, manifest_path, []
        )
        self.assertEqual(len(appendix_batch["items"]), len(decisions))
        appendix_payload = []
        for index, (item, decision) in enumerate(zip(appendix_batch["items"], decisions)):
            payload = {
                "appendix_id": item["appendix_id"],
                "decision": decision,
                "reason": "逐项核验招标附表。",
            }
            if decision == "include":
                payload.update(
                    {
                        "node_id": f"ADD-APPENDIX-{index + 1}",
                        "parent_id": "ADD-APPENDIX",
                    }
                )
            appendix_payload.append(payload)
        appendix_request = {
            "batch_token": appendix_batch["batch_token"],
            "items": appendix_payload,
        }
        if "include" in decisions:
            appendix_request["root_addition"] = {
                "node_id": "ADD-APPENDIX",
                "parent_id": None,
                "number": "第2章",
                "title": "技术附表",
                "reason": "招标文件包含独立附表。",
            }
        outline_runner.dispatch_command(
            "appendix-decision-batch",
            manifest,
            manifest_path,
            [
                json.dumps(
                    appendix_request,
                    ensure_ascii=False,
                )
            ],
        )
        evidence_id = str(appendix_batch["items"][0].get("evidence_id") or "")
        if evidence_id:
            outline_runner.dispatch_command(
                "read", manifest, manifest_path, [evidence_id]
            )
        if mutate_state is not None:
            state_path = root / "outline_decision_state.json"
            state = json_load(state_path)
            mutate_state(state, appendix_batch["items"])
            state_path.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
        complete_outline_review(outline_runner, manifest, manifest_path)
        outline_runner.dispatch_command("decisions", manifest, manifest_path, [])
        outline_runner.compose_manifest(manifest, manifest_path)
        return outline_runner, manifest, manifest_path



    def _make_hidden_numbered_custom_style(self, doc, style_name: str, base_name: str, num_id: str):
        """构造带隐藏自动编号的自定义标题样式（如"标题6-标书" basedOn Heading 6）。"""
        custom = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        custom.base_style = doc.styles[base_name]
        p_pr = custom.element.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), num_id)
        num_pr.append(ilvl)
        num_pr.append(nid)
        p_pr.append(num_pr)
        return custom



    def _add_paragraph_numpr(self, para, num_id: str) -> None:
        p_pr = para._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), num_id)
        num_pr.append(ilvl)
        num_pr.append(nid)
        p_pr.append(num_pr)
