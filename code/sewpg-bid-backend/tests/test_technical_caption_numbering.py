from __future__ import annotations

import json
import struct
import tempfile
import unittest
import zlib
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from app.document_processing.technical_document import captioning as runner
from app.document_processing.technical_document.captioning import captions


def _write_png(path: Path) -> Path:
    """造一张极小的 PNG，靠 add_picture 的显式宽高控制是否算"可编号大图"。"""
    if path.exists():
        return path

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    width = height = 8
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    return path


def _add_picture(doc, png: Path, *, width_cm: float = 6.0, height_cm: float = 4.0):
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_picture(str(png), width=Cm(width_cm), height=Cm(height_cm))
    return paragraph


def _para_texts(path: Path) -> list[str]:
    """深度取文：`paragraph.text` 读不到 w:fldSimple 里的编号数字。"""
    doc = Document(str(path))
    texts = []
    for paragraph in doc.paragraphs:
        text = "".join(node.text or "" for node in paragraph._p.iter(qn("w:t"))).strip()
        if text:
            texts.append(text)
    return texts


class CaptionNumberingTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.png = _write_png(self.tmp / "pic.png")
        self.addCleanup(self._tmp.cleanup)

    def _run(self, source: Path, **manifest_extra) -> dict:
        output = self.tmp / f"{source.stem}_captioned.docx"
        manifest_path = self.tmp / f"{source.stem}_manifest.json"
        manifest = {
            "schemaVersion": "bid-tech-caption-number-manifest-v1",
            "inputFile": str(source),
            "outputFile": str(output),
            **manifest_extra,
        }
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        return runner.run_manifest(manifest_path)

    def _build_body(self, path: Path) -> Path:
        """造一份方案 B 形态的成稿：目录/前言无章号 H1，正文 H1 带「第X章」文本编号。"""
        doc = Document()
        doc.add_paragraph("目录", style="Heading 1")
        doc.add_paragraph("前言  投标说明函", style="Heading 1")
        doc.add_paragraph("这是投标说明函正文。")

        doc.add_paragraph("第一章  标前概述", style="Heading 1")
        doc.add_paragraph("风场基本参数一览表")
        table = doc.add_table(rows=2, cols=3)
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                cell.text = f"数据{index}"
        _add_picture(doc, self.png)
        doc.add_paragraph("机组外形示意图")

        doc.add_paragraph("第二章  技术方案", style="Heading 1")
        _add_picture(doc, self.png)
        doc.add_paragraph("叶片关键尺寸")
        table2 = doc.add_table(rows=2, cols=2)
        for row in table2.rows:
            for index, cell in enumerate(row.cells):
                cell.text = f"值{index}"
        doc.save(str(path))
        return path

    def test_chapter_number_comes_from_heading_text(self) -> None:
        """章号取 H1 文本里的「第X章」，目录/前言不参与计数、其下不编号。"""
        source = self._build_body(self.tmp / "body.docx")
        result = self._run(source)

        self.assertEqual(result["status"], "completed")
        texts = _para_texts(Path(result["outputFile"]))
        self.assertIn("表1-1 风场基本参数一览表", texts)
        self.assertIn("图1-1 机组外形示意图", texts)
        # 第二章的表沿用第二章章号，且章内序号重新从 1 开始
        self.assertIn("表2-1 叶片关键尺寸", texts)
        self.assertEqual(result["summary"]["chapterCount"], 2)

    def test_figure_caption_yields_to_following_table(self) -> None:
        """夹在图和表之间的标题属于表：图注向下找候选时让给表格，自己另插一条。"""
        source = self._build_body(self.tmp / "arbitration.docx")
        result = self._run(source)
        texts = _para_texts(Path(result["outputFile"]))

        self.assertIn("表2-1 叶片关键尺寸", texts)
        self.assertIn("图2-1", texts)
        self.assertNotIn("图2-1 叶片关键尺寸", texts)

    def test_existing_label_is_stripped_before_renumbering(self) -> None:
        doc = Document()
        doc.add_paragraph("第三章  设备制造", style="Heading 1")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "项"
        table.rows[0].cells[1].text = "值"
        doc.element.body.insert(
            list(doc.element.body).index(table._tbl),
            Document().add_paragraph("表 11 制造基地产能总表")._p,
        )
        source = self.tmp / "old_label.docx"
        doc.save(str(source))

        result = self._run(source)
        texts = _para_texts(Path(result["outputFile"]))
        self.assertIn("表3-1 制造基地产能总表", texts)

    def test_seq_field_is_word_native(self) -> None:
        """SEQ 计数器名必须等于可见前缀，否则 Word 的交叉引用对话框里列不出来。"""
        source = self._build_body(self.tmp / "field.docx")
        result = self._run(source)

        doc = Document(result["outputFile"])
        instrs = [node.get(qn("w:instr")) for node in doc.element.body.iter(qn("w:fldSimple"))]
        self.assertTrue(instrs)
        self.assertIn(" SEQ 表 \\* ARABIC \\s 1 ", instrs)
        self.assertIn(" SEQ 图 \\* ARABIC \\s 1 ", instrs)
        # 域里要写缓存值，未按 F9 时显示也得是对的
        for field in doc.element.body.iter(qn("w:fldSimple")):
            cached = "".join(node.text or "" for node in field.iter(qn("w:t")))
            self.assertTrue(cached.isdigit(), f"SEQ 域缺少缓存值：{cached!r}")
        self.assertEqual(result["summary"]["numberingMode"], "seq_field")

    def test_update_fields_on_open_is_enabled(self) -> None:
        source = self._build_body(self.tmp / "update_fields.docx")
        result = self._run(source)

        settings = Document(result["outputFile"]).settings.element
        nodes = settings.findall(qn("w:updateFields"))
        self.assertEqual([node.get(qn("w:val")) for node in nodes], ["true"])

    def test_plain_mode_writes_text_number(self) -> None:
        source = self._build_body(self.tmp / "plain.docx")
        result = self._run(source, useWordFields=False)

        doc = Document(result["outputFile"])
        self.assertEqual(list(doc.element.body.iter(qn("w:fldSimple"))), [])
        self.assertIn("表1-1 风场基本参数一览表", _para_texts(Path(result["outputFile"])))
        self.assertEqual(result["summary"]["numberingMode"], "plain")

    def test_adjacent_image_paragraphs_share_one_caption(self) -> None:
        doc = Document()
        doc.add_paragraph("第一章  概述", style="Heading 1")
        _add_picture(doc, self.png)
        _add_picture(doc, self.png)
        _add_picture(doc, self.png)
        source = self.tmp / "grouped.docx"
        doc.save(str(source))

        result = self._run(source)
        self.assertEqual(result["summary"]["figureCount"], 1)
        self.assertIn("图1-1", _para_texts(Path(result["outputFile"])))

    def test_small_images_are_not_numbered(self) -> None:
        """宽<3cm 或高<2cm 的是图标/印章，不该占用图号。"""
        doc = Document()
        doc.add_paragraph("第一章  概述", style="Heading 1")
        _add_picture(doc, self.png, width_cm=1.0, height_cm=1.0)
        source = self.tmp / "icon.docx"
        doc.save(str(source))

        result = self._run(source)
        self.assertEqual(result["status"], "skipped")
        self.assertEqual(result["outputFile"], "")
        self.assertFalse((self.tmp / "icon_captioned.docx").exists())

    def test_layout_table_is_skipped(self) -> None:
        """含图单元格占比高且几乎没文字的是图片排版壳，不计表号。"""
        doc = Document()
        doc.add_paragraph("第一章  概述", style="Heading 1")
        table = doc.add_table(rows=1, cols=2)
        for cell in table.rows[0].cells:
            cell.paragraphs[0].add_run().add_picture(str(self.png), width=Cm(6), height=Cm(4))
        source = self.tmp / "layout_table.docx"
        doc.save(str(source))

        result = self._run(source)
        self.assertEqual(result["summary"]["tableCount"], 0)

    def test_headings_are_never_rewritten_as_captions(self) -> None:
        """`图表标题` 这类样式名含"标题"二字，Heading 判断必须精确匹配，否则整章题注全漏。"""
        numberer = captions.CaptionNumberer()
        self.assertTrue(numberer.is_heading_style("Heading 1"))
        self.assertTrue(numberer.is_heading_style("标题 2"))
        self.assertFalse(numberer.is_heading_style("图表标题"))
        self.assertFalse(numberer.is_heading_style("题注"))

    def test_chapter_number_falls_back_to_counting(self) -> None:
        """整篇 H1 都没有「第X章」文本编号时退回顺序计数，并如实报 warning。"""
        doc = Document()
        doc.add_paragraph("产品介绍", style="Heading 1")
        _add_picture(doc, self.png)
        source = self.tmp / "fallback.docx"
        doc.save(str(source))

        result = self._run(source)
        codes = [warning["code"] for warning in result["warnings"]]
        self.assertIn("caption_chapter_fallback", codes)
        self.assertIn("图1-1", _para_texts(Path(result["outputFile"])))

    def test_stale_manual_label_is_reported(self) -> None:
        """素材自带、写在图片上方的手写编号改写不到，必须报出来而不是静默留着。"""
        doc = Document()
        doc.add_paragraph("第二章  技术方案", style="Heading 1")
        doc.add_paragraph("图3-7 叶片结构")
        _add_picture(doc, self.png)
        source = self.tmp / "stale.docx"
        doc.save(str(source))

        result = self._run(source)
        codes = [warning["code"] for warning in result["warnings"]]
        self.assertIn("caption_number_gap", codes)

    def test_chinese_number_conversion(self) -> None:
        self.assertEqual(captions.chinese_number_to_int("一"), 1)
        self.assertEqual(captions.chinese_number_to_int("十"), 10)
        self.assertEqual(captions.chinese_number_to_int("十二"), 12)
        self.assertEqual(captions.chinese_number_to_int("二十三"), 23)
        self.assertEqual(captions.chinese_number_to_int("7"), 7)
        self.assertIsNone(captions.chinese_number_to_int("附录"))

    def test_output_must_differ_from_input(self) -> None:
        source = self._build_body(self.tmp / "same.docx")
        with self.assertRaises(ValueError):
            captions.number_captions(source, source)

    def test_caption_style_follows_heading_style_contract(self) -> None:
        """题注字体/字号取 heading_style.json 的 caption 段，不另立默认值。"""
        spec_path = self.tmp / "style.json"
        spec_path.write_text(
            json.dumps({"caption": {"zh_font": "楷体", "en_font": "Arial", "size_pt": 9, "line_spacing": 1.0}}),
            encoding="utf-8",
        )
        source = self._build_body(self.tmp / "styled.docx")
        result = self._run(source, styleSpecPath=str(spec_path))

        doc = Document(result["outputFile"])
        caption_paragraph = next(
            paragraph for paragraph in doc.paragraphs
            if "".join(node.text or "" for node in paragraph._p.iter(qn("w:t"))).startswith("表1-")
        )
        style = caption_paragraph.style
        self.assertEqual(style.font.size.pt, 9)
        self.assertEqual(style.element.get_or_add_rPr().get_or_add_rFonts().get(qn("w:eastAsia")), "楷体")


class CaptionNumberStepTest(unittest.TestCase):
    """S4 链路里的题注编号环节：三种状态都不许阻断出稿。"""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

        from app.services import tech_assembly

        self.tech_assembly = tech_assembly
        self.png = _write_png(self.tmp / "pic.png")

        doc = Document()
        doc.add_paragraph("第一章  概述", style="Heading 1")
        doc.add_paragraph("风场基本参数一览表")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "项"
        table.rows[0].cells[1].text = "值"
        self.source = self.tmp / "assembled.docx"
        doc.save(str(self.source))

    def _run_step(self):
        stages: list[str] = []
        step = self.tech_assembly._run_tech_caption_number_step(
            input_path=self.source,
            output_path=self.tmp / "assembled_captioned.docx",
            work_dir=self.tmp,
            progress_callback=lambda stage, _meta=None: stages.append(stage),
        )
        return step, stages

    def test_step_reports_completed_and_writes_manifest(self) -> None:
        step, stages = self._run_step()

        self.assertEqual(step["status"], "completed")
        self.assertTrue(Path(step["outputFile"]).exists())
        self.assertEqual(step["summary"]["tableCount"], 1)
        self.assertEqual(stages, ["calling_caption_number", "caption_number_completed"])

        manifest = json.loads((self.tmp / "tech_caption_number_input.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["schemaVersion"], "bid-tech-caption-number-manifest-v1")
        self.assertFalse(manifest["highlight"])
        self.assertTrue(Path(manifest["styleSpecPath"]).exists())

    def test_step_degrades_without_blocking(self) -> None:
        """题注编号失败时如实标 failed 并给出 warning，由调用方沿用组装原稿。"""
        original = self.tech_assembly._run_local_caption_number
        self.tech_assembly._run_local_caption_number = lambda *_args, **_kwargs: (_ for _ in ()).throw(
            RuntimeError("boom")
        )
        self.addCleanup(setattr, self.tech_assembly, "_run_local_caption_number", original)

        step, stages = self._run_step()
        self.assertEqual(step["status"], "failed")
        self.assertEqual(step["outputFile"], "")
        self.assertEqual([warning["code"] for warning in step["warnings"]], ["caption_number_failed"])
        self.assertEqual(stages, ["calling_caption_number", "caption_number_failed"])


if __name__ == "__main__":
    unittest.main()
